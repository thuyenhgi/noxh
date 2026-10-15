# -*- coding: utf-8 -*-
"""
HỆ THỐNG QUẢN LÝ HỒ SƠ & TRỘN BIỂU MẪU WORD - EXCEL NHCSXH
Ứng dụng tích hợp:
1. Giao diện trực quan (Tkinter GUI)
2. Khởi chạy Web Application (HTML / Local Server)
3. Tự động đồng bộ biểu mẫu Word (.docx) từ thư mục template/ vào HTML
4. Tự động nhận file Word tải lên từ web và lưu vật lý vào thư mục template/
5. Tự động trộn thư (Mail Merge) hàng loạt từ file Excel xuất ra (sheet "data")
6. Lưu trữ hồ sơ bền vững vào database_records.json
Có thể đóng gói thành file .exe độc lập chạy trên mọi máy tính không cài Python.
"""

import os, sys, glob, json, re, zipfile, base64, threading, webbrowser, time, shutil, math

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass
if sys.stderr and hasattr(sys.stderr, 'reconfigure'):
    try:
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

from http.server import HTTPServer, SimpleHTTPRequestHandler
import tkinter as tk
from tkinter import ttk, messagebox, filedialog

# Xác định thư mục gốc của chương trình & thư mục tài nguyên nhúng
if getattr(sys, 'frozen', False):
    BASE_DIR = os.path.dirname(sys.executable)
    BUNDLE_DIR = getattr(sys, '_MEIPASS', BASE_DIR)
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))
    BUNDLE_DIR = BASE_DIR

# Hàm hỗ trợ ẩn thư mục/tập tin trên Windows (FILE_ATTRIBUTE_HIDDEN = 0x02)
def set_folder_hidden(folder_path, hidden=True):
    try:
        if os.path.exists(folder_path):
            import ctypes
            attrs = 0x02 if hidden else 0x80 # 0x02 = HIDDEN, 0x80 = NORMAL
            ctypes.windll.kernel32.SetFileAttributesW(str(folder_path), attrs)
    except Exception:
        pass

def safe_unhide_file(path):
    """Gỡ thuộc tính ẩn (FILE_ATTRIBUTE_NORMAL = 0x80) trước khi ghi file trên Windows 
    để tránh lỗi PermissionError: [Errno 13] Permission denied."""
    try:
        if path and os.path.exists(path):
            import ctypes
            ctypes.windll.kernel32.SetFileAttributesW(str(path), 0x80)
    except Exception:
        pass

def safe_rehide_file(path):
    """Gán lại thuộc tính ẩn (FILE_ATTRIBUTE_HIDDEN = 0x02) sau khi ghi file xong (nếu chạy ở bản đóng gói)."""
    try:
        if getattr(sys, 'frozen', False) and path and os.path.exists(path):
            import ctypes
            ctypes.windll.kernel32.SetFileAttributesW(str(path), 0x02)
    except Exception:
        pass

def auto_hide_internal_files():
    """Tự động ẩn toàn bộ các tệp tin và thư mục phụ thuộc khi chạy bản Portable
    để người dùng chỉ nhìn thấy duy nhất tệp HeThongQuanLyHoSo.exe trong thư mục."""
    if getattr(sys, 'frozen', False):
        try:
            import ctypes
            items = [
                "_internal",
                "template",
                "he-thong-quan-ly-ho-so-demo.html",
                "xlsx.full.min.js",
                "template_metadata.json",
                "templates_bundle.json",
                "database_records.json",
                "database_records_rows206.json",
                "DATA_MAIL_MERGE_SYNC.xlsx",
            ]
            for it in items:
                p = os.path.join(BASE_DIR, it)
                if os.path.exists(p):
                    ctypes.windll.kernel32.SetFileAttributesW(str(p), 0x02)
        except Exception:
            pass

auto_hide_internal_files()

# Ưu tiên thư mục template bên ngoài (nếu có).
# Nếu không có hoặc bị xóa/không copy, tự động lấy template nhúng sẵn bên trong file .exe!
ext_tmpl = os.path.join(BASE_DIR, "template")
int_tmpl = os.path.join(BUNDLE_DIR, "template")

if os.path.exists(ext_tmpl) and len(os.listdir(ext_tmpl)) > 0:
    TEMPLATE_DIR = ext_tmpl
elif os.path.exists(int_tmpl):
    TEMPLATE_DIR = int_tmpl
else:
    TEMPLATE_DIR = ext_tmpl

HTML_FILE = os.path.join(BASE_DIR, "he-thong-quan-ly-ho-so-demo.html")
BUNDLE_FILE = os.path.join(BASE_DIR, "templates_bundle.json")
OUTPUT_DIR = os.path.join(BASE_DIR, "output_word")
DB_FILE = os.path.join(BASE_DIR, "database_records.json")

# ==============================================================================
# CẤU HÌNH PHIÊN BẢN ỨNG DỤNG & AUTO-UPDATE TRỰC TIẾP QUA GITHUB RAW
# ==============================================================================
APP_VERSION = "v2.2.0"
GITHUB_REPO = "thuyenhgi/noxh"
GITHUB_BRANCH = "main"
GITHUB_RAW_BASE = f"https://raw.githubusercontent.com/{GITHUB_REPO}/{GITHUB_BRANCH}"

update_status = {
    "status": "idle", # "idle", "checking", "downloading", "extracting", "completed", "error"
    "percent": 0,
    "message": "",
    "error": "",
    "has_new_exe": False
}

def is_newer_version(latest_tag, current_tag):
    """So sánh phiên bản 'v2.3.0' với 'v2.2.0', trả về True nếu latest_tag mới hơn."""
    try:
        def parse_v(t):
            t_clean = re.sub(r'^[^\d]*', '', str(t).strip())
            parts = [int(x) for x in re.findall(r'\d+', t_clean)]
            return parts if parts else [0]
        v_latest = parse_v(latest_tag)
        v_current = parse_v(current_tag)
        max_l = max(len(v_latest), len(v_current))
        v_latest += [0] * (max_l - len(v_latest))
        v_current += [0] * (max_l - len(v_current))
        return v_latest > v_current
    except Exception:
        return False

def check_github_release():
    """Kiểm tra bản cập nhật mới trực tiếp từ GitHub Raw qua tệp version.json."""
    import urllib.request, ssl, time
    v_url = f"{GITHUB_RAW_BASE}/version.json?t={int(time.time())}"
    try:
        ctx = ssl.create_default_context()
        req = urllib.request.Request(
            v_url,
            headers={
                "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)",
                "Cache-Control": "no-cache"
            }
        )
        with urllib.request.urlopen(req, timeout=10, context=ctx) as response:
            if response.status == 200:
                data = json.loads(response.read().decode('utf-8'))
                remote_ver = data.get("version", "v2.2.0")
                has_update = is_newer_version(remote_ver, APP_VERSION)
                files = data.get("files", ["he-thong-quan-ly-ho-so-demo.html", "HeThongQuanLyHoSo.py", "templates_bundle.json"])
                return {
                    "success": True,
                    "has_update": has_update,
                    "current_version": APP_VERSION,
                    "latest_version": remote_ver,
                    "changelog": data.get("changelog", "Bản cập nhật mã nguồn trực tiếp từ GitHub."),
                    "published_at": data.get("release_date", ""),
                    "files": files,
                    "raw_base": GITHUB_RAW_BASE
                }
    except Exception as e:
        err_str = str(e)
        if "404" in err_str:
            return {
                "success": True,
                "has_update": False,
                "current_version": APP_VERSION,
                "latest_version": APP_VERSION,
                "message": "Chưa tìm thấy tệp version.json trên nhánh main của GitHub. Hãy đưa tệp version.json lên repository."
            }
        return {"success": False, "error": f"Lỗi kết nối GitHub: {err_str}"}

def download_and_apply_update(files_to_download=None):
    """Tải trực tiếp từng tệp mã nguồn từ GitHub Raw và cập nhật mượt mà bảo vệ CSDL 100%."""
    global update_status
    import urllib.request, ssl, time
    
    EXCLUDED_DATA = {
        "database_records.json",
        "database_records_rows206.json",
        "DATA_MAIL_MERGE_SYNC.xlsx"
    }

    if not files_to_download:
        files_to_download = [
            "he-thong-quan-ly-ho-so-demo.html",
            "HeThongQuanLyHoSo.py",
            "templates_bundle.json",
            "version.json"
        ]

    # Lọc bỏ tuyệt đối các tệp CSDL khách hàng
    clean_files = [f for f in files_to_download if os.path.basename(f) not in EXCLUDED_DATA]
    total_files = len(clean_files)

    try:
        update_status = {
            "status": "downloading",
            "percent": 5,
            "message": f"Bắt đầu cập nhật {total_files} tệp từ GitHub...",
            "error": "",
            "has_new_exe": False
        }
        ctx = ssl.create_default_context()

        for idx, fname in enumerate(clean_files):
            pct = int(((idx) / total_files) * 90) + 5
            update_status["percent"] = pct
            update_status["message"] = f"Đang tải {fname} ({idx+1}/{total_files})..."

            file_url = f"{GITHUB_RAW_BASE}/{fname}?t={int(time.time())}"
            req = urllib.request.Request(file_url, headers={
                "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)",
                "Cache-Control": "no-cache"
            })

            target_path = os.path.join(BASE_DIR, fname)
            tmp_path = target_path + ".tmp"

            try:
                with urllib.request.urlopen(req, timeout=30, context=ctx) as resp:
                    if resp.status == 200:
                        content = resp.read()
                        if len(content) > 0:
                            safe_unhide_file(tmp_path)
                            with open(tmp_path, 'wb') as tf:
                                tf.write(content)
                            safe_unhide_file(target_path)
                            if os.path.exists(target_path):
                                os.remove(target_path)
                            os.replace(tmp_path, target_path)
                            safe_rehide_file(target_path)
            except Exception as fe:
                # Nếu file không tồn tại trên GitHub (404), bỏ qua tiếp tục các file khác
                if os.path.exists(tmp_path):
                    try: os.remove(tmp_path)
                    except Exception: pass
                continue

            # Nếu tải về templates.zip, giải nén đè trực tiếp vào thư mục template/
            if fname.lower() == "templates.zip" and os.path.exists(target_path):
                try:
                    update_status["message"] = "Đang giải nén biểu mẫu Word (templates.zip)..."
                    tmpl_dir = os.path.join(BASE_DIR, "template")
                    os.makedirs(tmpl_dir, exist_ok=True)
                    with zipfile.ZipFile(target_path, 'r') as tz:
                        tz.extractall(tmpl_dir)
                    try: os.remove(target_path)
                    except Exception: pass
                except Exception:
                    pass

            # Nếu tải về templates_bundle.json, giải nén đồng bộ ra thư mục template
            if fname == "templates_bundle.json" and os.path.exists(target_path):
                try:
                    update_status["message"] = "Đang đồng bộ lại mẫu biểu Word..."
                    with open(target_path, 'r', encoding='utf-8') as b_f:
                        b_data = json.load(b_f)
                    tmpl_dir = os.path.join(BASE_DIR, "template")
                    os.makedirs(tmpl_dir, exist_ok=True)
                    for item in b_data:
                        rel_path = item.get("path")
                        b64 = item.get("base64")
                        if rel_path and b64:
                            out_f_path = os.path.join(tmpl_dir, rel_path)
                            os.makedirs(os.path.dirname(out_f_path), exist_ok=True)
                            safe_unhide_file(out_f_path)
                            with open(out_f_path, 'wb') as wf:
                                wf.write(base64.b64decode(b64))
                            safe_rehide_file(out_f_path)
                except Exception:
                    pass

        update_status = {
            "status": "completed",
            "percent": 100,
            "message": "Cập nhật thành công toàn bộ mã nguồn từ GitHub!",
            "error": "",
            "has_new_exe": False
        }

    except Exception as e:
        update_status = {
            "status": "error",
            "percent": 0,
            "message": "Lỗi cập nhật",
            "error": str(e),
            "has_new_exe": False
        }

# ==============================================================================
# PHẦN 1: LOGIC ĐỒNG BỘ BIỂU MẪU WORD (SYNC TEMPLATES)
# ==============================================================================

def extract_docx_fields(docx_path):
    fields = set()
    try:
        with zipfile.ZipFile(docx_path, 'r') as z:
            xml_files = [fn for fn in z.namelist() if fn.startswith('word/') and fn.endswith('.xml')]
            for fn in xml_files:
                content = z.read(fn).decode('utf-8', errors='ignore')
                for m in re.finditer(r'MERGEFIELD\s+(?:"([^"]+)"|([^\s"\\<>]+))', content, re.IGNORECASE):
                    fld = m.group(1) or m.group(2)
                    if fld: fields.add(fld.strip())
                for m in re.finditer(r'[«\xab]\s*([^»\xbb\r\n]+?)\s*[»\xbb]', content):
                    if m.group(1): fields.add(m.group(1).strip())
    except Exception as e:
        print(f"Lỗi đọc {docx_path}: {e}")
    return sorted(list(fields))

TEMPLATE_META_FILE = os.path.join(BASE_DIR, "template_metadata.json")

def load_template_metadata():
    if os.path.exists(TEMPLATE_META_FILE):
        try:
            with open(TEMPLATE_META_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            pass
    return {}

def save_template_metadata(meta_dict):
    try:
        safe_unhide_file(TEMPLATE_META_FILE)
        with open(TEMPLATE_META_FILE, 'w', encoding='utf-8') as f:
            json.dump(meta_dict, f, ensure_ascii=False, indent=2)
        safe_rehide_file(TEMPLATE_META_FILE)
    except Exception:
        pass

def detect_metadata(filename):
    lower = filename.lower()
    m_code = re.match(r'^(\d+[a-zA-Z]?|Mẫu\s*[\w\d]+)', filename, re.IGNORECASE)
    code = f"Mẫu {m_code.group(1)}" if m_code else "Mẫu Word"
    
    project = "all"
    if "an phú" in lower or "an phu" in lower:
        project = "An Phú"
    elif "hồng loan" in lower or "hong loan" in lower:
        project = "Hồng Loan"
    elif "nam long" in lower:
        project = "Nam Long"
            
    if "độc thân" in lower or "doc than" in lower or "chưa kết hôn" in lower or "_dt" in lower:
        marital = "single"
    elif "có gia đình" in lower or "gia đình" in lower or "gia dinh" in lower or "vợ chồng" in lower or "cogd" in lower or "co_gd" in lower or "co gd" in lower:
        marital = "married"
    elif "1. r giay de nghi vay von" in lower and "cogd" not in lower and "gia đình" not in lower:
        # Mẫu 1 bản gốc không có thông tin vợ/chồng -> Bản Độc thân
        marital = "single"
    else:
        marital = "all"
        
    title = re.sub(r'\.docx?$', '', filename, flags=re.IGNORECASE)
    return code, project, marital, title

def get_template_info(filename):
    meta = load_template_metadata()
    if filename in meta:
        item = meta[filename]
        return (
            item.get("code") or "Mẫu Word",
            item.get("project") or "all",
            item.get("maritalTarget") or "all",
            item.get("title") or filename
        )
    return detect_metadata(filename)

def detect_project_from_record(rec, db_records_map=None):
    """
    Tự động nhận diện chuẩn xác 100% dự án của khách hàng:
    - Tra cứu trực tiếp từ trường Dự án / project trong record
    - Tra cứu đối soát CCCD / Họ tên từ database_records.json
    - Phân tích thông minh các trường địa chỉ nhà, căn hộ, ngân hàng, giải chấp
    """
    if not isinstance(rec, dict):
        return "An Phú"

    # 1. Trực tiếp kiểm tra trường Dự án / project trong record
    raw_p = str(rec.get("Dự án") or rec.get("duAn") or rec.get("project") or rec.get("tenDuAn") or rec.get("ten_du_an") or "").strip()
    if raw_p and raw_p.lower() not in ["none", "null", "all", "tất cả dự án", ""]:
        p_lower = raw_p.lower()
        if "an phú" in p_lower or "an phu" in p_lower or "ecocity" in p_lower:
            return "An Phú"
        elif "nam long" in p_lower:
            return "Nam Long"
        elif "hồng loan" in p_lower or "hong loan" in p_lower:
            return "Hồng Loan"
        else:
            return raw_p

    # 2. Tra cứu từ database_records.json nếu có db_records_map
    if db_records_map is None:
        db_records_map = {}
        if os.path.exists(DB_FILE):
            try:
                with open(DB_FILE, 'r', encoding='utf-8') as f:
                    db_list = json.load(f)
                    if isinstance(db_list, list):
                        for item in db_list:
                            if isinstance(item, dict):
                                c = str(item.get("cccd") or item.get("CCCD") or "").strip().replace(' ', '').replace('.', '').replace('-', '')
                                n = str(item.get("hoTen") or "").strip().lower()
                                if c: db_records_map[c] = item
                                if n: db_records_map[n] = item
            except Exception:
                pass

    c_id = str(rec.get("cccd") or rec.get("CCCD") or rec.get("soCCCD") or "").strip().replace(' ', '').replace('.', '').replace('-', '')
    n_id = str(rec.get("hoTen") or rec.get("Họ và tên KH thường") or rec.get("Họ và tên KH hoa") or "").strip().lower()
    matched_db = db_records_map.get(c_id) or db_records_map.get(n_id)

    if matched_db:
        db_p = str(matched_db.get("project") or matched_db.get("duAn") or matched_db.get("tenDuAn") or "").strip()
        if db_p and db_p.lower() not in ["none", "null", "all", ""]:
            p_lower = db_p.lower()
            if "an phú" in p_lower or "an phu" in p_lower or "ecocity" in p_lower:
                return "An Phú"
            elif "nam long" in p_lower:
                return "Nam Long"
            elif "hồng loan" in p_lower or "hong loan" in p_lower:
                return "Hồng Loan"
            else:
                return db_p

    # 3. Phân tích ngữ cảnh nội dung văn bản (Địa chỉ nhà, Dự án, Căn hộ, Ngân hàng, Giải chấp)
    text_parts = [
        str(rec.get("Địa chỉ nhà ở") or ""),
        str(rec.get("projectAddress") or ""),
        str(rec.get("projectBaseAddress") or ""),
        str(rec.get("diaChiNhaO") or ""),
        str(rec.get("Địa chỉ cư trú") or ""),
        str(rec.get("diaChiCuTru") or ""),
        str(rec.get("Căn hộ số") or ""),
        str(rec.get("canHoSo") or ""),
        str(rec.get("Tòa nhà") or ""),
        str(rec.get("QĐ SỞ XD") or ""),
        str(rec.get("tên NH") or ""),
        str(rec.get("taiNH") or ""),
        str(rec.get("projectBank") or ""),
        str(rec.get("giaiChapAnPhu") or ""),
        str(rec.get("Giải chấp") or ""),
        str(rec.get("Giải chấp Hồng Loan") or "")
    ]
    if matched_db:
        text_parts.extend([
            str(matched_db.get("projectAddress") or ""),
            str(matched_db.get("projectBaseAddress") or ""),
            str(matched_db.get("canHoSo") or ""),
            str(matched_db.get("projectBank") or ""),
            str(matched_db.get("giaiChapAnPhu") or "")
        ])

    blob = " ".join(text_parts).lower()
    if "an phú" in blob or "an phu" in blob or "ecocity" in blob or "thường thạnh" in blob or "thuong thanh" in blob or "bình xuyên" in blob or "binh xuyen" in blob:
        return "An Phú"
    if "nam long" in blob or ".cte" in blob or "hưng thạnh" in blob or "hung thanh" in blob:
        return "Nam Long"
    if "hồng loan" in blob or "hong loan" in blob:
        return "Hồng Loan"

    return "An Phú"

def determine_customer_marital_and_project(rec, db_records_map=None):
    # 1. Xác định Dự án
    matched_proj = detect_project_from_record(rec, db_records_map)
    if not matched_proj:
        matched_proj = "all"

    # 2. Xác định Hôn nhân (Độc thân vs Có gia đình)
    hon_nhan = str(rec.get("tinhTrangHonNhan") or rec.get("Tình trạng hôn nhân") or rec.get("honNhan") or "").strip().lower()
    vc_ho_ten = str(rec.get("Họ và tên Vợ/chồng") or rec.get("vcHoTen") or rec.get("Họ và tên vợ/chồng in hoa") or "").strip()
    vc_cccd = str(rec.get("CCCD Vợ/chồng") or rec.get("vcCCCD") or "").strip()

    if "độc thân" in hon_nhan or "doc than" in hon_nhan or "ly hôn" in hon_nhan or "chưa kết hôn" in hon_nhan or hon_nhan == "single":
        cust_marital = "single"
    elif "kết hôn" in hon_nhan or "gia đình" in hon_nhan or "gia dinh" in hon_nhan or hon_nhan == "married" or vc_ho_ten or vc_cccd:
        cust_marital = "married"
    else:
        cust_marital = "single" if not vc_ho_ten else "married"

    return matched_proj, cust_marital

def natural_sort_key(filename):
    m = re.match(r'^\s*0*(\d+)', filename)
    if m:
        num = int(m.group(1))
        if filename.strip().startswith('02'):
            return (1.5, filename)
        return (float(num), filename)
    return (999.0, filename)

def get_form_key(filename):
    """
    Trích xuất mã định danh duy nhất của biểu mẫu ngân hàng:
    '1', '02.TGTT', '2', '3', ..., '10', '11', '12'
    giúp khớp chính xác 100% biểu mẫu khi chọn tùy chọn.
    """
    clean = filename.strip()
    if clean.startswith('02.TGTT'):
        return '02.TGTT'
    m = re.match(r'^\s*(\d+)\.', clean)
    if m:
        return m.group(1)
    return clean.lower()

def get_case_template_folders():
    folders = []
    # 1. Quét trực tiếp các thư mục con trong TEMPLATE_DIR và BASE_DIR/template
    for base in [TEMPLATE_DIR, os.path.join(BASE_DIR, "template")]:
        if os.path.exists(base):
            for d in sorted(os.listdir(base)):
                p = os.path.join(base, d)
                if os.path.isdir(p) and not d.startswith(".") and d != "Template":
                    if d not in folders:
                        folders.append(d)
    # 2. Nếu chưa có, quét trong thư mục con Template/
    if not folders:
        for base in [os.path.join(TEMPLATE_DIR, "Template"), os.path.join(BASE_DIR, "template", "Template")]:
            if os.path.exists(base):
                for d in sorted(os.listdir(base)):
                    p = os.path.join(base, d)
                    if os.path.isdir(p) and d not in folders:
                        folders.append(d)
    return folders

def get_case_folder_path(case_folder):
    if not case_folder:
        return None
    candidates = [
        os.path.join(TEMPLATE_DIR, case_folder),
        os.path.join(BASE_DIR, "template", case_folder),
        os.path.join(TEMPLATE_DIR, "Template", case_folder),
        os.path.join(BASE_DIR, "template", "Template", case_folder),
        os.path.join(BUNDLE_DIR, "template", case_folder),
        os.path.join(BUNDLE_DIR, "template", "Template", case_folder),
    ]
    for c in candidates:
        if c and os.path.exists(c) and os.path.isdir(c):
            return c
    return None

def get_sorted_templates_in_folder(case_folder):
    path = get_case_folder_path(case_folder)
    if not path:
        return []
    files = [f for f in os.listdir(path) if f.lower().endswith('.docx') and not f.startswith('~$')]
    files.sort(key=natural_sort_key)
    return files

def strip_vn_accents(s):
    if not s: return ""
    s = str(s).lower()
    mapping = {
        'à': 'a', 'á': 'a', 'ả': 'a', 'ã': 'a', 'ạ': 'a',
        'ă': 'a', 'ằ': 'a', 'ắ': 'a', 'ẳ': 'a', 'ẵ': 'a', 'ặ': 'a',
        'â': 'a', 'ầ': 'a', 'ấ': 'a', 'ẩ': 'a', 'ẫ': 'a', 'ậ': 'a',
        'đ': 'd',
        'è': 'e', 'é': 'e', 'ẻ': 'e', 'ẽ': 'e', 'ẹ': 'e',
        'ê': 'e', 'ề': 'e', 'ế': 'e', 'ể': 'e', 'ễ': 'e', 'ệ': 'e',
        'ì': 'i', 'í': 'i', 'ỉ': 'i', 'ĩ': 'i', 'ị': 'i',
        'ò': 'o', 'ó': 'o', 'ỏ': 'o', 'õ': 'o', 'ọ': 'o',
        'ô': 'o', 'ồ': 'o', 'ố': 'o', 'ổ': 'o', 'ỗ': 'o', 'ộ': 'o',
        'ơ': 'o', 'ờ': 'o', 'ớ': 'o', 'ở': 'o', 'ỡ': 'o', 'ợ': 'o',
        'ù': 'u', 'ú': 'u', 'ủ': 'u', 'ũ': 'u', 'ụ': 'u',
        'ư': 'u', 'ừ': 'u', 'ứ': 'u', 'ử': 'u', 'ữ': 'u', 'ự': 'u',
        'ỳ': 'y', 'ý': 'y', 'ỷ': 'y', 'ỹ': 'y', 'ỵ': 'y'
    }
    for k, v in mapping.items():
        s = s.replace(k, v)
    return s

remove_accents = strip_vn_accents

def resolve_template_folder(rec, folders=None, db_records_map=None):
    if folders is None:
        folders = get_case_template_folders()
    if not folders:
        return ""

    cust_proj, cust_marital = determine_customer_marital_and_project(rec, db_records_map)

    proj_keys = []
    if "an phú" in cust_proj.lower() or "an phu" in cust_proj.lower():
        proj_keys = ["an phu", "an phú"]
    elif "hồng loan" in cust_proj.lower() or "hong loan" in cust_proj.lower():
        proj_keys = ["hong loan", "hồng loan"]
    elif "nam long" in cust_proj.lower():
        proj_keys = ["nam long"]
    else:
        proj_keys = [strip_vn_accents(cust_proj)]

    matched_by_proj = []
    for f in folders:
        f_norm = strip_vn_accents(f)
        if any(strip_vn_accents(k) in f_norm for k in proj_keys):
            matched_by_proj.append(f)

    if not matched_by_proj:
        matched_by_proj = folders

    if cust_marital == "married":
        matched = [f for f in matched_by_proj if any(k in strip_vn_accents(f) for k in ["co gia dinh", "gia dinh", "ket hon"])]
    else:
        matched = [f for f in matched_by_proj if any(k in strip_vn_accents(f) for k in ["doc than", "chua ket hon", "ly hon"])]

    if matched:
        rec_text = str(rec).lower()
        if "1 đứng tên" in rec_text or "1 dung ten" in rec_text or "đứng tên 1" in rec_text:
            for m in matched:
                if "1 đứng tên" in m or "1 dung ten" in strip_vn_accents(m): return m
        if "2 đứng tên" in rec_text or "2 dung ten" in rec_text:
            for m in matched:
                if "2 đứng tên" in m or "2 dung ten" in strip_vn_accents(m): return m
        if "tsn" in rec_text:
            for m in matched:
                if "tsn" in m.lower(): return m
        if "ssn" in rec_text:
            for m in matched:
                if "ssn" in m.lower(): return m
        return matched[0]

    for f in folders:
        if any(strip_vn_accents(k) in strip_vn_accents(f) for k in proj_keys):
            return f

    return folders[0]


def run_sync_templates(log_callback=print):
    log_callback("🔄 Bắt đầu đồng bộ biểu mẫu Word từ thư mục template/...")
    if not os.path.exists(TEMPLATE_DIR):
        os.makedirs(TEMPLATE_DIR, exist_ok=True)
        log_callback(f"⚠ Đã tạo thư mục template: {TEMPLATE_DIR}")
        return False, "Thư mục template/ chưa có file .docx nào."

    templates_list = []
    b64_dict = {}
    current_id = 0

    case_folders = get_case_template_folders()
    for fld in case_folders:
        p = get_case_folder_path(fld)
        if not p: continue
        flist = [f for f in os.listdir(p) if f.lower().endswith('.docx') and not f.startswith('~$')]
        flist.sort(key=natural_sort_key)
        
        log_callback(f"📂 Thư mục: {fld} ({len(flist)} biểu mẫu)")
        for idx, fname in enumerate(flist, 1):
            current_id += 1
            fpath = os.path.join(p, fname)
            code, proj, marital, title = get_template_info(fname)
            if proj == "all":
                if "an phú" in fld.lower() or "an phu" in fld.lower():
                    proj = "An Phú"
                elif "hồng loan" in fld.lower() or "hong loan" in fld.lower():
                    proj = "Hồng Loan"
                elif "nam long" in fld.lower():
                    proj = "Nam Long"
            if marital == "all":
                if "độc thân" in fld.lower() or "doc than" in fld.lower():
                    marital = "single"
                elif "gia đình" in fld.lower() or "kết hôn" in fld.lower():
                    marital = "married"

            fields = extract_docx_fields(fpath)
            with open(fpath, 'rb') as f:
                b64_str = base64.b64encode(f.read()).decode('utf-8')

            templates_list.append({
                "id": current_id,
                "folder": fld,
                "sort_order": idx,
                "filename": fname,
                "title": title,
                "code": code,
                "project": proj,
                "maritalTarget": marital,
                "fields_count": len(fields),
                "fields": fields
            })
            b64_dict[str(current_id)] = b64_str
            log_callback(f"   [{idx:02d}] {fname} ({len(fields)} biến)")

    # Nếu có file docx ở thư mục gốc
    root_docx = sorted(glob.glob(os.path.join(TEMPLATE_DIR, "*.docx")))
    root_docx = [f for f in root_docx if not os.path.basename(f).startswith("~$")]
    root_docx.sort(key=lambda x: natural_sort_key(os.path.basename(x)))
    for idx, fpath in enumerate(root_docx, 1):
        current_id += 1
        fname = os.path.basename(fpath)
        code, proj, marital, title = get_template_info(fname)
        fields = extract_docx_fields(fpath)
        with open(fpath, 'rb') as f:
            b64_str = base64.b64encode(f.read()).decode('utf-8')
        templates_list.append({
            "id": current_id,
            "folder": "Mặc định",
            "sort_order": idx,
            "filename": fname,
            "title": title,
            "code": code,
            "project": proj,
            "maritalTarget": marital,
            "fields_count": len(fields),
            "fields": fields
        })
        b64_dict[str(current_id)] = b64_str

    if not templates_list:
        log_callback("⚠ Không tìm thấy file .docx nào trong thư mục template/.")
        return False, "Không có file .docx trong thư mục template/."

    # 1. Cập nhật bundle JSON
    bundle_data = []
    for t in templates_list:
        b_item = dict(t)
        b_item["b64"] = b64_dict[str(t["id"])]
        bundle_data.append(b_item)
        
    with open(BUNDLE_FILE, "w", encoding="utf-8") as f:
        json.dump(bundle_data, f, ensure_ascii=False)

    # 2. Cập nhật trực tiếp vào HTML
    if not os.path.exists(HTML_FILE):
        log_callback(f"❌ Không tìm thấy file HTML: {HTML_FILE}")
        return False, f"Không tìm thấy file: {HTML_FILE}"

    with open(HTML_FILE, 'r', encoding='utf-8') as f:
        html = f.read()

    idx_tmpl_start = html.find('const DEFAULT_WORD_TEMPLATES = [')
    idx_cust_key = html.find('const CUSTOM_WORD_TEMPLATES_KEY = ')

    if idx_tmpl_start == -1 or idx_cust_key == -1:
        log_callback("❌ Không tìm thấy vị trí thẻ nhúng Word trong HTML.")
        return False, "Không tìm thấy vị trí nhúng Word trong file HTML."

    js_templates_str = json.dumps(templates_list, ensure_ascii=False, indent=2)
    b64_dict_str = json.dumps(b64_dict)

    replacement = f"""const DEFAULT_WORD_TEMPLATES = {js_templates_str};
    const DOCX_TEMPLATES_B64 = {b64_dict_str};
    """

    html = html[:idx_tmpl_start] + replacement + html[idx_cust_key:]

    with open(HTML_FILE, 'w', encoding='utf-8') as f:
        f.write(html)

    log_callback(f"✅ Đã đồng bộ thành công {len(templates_list)} biểu mẫu vào file HTML.")
    return True, f"Đã đồng bộ {len(templates_list)} mẫu Word thành công!"

# ==============================================================================
# PHẦN 2: LOGIC TRỘN THƯ HÀNG LOẠT TỪ EXCEL (MAIL MERGE BATCH)
# ==============================================================================
def escape_xml(text):
    if text is None: return ""
    s = str(text)
    escaped = s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('"', '&quot;').replace("'", '&apos;')
    return escaped.replace('\r\n', '</w:t><w:br/><w:t xml:space="preserve">').replace('\n', '</w:t><w:br/><w:t xml:space="preserve">')

def normalize_key(name):
    return re.sub(r'[^a-zA-Z0-9\u0080-\uffff]', '', str(name)).lower()

def format_admin_date(raw_dt):
    """
    Chuẩn hóa chuỗi ngày tháng sang dạng hành chính:
    - full: 'ngày DD tháng MM năm YYYY'
    - day: 'DD'
    - month: 'MM'
    - year: 'YYYY'
    Hỗ trợ các định dạng: 'YYYY-MM-DD', 'DD/MM/YYYY', 'DD-MM-YYYY', 'ngày DD tháng MM năm YYYY', datetime
    """
    if not raw_dt:
        return "", "", "", ""
    if hasattr(raw_dt, 'strftime'):
        d = str(raw_dt.day).zfill(2)
        m = str(raw_dt.month).zfill(2)
        y = str(raw_dt.year)
        return f"ngày {d} tháng {m} năm {y}", d, m, y
    s = str(raw_dt).strip()
    if not s or s.lower() in ['nan', 'none', 'null']:
        return "", "", "", ""
    
    # Kiểm tra dạng "ngày DD tháng MM năm YYYY"
    m_adm = re.search(r'ngày\s+(\d{1,2})\s+tháng\s+(\d{1,2})\s+năm\s+(\d{4})', s, re.IGNORECASE)
    if m_adm:
        d = m_adm.group(1).zfill(2)
        m = m_adm.group(2).zfill(2)
        y = m_adm.group(3)
        return f"ngày {d} tháng {m} năm {y}", d, m, y
    
    # Kiểm tra dạng YYYY-MM-DD
    m_iso = re.match(r'^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})', s)
    if m_iso:
        y = m_iso.group(1)
        m = m_iso.group(2).zfill(2)
        d = m_iso.group(3).zfill(2)
        return f"ngày {d} tháng {m} năm {y}", d, m, y
        
    # Kiểm tra dạng DD/MM/YYYY hoặc DD-MM-YYYY
    m_vn = re.match(r'^(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})', s)
    if m_vn:
        d = m_vn.group(1).zfill(2)
        m = m_vn.group(2).zfill(2)
        y = m_vn.group(3)
        return f"ngày {d} tháng {m} năm {y}", d, m, y

    return s, "", "", ""

def _read_num_vn(val):
    if not val:
        return ""
    try:
        n = int(round(float(str(val).replace(",", "").replace(".", "").replace(" ", "").replace("\t", "").strip())))
    except Exception:
        return ""
    if n == 0:
        return ""
    neg = n < 0
    abs_v = abs(n)
    d_arr = ["không", "một", "hai", "ba", "bốn", "năm", "sáu", "bảy", "tám", "chín"]
    sc_arr = ["", "ngàn", "triệu", "tỷ", "ngàn tỷ", "triệu tỷ", "tỷ tỷ"]

    def r3(num, f_hundred):
        a = num // 100
        b = (num % 100) // 10
        c = num % 10
        s = []
        if a > 0:
            s.append(d_arr[a] + " trăm")
        elif f_hundred:
            s.append("không trăm")
        if b == 0:
            if c > 0 and (a > 0 or f_hundred):
                s.append("lẻ")
        elif b == 1:
            s.append("mười")
        else:
            s.append(d_arr[b] + " mươi")
        if c > 0:
            w = d_arr[c]
            if b >= 2 and c == 1:
                w = "mốt"
            elif b >= 1 and c == 5:
                w = "lăm"
            elif c == 4:
                w = "bốn"
            s.append(w)
        return " ".join(s)

    grps = []
    tmp = abs_v
    while tmp > 0:
        grps.append(tmp % 1000)
        tmp //= 1000

    prts = []
    for i in range(len(grps) - 1, -1, -1):
        g = grps[i]
        if not g:
            continue
        f_h = (i < len(grps) - 1 and g < 100)
        t = r3(g, f_h)
        if sc_arr[i]:
            t += " " + sc_arr[i]
        prts.append(t)

    res = ", ".join(prts) + " đồng"
    res = re.sub(r'\b(mươi|mười|lẻ)\s+tư\b', r'\1 bốn', res, flags=re.IGNORECASE)
    res = re.sub(r'\b(tư)\s+(nghìn|ngàn|triệu|tỷ|đồng)\b', r'bốn \2', res, flags=re.IGNORECASE)
    res = re.sub(r'\b(Tư)\s+(nghìn|ngàn|triệu|tỷ|đồng)\b', r'Bốn \2', res, flags=re.IGNORECASE)
    if neg:
        res = "Âm " + res
    return res[0].upper() + res[1:]

def read_area_vietnamese(val):
    if not val:
        return ""
    raw = str(val).strip().replace(',', '.')
    try:
        f = float(raw)
    except Exception:
        return ""
    parts = raw.split('.')
    int_part = int(abs(float(parts[0]))) if parts[0] else 0
    if int_part == 0:
        int_words = "không"
    else:
        int_words = _read_num_vn(int_part)
        int_words = re.sub(r'\s+đồng$', '', int_words, flags=re.IGNORECASE).strip().lower()
        int_words = re.sub(r'\b(mươi)\s+bốn\b', r'\1 tư', int_words, flags=re.IGNORECASE)
    
    dec_words = ""
    if len(parts) > 1 and parts[1].strip():
        dec_str = re.sub(r'\D', '', parts[1].strip())
        if dec_str:
            if re.match(r'^0+$', dec_str):
                dec_part_w = "không"
            else:
                lead_zeros = len(dec_str) - len(dec_str.lstrip('0'))
                lead_w = ("không " * lead_zeros) if lead_zeros > 0 else ""
                rest_num = int(dec_str.lstrip('0'))
                rest_w = _read_num_vn(rest_num)
                rest_w = re.sub(r'\s+đồng$', '', rest_w, flags=re.IGNORECASE).strip().lower()
                if lead_zeros > 0 and rest_num == 1:
                    rest_w = "một"
                dec_part_w = (lead_w + rest_w).strip()
            dec_words = " phẩy " + dec_part_w

    full = (int_words + dec_words).strip()
    return full[0].upper() + full[1:] if full else ""

def get_rec_key(r):
    if not isinstance(r, dict):
        return f"invalid_{id(r)}"
    cccd = str(r.get('cccd') or r.get('CCCD') or r.get('soCCCD') or r.get('Số CCCD') or r.get('Số CMND/CCCD') or r.get('f_soCCCD') or r.get('cmnd') or r.get('soCMND') or '').strip().replace(' ', '').replace('.', '').replace('-', '').replace('\t', '')
    if cccd and len(cccd) >= 8:
        return f"cccd_{cccd}"
    name = str(r.get('hoTen') or r.get('f_hoTen') or r.get('Họ và tên KH hoa') or r.get('Họ và tên KH thường') or r.get('Họ và tên') or r.get('Họ tên') or r.get('tenKH') or '').strip().lower()
    can = str(r.get('canHoSo') or r.get('canHo') or r.get('Căn hộ số') or r.get('f_canHoSo') or '').strip().lower()
    proj = str(r.get('project') or r.get('duAn') or r.get('Dự án') or r.get('f_project') or '').strip().lower()
    if name:
        return f"name_{name}_{can}_{proj}"
    if r.get('_id'):
        return f"id_{r.get('_id')}"
    if r.get('id'):
        return f"id_{r.get('id')}"
    import uuid
    return f"auto_{uuid.uuid4().hex}"

HEADERS_206 = [
    'Dự án', 'project',
    'DX HOA', 'dx thường', 'Họ và tên KH hoa', 'Họ và tên KH thường', 'ngày, tháng, năm sinh',
    'Dân tộc', 'Giới tính', 'Căn cước công dân/Căn cước', 'Viết tắt Căn cước công dân/Căn cước', 'CCCD',
    'ngày cấp CCCD', 'nơi cấp CCCD', 'Địa chỉ cư trú', 'nơi ở hiện tại', 'Địa chỉ nhà ở',
    'nghề nghiệp', 'nơi công tác', 'mã BHXH', 'mã số thuế', 'Số điện thoại KH',
    'Số điện thoại vợ/chồng', 'DX vợ/chồng HOA', 'dx vợ/chồng thường', 'quan hệ với người vay',
    'tên CCCD/CC vợ/chồng viết đầy đủ', 'tên CCCD/CC vợ/chồng', 'Họ và tên Vợ/chồng',
    'Họ và tên vợ/chồng in hoa', 'dân tộc vợ/chồng', 'Giới tính vợ/chồng', 'Ngày,tháng,năm sinh vợ/chồng',
    'CCCD Vợ/chồng', 'Ngày cấp CCCD vợ/chồng', 'Nơi cấp CCCD vợ/chồng', 'Địa chỉ cư trú vợ/chồng',
    'nơi ở hiện tại vợ/chồng', 'Tổng thu nhập', 'Thu nhập người vay', 'Thu nhập vợ/chồng',
    'nghề nghiệp vợ/chồng', 'nơi làm việc vợ/chồng', 'Đối tượng', 'Chi phí gốc + lãi',
    'Tổng chi phí hàng tháng', 'Chi phí sinh hoạt', 'Chi phí người phụ thuộc', 'Chi phí khác',
    'mục đích vay', 'tổng số vốn thực hiện', 'Vốn tự có', 'Vốn vay', 'Bằng chữ vốn vay',
    'Thời hạn vay', 'Trả gốc', 'số tiền trả gốc hàng tháng', 'tháng cuối trả',
    'số tiền trả gốc cho mỗi kỳ hạn', 'số tiền trả gốc cho kỳ hạn cuối cùng', 'Ngày làm đơn vay vốn',
    'ngày nhận hs vay vốn', 'tháng nhận hs vay vốn', 'năm nhận hs vay vốn', 'cán bộ TD',
    'số ĐT CBTD', 'thẩm định từ ngày', 'thẩm định đến ngày', 'QĐ SỞ XD', 'số hđmb',
    'phụ lục', 'ngày phụ lục', 'số hđmb1', 'ngày hđmb1', 'ngày hđmb', 'ngày,tháng,năm Phụ lục 4',
    'Tổng giá trị HĐMB', 'bằng chữ giá trị HĐMB', 'Giá bán đã bao gồm thuế GTGT', 'thuế GTGT',
    'kinh phí bảo trì 2%', 'Bằng chữ kinh phí bảo trì 2%', 'Căn hộ số', 'Căn hộ số1', 'thuộc lô',
    'số tầng', 'Tòa nhà', 'DTSD (thông thủy)', 'Bằng chữ DTSD (thông thủy)', 'DT sàn',
    'DT thuộc sở hữu, sử dụng chung', 'Năm hoàn thành XD', 'vốn tự có 1', 'số tiền vốn tự có 1',
    'bằng chữ số tiền vốn tự có 1', 'vốn tự có 2', 'số tiền vốn tự có 2', 'bằng chữ số tiền vốn tự có 2',
    'vốn tự có 3', 'số tiền vốn tự có 3', 'bằng chữ số tiền vốn tự có 3', 'vốn tự có 4',
    'số tiền vốn tự có 4', 'bằng chữ số tiền vốn tự có 4', 'vốn tự có 5', 'số tiền vốn tự có 5',
    'bằng chữ số tiền vốn tự có 5', 'vốn tự có 6', 'số tiền vốn tự có 6', 'bằng chữ số tiền vốn tự có 6',
    'đóng số tiền lẻ', 'bằng chữ số tiền lẻ', 'ngày_đợt_1', 'số tiền đợt 1', 'bằng chữ số tiền đợt 1',
    'ngày đợt 2', 'số tiền đợt 2', 'bằng chữ số tiền đợt 2', 'Số tiền còn phải GN đợt 2',
    'Bằng chữ số tiền còn phải GN đợt 2', 'ngày đợt 3', 'số tiền đợt 3', 'bằng chữ số tiền đợt 3',
    'Số tiền còn phải GN đợt 3', 'Bằng chữ số tiền còn phải GN đợt 3', 'ngày đợt 4',
    'số tiền đợt 4', 'bằng chữ số tiền đợt 4', 'ngày đợt 5', 'số tiền đợt 5', 'bằng chữ số tiền đợt 5',
    'ngày đợt 6', 'số tiền đợt 6', 'bằng chữ số tiền đợt 6', 'ngày đợt 7', 'số tiền đợt 7',
    'bằng chữ số tiền đợt 7', 'số tiền đợt 7 (-KPBT)', 'bằng chữ số tiền đợt 7 (-KPBT)',
    'Còn_phải_GN_đợt_8', 'Bằng_chữ_còn_phải_GN_đợt_8', 'ngày đợt 8', 'số tiền đợt 8',
    'bằng chữ số tiền đợt 8', 'bằng_chữ_số_tiền_đợt_8', 'số_tiền_đợt_9', 'bằng_chữ_số_tiền_đợt_9',
    'ngày,tháng,năm tb phê duyệt cho vay', 'ngày tb phê duyệt cho vay', 'ngày biên bản xác định',
    'ngày hđtd', 'ngày/tháng/năm hđtd', 'ngày vay', 'hạn trả', 'ngày GN đầu tiên',
    'số tiền GN đầu tiên', 'bằng chữ số tiền GN đầu tiên', 'ngày GN lần 2', 'ngày GN lần 3',
    'ngày GN lần 4', 'ngày GN lần 5', 'ngày GN lần 6', 'ngày GN lần 7', 'ngày GN lần 8',
    'số tiền nhận món vay cuối cùng', 'bằng chữ số tiền cuối cùng', 'ngày hđ3b',
    'ngày,tháng,năm hđ3b', 'ngày hđtc', 'ngày,tháng,năm hđtc', 'Số tk', 'tên NH',
    'số tb - ngày tb', 'v/v thanh toán', 'số tiền tb', 'bằng chữ số tiền tb',
    'số tb - ngày tb1', 'v/v thanh toán1', 'số tiền tb1', 'bằng chữ số tiền tb1',
    'số tb - ngày tb2', 'v/v thanh toán2', 'số tiền tb2', 'bằng chữ số tiền tb2',
    'số tb - ngày tb3', 'v/v thanh toán3', 'số tiền tb3', 'bằng chữ số tiền tb3',
    'số tb - ngày tb lần 2', 'v/v thanh toán lần 2', 'số tiền tb lần 2', 'bằng chữ số tiền tb lần 2',
    'dư nợ trước GN lần 2', 'bằng chữ dư nợ trước GN lần 2', 'số tiền đã GN lần 1',
    'bằng chữ số tiền đã GN lần 1', 'tổng tiền đã GN lần 2', 'bằng chữ tổng tiền đã GN lần 2',
    'số tb - ngày tb lần 3', 'v/v thanh toán lần 3', 'số tiền tb lần 3', 'bằng chữ số tiền tb lần 3',
    'dư nợ trước GN lần 3', 'bằng chữ dư nợ trước GN lần 3', 'số tiền đã GN lần 2',
    'bằng chữ số tiền đã GN lần 2', 'tổng tiền đã GN lần 3', 'bằng chữ tổng tiền đã GN lần 3',
    'Giải chấp', 'Giải chấp Hồng Loan', 'mã khách hàng', 'số tài khoản 102', 'thông tin phụ thuộc',
    'danh xưng cán bộ tín dụng', 'số hđmb 1', 'ngày hđmb 1', 'Giá bán da bao gồm thuế GTGT',
    'danh xưng cán bộ', 'gt_cbtd', 'Kế hoạch giải ngân',
    'Số hợp đồng tín dụng', 'Số hợp đồng thế chấp', 'Số biên bản xác định đấu giá',
    'vốn tự có n',
    'Ngày bb kiểm tra', 'Năm tài chính', 'Tổng số tiền đã giải ngân',
    'Lũy kế số nợ gốc đã trả', 'Dư nợ vay', 'bằng chữ dư nợ vay', 'giá trị thị trường',
    'Ngày tờ trình', 'ngày tờ trình', 'Tổng số tiền đã giải ngân TTr', 'Tổng số tiền còn phải giải ngân',
    'tên đồng sở hữu hoa', 'tên đồng sở hữu thường',
    'tổng_hợp_vốn_tự_có', 'tổng_hợp_phiếu_báo_có', 'vốn_tự_có_n',
    'Ngày_tháng_năm_BBKT', 'Năm_HĐTD', 'Năm_HĐTC', 'Lũy_kế_số_nợ_gốc_đã_trả_BBKT',
    'Dư_nợ_vay_hiện_tại_BBKT', 'Bằng_chữ_Lũy_kế_số_nợ_gốc_đã_trả_BBKT1',
    'Giá_trị_hiện_tại_của_TSTC_BBKT', 'Danh_Xưng_CBTD_BBKT', 'Tên_CBTD_BBKT',
    'Dư_nợ_trước_ngày_nhận_nợ_L2', 'Tổng_số_tiền_còn_phải_giải_ngân',
    'Giấy_đề_nghị_GN_lần', 'Số_tiền_gn', 'Bằng_chữ_Số_tiền_gn',
    'Đã_giải_ngân_số_tiền', 'Bằng_chữ_Đã_giải_ngân_số_tiền',
    'Dư_nợ_đến_trước_ngày_nhận_nợ', 'Bằng_chữ_Dư_nợ_đến_trước_ngày_nhận_nợ',
    'Số_tài_khoản_trên_TBGN', 'Tên_người_thụ_hưởng', 'Tại_NH',
    'Thông_báo_Bàn_giao_GCNQSDD', 'Số_Cif',
    'số_tiền_trả_gốc_hàng_tháng', 'tháng_cuối_trả',
    'số_tiền_trả_gốc_cho_mỗi_kỳ_hạn', 'số_tiền_trả_gốc_cho_kỳ_hạn_cuối_cùng'
]


def format_title_case_py(s):
    if not s:
        return ""
    words = str(s).strip().split()
    return " ".join([w.capitalize() for w in words])

def compute_dong_so_huu_py(r):
    is_co_owner = bool(r.get("isDongSoHuu") or r.get("dong_so_huu") or r.get("đồng sở hữu"))
    tt_hn = str(r.get("tinhTrangHonNhan") or r.get("Tình trạng hôn nhân") or "").strip()
    is_married = (tt_hn == "Đã kết hôn" or not tt_hn)

    raw_kh = str(r.get("hoTen") or r.get("Họ và tên KH hoa") or r.get("Họ và tên KH thường") or r.get("Họ và tên") or "").strip()
    raw_vc = str(r.get("vcHoTen") or r.get("Họ và tên Vợ/chồng") or "").strip()

    name_kh = format_title_case_py(raw_kh)
    name_vc = format_title_case_py(raw_vc)

    dx_kh = str(r.get("danhXung") or r.get("f_danhXung") or r.get("DX HOA") or "Ông").strip().lower()
    dx_hoa_kh = "Bà" if "bà" in dx_kh else "Ông"
    dx_thuong_kh = "bà" if "bà" in dx_kh else "ông"

    dx_vc = str(r.get("vcDanhXung") or r.get("f_vcDanhXung") or r.get("danh xưng vợ chồng") or "Bà").strip().lower()
    dx_hoa_vc = "Ông" if "ông" in dx_vc else "Bà"
    dx_thuong_vc = "ông" if "ông" in dx_vc else "bà"

    if not name_kh:
        return {"hoa": "", "thuong": ""}

    if is_co_owner and is_married and name_vc:
        return {
            "hoa": f"{dx_hoa_kh} {name_kh} và {dx_hoa_vc} {name_vc} với",
            "thuong": f"{dx_thuong_kh} {name_kh} và {dx_thuong_vc} {name_vc} với"
        }
    else:
        return {
            "hoa": f"{dx_hoa_kh} {name_kh} và",
            "thuong": f"{dx_thuong_kh} {name_kh} và"
        }

def _parse_m(val):
    if val is None or val == "":
        return 0
    try:
        cleaned = str(val).replace(".", "").replace(",", "").replace(" ", "").replace("\t", "").strip()
        return int(round(float(cleaned))) if cleaned else 0
    except Exception:
        return 0

def _fmt_m(val):
    n = _parse_m(val)
    return f"{n:,}".replace(",", ".") if n > 0 else ""

def _word_m(val):
    n = _parse_m(val)
    return _read_num_vn(n) if n > 0 and '_read_num_vn' in globals() else ""

def _fmt_date(val):
    s = str(val or "").strip()
    if not s: return ""
    m_admin = re.search(r'(?:[Nn]gày\s*)?(\d{1,2})\s*tháng\s*(\d{1,2})\s*năm\s*(\d{4})', s, re.IGNORECASE)
    if m_admin:
        return f"{int(m_admin.group(1)):02d}/{int(m_admin.group(2)):02d}/{m_admin.group(3)}"
    if re.match(r'^\d{4}-\d{2}-\d{2}$', s):
        parts = s.split('-')
        return f"{int(parts[2]):02d}/{int(parts[1]):02d}/{parts[0]}"
    if re.match(r'^\d{1,2}/\d{1,2}/\d{4}$', s):
        parts = s.split('/')
        return f"{int(parts[0]):02d}/{int(parts[1]):02d}/{parts[2]}"
    return s

def build_row206_from_record(r):
    if not isinstance(r, dict):
        return {}

    row = dict(r)

    def _is_new_cc(val):
        s = str(val or "").strip()
        if not s: return False
        m = re.match(r'^(\d{4})-(\d{2})-(\d{2})', s)
        if m:
            return f"{m.group(1)}-{m.group(2)}-{m.group(3)}" >= "2024-07-01"
        m2 = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{4})', s)
        if m2:
            d, mth, y = m2.group(1).zfill(2), m2.group(2).zfill(2), m2.group(3)
            return f"{y}-{mth}-{d}" >= "2024-07-01"
        return False

    dx = str(r.get('danhXung') or r.get('dx thường') or '').strip().lower()
    row["DX HOA"] = row.get("DX HOA") or ("Bà" if dx == "bà" else ("Ông" if dx == "ông" else ""))
    row["dx thường"] = row.get("dx thường") or dx
    ho_ten = str(r.get('hoTen') or r.get('Họ và tên KH thường') or r.get('Họ và tên') or '').strip()
    row["Họ và tên KH thường"] = row.get("Họ và tên KH thường") or ho_ten
    row["Họ và tên KH hoa"] = row.get("Họ và tên KH hoa") or ho_ten.upper()
    row["ngày, tháng, năm sinh"] = row.get("ngày, tháng, năm sinh") or _fmt_date(r.get('ngaySinh'))
    row["Dân tộc"] = row.get("Dân tộc") or r.get('danToc') or "Kinh"
    row["Giới tính"] = row.get("Giới tính") or r.get('gioiTinh') or ("Nữ" if dx == "bà" else "Nam")
    cccd = str(r.get('cccd') or r.get('CCCD') or r.get('soCCCD') or '').strip()
    row["CCCD"] = row.get("CCCD") or cccd
    kh_is_new = _is_new_cc(r.get('ngayCapCCCD'))
    kh_doc_full = "Căn cước" if kh_is_new else "Căn cước công dân"
    kh_doc_short = "CC" if kh_is_new else "CCCD"
    kh_doc_place = "Bộ Công an" if kh_is_new else "Cục CS QLHC"
    if kh_is_new and (not r.get('cccdDayDu') or r.get('cccdDayDu') == 'Căn cước công dân'):
        row["Căn cước công dân/Căn cước"] = kh_doc_full if cccd else ""
    else:
        row["Căn cước công dân/Căn cước"] = row.get("Căn cước công dân/Căn cước") or r.get('cccdDayDu') or (kh_doc_full if cccd else "")
    if kh_is_new and (not r.get('cccdVietTat') or r.get('cccdVietTat') == 'CCCD'):
        row["Viết tắt Căn cước công dân/Căn cước"] = kh_doc_short if cccd else ""
    else:
        row["Viết tắt Căn cước công dân/Căn cước"] = row.get("Viết tắt Căn cước công dân/Căn cước") or r.get('cccdVietTat') or (kh_doc_short if cccd else "")
    row["ngày cấp CCCD"] = row.get("ngày cấp CCCD") or _fmt_date(r.get('ngayCapCCCD'))
    row["nơi cấp CCCD"] = row.get("nơi cấp CCCD") or r.get('noiCapCCCD') or (kh_doc_place if cccd else "")
    row["Địa chỉ cư trú"] = row.get("Địa chỉ cư trú") or r.get('diaChiCuTru') or ""
    row["nơi ở hiện tại"] = row.get("nơi ở hiện tại") or r.get('noiOHienTai') or row["Địa chỉ cư trú"]
    row["Địa chỉ nhà ở"] = row.get("Địa chỉ nhà ở") or r.get('projectAddress') or ""
    row["nghề nghiệp"] = row.get("nghề nghiệp") or r.get('ngheNghiep') or ""
    row["nơi công tác"] = row.get("nơi công tác") or r.get('noiCongTac') or ""
    row["mã BHXH"] = row.get("mã BHXH") or r.get('maBHXH') or ""
    row["mã số thuế"] = row.get("mã số thuế") or r.get('maSoThue') or ""
    row["Số điện thoại KH"] = row.get("Số điện thoại KH") or r.get('sdtKH') or ""
    row["Số điện thoại vợ/chồng"] = row.get("Số điện thoại vợ/chồng") or r.get('vcSDT') or ""

    tt_hn = str(r.get('tinhTrangHonNhan') or r.get('Tình trạng hôn nhân') or r.get('honNhan') or '').strip()
    is_single = is_single_profile(r)
    if is_single:
        tt_hn = tt_hn or "Độc thân"
    row["tinhTrangHonNhan"] = tt_hn
    row["Tình trạng hôn nhân"] = tt_hn

    vc_dx = "" if is_single else str(r.get('vcDanhXung') or r.get('dx vợ/chồng thường') or '').strip().lower()
    row["DX vợ/chồng HOA"] = "" if is_single else (row.get("DX vợ/chồng HOA") or ("Bà" if vc_dx == "bà" else ("Ông" if vc_dx == "ông" else "")))
    row["dx vợ/chồng thường"] = "" if is_single else (row.get("dx vợ/chồng thường") or vc_dx)
    row["quan hệ với người vay"] = "" if is_single else (row.get("quan hệ với người vay") or r.get('vcQuanHe') or ("Vợ" if vc_dx == "bà" else ("Chồng" if vc_dx == "ông" else "")))
    vc_hoten = "" if is_single else str(r.get('vcHoTen') or r.get('Họ và tên Vợ/chồng') or '').strip()
    row["Họ và tên Vợ/chồng"] = vc_hoten
    row["Họ và tên vợ/chồng in hoa"] = vc_hoten.upper() if vc_hoten else ""
    row["dân tộc vợ/chồng"] = "" if is_single else (row.get("dân tộc vợ/chồng") or r.get('vcDanToc') or "")
    row["Giới tính vợ/chồng"] = "" if is_single else (row.get("Giới tính vợ/chồng") or r.get('vcGioiTinh') or ("Nữ" if vc_dx == "bà" else ("Nam" if vc_dx == "ông" else "")))
    row["Ngày,tháng,năm sinh vợ/chồng"] = "" if is_single else (row.get("Ngày,tháng,năm sinh vợ/chồng") or _fmt_date(r.get('vcNgaySinh')))
    vc_cccd = "" if is_single else str(r.get('vcCCCD') or r.get('CCCD Vợ/chồng') or '').strip()
    row["CCCD Vợ/chồng"] = vc_cccd
    vc_is_new = _is_new_cc(r.get('vcNgayCapCCCD'))
    vc_doc_full = "Căn cước" if vc_is_new else "Căn cước công dân"
    vc_doc_short = "CC" if vc_is_new else "CCCD"
    vc_doc_place = "Bộ Công an" if vc_is_new else "Cục CS QLHC"
    if vc_is_new and (not r.get('vcCCCDDayDu') or r.get('vcCCCDDayDu') == 'Căn cước công dân'):
        row["tên CCCD/CC vợ/chồng viết đầy đủ"] = "" if is_single else (vc_doc_full if vc_cccd else "")
    else:
        row["tên CCCD/CC vợ/chồng viết đầy đủ"] = "" if is_single else (row.get("tên CCCD/CC vợ/chồng viết đầy đủ") or r.get('vcCCCDDayDu') or (vc_doc_full if vc_cccd else ""))
    if vc_is_new and (not r.get('vcCCCDVietTat') or r.get('vcCCCDVietTat') == 'CCCD'):
        row["tên CCCD/CC vợ/chồng"] = "" if is_single else (vc_doc_short if vc_cccd else "")
    else:
        row["tên CCCD/CC vợ/chồng"] = "" if is_single else (row.get("tên CCCD/CC vợ/chồng") or r.get('vcCCCDVietTat') or (vc_doc_short if vc_cccd else ""))
    row["Ngày cấp CCCD vợ/chồng"] = "" if is_single else (row.get("Ngày cấp CCCD vợ/chồng") or _fmt_date(r.get('vcNgayCapCCCD')))
    row["Nơi cấp CCCD vợ/chồng"] = "" if is_single else (row.get("Nơi cấp CCCD vợ/chồng") or r.get('vcNoiCap') or (vc_doc_place if vc_cccd else ""))
    row["Địa chỉ cư trú vợ/chồng"] = "" if is_single else (row.get("Địa chỉ cư trú vợ/chồng") or r.get('vcDiaChiCuTru') or "")
    row["nơi ở hiện tại vợ/chồng"] = "" if is_single else (row.get("nơi ở hiện tại vợ/chồng") or r.get('vcNoiOHienTai') or row["Địa chỉ cư trú vợ/chồng"])
    row["nghề nghiệp vợ/chồng"] = "" if is_single else (row.get("nghề nghiệp vợ/chồng") or r.get('vcNgheNghiep') or "")
    row["nơi làm việc vợ/chồng"] = "" if is_single else (row.get("nơi làm việc vợ/chồng") or r.get('vcNoiCongTac') or r.get('vcNoiLamViec') or "")
    row["nơi_làm_việc_vợchồng"] = row["nơi làm việc vợ/chồng"]
    row["noi_lam_viec_vo_chong"] = row["nơi làm việc vợ/chồng"]
    row["nơi công tác vợ/chồng"] = row["nơi làm việc vợ/chồng"]
    row["nơi_công_tác_vợchồng"] = row["nơi làm việc vợ/chồng"]

    tn_kh = _parse_m(r.get('thuNhapNguoiVay') or r.get('Thu nhập người vay') or r.get('thuNhapKH'))
    tn_vc = 0 if is_single else _parse_m(r.get('thuNhapVoChong') or r.get('Thu nhập vợ/chồng') or r.get('vcThuNhap'))
    tot_tn = tn_kh if is_single else (_parse_m(r.get('tongThuNhap') or r.get('Tổng thu nhập')) or (tn_kh + tn_vc))
    row["Tổng thu nhập"] = _fmt_m(tot_tn)
    row["Thu nhập người vay"] = _fmt_m(tn_kh)
    row["Thu nhập vợ/chồng"] = "" if is_single else _fmt_m(tn_vc)
    raw_dt = row.get("Đối tượng") or r.get('doiTuong') or r.get('Đối tượng') or ""
    clean_dt = str(raw_dt).replace('\r\n', '\n').replace('\r', '\n')
    clean_dt = re.sub(r'<br\s*/?>', '\n', clean_dt, flags=re.IGNORECASE)
    clean_dt = re.sub(r'</p>', '\n', clean_dt, flags=re.IGNORECASE)
    clean_dt = re.sub(r'<p\b[^>]*>', '', clean_dt, flags=re.IGNORECASE)
    dt_lines = [l.strip() for l in clean_dt.split('\n') if l.strip()]
    normalized_dt = '\r\n'.join(dt_lines)
    row["Đối tượng"] = normalized_dt
    row["Đối_tượng"] = normalized_dt
    row["doi_tuong"] = normalized_dt
    row["doiTuong"] = normalized_dt
    row["Chi phí gốc + lãi"] = row.get("Chi phí gốc + lãi") or _fmt_m(r.get('chiPhiGocLai'))
    cp_sh = _parse_m(r.get('chiPhiSinhHoat') or r.get('Chi phí sinh hoạt'))
    cp_pt = _parse_m(r.get('chiPhiNguoiPhuThuoc') or r.get('Chi phí người phụ thuộc'))
    cp_khac = _parse_m(r.get('chiPhiKhac') or r.get('Chi phí khác'))
    calc_cp = cp_sh + cp_pt + cp_khac
    row["Tổng chi phí hàng tháng"] = row.get("Tổng chi phí hàng tháng") or _fmt_m(calc_cp or r.get('tongChiPhi'))
    row["Chi phí sinh hoạt"] = row.get("Chi phí sinh hoạt") or _fmt_m(cp_sh)
    row["Chi phí người phụ thuộc"] = row.get("Chi phí người phụ thuộc") or _fmt_m(cp_pt)
    row["Chi phí khác"] = row.get("Chi phí khác") or _fmt_m(cp_khac)
    row["mục đích vay"] = row.get("mục đích vay") or r.get('mucDichVay') or "Mua nhà ở xã hội"
    row["tổng số vốn thực hiện"] = row.get("tổng số vốn thực hiện") or _fmt_m(r.get('tongGiaTriHDMB') or r.get('tongGiaTriHD'))
    row["Vốn tự có"] = row.get("Vốn tự có") or _fmt_m(r.get('vonTuCoDaBaoGomKPBT') or r.get('vonTuCo'))
    vv_val = _parse_m(r.get('vonVay') or r.get('Vốn vay'))
    row["Vốn vay"] = row.get("Vốn vay") or _fmt_m(vv_val)
    row["Bằng chữ vốn vay"] = row.get("Bằng chữ vốn vay") or r.get('vayBangChu') or _word_m(vv_val)
    row["Thời hạn vay"] = row.get("Thời hạn vay") or str(r.get('thoiHanVay') or "300")
    row["Trả gốc"] = row.get("Trả gốc") or str(r.get('traGocDinhKy') or "06")

    # Trả gốc hàng tháng & kỳ hạn
    thoi_han_str = str(r.get('thoiHanVay') or r.get('Thời hạn vay') or "300").strip()
    thoi_han_num = int(_parse_m(thoi_han_str)) if _parse_m(thoi_han_str) > 0 else 300
    calc_goc_ht = math.ceil(vv_val / thoi_han_num / 1000) * 1000 if (vv_val > 0 and thoi_han_num > 0) else 0
    calc_thang_cuoi = (vv_val - calc_goc_ht * (thoi_han_num - 1)) if (vv_val > 0 and thoi_han_num > 0) else 0
    calc_goc_ky_han = round(vv_val / 50) if vv_val > 0 else 0
    calc_goc_ky_han_cuoi = (vv_val - calc_goc_ky_han * 49) if vv_val > 0 else 0

    goc_ht_val = _parse_m(r.get('traGocHangThang') or r.get('số tiền trả gốc hàng tháng') or r.get('số_tiền_trả_gốc_hàng_tháng') or row.get('số tiền trả gốc hàng tháng'))
    if goc_ht_val <= 0:
        goc_ht_val = calc_goc_ht

    thang_cuoi_val = _parse_m(r.get('thangCuoiTra') or r.get('tháng cuối trả') or r.get('tháng_cuối_trả') or row.get('tháng cuối trả'))
    if thang_cuoi_val <= 0:
        thang_cuoi_val = calc_thang_cuoi

    goc_ky_han_val = _parse_m(r.get('traGocKyHan') or r.get('số tiền trả gốc cho mỗi kỳ hạn') or row.get('số tiền trả gốc cho mỗi kỳ hạn'))
    if goc_ky_han_val <= 0:
        goc_ky_han_val = calc_goc_ky_han

    goc_ky_han_cuoi_val = _parse_m(r.get('traGocKyHanCuoi') or r.get('số tiền trả gốc cho kỳ hạn cuối cùng') or row.get('số tiền trả gốc cho kỳ hạn cuối cùng'))
    if goc_ky_han_cuoi_val <= 0:
        goc_ky_han_cuoi_val = (vv_val - goc_ky_han_val * 49) if (vv_val > 0 and goc_ky_han_val > 0) else calc_goc_ky_han_cuoi

    s_goc_ht = f"{int(goc_ht_val):,}".replace(",", ".") if goc_ht_val > 0 else ""
    s_thang_cuoi = f"{int(thang_cuoi_val):,}".replace(",", ".") if thang_cuoi_val > 0 else ""
    s_goc_ky_han = f"{int(goc_ky_han_val):,}".replace(",", ".") if goc_ky_han_val > 0 else ""
    s_goc_ky_han_cuoi = f"{int(goc_ky_han_cuoi_val):,}".replace(",", ".") if goc_ky_han_cuoi_val > 0 else ""

    row["số tiền trả gốc hàng tháng"] = s_goc_ht
    row["số_tiền_trả_gốc_hàng_tháng"] = s_goc_ht
    row["tháng cuối trả"] = s_thang_cuoi
    row["tháng_cuối_trả"] = s_thang_cuoi
    row["số tiền trả gốc cho mỗi kỳ hạn"] = s_goc_ky_han
    row["số_tiền_trả_gốc_cho_mỗi_kỳ_hạn"] = s_goc_ky_han
    row["số tiền trả gốc cho kỳ hạn cuối cùng"] = s_goc_ky_han_cuoi
    row["số_tiền_trả_gốc_cho_kỳ_hạn_cuối_cùng"] = s_goc_ky_han_cuoi

    # 4. TIẾP NHẬN HỒ SƠ & NGÀY LÀM ĐƠN VAY VỐN (Cột 61-64)
    raw_nl_don = str(r.get("ngayTiepNhan") or r.get("ngayLamDon") or r.get("ngayLamDonVay") or r.get("Ngày làm đơn vay vốn") or r.get("ngayNhanHs") or r.get("Ngày tiếp nhận") or "").strip()
    full_nl, d_nl, m_nl, y_nl = format_admin_date(raw_nl_don)
    if not d_nl and r.get("ngay"):
        d_nl = str(r.get("ngay")).strip().zfill(2)
    if not m_nl and r.get("thang"):
        m_nl = str(r.get("thang")).strip().zfill(2)
    if not y_nl and r.get("nam"):
        y_nl = str(r.get("nam")).strip()
    if not full_nl and d_nl and m_nl and y_nl:
        full_nl = f"ngày {d_nl} tháng {m_nl} năm {y_nl}"

    row["Ngày làm đơn vay vốn"] = full_nl
    row["Ngày_làm_đơn_vay_vốn"] = full_nl
    row["ngay_lam_don_vay_von"] = full_nl
    row["ngày_làm_đơn_vay_vốn"] = full_nl
    row["ngày làm đơn vay vốn"] = full_nl
    row["NGÀY_LÀM_ĐƠN_VAY_VỐN"] = full_nl

    row["ngày nhận hs vay vốn"] = d_nl
    row["ngày_nhận_hs_vay_vốn"] = d_nl
    row["ngay_nhan_hs_vay_von"] = d_nl
    row["tháng nhận hs vay vốn"] = m_nl
    row["tháng_nhận_hs_vay_vốn"] = m_nl
    row["thang_nhan_hs_vay_von"] = m_nl
    row["năm nhận hs vay vốn"] = y_nl
    row["năm_nhận_hs_vay_vốn"] = y_nl
    row["nam_nhan_hs_vay_von"] = y_nl

    can_ho = str(r.get('canHoSo') or r.get('Căn hộ số') or r.get('canHo') or '').strip()
    row["Căn hộ số"] = row.get("Căn hộ số") or can_ho
    row["Căn hộ số1"] = row.get("Căn hộ số1") or can_ho
    detected_p = detect_project_from_record(r)
    row["Dự án"] = detected_p or row.get("Dự án") or r.get('project') or r.get('duAn') or ""
    row["project"] = row["Dự án"]
    row["duAn"] = row["Dự án"]
    row["thuộc lô"] = row.get("thuộc lô") or r.get('lo') or ""
    row["số tầng"] = row.get("số tầng") or r.get('tang') or ""
    row["Tòa nhà"] = row.get("Tòa nhà") or r.get('toaNha') or ""
    dt_tt_val = row.get("DTSD (thông thủy)") or r.get('dtThongThuy') or ""
    row["DTSD (thông thủy)"] = dt_tt_val
    row["Bằng chữ DTSD (thông thủy)"] = row.get("Bằng chữ DTSD (thông thủy)") or r.get('bangChuDTThongThuy') or read_area_vietnamese(dt_tt_val)
    row["DT sàn"] = row.get("DT sàn") or r.get('dtSan') or ""
    row["DT thuộc sở hữu, sử dụng chung"] = row.get("DT thuộc sở hữu, sử dụng chung") or r.get('dtChung') or ""
    row["Năm hoàn thành XD"] = row.get("Năm hoàn thành XD") or r.get('namHoanThanh') or ""
    row["Tổng giá trị HĐMB"] = row.get("Tổng giá trị HĐMB") or _fmt_m(r.get('tongGiaTriHDMB') or r.get('tongGiaTriHD'))
    row["bằng chữ giá trị HĐMB"] = row.get("bằng chữ giá trị HĐMB") or _word_m(r.get('tongGiaTriHDMB') or r.get('tongGiaTriHD'))
    row["Giá bán đã bao gồm thuế GTGT"] = row.get("Giá bán đã bao gồm thuế GTGT") or _fmt_m(r.get('giaBanBaoGomGTGT') or r.get('giaBanNha'))
    row["Giá bán da bao gồm thuế GTGT"] = row.get("Giá bán da bao gồm thuế GTGT") or row["Giá bán đã bao gồm thuế GTGT"]
    row["thuế GTGT"] = row.get("thuế GTGT") or _fmt_m(r.get('thueGTGT'))
    row["kinh phí bảo trì 2%"] = row.get("kinh phí bảo trì 2%") or _fmt_m(r.get('kinhPhiBaoTri'))
    row["Bằng chữ kinh phí bảo trì 2%"] = row.get("Bằng chữ kinh phí bảo trì 2%") or _word_m(r.get('kinhPhiBaoTri'))
    row["số hđmb"] = row.get("số hđmb") or r.get('soHDMB') or ""
    row["phụ lục"] = row.get("phụ lục") or r.get('phuLuc') or ""
    pl_admin_date = r.get('ngayThangNamPhuLuc') or _fmt_date(r.get('ngayPhuLuc')) or ""
    row["ngày phụ lục"] = row.get("ngày phụ lục") or pl_admin_date
    row["ngay_phu_luc"] = row["ngày phụ lục"]
    row["ngàythángnăm_phụ_lục"] = row["ngày phụ lục"]
    row["ngày_phụ_lục"] = row["ngày phụ lục"]

    # Checkbox giới tính Nam / Nữ (Kiểu 1: Nam ☑  Nữ ☐ hoặc Nam ☐  Nữ ☑)
    gt_cust = str(r.get('gioiTinh') or r.get('Giới tính') or '').strip().lower()
    is_cust_m = (gt_cust == 'nam' or str(r.get('dx') or r.get('dx_hoa') or '').strip().lower() in ['ông', 'ong'])
    row["check_nam"] = "☑" if is_cust_m else "☐"
    row["check_nu"] = "☐" if is_cust_m else "☑"
    row["chk_nam"] = row["check_nam"]
    row["chk_nu"] = row["check_nu"]
    row["giới_tính_checkbox"] = f"Nam  {row['check_nam']}     Nữ  {row['check_nu']}"
    row["gioi_tinh_checkbox"] = row["giới_tính_checkbox"]

    row["số hđmb1"] = row.get("số hđmb1") or r.get('soHDMB1') or ""
    row["số hđmb 1"] = row.get("số hđmb 1") or row["số hđmb1"]
    row["ngày hđmb1"] = row.get("ngày hđmb1") or r.get('ngayHDMB1') or ""
    row["ngày hđmb 1"] = row.get("ngày hđmb 1") or row["ngày hđmb1"]
    row["ngày hđmb"] = row.get("ngày hđmb") or _fmt_date(r.get('ngayHDMB'))
    row["cán bộ TD"] = row.get("cán bộ TD") or r.get('canBo') or ""
    row["số ĐT CBTD"] = row.get("số ĐT CBTD") or r.get('sdtCanBo') or ""
    row["danh xưng cán bộ tín dụng"] = row.get("danh xưng cán bộ tín dụng") or r.get('dxCanBo') or "Ông"
    row["thẩm định từ ngày"] = row.get("thẩm định từ ngày") or _fmt_date(r.get('thamDinhTu'))
    row["thẩm định đến ngày"] = row.get("thẩm định đến ngày") or _fmt_date(r.get('thamDinhDen'))
    row["QĐ SỞ XD"] = row.get("QĐ SỞ XD") or r.get('qdSoXD') or ""

    for i in range(1, 7):
        k_vtc = f"vốn tự có {i}"
        row[k_vtc] = row.get(k_vtc) or r.get(f"baoCo{i}") or r.get(f"loaiBaoCo{i}") or ""
        val_vtc = r.get(f"vonTuCo{i}") or r.get(f"số tiền vốn tự có {i}")
        row[f"số tiền vốn tự có {i}"] = row.get(f"số tiền vốn tự có {i}") or _fmt_m(val_vtc)
        row[f"bằng chữ số tiền vốn tự có {i}"] = row.get(f"bằng chữ số tiền vốn tự có {i}") or _word_m(val_vtc)

    row["đóng số tiền lẻ"] = row.get("đóng số tiền lẻ") or _fmt_m(r.get('dongTienLe'))
    row["bằng chữ số tiền lẻ"] = row.get("bằng chữ số tiền lẻ") or _word_m(r.get('dongTienLe'))

    for i in range(1, 9):
        d_val = r.get(f"date{i}") or r.get(f"desc{i}") or r.get(f"ngay_dot_{i}") or r.get(f"ngày đợt {i}") or ""
        if i == 1:
            row["ngày_đợt_1"] = row.get("ngày_đợt_1") or d_val
        else:
            row[f"ngày đợt {i}"] = row.get(f"ngày đợt {i}") or d_val

        a_val = r.get(f"amt{i}") or r.get(f"số tiền đợt {i}")
        row[f"số tiền đợt {i}"] = row.get(f"số tiền đợt {i}") or _fmt_m(a_val)
        row[f"bằng chữ số tiền đợt {i}"] = row.get(f"bằng chữ số tiền đợt {i}") or _word_m(a_val)
        if i == 8:
            row["bằng_chữ_số_tiền_đợt_8"] = row.get("bằng_chữ_số_tiền_đợt_8") or row[f"bằng chữ số tiền đợt 8"]

        manual_remains = r.get("_manualRemains") or {}
        if not isinstance(manual_remains, dict):
            manual_remains = {}
        is_man = bool(manual_remains.get(i) or manual_remains.get(str(i)))

        numAmt2 = _parse_m(r.get("amt2") or r.get("số tiền đợt 2"))
        dtl = _parse_m(r.get("dongTienLe") or r.get("đóng số tiền lẻ"))
        auto_rem2 = max(0, numAmt2 - dtl) if numAmt2 > 0 else 0

        if i == 2:
            rem2_v = r.get("remain2") if is_man else auto_rem2
            row["Số tiền còn phải GN đợt 2"] = _fmt_m(rem2_v) if rem2_v else ""
            row["Bằng chữ số tiền còn phải GN đợt 2"] = _word_m(rem2_v) if rem2_v else ""
        elif i == 3:
            # - Số tiền còn phải giải ngân đợt 3 = số tiền giải ngân đợt 3 = numAmt3
            rem3_v = r.get("remain3") if is_man else (a_val if a_val else r.get("remain3"))
            row["Số tiền còn phải GN đợt 3"] = _fmt_m(rem3_v) if rem3_v else ""
            row["Bằng chữ số tiền còn phải GN đợt 3"] = _word_m(rem3_v) if rem3_v else ""
        elif i in [4, 5, 6, 7]:
            rem_k = r.get(f"remain{i}") if is_man else ""
            row[f"Số tiền còn phải GN đợt {i}"] = _fmt_m(rem_k) if rem_k else ""
            row[f"Bằng chữ số tiền còn phải GN đợt {i}"] = _word_m(rem_k) if rem_k else ""
        elif i == 8:
            rem8_v = r.get("remain8") if is_man else (r.get("remain8") or a_val)
            row["Còn_phải_GN_đợt_8"] = _fmt_m(rem8_v) if rem8_v else ""
            row["Bằng_chữ_còn_phải_GN_đợt_8"] = _word_m(rem8_v) if rem8_v else ""

    gn1_v = r.get('soTienGNDauTien') or r.get('số tiền GN đầu tiên')
    if not gn1_v:
        dtl_v = _parse_m(r.get("dongTienLe") or r.get("đóng số tiền lẻ"))
        p_name = (r.get("project") or r.get("Dự án") or "").lower()
        if "an phú" in p_name or "an phu" in p_name:
            # (G106 + G110 + G114 + G118 + G122) - I103
            s_2to6 = sum([_parse_m(row.get(f"số tiền đợt {k}")) for k in range(2, 7)])
            gn1_v = max(0, int(round(s_2to6 - dtl_v)))
        else:
            # G126 - I103 (đợt 7 - đóng tiền lẻ)
            amt7_v = _parse_m(row.get("số tiền đợt 7"))
            gn1_v = max(0, int(round(amt7_v - dtl_v)))
        if not gn1_v:
            gn1_v = _parse_m(row.get("số tiền đợt 1"))

    row["số tiền GN đầu tiên"] = row.get("số tiền GN đầu tiên") or _fmt_m(gn1_v)
    row["bằng chữ số tiền GN đầu tiên"] = row.get("bằng chữ số tiền GN đầu tiên") or _word_m(gn1_v)
    row["ngày GN đầu tiên"] = row.get("ngày GN đầu tiên") or _fmt_date(r.get('ngayGNDauTien'))

    # Kế hoạch giải ngân (Mục 3 HĐTD)
    kh_gn = r.get("keHoachGiaiNgan") or r.get("Kế hoạch giải ngân") or r.get("kế_hoạch_giải_ngân")
    if not kh_gn:
        active_loan_cols = []
        for col in range(2, 9):
            p_val = _parse_m(row.get(f"pct{col}") or row.get(f"tỷ lệ đợt {col}"))
            a_val = _parse_m(row.get(f"số tiền đợt {col}"))
            d_val = str(row.get(f"ngày đợt {col}") or r.get(f"date{col}") or r.get(f"desc{col}") or "").strip()
            if p_val > 0 or a_val > 0:
                active_loan_cols.append(col)
        
        cols = active_loan_cols if active_loan_cols else [2, 3, 4, 8]
        m = len(cols)
        lines = []

        def _get_money_connector(content):
            if not content: return ", số tiền "
            if re.search(r'(?:với\s+)?số\s+tiền(?:\s+tương\s+đương)?(?:\s+là)?\s*[:]?\s*$', content, flags=re.IGNORECASE):
                return "" if content.endswith(" ") else " "
            if re.search(r'[,]\s*$', content):
                return "số tiền là " if content.endswith(" ") else " số tiền là "
            if re.search(r'[:]\s*$', content):
                return "" if content.endswith(" ") else " "
            return ", số tiền "

        def _fmt_sch_cond(val, fallback):
            s = str(val or "").strip()
            if not s: return fallback
            if re.match(r'^\d{4}-\d{2}-\d{2}$', s):
                parts = s.split('-')
                return f"Ngày {parts[2]}/{parts[1]}/{parts[0]}"
            if re.match(r'^\d{1,2}/\d{1,2}/\d{2,4}$', s):
                return f"Ngày {s}"
            return s

        is_hl_or_nl = ("hồng loan" in str(r.get("project") or row.get("Dự án") or "").lower() or 
                       "nam long" in str(r.get("project") or row.get("Dự án") or "").lower())
        maint_v = _parse_m(row.get("kinh phí bảo trì 2%"))

        for idx, col in enumerate(cols):
            lan = idx + 1
            is_last = (idx == m - 1)
            label_prefix = f"+ Lần {lan}:"

            if idx == 0:
                raw_d1 = r.get("ngayGNDauTien") or r.get("ngayVay") or row.get("ngày GN đầu tiên") or row.get("ngày vay") or r.get("date2")
                d1_str = _fmt_sch_cond(raw_d1, "«ngày_GN_đầu_tiên»")
                if not re.match(r'^ngày\s+', d1_str, flags=re.IGNORECASE):
                    d1_str = f"Ngày {d1_str}"
                rem2_v = _parse_m(row.get("Số tiền còn phải GN đợt 2") or row.get("số tiền GN đầu tiên") or r.get("soTienGNDauTien") or row.get("số tiền đợt 2"))
                rem2_s = _fmt_m(rem2_v) if rem2_v else "«số_tiền_GN_đầu_tiên»"
                rem2_w = _word_m(rem2_v) if rem2_v else "«bằng_chữ_số_tiền_GN_đầu_tiên»"
                lines.append(f"{label_prefix} {d1_str}, số tiền {rem2_s} đồng (Bằng chữ: {rem2_w})")
            elif col == 7:
                raw_d7 = r.get("desc7") or r.get("date7") or r.get("desc5") or r.get("date5") or row.get("nội dung đợt 7") or row.get("ngày đợt 7")
                fallback_tag7 = "«nội dung đợt 5»" if (lan == 4 and m == 5) else ("«nội dung đợt 4»" if (lan == 3 and m == 4) else "«nội dung đợt 7»")
                d7_s = _fmt_sch_cond(raw_d7, fallback_tag7)
                a7_raw = _parse_m(row.get("số tiền đợt 7") or row.get("số tiền đợt 5"))
                a7_no_kpbt = a7_raw if is_hl_or_nl else (a7_raw - maint_v if a7_raw > maint_v else a7_raw)
                amt_tag7 = "«số_tiền_đợt_5»" if (lan == 4 and m == 5) else ("«số_tiền_đợt_4»" if (lan == 3 and m == 4) else "«số_tiền_đợt_7_KPBT»")
                amt_words_tag7 = "«bằng_chữ_số_tiền_đợt_5»" if (lan == 4 and m == 5) else ("«bằng_chữ_số_tiền_đợt_4»" if (lan == 3 and m == 4) else "«bằng_chữ_số_tiền_đợt_7_KPBT»")
                a7_s = _fmt_m(a7_no_kpbt) if a7_no_kpbt else amt_tag7
                a7_w = _word_m(a7_no_kpbt) if a7_no_kpbt else amt_words_tag7
                conn7 = _get_money_connector(d7_s)
                lines.append(f"{label_prefix} {d7_s}{conn7}{a7_s} đồng (Bằng chữ: {a7_w})")
            else:
                raw_d = r.get(f"desc{col}") or r.get(f"date{col}") or row.get(f"nội dung đợt {col}") or row.get(f"ngày đợt {col}")
                fallback_tag = f"«nội dung đợt {col}»"
                d_s = _fmt_sch_cond(raw_d, fallback_tag)
                a_v = _parse_m(row.get(f"số tiền đợt {col}"))
                a_s = _fmt_m(a_v) if a_v else f"«số_tiền_đợt_{col}»"
                a_w = _word_m(a_v) if a_v else f"«bằng_chữ_số_tiền_đợt_{col}»"
                conn = _get_money_connector(d_s)
                lines.append(f"{label_prefix} {d_s}{conn}{a_s} đồng (Bằng chữ: {a_w})")

        for idx in range(len(lines)):
            lines[idx] = re.sub(r'[.;\s]+$', '', lines[idx])
            if idx < len(lines) - 1:
                lines[idx] += ";"
            else:
                lines[idx] += "."

        kh_gn = "\n".join(lines)

    row["Kế hoạch giải ngân"] = kh_gn
    row["kế_hoạch_giải_ngân"] = kh_gn
    row["ke_hoach_giai_ngan"] = kh_gn
    row["kế hoạch giải ngân"] = kh_gn

    # 13. SỐ HỢP ĐỒNG TÍN DỤNG, SỐ HỢP ĐỒNG THẾ CHẤP, SỐ BIÊN BẢN XÁC ĐỊNH ĐẤU GIÁ
    from datetime import datetime
    cccd_clean = str(row.get("CCCD") or r.get("cccd") or "").replace(" ", "").strip()
    nam_hs = str(r.get("namNhanHs") or r.get("năm nhận hs vay vốn") or row.get("năm nhận hs vay vốn") or "").strip()
    if not nam_hs:
        raw_dt = str(r.get("ngayTiepNhan") or r.get("ngayNhanHs") or r.get("ngayVay") or row.get("ngày vay") or "").strip()
        m_y = re.search(r'\b(20\d\d)\b', raw_dt)
        nam_hs = m_y.group(1) if m_y else str(datetime.now().year)

    so_hdtd = str(r.get("soHDTD") or row.get("Số hợp đồng tín dụng") or "").strip()
    if not so_hdtd and cccd_clean:
        so_hdtd = f"{cccd_clean}/{nam_hs}/HĐTD"

    so_hdtc = str(r.get("soHDTC") or row.get("Số hợp đồng thế chấp") or "").strip()
    if not so_hdtc and cccd_clean:
        so_hdtc = f"{cccd_clean}/{nam_hs}/HĐTTC"

    so_bbxd = str(r.get("soBienBanXacDinh") or r.get("soBBXD") or r.get("soBBDH") or row.get("Số biên bản xác định đấu giá") or "").strip()
    if not so_bbxd and cccd_clean:
        so_bbxd = f"{cccd_clean}/{nam_hs}/BBĐH"

    row["Số hợp đồng tín dụng"] = so_hdtd
    row["số_hđtd"] = so_hdtd
    row["so_hdtd"] = so_hdtd
    row["số hợp đồng tín dụng"] = so_hdtd
    row["Số_HĐTD"] = so_hdtd
    row["số_hợp_đồng_tín_dụng"] = so_hdtd

    row["Số hợp đồng thế chấp"] = so_hdtc
    row["số_hđtc"] = so_hdtc
    row["so_hdtc"] = so_hdtc
    row["số hợp đồng thế chấp"] = so_hdtc
    row["Số_HĐTC"] = so_hdtc
    row["số_hợp_đồng_thế_chấp"] = so_hdtc

    row["Số biên bản xác định đấu giá"] = so_bbxd
    row["số_biên_bản_xác_định"] = so_bbxd
    row["so_bbdh"] = so_bbxd
    row["so_bbxd"] = so_bbxd
    row["số_biên_bản_đấu_giá"] = so_bbxd
    row["số biên bản xác định"] = so_bbxd
    row["Số_biên_bản_xác_định"] = so_bbxd
    row["số_biên_bản_xác_định_đấu_giá"] = so_bbxd

    # 13a. NGÀY THÔNG BÁO PHÊ DUYỆT CHO VAY VÀ NGÀY BIÊN BẢN XÁC ĐỊNH
    raw_tb_pd = str(r.get("ngayThangNamThongBao") or r.get("ngayThongBao") or r.get("ngày,tháng,năm tb phê duyệt cho vay") or r.get("ngày tb phê duyệt cho vay") or r.get("thamDinhDen") or r.get("thẩm định đến ngày") or "").strip()
    full_tb_pd, d_tb_pd, m_tb_pd, y_tb_pd = format_admin_date(raw_tb_pd)
    fmt_tb_pd = _fmt_date(raw_tb_pd)
    if not full_tb_pd and fmt_tb_pd:
        m_parts = fmt_tb_pd.split('/')
        if len(m_parts) == 3:
            full_tb_pd = f"ngày {m_parts[0]} tháng {m_parts[1]} năm {m_parts[2]}"

    row["ngày,tháng,năm tb phê duyệt cho vay"] = full_tb_pd or fmt_tb_pd
    row["ngàythángnăm_tb_phê_duyệt_cho_vay"] = full_tb_pd or fmt_tb_pd
    row["ngày_tháng_năm_tb_phê_duyệt_cho_vay"] = full_tb_pd or fmt_tb_pd
    row["ngaythangnam_tb_phe_duyet_cho_vay"] = full_tb_pd or fmt_tb_pd
    row["Ngày__tháng__năm_TB_phê_duyệt_cho_vay"] = full_tb_pd or fmt_tb_pd
    row["Ngày_tháng_năm_tb_phê_duyệt_cho_vay"] = full_tb_pd or fmt_tb_pd
    row["ngày tb phê duyệt cho vay"] = fmt_tb_pd or full_tb_pd
    row["ngày_tb_phê_duyệt_cho_vay"] = fmt_tb_pd or full_tb_pd
    row["ngay_tb_phe_duyet_cho_vay"] = fmt_tb_pd or full_tb_pd
    row["ngayThongBao"] = fmt_tb_pd or full_tb_pd
    row["ngayThangNamThongBao"] = full_tb_pd or fmt_tb_pd

    raw_bb_xd = str(r.get("ngayBienBanXacDinh") or r.get("ngày biên bản xác định") or r.get("thamDinhDen") or raw_tb_pd or "").strip()
    full_bb_xd, d_bb_xd, m_bb_xd, y_bb_xd = format_admin_date(raw_bb_xd)
    fmt_bb_xd = _fmt_date(raw_bb_xd)
    row["ngày biên bản xác định"] = full_bb_xd or fmt_bb_xd
    row["ngày_biên_bản_xác_định"] = full_bb_xd or fmt_bb_xd
    row["ngay_bien_ban_xac_dinh"] = full_bb_xd or fmt_bb_xd
    row["ngayBienBanXacDinh"] = full_bb_xd or fmt_bb_xd

    # 13b. NGÀY HỢP ĐỒNG TÍN DỤNG, NGÀY HỢP ĐỒNG THẾ CHẤP, NGÀY HỢP ĐỒNG 3 BÊN, NGÀY VAY, HẠN TRẢ
    ngay_vay_raw = str(r.get("ngayVay") or r.get("ngày vay") or row.get("ngày vay") or r.get("date1") or r.get("ngayNhanNo") or "").strip()
    ngay_hdtd_raw = str(r.get("ngayHDTD") or r.get("ngày hđtd") or r.get("ngày_hđtd") or row.get("ngày hđtd") or ngay_vay_raw or "").strip()
    full_hdtd, d_hdtd, m_hdtd, y_hdtd = format_admin_date(ngay_hdtd_raw)
    fmt_hdtd = _fmt_date(ngay_hdtd_raw)

    row["ngày hđtd"] = fmt_hdtd
    row["ngày_hđtd"] = fmt_hdtd
    row["ngay_hdtd"] = fmt_hdtd
    row["ngayHDTD"] = fmt_hdtd
    row["ngày/tháng/năm hđtd"] = full_hdtd or fmt_hdtd
    row["ngàythángnăm_hđtd"] = full_hdtd or fmt_hdtd
    row["ngày_tháng_năm_hđtd"] = full_hdtd or fmt_hdtd
    row["ngaythangnam_hdtd"] = full_hdtd or fmt_hdtd
    row["Ngày__tháng__năm_HĐTD"] = full_hdtd or fmt_hdtd
    row["Ngày_tháng_năm_HĐTD"] = full_hdtd or fmt_hdtd
    y_td_calc = y_hdtd or (d_hdtd if len(d_hdtd) == 4 else "") or nam_hs
    if y_td_calc:
        row["Năm_HĐTD"] = y_td_calc
        row["Năm HĐTD"] = y_td_calc
        row["năm_hđtd"] = y_td_calc

    # HĐTC
    ngay_hdtc_raw = str(r.get("ngayHDTC") or r.get("ngày hđtc") or r.get("ngày_hđtc") or row.get("ngày hđtc") or ngay_hdtd_raw or "").strip()
    full_hdtc, d_hdtc, m_hdtc, y_hdtc = format_admin_date(ngay_hdtc_raw)
    fmt_hdtc = _fmt_date(ngay_hdtc_raw)
    row["ngày hđtc"] = fmt_hdtc
    row["ngày_hđtc"] = fmt_hdtc
    row["ngay_hdtc"] = fmt_hdtc
    row["ngayHDTC"] = fmt_hdtc
    row["ngày,tháng,năm hđtc"] = full_hdtc or fmt_hdtc
    row["ngàythángnăm_hđtc"] = full_hdtc or fmt_hdtc
    row["ngày_tháng_năm_hđtc"] = full_hdtc or fmt_hdtc
    row["ngaythangnam_hdtc"] = full_hdtc or fmt_hdtc
    row["Ngày__tháng__năm_HĐTC"] = full_hdtc or fmt_hdtc
    row["Ngày_tháng_năm_HĐTC"] = full_hdtc or fmt_hdtc
    y_tc_calc = y_hdtc or y_td_calc
    if y_tc_calc:
        row["Năm_HĐTC"] = y_tc_calc
        row["Năm HĐTC"] = y_tc_calc
        row["năm_hđtc"] = y_tc_calc

    # HĐ3B
    ngay_hd3b_raw = str(r.get("ngayHD3B") or r.get("ngày hđ3b") or r.get("ngày_hđ3b") or row.get("ngày hđ3b") or ngay_hdtd_raw or "").strip()
    full_hd3b, d_hd3b, m_hd3b, y_hd3b = format_admin_date(ngay_hd3b_raw)
    fmt_hd3b = _fmt_date(ngay_hd3b_raw)
    row["ngày hđ3b"] = fmt_hd3b
    row["ngày_hđ3b"] = fmt_hd3b
    row["ngay_hd3b"] = fmt_hd3b
    row["ngayHD3B"] = fmt_hd3b
    row["ngày,tháng,năm hđ3b"] = full_hd3b or fmt_hd3b
    row["ngàythángnăm_hđ3b"] = full_hd3b or fmt_hd3b
    row["ngày_tháng_năm_hđ3b"] = full_hd3b or fmt_hd3b
    row["ngaythangnam_hd3b"] = full_hd3b or fmt_hd3b

    # Ngày vay & Hạn trả & Ngày GN đầu tiên
    fmt_ngay_vay = _fmt_date(ngay_vay_raw)
    row["ngày vay"] = fmt_ngay_vay
    row["ngày_vay"] = fmt_ngay_vay
    row["ngay_vay"] = fmt_ngay_vay
    row["ngayVay"] = fmt_ngay_vay

    han_tra_raw = str(r.get("hanTra") or r.get("hạn trả") or r.get("hạn_trả") or row.get("hạn trả") or "").strip()
    if not han_tra_raw and ngay_vay_raw:
        try:
            m_ymd = re.match(r'^(\d{4})-(\d{2})-(\d{2})', ngay_vay_raw)
            if m_ymd:
                y_v, m_v, d_v = int(m_ymd.group(1)), int(m_ymd.group(2)), int(m_ymd.group(3))
                y_ht = y_v + 25
                han_tra_raw = f"{d_v:02d}/{m_v:02d}/{y_ht}"
        except Exception:
            pass
    fmt_han_tra = _fmt_date(han_tra_raw)
    row["hạn trả"] = fmt_han_tra
    row["hạn_trả"] = fmt_han_tra
    row["han_tra"] = fmt_han_tra

    gn1_raw = str(r.get("ngayGNDauTien") or r.get("ngày GN đầu tiên") or r.get("ngày_gn_đầu_tiên") or row.get("ngày GN đầu tiên") or ngay_vay_raw or "").strip()
    fmt_gn1 = _fmt_date(gn1_raw)
    row["ngày GN đầu tiên"] = fmt_gn1
    row["ngày_gn_đầu_tiên"] = fmt_gn1
    row["ngay_gn_dau_tien"] = fmt_gn1

    # 14. VỐN TỰ CÓ N (BỎ 2% KPBT)
    vtc_n_val = str(r.get("vonTuCoN") or r.get("vốn tự có n") or r.get("vốn_tự_có_n") or "").strip()
    if not vtc_n_val:
        lines_n = []
        for bi in range(1, 9):
            bc_v = str(r.get(f"vốn tự có {bi}") or r.get(f"vốn_tự_có_{bi}") or r.get(f"baoCo{bi}") or "").strip()
            bc_v = re.sub(r'[;,\s]+$', '', bc_v)
            t_v = _parse_m(r.get(f"số tiền vốn tự có {bi}") or r.get(f"số_tiền_vốn_tự_có_{bi}") or r.get(f"vonTuCo{bi}") or 0)
            if bc_v or t_v > 0:
                t_str = f"{int(t_v):,}".replace(",", ".") if t_v > 0 else f"«số_tiền_vốn_tự_có_{bi}»"
                b_str = bc_v if bc_v else f"«vốn_tự_có_{bi}»"
                lines_n.append(f"- {b_str}, số tiền: {t_str} đồng.")
        if not lines_n:
            lines_n.append("- «vốn_tự_có_1», số tiền: «số_tiền_vốn_tự_có_1» đồng.")
        vtc_n_val = "\n".join(lines_n)

    row["vốn tự có n"] = vtc_n_val
    row["vốn_tự_có_n"] = vtc_n_val
    row["Vốn tự có n"] = vtc_n_val
    row["Vốn_tự_có_n"] = vtc_n_val
    row["von_tu_co_n"] = vtc_n_val

    # 14b. GIẢI CHẤP
    gc_ap = str(r.get("giaiChapAnPhu") or r.get("Giải chấp") or r.get("giaiChap") or r.get("giải chấp") or "").strip()
    gc_hl = str(r.get("giaiChapHongLoan") or r.get("Giải chấp Hồng Loan") or r.get("giải chấp hồng loan") or "").strip()
    row["Giải chấp"] = gc_ap
    row["Giải_chấp"] = gc_ap
    row["giải_chấp"] = gc_ap
    row["giải chấp"] = gc_ap
    row["Giải chấp Hồng Loan"] = gc_hl
    row["Giải_chấp_Hồng_Loan"] = gc_hl
    row["giải_chấp_hồng_loan"] = gc_hl
    row["giải chấp hồng loan"] = gc_hl

    # 15. BIÊN BẢN KIỂM TRA (PHẦN 10)
    bbkt_ngay = str(r.get("ngayBBKT") or r.get("Ngày bb kiểm tra") or r.get("Ngày_tháng_năm_BBKT") or "").strip()
    bbkt_nam = str(r.get("namTaiChinh") or r.get("Năm tài chính") or r.get("Năm_HĐTD") or "").strip()
    if not bbkt_nam:
        raw_dt = str(r.get("ngayTiepNhan") or r.get("ngayNhanHs") or r.get("ngayVay") or r.get("ngày vay") or "").strip()
        m_y = re.search(r'\b(20\d\d)\b', raw_dt)
        bbkt_nam = m_y.group(1) if m_y else str(datetime.now().year)

    bbkt_tong_gn = _parse_m(r.get("tongTienDaGiaiNganBBKT") or r.get("Tổng số tiền đã giải ngân") or r.get("Vốn vay") or r.get("vonVay") or 0)
    if bbkt_tong_gn <= 0 and r.get("amt1"):
        bbkt_tong_gn = _parse_m(r.get("amt1"))

    bbkt_goc_tra = _parse_m(r.get("luyKeGocDaTraBBKT") or r.get("Lũy kế số nợ gốc đã trả") or 0)

    if r.get("duNoVayBBKT") is not None and str(r.get("duNoVayBBKT")).strip() != "":
        bbkt_du_no = _parse_m(r.get("duNoVayBBKT"))
    elif r.get("Dư nợ vay") is not None and str(r.get("Dư nợ vay")).strip() != "":
        bbkt_du_no = _parse_m(r.get("Dư nợ vay"))
    else:
        bbkt_du_no = max(0.0, bbkt_tong_gn - bbkt_goc_tra)

    bbkt_chu_du_no = str(r.get("bangChuDuNoVayBBKT") or r.get("bằng chữ dư nợ vay") or "").strip()
    if not bbkt_chu_du_no and bbkt_du_no > 0:
        bbkt_chu_du_no = _word_m(bbkt_du_no)

    bbkt_gttt = str(r.get("giaTriThiTruongBBKT") or r.get("giá trị thị trường") or r.get("Giá_trị_hiện_tại_của_TSTC_BBKT") or "").strip()

    row["Ngày bb kiểm tra"] = bbkt_ngay
    row["Năm tài chính"] = bbkt_nam
    row["Tổng số tiền đã giải ngân"] = f"{int(bbkt_tong_gn):,}".replace(",", ".") if bbkt_tong_gn > 0 else ""
    row["Lũy kế số nợ gốc đã trả"] = f"{int(bbkt_goc_tra):,}".replace(",", ".") if bbkt_goc_tra > 0 else "0"
    row["Dư nợ vay"] = f"{int(bbkt_du_no):,}".replace(",", ".") if bbkt_du_no > 0 else "0"
    row["bằng chữ dư nợ vay"] = bbkt_chu_du_no
    row["giá trị thị trường"] = bbkt_gttt

    # Aliases
    row["ngày_bb_kiểm_tra"] = bbkt_ngay
    row["ngay_bb_kiem_tra"] = bbkt_ngay
    row["Ngày_tháng_năm_BBKT"] = bbkt_ngay
    row["Năm_HĐTD"] = bbkt_nam
    row["năm_tài_chính"] = bbkt_nam
    row["nam_tai_chinh"] = bbkt_nam
    row["tổng_số_tiền_đã_giải_ngân"] = row["Tổng số tiền đã giải ngân"]
    row["tong_so_tien_da_giai_ngan"] = row["Tổng số tiền đã giải ngân"]
    row["Lũy_kế_số_nợ_gốc_đã_trả_BBKT"] = row["Lũy kế số nợ gốc đã trả"]
    row["lũy_kế_số_nợ_gốc_đã_trả"] = row["Lũy kế số nợ gốc đã trả"]
    row["luy_ke_so_no_goc_da_tra"] = row["Lũy kế số nợ gốc đã trả"]
    row["Dư_nợ_vay_hiện_tại_BBKT"] = row["Dư nợ vay"]
    row["dư_nợ_vay"] = row["Dư nợ vay"]
    row["du_no_vay"] = row["Dư nợ vay"]
    row["Bằng_chữ_Lũy_kế_số_nợ_gốc_đã_trả_BBKT1"] = bbkt_chu_du_no
    row["bằng_chữ_dư_nợ_vay"] = bbkt_chu_du_no
    row["bang_chu_du_no_vay"] = bbkt_chu_du_no
    row["Giá_trị_hiện_tại_của_TSTC_BBKT"] = bbkt_gttt
    row["giá_trị_thị_trường"] = bbkt_gttt
    row["gia_tri_thi_truong"] = bbkt_gttt

    row["Số_Căn_hộ_BBKT"] = str(row.get("Căn hộ số") or r.get("canHoSo") or "").strip()
    dx_cb_val = str(row.get("danh xưng cán bộ tín dụng") or r.get("dxCanBo") or r.get("danhXungCanBo") or r.get("dx_cán_bộ") or "Ông").strip()
    ten_cb_val = str(row.get("cán bộ TD") or r.get("canBo") or r.get("tenCanBo") or r.get("cán_bộ_tín_dụng") or "").strip()
    row["Danh_Xưng_CBTD_BBKT"] = dx_cb_val
    row["Tên_CBTD_BBKT"] = ten_cb_val
    row["dx_cán_bộ"] = dx_cb_val
    row["dx_can_bo"] = dx_cb_val
    row["danh_xưng_cán_bộ"] = dx_cb_val
    row["Danh_xưng_cán_bộ"] = dx_cb_val
    row["danh_xưng_cán_bộ_tín_dụng"] = dx_cb_val
    row["cán_bộ_tín_dụng"] = ten_cb_val
    row["cán_bộ_td"] = ten_cb_val
    row["cán bộ TD"] = ten_cb_val
    row["Tên_cán_bộ_tín_dụng"] = ten_cb_val
    if not row.get("Địa_chỉ_cư_trú"):
        row["Địa_chỉ_cư_trú"] = str(row.get("Địa chỉ cư trú") or r.get("diaChiCuTru") or r.get("diaChi") or "").strip()

    # 16. TỜ TRÌNH ĐỀ NGHỊ GIẢI NGÂN (PHẦN 11 - MẪU 16/NƠXH - MẪU 18)
    raw_ttr_date = str(r.get("ngayToTrinh") or r.get("Ngày tờ trình") or r.get("ngày tờ trình") or "").strip()
    ttr_hoa = ""
    ttr_thuong = ""
    if raw_ttr_date:
        m_full = re.search(r'(?:[Nn]gày\s*)?(\d{1,2})\s*tháng\s*(\d{1,2})\s*năm\s*(\d{4})', raw_ttr_date, re.IGNORECASE)
        if m_full:
            d, m, y = m_full.group(1), m_full.group(2), m_full.group(3)
            ttr_hoa = f"Ngày {d} tháng {m} năm {y}"
            ttr_thuong = f"ngày {d} tháng {m} năm {y}"
        elif re.match(r'^\d{4}-\d{2}-\d{2}$', raw_ttr_date):
            parts = raw_ttr_date.split('-')
            d, m, y = str(int(parts[2])), str(int(parts[1])), parts[0]
            ttr_hoa = f"Ngày {d} tháng {m} năm {y}"
            ttr_thuong = f"ngày {d} tháng {m} năm {y}"
        elif re.match(r'^\d{1,2}/\d{1,2}/\d{4}$', raw_ttr_date):
            parts = raw_ttr_date.split('/')
            d, m, y = str(int(parts[0])), str(int(parts[1])), parts[2]
            ttr_hoa = f"Ngày {d} tháng {m} năm {y}"
            ttr_thuong = f"ngày {d} tháng {m} năm {y}"
        else:
            ttr_hoa = raw_ttr_date
            ttr_thuong = raw_ttr_date[:1].lower() + raw_ttr_date[1:] if raw_ttr_date else ""

    # Tổng số tiền đã giải ngân TTr:
    ttr_tong_da_gn = 0
    if r.get("tongTienDaGiaiNganTTr") is not None and str(r.get("tongTienDaGiaiNganTTr")).strip() != "":
        ttr_tong_da_gn = _parse_m(r.get("tongTienDaGiaiNganTTr"))
    elif r.get("Tổng số tiền đã giải ngân TTr") is not None and str(r.get("Tổng số tiền đã giải ngân TTr")).strip() != "":
        ttr_tong_da_gn = _parse_m(r.get("Tổng số tiền đã giải ngân TTr"))
    else:
        amt1_val = _parse_m(r.get("soTienGNDauTien") or r.get("amt1") or row.get("số tiền GN đầu tiên") or row.get("số tiền đợt 1"))
        running_tot = amt1_val
        for bi in range(2, 9):
            has_dot = bool(r.get(f"tb{bi}") or r.get(f"date{bi}") or r.get(f"soTienTB{bi}") or r.get(f"số thông báo đợt {bi}") or r.get(f"số tb - ngày tb{bi}"))
            dot_amt = _parse_m(r.get(f"amt{bi}") or r.get(f"soTienTB{bi}") or row.get(f"số tiền đợt {bi}"))
            if has_dot and dot_amt > 0:
                running_tot += dot_amt
        ttr_tong_da_gn = running_tot

    von_vay_val = _parse_m(row.get("Vốn vay") or r.get("vonVay") or r.get("Vốn_vay") or 0)
    if r.get("tongTienConPhaiGNTTr") is not None and str(r.get("tongTienConPhaiGNTTr")).strip() != "":
        ttr_con_phai_gn = _parse_m(r.get("tongTienConPhaiGNTTr"))
    elif r.get("Tổng số tiền còn phải giải ngân") is not None and str(r.get("Tổng số tiền còn phải giải ngân")).strip() != "":
        ttr_con_phai_gn = _parse_m(r.get("Tổng số tiền còn phải giải ngân"))
    else:
        ttr_con_phai_gn = max(0, von_vay_val - ttr_tong_da_gn) if von_vay_val > 0 else 0

    s_ttr_tong_da_gn = f"{int(ttr_tong_da_gn):,}".replace(",", ".") if ttr_tong_da_gn > 0 else "0"
    s_ttr_con_phai_gn = f"{int(ttr_con_phai_gn):,}".replace(",", ".") if ttr_con_phai_gn > 0 else "0"

    row["Ngày tờ trình"] = ttr_hoa
    row["ngày tờ trình"] = ttr_thuong
    row["Tổng số tiền đã giải ngân TTr"] = s_ttr_tong_da_gn
    row["Tổng số tiền còn phải giải ngân"] = s_ttr_con_phai_gn

    row["Ngày_tờ_trình"] = ttr_hoa
    row["ngày_tờ_trình"] = ttr_thuong
    row["Ngay_to_trinh"] = ttr_hoa
    row["ngay_to_trinh"] = ttr_thuong
    row["Dư_nợ_trước_ngày_nhận_nợ_L2"] = s_ttr_tong_da_gn
    row["dư_nợ_trước_ngày_nhận_nợ_l2"] = s_ttr_tong_da_gn
    row["Tổng_số_tiền_còn_phải_giải_ngân"] = s_ttr_con_phai_gn
    row["tong_so_tien_con_phai_giai_ngan"] = s_ttr_con_phai_gn

    # 16b. CÒN PHẢI GIẢI NGÂN VÀ LŨY KẾ THEO TỪNG ĐỢT (ĐỢT 1 ĐẾN 8)
    _cum_gn = 0
    for ki in range(1, 9):
        amt_ki_val = _parse_m(r.get(f"amt{ki}") or row.get(f"số tiền đợt {ki}") or 0)
        _cum_gn += amt_ki_val
        con_lai_ki = max(0, von_vay_val - _cum_gn) if von_vay_val > 0 else 0
        s_con_lai_ki = _fmt_m(con_lai_ki) if con_lai_ki > 0 else "0"
        w_con_lai_ki = _word_m(con_lai_ki) if con_lai_ki > 0 else "không đồng"
        row[f"Số_tiền_còn_phải_GN_đợt_{ki}"] = s_con_lai_ki
        row[f"số_tiền_còn_phải_gn_đợt_{ki}"] = s_con_lai_ki
        row[f"Còn_phải_GN_đợt_{ki}"] = s_con_lai_ki
        row[f"còn_phải_gn_đợt_{ki}"] = s_con_lai_ki
        row[f"Bằng_chữ_số_tiền_còn_phải_GN_đợt_{ki}"] = w_con_lai_ki
        row[f"bằng_chữ_số_tiền_còn_phải_gn_đợt_{ki}"] = w_con_lai_ki
        row[f"bằng_chữ_còn_phải_gn_đợt_{ki}"] = w_con_lai_ki

    # 17. ĐỒNG SỞ HỮU
    dsh_py = compute_dong_so_huu_py(r)
    row["tên đồng sở hữu hoa"] = dsh_py.get("hoa", "")
    row["tên đồng sở hữu thường"] = dsh_py.get("thuong", "")
    row["tên_đồng_sở_hữu_thường"] = dsh_py.get("thuong", "")
    row["đồng_sở_hữu_hoa"] = dsh_py.get("hoa", "")
    row["đồng_sở_hữu_thường"] = dsh_py.get("thuong", "")

    # 17b. GIẤY TỜ XÁC NHẬN THU NHẬP / LƯƠNG (MẪU 11 & MẪU 3)
    xn_luong = str(r.get("xacNhanLuong") or r.get("xác nhận lương") or r.get("xác_nhận_lương") or r.get("xác nhận thu nhập") or r.get("xác_nhận_thu_nhập") or "").strip()
    if not xn_luong:
        tt_hn = str(r.get("tinhTrangHonNhan") or r.get("Tình trạng hôn nhân") or "").strip()
        is_single = (tt_hn == "Độc thân" or tt_hn == "Ly hôn")
        raw_kh = str(r.get("hoTen") or r.get("Họ và tên KH thường") or r.get("Họ và tên KH hoa") or r.get("Họ và tên") or "").strip()
        raw_vc = str(r.get("vcHoTen") or r.get("Họ và tên Vợ/chồng") or "").strip()
        name_kh = format_title_case_py(raw_kh) if 'format_title_case_py' in globals() else raw_kh.title()
        name_vc = format_title_case_py(raw_vc) if 'format_title_case_py' in globals() else raw_vc.title()
        if name_kh:
            if is_single or not name_vc:
                xn_luong = f"Xác nhận thu nhập ({name_kh})"
            else:
                xn_luong = f"Xác nhận thu nhập ({name_kh}), Kê khai thu nhập ({name_vc})"
    for k_xn in ["xác_nhận_lương", "xác nhận lương", "xác_nhận_thu_nhập", "xác nhận thu nhập", "xac_nhan_luong", "xac_nhan_thu_nhap", "Xác_nhận_lương", "Xác_nhận_thu_nhập"]:
        row[k_xn] = xn_luong

    # 17c. MÃ KHÁCH HÀNG, SỐ TÀI KHOẢN 102 & THÔNG TIN NGƯỜI PHỤ THUỘC (CỘT HA, HB, HC)
    ma_kh_val = str(r.get("maKH") or r.get("mã khách hàng") or r.get("mã_khách_hàng") or r.get("ma_kh") or r.get("Mã khách hàng") or "").strip()
    for k_makh in ["mã khách hàng", "mã_khách_hàng", "Mã_khách_hàng", "mã_kh", "Mã_KH", "ma_kh", "makh", "maKH"]:
        row[k_makh] = ma_kh_val

    stk102_val = str(r.get("soTK102") or r.get("số tài khoản 102") or r.get("số_tài_khoản_102") or r.get("so_tk_102") or r.get("Số tài khoản 102") or "").strip()
    for k_stk in ["số tài khoản 102", "số_tài_khoản_102", "Số_tài_khoản_102", "số_tk_102", "Số_TK_102", "stk_102", "STK_102", "so_tk_102", "soTK102"]:
        row[k_stk] = stk102_val

    ttpt_val = str(r.get("thongTinPhuThuoc") or r.get("thông tin phụ thuộc") or r.get("thông_tin_phụ_thuộc") or r.get("người_phụ_thuộc") or r.get("người phụ thuộc") or r.get("thong_tin_phu_thuoc") or "").strip()
    for k_pt in ["thông tin phụ thuộc", "thông_tin_phụ_thuộc", "Thông_tin_phụ_thuộc", "thong_tin_phu_thuoc", "người_phụ_thuộc", "người phụ thuộc", "Người_phụ_thuộc", "thongTinPhuThuoc"]:
        row[k_pt] = ttpt_val

    # 18. TỔNG HỢP VỐN TỰ CÓ & PHIẾU BÁO CÓ
    th_vtc_val = str(r.get("tongHopVonTuCo") or r.get("tổng_hợp_vốn_tự_có") or r.get("tổng hợp vốn tự có") or "").strip()
    if not th_vtc_val:
        lines_vtc = []
        order_idx = 1
        for bi in range(1, 9):
            bc_v = str(r.get(f"vốn tự có {bi}") or r.get(f"vốn_tự_có_{bi}") or r.get(f"baoCo{bi}") or "").strip()
            bc_v = re.sub(r'(số giao dịch|chứng từ:|toán số:)\s*([A-Za-z0-9])', r'\1 \2', bc_v, flags=re.IGNORECASE)
            bc_v = re.sub(r'[;,\s]+$', '', bc_v)
            t_v = _parse_m(r.get(f"số tiền vốn tự có {bi}") or r.get(f"số_tiền_vốn_tự_có_{bi}") or r.get(f"vonTuCo{bi}") or 0)
            if bc_v or t_v > 0:
                t_str = f"{int(t_v):,}".replace(",", ".") if t_v > 0 else f"«số_tiền_vốn_tự_có_{bi}»"
                b_str = bc_v if bc_v else f"«vốn_tự_có_{bi}»"
                lines_vtc.append(f"({order_idx}) {b_str}, số tiền: {t_str} đồng.")
                order_idx += 1
        if not lines_vtc:
            lines_vtc.append("(1) «vốn_tự_có_1», số tiền: «số_tiền_vốn_tự_có_1» đồng.")
            order_idx = 2
        kpbt_v = _parse_m(r.get("kinh phí bảo trì 2%") or r.get("kinh_phí_bảo_trì_2") or r.get("kinhPhiBaoTri") or 0)
        kpbt_s = f"{int(kpbt_v):,}".replace(",", ".") if kpbt_v > 0 else "«kinh_phí_bảo_trì_2»"
        lines_vtc.append(f"({order_idx}) Số tiền thanh toán 2% kinh phí bảo trì {kpbt_s} đồng. Bên mua tự thanh toán bằng vốn tự có cho Bên Bán trong thời hạn từ ngày Bên Bán gửi thông báo bàn giao Căn hộ đến trước thời điểm ký Biên bản bàn giao căn hộ.")
        th_vtc_val = "\n".join(lines_vtc)

    row["tổng_hợp_vốn_tự_có"] = th_vtc_val
    row["tổng hợp vốn tự có"] = th_vtc_val
    row["Tổng_hợp_vốn_tự_có"] = th_vtc_val
    row["tong_hop_von_tu_co"] = th_vtc_val
    row["tongHopVonTuCo"] = th_vtc_val

    # 18b. CHỨNG TỪ VỐN TỰ CÓ ĐỢT 1 ĐẾN 8
    vtc1_raw = str(r.get("noiDungChungTuDot1") or r.get("chứng_từ_vốn_tự_có_đợt_1") or r.get("chứng từ vốn tự có đợt 1") or "").strip()
    if not vtc1_raw:
        bc1_v = str(r.get("vốn tự có 1") or r.get("vốn_tự_có_1") or r.get("baoCo1") or "").strip()
        t1_v = _parse_m(r.get("số tiền vốn tự có 1") or r.get("số_tiền_vốn_tự_có_1") or r.get("vonTuCo1") or 0)
        if bc1_v and t1_v > 0:
            vtc1_raw = f"{bc1_v}, số tiền {_fmt_m(t1_v)} đồng"
        elif bc1_v:
            vtc1_raw = bc1_v
        elif t1_v > 0:
            vtc1_raw = f"Số tiền {_fmt_m(t1_v)} đồng"
    row["chứng_từ_vốn_tự_có_đợt_1"] = vtc1_raw
    row["chứng từ vốn tự có đợt 1"] = vtc1_raw
    row["chung_tu_von_tu_co_dot_1"] = vtc1_raw
    row["chứng_từ_đợt_1"] = vtc1_raw

    vtc2_raw = str(r.get("noiDungChungTuDot2") or r.get("chứng_từ_vốn_tự_có_đợt_2") or r.get("chứng từ vốn tự có đợt 2") or "").strip()
    if not vtc2_raw:
        bc2_v = str(r.get("vốn tự có 2") or r.get("vốn_tự_có_2") or r.get("baoCo2") or "").strip()
        t2_v = _parse_m(r.get("số tiền vốn tự có 2") or r.get("số_tiền_vốn_tự_có_2") or r.get("vonTuCo2") or 0)
        if bc2_v and t2_v > 0:
            vtc2_raw = f"{bc2_v}, số tiền {_fmt_m(t2_v)} đồng"
        elif bc2_v:
            vtc2_raw = bc2_v
        elif t2_v > 0:
            vtc2_raw = f"Số tiền {_fmt_m(t2_v)} đồng"
    row["chứng_từ_vốn_tự_có_đợt_2"] = vtc2_raw
    row["chứng từ vốn tự có đợt 2"] = vtc2_raw
    row["chung_tu_von_tu_co_dot_2"] = vtc2_raw
    row["chứng_từ_đợt_2"] = vtc2_raw

    for di in range(3, 9):
        vtcd_raw = str(r.get(f"noiDungChungTuDot{di}") or r.get(f"chứng_từ_vốn_tự_có_đợt_{di}") or r.get(f"chứng từ vốn tự có đợt {di}") or "").strip()
        if not vtcd_raw:
            bcd_v = str(r.get(f"vốn tự có {di}") or r.get(f"vốn_tự_có_{di}") or r.get(f"baoCo{di}") or "").strip()
            td_v = _parse_m(r.get(f"số tiền vốn tự có {di}") or r.get(f"số_tiền_vốn_tự_có_{di}") or r.get(f"vonTuCo{di}") or 0)
            if bcd_v and td_v > 0:
                vtcd_raw = f"{bcd_v}, số tiền {_fmt_m(td_v)} đồng"
            elif bcd_v:
                vtcd_raw = bcd_v
            elif td_v > 0:
                vtcd_raw = f"Số tiền {_fmt_m(td_v)} đồng"
        row[f"chứng_từ_vốn_tự_có_đợt_{di}"] = vtcd_raw
        row[f"chứng từ vốn tự có đợt {di}"] = vtcd_raw
        row[f"chung_tu_von_tu_co_dot_{di}"] = vtcd_raw

    p_name_bc = str(r.get("project") or r.get("dự án") or r.get("duAn") or r.get("du_an") or row.get("Dự án") or row.get("dự án") or row.get("project") or r.get("tenDuAn") or "").strip().lower()
    is_an_phu_2026_bc = "an phú" in p_name_bc or "an phu" in p_name_bc
    kpbt_val_bc = _parse_m(row.get("kinh phí bảo trì 2%") or r.get("kinhPhiBaoTri") or 0)
    kpbt_str_bc = _fmt_m(kpbt_val_bc) if kpbt_val_bc > 0 else (str(row.get("kinh phí bảo trì 2%") or r.get("kinhPhiBaoTri") or "").strip() or "«kinh phí bảo trì 2%»")

    th_bc_val = str(r.get("tongHopPhieuBaoCo") or r.get("tổng_hợp_phiếu_báo_có") or r.get("tổng hợp phiếu báo có") or "").strip()
    if not th_bc_val or (is_an_phu_2026_bc and "kinh phí bảo trì 2%" not in th_bc_val):
        bc_list = []
        for bi in range(1, 9):
            bc_raw = str(r.get(f"vốn tự có {bi}") or r.get(f"vốn_tự_có_{bi}") or r.get(f"baoCo{bi}") or row.get(f"vốn tự có {bi}") or "").strip()
            bc_clean = re.sub(r'(số giao dịch|chứng từ:|toán số:)\s*([A-Za-z0-9])', r'\1 \2', bc_raw, flags=re.IGNORECASE)
            bc_clean = re.sub(r'[;,\s]+$', '', bc_clean)
            t_val = _parse_m(r.get(f"số tiền vốn tự có {bi}") or r.get(f"số_tiền_vốn_tự_có_{bi}") or r.get(f"vonTuCo{bi}") or row.get(f"số tiền vốn tự có {bi}") or 0)
            if bc_clean or t_val > 0:
                t_str = f"{int(t_val):,}".replace(",", ".") if t_val > 0 else f"«số_tiền_vốn_tự_có_{bi}»"
                b_str = bc_clean if bc_clean else f"«vốn_tự_có_{bi}»"
                bc_list.append({'index': bi, 'amount': t_val, 'bc_val': b_str, 'v_str': t_str, 'v_words': _word_m(t_val) if t_val > 0 else f"«bằng_chữ_số_tiền_vốn_tự_có_{bi}»"})

        amt1_val = _parse_m(row.get("số tiền đợt 1") or r.get("amt1") or row.get("số tiền GN đầu tiên") or r.get("soTienGNDauTien") or 0)
        p1_vouchers = []
        sum_p1 = 0
        next_v_idx = 0
        if bc_list:
            if amt1_val > 0:
                for k_v, v_item in enumerate(bc_list):
                    p1_vouchers.append(v_item)
                    sum_p1 += v_item['amount']
                    next_v_idx = k_v + 1
                    if sum_p1 >= amt1_val:
                        break
            else:
                p1_vouchers = bc_list[:2] if len(bc_list) > 1 else [bc_list[0]]
                next_v_idx = len(p1_vouchers)
                sum_p1 = sum(v['amount'] for v in p1_vouchers)

        phieu_d1_txt = ""
        so_du_d1 = 0
        if p1_vouchers:
            items_p1 = [f"số tiền {v['v_str']} đồng (Bằng chữ: {v['v_words']}) theo: {v['bc_val']}" for v in p1_vouchers]
            if len(items_p1) == 1:
                base_p1_str = items_p1[0]
            elif len(items_p1) == 2:
                base_p1_str = f"{items_p1[0]} và {items_p1[1]}"
            else:
                base_p1_str = f"{', '.join(items_p1[:-1])} và {items_p1[-1]}"

            if amt1_val > 0 and sum_p1 > amt1_val:
                so_du_d1 = sum_p1 - amt1_val
                amt1_s = _fmt_m(amt1_val)
                amt1_w = _word_m(amt1_val)
                s_du_s = _fmt_m(so_du_d1)
                s_du_w = _word_m(so_du_d1)
                phieu_d1_txt = f"Bên mua đã nộp {base_p1_str}, trong đó số tiền của đợt 1 là {amt1_s} đồng (Bằng chữ: {amt1_w}) và số tiền của đợt 2 là {s_du_s} đồng (Bằng chữ: {s_du_w})."
            else:
                phieu_d1_txt = f"Bên mua đã nộp {base_p1_str}."

        items_p2 = []
        if so_du_d1 > 0:
            items_p2.append(f"số tiền {_fmt_m(so_du_d1)} đồng (Bằng chữ: {_word_m(so_du_d1)})")
        for k in range(next_v_idx, len(bc_list)):
            v = bc_list[k]
            items_p2.append(f"số tiền {v['v_str']} đồng (Bằng chữ: {v['v_words']}) theo: {v['bc_val']}")

        phieu_d2_txt = ""
        if len(items_p2) == 1:
            phieu_d2_txt = f"Bên mua đã nộp {items_p2[0]}. "
        elif len(items_p2) == 2:
            phieu_d2_txt = f"Bên mua đã nộp {items_p2[0]} và {items_p2[1]}. "
        elif len(items_p2) > 2:
            phieu_d2_txt = f"Bên mua đã nộp {', '.join(items_p2[:-1])} và {items_p2[-1]}. "

        def _parse_pct(val):
            if val is None or val == "": return 0
            try:
                cleaned = str(val).replace('%', '').replace(',', '.').strip()
                f_val = float(cleaned) if cleaned else 0
                return (f_val * 100) if (0 < f_val <= 1.0) else f_val
            except Exception:
                return 0

        pcts_m, amts_m = {}, {}
        active_c = []
        for c in range(1, 9):
            pv = _parse_pct(r.get(f"pct{c}") or row.get(f"tỷ lệ đợt {c}") or 0)
            av = _parse_m(r.get(f"amt{c}") or row.get(f"số tiền đợt {c}") or 0)
            rd = str(r.get(f"date{c}") or r.get(f"desc{c}") or row.get(f"ngày đợt {c}") or "").strip()
            pcts_m[c] = pv
            amts_m[c] = av
            if pv > 0 or av > 0 or (c == 1 and (pv > 0 or av > 0 or bool(rd))):
                active_c.append(c)
        if not any(pcts_m[c] > 0 for c in range(1, 9)):
            active_c = [1, 2, 7, 8]

        out_lines_bc = []
        cur_stt_bc = 1
        penultimate_c = active_c[-2] if len(active_c) >= 2 else -1

        for col in active_c:
            p_val_c = pcts_m.get(col, 0)
            a_val_c = amts_m.get(col, 0)
            raw_c = str(r.get(f"desc{col}") or r.get(f"nội dung đợt {col}") or r.get(f"nội_dung_đợt_{col}") or r.get(f"desc{cur_stt_bc}") or r.get(f"nội dung đợt {cur_stt_bc}") or r.get(f"nội_dung_đợt_{cur_stt_bc}") or r.get(f"date{col}") or r.get(f"ngày đợt {col}") or r.get(f"ngày_đợt_{col}") or r.get(f"date{cur_stt_bc}") or r.get(f"ngày đợt {cur_stt_bc}") or r.get(f"ngày_đợt_{cur_stt_bc}") or "").strip()

            p_s = f"{int(round(p_val_c))}%" if p_val_c > 0 else f"«%_đợt_{col}»"
            a_s = _fmt_m(a_val_c) if a_val_c > 0 else f"«số_tiền_đợt_{col}»"
            a_w = _word_m(a_val_c) if a_val_c > 0 else f"«bằng_chữ_số_tiền_đợt_{col}»"

            clean_c = raw_c
            if not clean_c:
                clean_c = f"«nội_dung_đợt_{cur_stt_bc}»"
                conn_bc = " với số tiền "
            else:
                if re.search(r'(?:với\s+)?số\s+tiền(?:\s+tương\s+đương)?(?:\s+là)?\s*[:]?\s*$', clean_c, flags=re.IGNORECASE):
                    conn_bc = "" if clean_c.endswith(" ") else " "
                elif re.search(r'[,]\s*$', clean_c):
                    conn_bc = "số tiền là " if clean_c.endswith(" ") else " số tiền là "
                elif re.search(r'[:]\s*$', clean_c):
                    conn_bc = "" if clean_c.endswith(" ") else " "
                else:
                    conn_bc = " với số tiền "

            line_bc = f"+ Đợt {cur_stt_bc}: {clean_c}{conn_bc}{a_s} đồng (Bằng chữ: {a_w})."
            if col == 1 and phieu_d1_txt:
                line_bc += f" {phieu_d1_txt}"
            elif col == 2 and phieu_d2_txt:
                line_bc += f" {phieu_d2_txt.strip()}"

            # Đối với An Phú: Đợt kế cuối (hoặc col == 7) nối thêm câu KPBT 2%
            is_ap_dot = is_an_phu_2026_bc or (len(active_c) == 4 and active_c[0] == 1 and active_c[1] == 2 and active_c[2] == 7 and active_c[3] == 8)
            if is_ap_dot and (col == penultimate_c or col == 7):
                line_bc += f" Trong đó, Bên mua tự thanh toán bằng vốn tự có kinh phí bảo trì 2% số tiền là {kpbt_str_bc} đồng cho Bên bán khi có thông báo bàn giao căn hộ và trước khi nhận bàn giao Căn hộ."

            out_lines_bc.append(re.sub(r'\s+', ' ', line_bc).strip())
            cur_stt_bc += 1

        for idx_b in range(len(out_lines_bc)):
            out_lines_bc[idx_b] = re.sub(r'[.;\s]+$', '', out_lines_bc[idx_b])
            if idx_b < len(out_lines_bc) - 1:
                out_lines_bc[idx_b] += ";"
            else:
                out_lines_bc[idx_b] += "."

        th_bc_val = "\n".join(out_lines_bc)

    # Đảm bảo nếu th_bc_val đã có nhưng chưa có câu KPBT 2% đối với An Phú thì bổ sung vào đợt kế cuối
    if (is_an_phu_2026_bc or "bình xuyên" in str(r).lower() or "113002951528" in str(r)) and th_bc_val and "kinh phí bảo trì 2%" not in th_bc_val:
        th_lines = th_bc_val.split("\n")
        if len(th_lines) >= 2:
            target_idx = len(th_lines) - 2
            line_target = re.sub(r'[.;\s]+$', '', th_lines[target_idx])
            th_lines[target_idx] = line_target + f" Trong đó, Bên mua tự thanh toán bằng vốn tự có kinh phí bảo trì 2% số tiền là {kpbt_str_bc} đồng cho Bên bán khi có thông báo bàn giao căn hộ và trước khi nhận bàn giao Căn hộ.;"
            th_bc_val = "\n".join(th_lines)

    row["tổng_hợp_phiếu_báo_có"] = th_bc_val
    row["tổng hợp phiếu báo có"] = th_bc_val
    row["Tổng_hợp_phiếu_báo_có"] = th_bc_val
    row["tong_hop_phieu_bao_co"] = th_bc_val
    row["tongHopPhieuBaoCo"] = th_bc_val

    # 19. MẪU 15 & 16: GIẤY ĐỀ NGHỊ GIẢI NGÂN LẦN N & BIÊN BẢN GIAO NHẬN HỒ SƠ N
    lan_gn = str(r.get("lanGiaiNgan") or r.get("soLanGiaiNgan") or r.get("Giấy_đề_nghị_GN_lần") or "2").strip()
    row["Giấy_đề_nghị_GN_lần"] = lan_gn
    row["giấy_đề_nghị_gn_lần"] = lan_gn

    t_gn_dot = _parse_m(r.get("soTienGiaiNgan") or r.get("amt2") or r.get("soTienTB2") or row.get("số tiền đợt 2") or 0)
    row["Số_tiền_gn"] = f"{int(t_gn_dot):,}".replace(",", ".") if t_gn_dot > 0 else ""
    row["số_tiền_gn"] = row["Số_tiền_gn"]
    row["Bằng_chữ_Số_tiền_gn"] = _word_m(t_gn_dot) if t_gn_dot > 0 else ""
    row["bằng_chữ_số_tiền_gn"] = row["Bằng_chữ_Số_tiền_gn"]

    t_da_gn = _parse_m(r.get("soTienDaGiaiNgan") or r.get("amt1") or row.get("số tiền đợt 1") or row.get("dư nợ trước GN lần 2") or 0)
    row["Đã_giải_ngân_số_tiền"] = f"{int(t_da_gn):,}".replace(",", ".") if t_da_gn > 0 else ""
    row["đã_giải_ngân_số_tiền"] = row["Đã_giải_ngân_số_tiền"]
    row["Bằng_chữ_Đã_giải_ngân_số_tiền"] = _word_m(t_da_gn) if t_da_gn > 0 else ""
    row["bằng_chữ_đã_giải_ngân_số_tiền"] = row["Bằng_chữ_Đã_giải_ngân_số_tiền"]

    t_du_no_truoc = _parse_m(r.get("duNoDenTruocNgayNhanNo") or row.get("dư nợ trước GN lần 2") or t_da_gn)
    row["Dư_nợ_đến_trước_ngày_nhận_nợ"] = f"{int(t_du_no_truoc):,}".replace(",", ".") if t_du_no_truoc > 0 else ""
    row["dư_nợ_đến_trước_ngày_nhận_nợ"] = row["Dư_nợ_đến_trước_ngày_nhận_nợ"]
    row["Bằng_chữ_Dư_nợ_đến_trước_ngày_nhận_nợ"] = _word_m(t_du_no_truoc) if t_du_no_truoc > 0 else ""
    row["bằng_chữ_dư_nợ_đến_trước_ngày_nhận_nợ"] = row["Bằng_chữ_Dư_nợ_đến_trước_ngày_nhận_nợ"]

    # 19b. TÀI KHOẢN VÀ NGÂN HÀNG THỤ HƯỞNG CHỦ ĐẦU TƯ
    p_name_cur = str(r.get("project") or r.get("dự án") or r.get("duAn") or row.get("Dự án") or "").strip().lower()
    default_acc = "113002951528"
    default_bank = "Ngân hàng TMCP Công thương Việt Nam - CN Bình Xuyên"
    if "nam long" in p_name_cur:
        default_acc = "7063020417"
        default_bank = "Phòng giao dịch Ngân hàng Chính sách xã hội Cái Răng"
    elif "hồng loan" in p_name_cur or "hong loan" in p_name_cur:
        default_acc = "7063261897"
        default_bank = "Phòng giao dịch Ngân hàng Chính sách xã hội Cái Răng"

    acc_cdu = str(r.get("projectAccount") or r.get("soTaiKhoan") or r.get("soTaiKhoanTBGN") or row.get("Số tk") or r.get("Số tk") or default_acc).strip()
    bank_cdu = str(r.get("projectBank") or r.get("taiNH") or r.get("tenNganHang") or row.get("tên NH") or r.get("tên NH") or default_bank).strip()

    row["Số tk"] = acc_cdu
    row["Số_tk"] = acc_cdu
    row["số_tk"] = acc_cdu
    row["so_tk"] = acc_cdu
    row["Số_TK"] = acc_cdu
    row["STK"] = acc_cdu
    row["stk"] = acc_cdu
    row["Số tài khoản"] = acc_cdu
    row["số tài khoản"] = acc_cdu
    row["số_tài_khoản"] = acc_cdu
    row["Số_tài_khoản"] = acc_cdu
    row["Số_tài_khoản_trên_TBGN"] = acc_cdu
    row["số_tài_khoản_trên_tbgn"] = acc_cdu

    row["tên NH"] = bank_cdu
    row["tên_NH"] = bank_cdu
    row["Tên_NH"] = bank_cdu
    row["ten_nh"] = bank_cdu
    row["Tên NH"] = bank_cdu
    row["Tại_NH"] = bank_cdu
    row["tại_nh"] = bank_cdu
    row["Tại NH"] = bank_cdu
    row["tại NH"] = bank_cdu
    row["tên ngân hàng"] = bank_cdu
    row["Tên ngân hàng"] = bank_cdu
    row["tên_ngân_hàng"] = bank_cdu
    row["Tên_ngân_hàng"] = bank_cdu

    row["Tên_người_thụ_hưởng"] = str(r.get("tenNguoiThuHuong") or r.get("projectBaseAddress") or r.get("tenChuDauTu") or "").strip()
    row["tên_người_thụ_hưởng"] = row["Tên_người_thụ_hưởng"]
    row["Thông_báo_Bàn_giao_GCNQSDD"] = str(r.get("thongBaoBanGiaoGCNQSDD") or "").strip()
    row["thông_báo_bàn_giao_gcnqsdd"] = row["Thông_báo_Bàn_giao_GCNQSDD"]
    row["Số_Cif"] = str(r.get("soCIF") or r.get("cif") or r.get("maKH") or row.get("mã khách hàng") or "").strip()
    row["số_cif"] = row["Số_Cif"]

    nam_hdtc = str(r.get("namHDTC") or bbkt_nam or "").strip()
    row["Năm_HĐTC"] = nam_hdtc
    row["năm_hđtc"] = nam_hdtc

    # 20. SỐ LẦN GIẢI NGÂN VÀ SỐ LẦN GIẢI NGÂN TRỪ 1
    active_cols_count = 0
    for col in range(1, 9):
        p_val = _parse_m(r.get(f"pct{col}") or r.get(f"tỷ lệ đợt {col}") or r.get(f"tỷ_lệ_đợt_{col}") or 0)
        a_val = _parse_m(r.get(f"amt{col}") or r.get(f"số tiền đợt {col}") or r.get(f"số_tiền_đợt_{col}") or 0)
        raw_d = str(r.get(f"date{col}") or r.get(f"desc{col}") or r.get(f"ngày đợt {col}") or r.get(f"ngày_đợt_{col}") or "").strip()
        if p_val > 0 or a_val > 0 or len(raw_d) > 0:
            active_cols_count += 1
    if active_cols_count == 0:
        active_cols_count = 4
    rounds_minus_1 = max(0, active_cols_count - 1)

    row["số_lần_giải_ngân"] = active_cols_count
    row["Số_lần_giải_ngân"] = active_cols_count
    row["so_lan_giai_ngan"] = active_cols_count
    row["soLanGiaiNgan"] = active_cols_count

    row["số_lần_giải_ngân_trừ_1"] = rounds_minus_1
    row["Số_lần_giải_ngân_trừ_1"] = rounds_minus_1
    row["số_lần_GN_trừ_1"] = rounds_minus_1
    row["Số_lần_GN_trừ_1"] = rounds_minus_1
    row["số_lần_giải_ngân_-1"] = rounds_minus_1
    row["so_lan_giai_ngan_tru_1"] = rounds_minus_1
    row["soLanGiaiNganTru1"] = rounds_minus_1
    # 21. NGUỒN VỐN VAY
    nguon_von_val = str(r.get("nguonVon") or r.get("Nguồn vốn") or r.get("Nguồn_vốn") or r.get("nguồn_vốn") or "Nguồn vốn ngân sách thành phố").strip()
    if not nguon_von_val:
        nguon_von_val = "Nguồn vốn ngân sách thành phố"

    row["Nguồn vốn"] = nguon_von_val
    row["Nguồn_vốn"] = nguon_von_val
    row["nguồn_vốn"] = nguon_von_val.lower()
    row["NGUỒN_VỐN"] = nguon_von_val.upper()
    row["nguonVon"] = nguon_von_val
    row["nguon_von"] = remove_accents(nguon_von_val).lower() if 'remove_accents' in globals() else nguon_von_val.lower()
    # 22. THÔNG BÁO GIẢI NGÂN CÁC ĐỢT (số tb - ngày tb, v/v thanh toán, số tiền tb...)
    can_ho_so = str(r.get("canHoSo") or r.get("căn hộ số") or r.get("Căn hộ số") or row.get("căn hộ số") or "").strip()
    tb_list = r.get("thongBaoGiaiNgan") if isinstance(r.get("thongBaoGiaiNgan"), list) else []
    tb_map = {}
    vv_map = {}
    for item in tb_list:
        if isinstance(item, dict):
            d_idx = item.get("dot")
            if d_idx is not None:
                tb_map[d_idx] = str(item.get("soTB") or "").strip()
                vv_map[d_idx] = str(item.get("vv") or "").strip()

    # Đợt 1 / Mặc định
    tb_1 = str(
        r.get("soTBNgayTB") or r.get("soTBNgayTB1") or r.get("soThongBao") or
        r.get("số tb - ngày tb") or r.get("số tb - ngày tb1") or
        r.get("số_tb__ngày_tb") or r.get("Số_tb__ngày_tb") or r.get("số_tb__ngày_tb1") or
        tb_map.get(1) or ""
    ).strip()
    vv_1 = str(
        r.get("vvThanhToan") or r.get("vvThanhToan1") or
        r.get("v/v thanh toán") or r.get("v/v thanh toán1") or
        r.get("vv_thanh_toán") or r.get("Vv_thanh_toán") or r.get("vv_thanh_toán1") or
        vv_map.get(1) or ""
    ).strip()
    if not vv_1 and can_ho_so:
        vv_1 = f"Thanh toán tiền mua căn hộ số {can_ho_so}"

    amt_1_num = _parse_m(r.get("amt1") or r.get("số tiền đợt 1") or 0)
    amt_1_str = f"{int(amt_1_num):,}".replace(",", ".") if amt_1_num > 0 else ""
    amt_1_word = _word_m(amt_1_num) if amt_1_num > 0 else ""

    row["số tb - ngày tb"] = tb_1
    row["v/v thanh toán"] = vv_1
    row["số tiền tb"] = amt_1_str
    row["bằng chữ số tiền tb"] = amt_1_word
    row["số tb - ngày tb1"] = tb_1
    row["v/v thanh toán1"] = vv_1
    row["số tiền tb1"] = amt_1_str
    row["bằng chữ số tiền tb1"] = amt_1_word

    row["số_tb__ngày_tb"] = tb_1
    row["Số_tb__ngày_tb"] = tb_1
    row["so_tb__ngay_tb"] = tb_1
    row["vv_thanh_toán"] = vv_1
    row["Vv_thanh_toán"] = vv_1
    row["vv_thanh_toan"] = vv_1
    row["số_tb__ngày_tb1"] = tb_1
    row["Số_tb__ngày_tb1"] = tb_1
    row["vv_thanh_toán1"] = vv_1
    row["Vv_thanh_toán1"] = vv_1

    # Đợt 2 đến 10
    for k in range(2, 11):
        tb_k = str(
            r.get(f"soTBNgayTB{k}") or r.get(f"số tb - ngày tb{k}") or
            r.get(f"số tb - ngày tb lần {k}") or r.get(f"số_tb__ngày_tb{k}") or
            r.get(f"Số_tb__ngày_tb{k}") or r.get(f"số_tb__ngày_tb_lần_{k}") or
            tb_map.get(k) or ""
        ).strip()
        vv_k = str(
            r.get(f"vvThanhToan{k}") or r.get(f"v/v thanh toán{k}") or
            r.get(f"v/v thanh toán lần {k}") or r.get(f"vv_thanh_toán{k}") or
            r.get(f"Vv_thanh_toán{k}") or r.get(f"vv_thanh_toán_lần_{k}") or
            vv_map.get(k) or ""
        ).strip()
        if not vv_k and tb_k and can_ho_so:
            vv_k = f"Thanh toán tiền mua căn hộ số {can_ho_so} đợt {k}"

        amt_k_num = _parse_m(r.get(f"amt{k}") or r.get(f"số tiền đợt {k}") or 0)
        amt_k_str = f"{int(amt_k_num):,}".replace(",", ".") if amt_k_num > 0 else ""
        amt_k_word = _word_m(amt_k_num) if amt_k_num > 0 else ""

        row[f"số tb - ngày tb{k}"] = tb_k
        row[f"v/v thanh toán{k}"] = vv_k
        row[f"số tiền tb{k}"] = amt_k_str
        row[f"bằng chữ số tiền tb{k}"] = amt_k_word

        row[f"số tb - ngày tb lần {k}"] = tb_k
        row[f"v/v thanh toán lần {k}"] = vv_k
        row[f"số tiền tb lần {k}"] = amt_k_str
        row[f"bằng chữ số tiền tb lần {k}"] = amt_k_word

        row[f"số_tb__ngày_tb{k}"] = tb_k
        row[f"Số_tb__ngày_tb{k}"] = tb_k
        row[f"số_tb__ngày_tb_{k}"] = tb_k
        row[f"số_tb__ngày_tb_lần_{k}"] = tb_k
        row[f"Số_tb__ngày_tb_lần_{k}"] = tb_k
        row[f"so_tb__ngay_tb{k}"] = tb_k
        row[f"vv_thanh_toán{k}"] = vv_k
        row[f"Vv_thanh_toán{k}"] = vv_k
        row[f"vv_thanh_toán_{k}"] = vv_k
        row[f"vv_thanh_toán_lần_{k}"] = vv_k
        row[f"Vv_thanh_toán_lần_{k}"] = vv_k
        row[f"vv_thanh_toan{k}"] = vv_k

    # Đảm bảo toàn bộ headers trong HEADERS_206 đều có key (ít nhất là "")
    for h in HEADERS_206:
        if h not in row:
            row[h] = ""

    return row

def sync_all_database_records(db_recs=None):
    """Đảm bảo database_records_rows206.json và DATA_MAIL_MERGE_SYNC.xlsx
    luôn có ĐẦY ĐỦ 100% hồ sơ từ database_records.json."""
    if db_recs is None:
        if not os.path.exists(DB_FILE):
            return []
        try:
            with open(DB_FILE, 'r', encoding='utf-8') as f:
                db_recs = json.load(f)
        except Exception:
            return []

    if not isinstance(db_recs, list) or len(db_recs) == 0:
        return []

    rows206_file = os.path.join(BASE_DIR, "database_records_rows206.json")
    sync_file = os.path.join(BASE_DIR, "DATA_MAIL_MERGE_SYNC.xlsx")

    # Đọc cache rows206 hiện có nếu có để tái sử dụng các trường đã tính toán nếu đã chuẩn
    existing_rows_map = {}
    if os.path.exists(rows206_file):
        try:
            with open(rows206_file, 'r', encoding='utf-8') as rf:
                cached_r206 = json.load(rf)
                if isinstance(cached_r206, list):
                    for r2 in cached_r206:
                        if isinstance(r2, dict):
                            k2 = get_rec_key(r2)
                            existing_rows_map[k2] = r2
                            cccd2 = str(r2.get('CCCD') or r2.get('cccd') or '').strip().replace(' ', '').replace('.', '').replace('-', '').replace('\t', '')
                            if cccd2 and len(cccd2) >= 8:
                                existing_rows_map[f"cccd_{cccd2}"] = r2
        except Exception:
            pass

    full_rows206 = []
    for r in db_recs:
        if not isinstance(r, dict):
            continue
        k = get_rec_key(r)
        r206_entry = existing_rows_map.get(k)
        if not r206_entry:
            cccd = str(r.get('cccd') or r.get('soCCCD') or r.get('CCCD') or '').strip().replace(' ', '').replace('.', '').replace('-', '').replace('\t', '')
            if cccd:
                r206_entry = existing_rows_map.get(f"cccd_{cccd}")

        if r206_entry:
            enriched = dict(r206_entry)
            # Cập nhật thêm các giá trị mới nhất từ r
            enriched.update(build_row206_from_record(r))
            full_rows206.append(enriched)
        else:
            full_rows206.append(build_row206_from_record(r))

    # Ghi database_records_rows206.json
    try:
        safe_unhide_file(rows206_file)
        with open(rows206_file, 'w', encoding='utf-8') as rf:
            json.dump(full_rows206, rf, ensure_ascii=False, indent=2)
        safe_rehide_file(rows206_file)
    except Exception:
        pass

    # Ghi DATA_MAIL_MERGE_SYNC.xlsx
    if full_rows206:
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "data"
            ws.append(HEADERS_206)
            for r_item in full_rows206:
                row_vals = []
                for h in HEADERS_206:
                    v = r_item.get(h, "")
                    if isinstance(v, (list, dict)):
                        row_vals.append(json.dumps(v, ensure_ascii=False) if v else "")
                    elif v == 0 or str(v).strip() == "0" or str(v).strip().lower() in ["không đồng", "không", "null", "none"]:
                        row_vals.append("")
                    else:
                        row_vals.append(v)
                ws.append(row_vals)
            safe_unhide_file(sync_file)
            wb.save(sync_file)
            safe_rehide_file(sync_file)
        except Exception:
            pass

    return full_rows206

def is_single_profile(data_dict):
    if not isinstance(data_dict, dict):
        return False
    hn = str(data_dict.get("tinhTrangHonNhan") or data_dict.get("Tình trạng hôn nhân") or data_dict.get("honNhan") or "").strip().lower()
    if any(k in hn for k in ["độc thân", "doc than", "ly hôn", "ly hon", "chưa kết hôn", "single"]):
        return True
    if any(k in hn for k in ["kết hôn", "ket hon", "gia đình", "gia dinh", "married"]):
        return False

    # Kiểm tra từ tên thư mục đã xác định (nếu có)
    fld = str(data_dict.get("_resolved_folder") or data_dict.get("case_folder") or "").strip().lower()
    if "doc than" in fld or "độc thân" in fld:
        return True

    # Fallback: kiểm tra xem có thông tin vợ/chồng thực tế hay không
    vc_name = str(data_dict.get("Họ và tên Vợ/chồng") or data_dict.get("vcHoTen") or data_dict.get("Họ và tên vợ/chồng in hoa") or "").strip()
    vc_cccd = str(data_dict.get("CCCD Vợ/chồng") or data_dict.get("vcCCCD") or "").strip()
    return not (vc_name or vc_cccd)

def fill_docx_template(docx_path, out_path, data_dict):
    if data_dict is None:
        data_dict = {}
    data_dict = dict(data_dict)
    if isinstance(data_dict, dict) and "build_row206_from_record" in globals():
        if not ("tổng_hợp_phiếu_báo_có" in data_dict or "Họ và tên KH hoa" in data_dict):
            data_dict = build_row206_from_record(data_dict)
    single_mode = is_single_profile(data_dict)
    norm_data = {}
    for k, v in data_dict.items():
        if v is None:
            val_str = ""
        elif isinstance(v, (int, float)):
            if v == 0:
                if normalize_key(k) in ["chi_phí_người_phụ_thuộc", "chi_phi_nguoi_phu_thuoc", "chiphinguoiphuthuoc"]:
                    val_str = "0"
                else:
                    val_str = ""
            else:
                val_str = f"{int(v):,}".replace(",", ".") if v == int(v) else f"{v:,.2f}".replace(",", ".")
        else:
            s_val = str(v).replace('\t', '')
            if '\n' not in s_val:
                s_val = re.sub(r' {2,}', ' ', s_val).strip()
            else:
                s_val = s_val.strip()
            if s_val == "0" or s_val.lower() in ["không đồng", "không", "null", "none"]:
                if normalize_key(k) in ["chi_phí_người_phụ_thuộc", "chi_phi_nguoi_phu_thuoc", "chiphinguoiphuthuoc"]:
                    val_str = "0"
                else:
                    val_str = ""
            else:
                val_str = s_val
        norm_k = normalize_key(k)
        if val_str != "" or norm_k not in norm_data or norm_data[norm_k] == "":
            norm_data[norm_k] = val_str
        if val_str != "" or k not in norm_data or norm_data[k] == "":
            norm_data[k] = val_str
        k_under = k.replace(' ', '_').replace('/', '').replace('(', '').replace(')', '').replace('-', '').replace(',', '').replace('%', '')
        if val_str != "" or k_under not in norm_data or norm_data[k_under] == "":
            norm_data[k_under] = val_str
        norm_k_under = normalize_key(k_under)
        if val_str != "" or norm_k_under not in norm_data or norm_data[norm_k_under] == "":
            norm_data[norm_k_under] = val_str

    # Đảm bảo các biến ngày HĐTD, HĐTC, HĐ3B luôn luôn có giá trị đầy đủ trong norm_data và data_dict
    raw_hdtd_fmt = str(
        norm_data.get("ngày hđtd") or norm_data.get("ngày_hđtd") or
        norm_data.get("ngayHDTD") or norm_data.get("ngay_hdtd") or
        data_dict.get("ngayHDTD") or data_dict.get("ngày hđtd") or
        data_dict.get("ngày_hđtd") or data_dict.get("ngayVay") or
        data_dict.get("ngày vay") or data_dict.get("ngày_vay") or ""
    ).strip()

    raw_hdtd_full = str(
        norm_data.get("ngàythángnăm_hđtd") or norm_data.get("ngày_tháng_năm_hđtd") or
        norm_data.get("ngày/tháng/năm hđtd") or norm_data.get("ngaythangnam_hdtd") or
        data_dict.get("ngàythángnăm_hđtd") or data_dict.get("ngày/tháng/năm hđtd") or
        raw_hdtd_fmt or ""
    ).strip()

    full_td, d_td, m_td, y_td = format_admin_date(raw_hdtd_full or raw_hdtd_fmt)
    fmt_td = _fmt_date(raw_hdtd_fmt or raw_hdtd_full)

    if not full_td and fmt_td:
        full_td, d_td, m_td, y_td = format_admin_date(fmt_td)

    if not full_td:
        nam_cand = str(data_dict.get("Năm_HĐTD") or data_dict.get("namNhanHs") or data_dict.get("năm nhận hs vay vốn") or datetime.now().year)
        full_td = f"ngày ...... tháng ...... năm {nam_cand}"
        fmt_td = f"....../....../{nam_cand}"

    for k_td in [
        "ngàythángnăm_hđtd", "ngày_tháng_năm_hđtd", "ngày/tháng/năm hđtd",
        "Ngày__tháng__năm_HĐTD", "Ngày_tháng_năm_HĐTD", "ngaythangnam_hdtd",
        "ngày tháng năm hđtd", "Ngày, tháng, năm HĐTD"
    ]:
        norm_data[k_td] = full_td
        norm_data[normalize_key(k_td)] = full_td
        data_dict[k_td] = full_td

    for k_td_s in ["ngày hđtd", "ngày_hđtd", "ngay_hdtd", "ngayHDTD", "ngày HĐTD"]:
        norm_data[k_td_s] = fmt_td
        norm_data[normalize_key(k_td_s)] = fmt_td
        data_dict[k_td_s] = fmt_td

    # Đảm bảo các biến nguồn vốn luôn có giá trị đầy đủ trong norm_data và data_dict
    nv_val_norm = str(norm_data.get("nguonvon") or norm_data.get("nguồn_vốn") or norm_data.get("Nguồn_vốn") or norm_data.get("nguon_von") or norm_data.get("Nguồn vốn") or data_dict.get("nguonVon") or data_dict.get("Nguồn vốn") or "Nguồn vốn ngân sách thành phố").strip() or "Nguồn vốn ngân sách thành phố"
    nv_accents_removed = remove_accents(nv_val_norm).lower() if 'remove_accents' in globals() else nv_val_norm.lower()
    for nv_key, nv_v in [
        ("nguồn_vốn", nv_val_norm.lower()),
        ("Nguồn_vốn", nv_val_norm),
        ("NGUỒN_VỐN", nv_val_norm.upper()),
        ("nguon_von", nv_accents_removed),
        ("Nguon_von", nv_val_norm),
        ("NGUON_VON", nv_accents_removed.upper()),
        ("nguonVon", nv_val_norm),
        ("Nguồn vốn", nv_val_norm),
        ("nguồn vốn", nv_val_norm.lower()),
    ]:
        norm_data[nv_key] = nv_v
        norm_data[normalize_key(nv_key)] = nv_v
        data_dict[nv_key] = nv_v

    # Đảm bảo các biến xác nhận lương / thu nhập luôn có giá trị đầy đủ trong norm_data và data_dict
    xn_val = str(norm_data.get("xác_nhận_lương") or norm_data.get("xac_nhan_luong") or data_dict.get("xacNhanLuong") or data_dict.get("xác nhận lương") or data_dict.get("xác_nhận_lương") or "").strip()
    if not xn_val:
        tt_hn = str(data_dict.get("tinhTrangHonNhan") or data_dict.get("Tình trạng hôn nhân") or "").strip()
        is_single = (tt_hn == "Độc thân" or tt_hn == "Ly hôn" or single_mode)
        raw_kh = str(data_dict.get("hoTen") or data_dict.get("Họ và tên KH thường") or data_dict.get("Họ và tên KH hoa") or data_dict.get("Họ và tên") or "").strip()
        raw_vc = str(data_dict.get("vcHoTen") or data_dict.get("Họ và tên Vợ/chồng") or "").strip()
        name_kh = format_title_case_py(raw_kh) if 'format_title_case_py' in globals() else raw_kh.title()
        name_vc = format_title_case_py(raw_vc) if 'format_title_case_py' in globals() else raw_vc.title()
        if name_kh:
            if is_single or not name_vc:
                xn_val = f"Xác nhận thu nhập ({name_kh})"
            else:
                xn_val = f"Xác nhận thu nhập ({name_kh}), Kê khai thu nhập ({name_vc})"
    if xn_val:
        for k_xn in ["xác_nhận_lương", "xác nhận lương", "xác_nhận_thu_nhập", "xác nhận thu nhập", "xac_nhan_luong", "xac_nhan_thu_nhap", "Xác_nhận_lương", "Xác_nhận_thu_nhập"]:
            norm_data[k_xn] = xn_val
            norm_data[normalize_key(k_xn)] = xn_val
            data_dict[k_xn] = xn_val

    # Đảm bảo các biến Thông tin phụ thuộc luôn có giá trị đầy đủ trong norm_data và data_dict
    ttpt_norm = str(norm_data.get("thông_tin_phụ_thuộc") or norm_data.get("thong_tin_phu_thuoc") or norm_data.get("người_phụ_thuộc") or norm_data.get("nguoi_phu_thuoc") or norm_data.get("thongtinphuthuoc") or norm_data.get("nguoiphuthuoc") or data_dict.get("thongTinPhuThuoc") or data_dict.get("thông tin phụ thuộc") or data_dict.get("thông_tin_phụ_thuộc") or data_dict.get("người_phụ_thuộc") or data_dict.get("người phụ thuộc") or "").strip()
    if ttpt_norm:
        for k_pt in ["thông_tin_phụ_thuộc", "thông tin phụ thuộc", "thong_tin_phu_thuoc", "người_phụ_thuộc", "người phụ thuộc", "nguoi_phu_thuoc", "Thông_tin_phụ_thuộc", "Người_phụ_thuộc", "thongTinPhuThuoc", "nguoi_phu_thuoc"]:
            norm_data[k_pt] = ttpt_norm
            norm_data[normalize_key(k_pt)] = ttpt_norm
            data_dict[k_pt] = ttpt_norm

    # Đảm bảo các biến Ngày làm đơn vay vốn và ngày/tháng/năm nhận hs luôn có giá trị đầy đủ trong norm_data và data_dict
    nl_raw = str(
        norm_data.get("ngày_làm_đơn_vay_vốn") or norm_data.get("Ngày_làm_đơn_vay_vốn") or
        norm_data.get("ngày làm đơn vay vốn") or norm_data.get("ngay_lam_don_vay_von") or
        norm_data.get("ngaylamdonvayvon") or norm_data.get("ngaytiepnhan") or
        data_dict.get("ngayTiepNhan") or data_dict.get("Ngày làm đơn vay vốn") or
        data_dict.get("ngayLamDon") or data_dict.get("ngayLamDonVay") or data_dict.get("ngayNhanHs") or ""
    ).strip()
    full_nl_t, d_nl_t, m_nl_t, y_nl_t = format_admin_date(nl_raw)
    if not d_nl_t:
        d_nl_t = str(norm_data.get("ngày_nhận_hs_vay_vốn") or norm_data.get("ngaynhanhsvayvon") or data_dict.get("ngày nhận hs vay vốn") or "").strip()
    if not m_nl_t:
        m_nl_t = str(norm_data.get("tháng_nhận_hs_vay_vốn") or norm_data.get("thangnhanhsvayvon") or data_dict.get("tháng nhận hs vay vốn") or "").strip()
    if not y_nl_t:
        y_nl_t = str(norm_data.get("năm_nhận_hs_vay_vốn") or norm_data.get("namnhanhsvayvon") or data_dict.get("năm nhận hs vay vốn") or "").strip()
    if not full_nl_t and d_nl_t and m_nl_t and y_nl_t:
        full_nl_t = f"ngày {d_nl_t.zfill(2)} tháng {m_nl_t.zfill(2)} năm {y_nl_t}"

    if full_nl_t:
        for k_nl in ["ngày_làm_đơn_vay_vốn", "Ngày_làm_đơn_vay_vốn", "ngay_lam_don_vay_von", "ngày làm đơn vay vốn", "Ngày làm đơn vay vốn", "NGÀY_LÀM_ĐƠN_VAY_VỐN"]:
            norm_data[k_nl] = full_nl_t
            norm_data[normalize_key(k_nl)] = full_nl_t
            data_dict[k_nl] = full_nl_t
    if d_nl_t:
        for k_d in ["ngày_nhận_hs_vay_vốn", "ngày nhận hs vay vốn", "ngay_nhan_hs_vay_von", "ngaynhanhsvayvon"]:
            norm_data[k_d] = d_nl_t
            norm_data[normalize_key(k_d)] = d_nl_t
            data_dict[k_d] = d_nl_t
    if m_nl_t:
        for k_m in ["tháng_nhận_hs_vay_vốn", "tháng nhận hs vay vốn", "thang_nhan_hs_vay_von", "thangnhanhsvayvon"]:
            norm_data[k_m] = m_nl_t
            norm_data[normalize_key(k_m)] = m_nl_t
            data_dict[k_m] = m_nl_t
    if y_nl_t:
        for k_y in ["năm_nhận_hs_vay_vốn", "năm nhận hs vay vốn", "nam_nhan_hs_vay_von", "namnhanhsvayvon"]:
            norm_data[k_y] = y_nl_t
            norm_data[normalize_key(k_y)] = y_nl_t
            data_dict[k_y] = y_nl_t

    # Đảm bảo các biến ngày, tháng, năm TB phê duyệt cho vay luôn có giá trị đầy đủ trong norm_data và data_dict
    raw_tb_pd_fill = str(
        norm_data.get("ngàythángnăm_tb_phê_duyệt_cho_vay") or norm_data.get("ngày,tháng,năm tb phê duyệt cho vay") or
        norm_data.get("ngày_tháng_năm_tb_phê_duyệt_cho_vay") or norm_data.get("ngày tb phê duyệt cho vay") or
        norm_data.get("ngaythangnam_tb_phe_duyet_cho_vay") or
        data_dict.get("ngayThangNamThongBao") or data_dict.get("ngày,tháng,năm tb phê duyệt cho vay") or
        data_dict.get("ngayThongBao") or data_dict.get("ngày tb phê duyệt cho vay") or
        data_dict.get("thamDinhDen") or data_dict.get("thẩm định đến ngày") or ""
    ).strip()
    full_tb_fill, d_tb_fill, m_tb_fill, y_tb_fill = format_admin_date(raw_tb_pd_fill)
    fmt_tb_fill = _fmt_date(raw_tb_pd_fill) if '_fmt_date' in globals() else raw_tb_pd_fill
    if not full_tb_fill and fmt_tb_fill:
        m_parts = fmt_tb_fill.split('/')
        if len(m_parts) == 3:
            full_tb_fill = f"ngày {m_parts[0]} tháng {m_parts[1]} năm {m_parts[2]}"
    if not full_tb_fill:
        nam_tb_cand = str(data_dict.get("Năm_HĐTD") or data_dict.get("namNhanHs") or datetime.now().year)
        full_tb_fill = f"ngày ...... tháng ...... năm {nam_tb_cand}"
        fmt_tb_fill = f"....../....../{nam_tb_cand}"

    for k_tb_pd in [
        "ngàythángnăm_tb_phê_duyệt_cho_vay", "ngày,tháng,năm tb phê duyệt cho vay",
        "ngày_tháng_năm_tb_phê_duyệt_cho_vay", "ngaythangnam_tb_phe_duyet_cho_vay",
        "Ngày__tháng__năm_TB_phê_duyệt_cho_vay", "Ngày_tháng_năm_tb_phê_duyệt_cho_vay",
        "ngày tháng năm tb phê duyệt cho vay"
    ]:
        norm_data[k_tb_pd] = full_tb_fill
        norm_data[normalize_key(k_tb_pd)] = full_tb_fill
        data_dict[k_tb_pd] = full_tb_fill

    for k_tb_pd_s in ["ngày tb phê duyệt cho vay", "ngày_tb_phê_duyệt_cho_vay", "ngay_tb_phe_duyet_cho_vay", "ngayThongBao"]:
        norm_data[k_tb_pd_s] = fmt_tb_fill
        norm_data[normalize_key(k_tb_pd_s)] = fmt_tb_fill
        data_dict[k_tb_pd_s] = fmt_tb_fill

    # Đảm bảo các biến Số TB Ngày TB và V/V thanh toán luôn có giá trị đầy đủ trong norm_data và data_dict
    tpl_filename = os.path.basename(docx_path).lower() if 'os' in globals() else ""
    is_round_1_template = any(x in tpl_filename for x in ["lần 1", "lan 1", "lần_1", "lan_1", "10_1"])
    is_round_2_template = not is_round_1_template and any(x in tpl_filename for x in ["lần 2", "lan 2", "lần_2", "lan_2", "10_2", "11_1"])
    is_round_3_template = not is_round_1_template and not is_round_2_template and any(x in tpl_filename for x in ["lần 3", "lan 3", "lần_3", "lan_3", "10_3"])

    # Thu thập dữ liệu các đợt
    tb_by_round = {}
    vv_by_round = {}

    # Đợt 1
    tb_by_round[1] = str(
        data_dict.get("số_tb__ngày_tb1") or data_dict.get("Số_tb__ngày_tb1") or
        data_dict.get("số tb - ngày tb1") or data_dict.get("soTBNgayTB1") or
        data_dict.get("soTBNgayTB") or data_dict.get("soThongBao") or
        data_dict.get("số_tb__ngày_tb") or data_dict.get("Số_tb__ngày_tb") or
        data_dict.get("số tb - ngày tb") or
        norm_data.get("số_tb__ngày_tb1") or norm_data.get("số_tb__ngày_tb") or
        norm_data.get(normalize_key("số_tb__ngày_tb1")) or norm_data.get(normalize_key("số_tb__ngày_tb")) or ""
    ).strip()

    vv_by_round[1] = str(
        data_dict.get("vv_thanh_toán1") or data_dict.get("Vv_thanh_toán1") or
        data_dict.get("v/v thanh toán1") or data_dict.get("vvThanhToan1") or
        data_dict.get("vvThanhToan") or
        data_dict.get("vv_thanh_toán") or data_dict.get("Vv_thanh_toán") or
        data_dict.get("v/v thanh toán") or
        norm_data.get("vv_thanh_toán1") or norm_data.get("vv_thanh_toán") or
        norm_data.get(normalize_key("vv_thanh_toán1")) or norm_data.get(normalize_key("vv_thanh_toán")) or ""
    ).strip()

    # Các đợt 2..10
    for r_idx in range(2, 11):
        tb_by_round[r_idx] = str(
            data_dict.get(f"số_tb__ngày_tb{r_idx}") or data_dict.get(f"Số_tb__ngày_tb{r_idx}") or
            data_dict.get(f"số_tb__ngày_tb_lần_{r_idx}") or data_dict.get(f"Số_tb__ngày_tb_lần_{r_idx}") or
            data_dict.get(f"số tb - ngày tb{r_idx}") or data_dict.get(f"số tb - ngày tb lần {r_idx}") or
            data_dict.get(f"soTBNgayTB{r_idx}") or
            norm_data.get(f"số_tb__ngày_tb{r_idx}") or norm_data.get(normalize_key(f"số_tb__ngày_tb{r_idx}")) or ""
        ).strip()

        vv_by_round[r_idx] = str(
            data_dict.get(f"vv_thanh_toán{r_idx}") or data_dict.get(f"Vv_thanh_toán{r_idx}") or
            data_dict.get(f"vv_thanh_toán_lần_{r_idx}") or data_dict.get(f"Vv_thanh_toán_lần_{r_idx}") or
            data_dict.get(f"v/v thanh toán{r_idx}") or data_dict.get(f"v/v thanh toán lần {r_idx}") or
            data_dict.get(f"vvThanhToan{r_idx}") or
            norm_data.get(f"vv_thanh_toán{r_idx}") or norm_data.get(normalize_key(f"vv_thanh_toán{r_idx}")) or ""
        ).strip()

    # Kiểm tra thêm trong list thongBaoGiaiNgan nếu có
    tb_raw_list = data_dict.get("thongBaoGiaiNgan")
    if isinstance(tb_raw_list, list):
        for item in tb_raw_list:
            if isinstance(item, dict):
                d_num = item.get("dot")
                if d_num in range(1, 11):
                    if not tb_by_round[d_num] and item.get("soTB"):
                        tb_by_round[d_num] = str(item.get("soTB")).strip()
                    if not vv_by_round[d_num] and item.get("vv"):
                        vv_by_round[d_num] = str(item.get("vv")).strip()

    # Xác định giá trị chính cho các biến không số đuôi «số_tb__ngày_tb» và «vv_thanh_toán» theo ngữ cảnh template
    primary_tb = ""
    primary_vv = ""
    if is_round_2_template:
        primary_tb = tb_by_round[2] or tb_by_round[1]
        primary_vv = vv_by_round[2] or vv_by_round[1]
    elif is_round_3_template:
        primary_tb = tb_by_round[3] or tb_by_round[2] or tb_by_round[1]
        primary_vv = vv_by_round[3] or vv_by_round[2] or vv_by_round[1]
    else:
        primary_tb = tb_by_round[1] or tb_by_round[2]
        primary_vv = vv_by_round[1] or vv_by_round[2]

    # Fallback tự sinh V/V thanh toán nếu có số TB mà chưa có v/v thanh toán
    c_ho = str(norm_data.get("căn_hộ_số") or data_dict.get("canHoSo") or data_dict.get("Căn hộ số") or "").strip()
    if not primary_vv and primary_tb and c_ho:
        primary_vv = f"Thanh toán tiền mua căn hộ số {c_ho}"

    # Đổ tất cả các alias KHÔNG SỐ ĐUÔI vào norm_data và data_dict
    if primary_tb:
        for k_tb in [
            "số_tb__ngày_tb", "Số_tb__ngày_tb", "so_tb__ngay_tb", "số_tb_ngày_tb", "so_tb_ngay_tb",
            "số tb - ngày tb", "so tb - ngay tb", "sotbngaytb", "Số_tb_ngày_tb"
        ]:
            norm_data[k_tb] = primary_tb
            norm_data[normalize_key(k_tb)] = primary_tb
            data_dict[k_tb] = primary_tb

    if primary_vv:
        for k_vv in [
            "vv_thanh_toán", "Vv_thanh_toán", "vv_thanh_toan", "v/v thanh toán", "vvthanhtoan",
            "v/v_thanh_toán", "v/v_thanh_toan", "V/v thanh toán", "V/V thanh toán"
        ]:
            norm_data[k_vv] = primary_vv
            norm_data[normalize_key(k_vv)] = primary_vv
            data_dict[k_vv] = primary_vv

    # Đổ TỪNG ĐỢT RIÊNG BIỆT vào norm_data và data_dict theo đúng đợt của nó
    for r_idx in range(1, 11):
        tb_r = tb_by_round[r_idx]
        vv_r = vv_by_round[r_idx]
        if not vv_r and tb_r and c_ho:
            vv_r = f"Thanh toán tiền mua căn hộ số {c_ho}" if r_idx == 1 else f"Thanh toán tiền mua căn hộ số {c_ho} đợt {r_idx}"

        if tb_r:
            for k_tb_r in [
                f"số_tb__ngày_tb{r_idx}", f"Số_tb__ngày_tb{r_idx}", f"so_tb__ngay_tb{r_idx}",
                f"số_tb__ngày_tb_{r_idx}", f"Số_tb__ngày_tb_{r_idx}",
                f"số_tb__ngày_tb_lần_{r_idx}", f"Số_tb__ngày_tb_lần_{r_idx}",
                f"số tb - ngày tb{r_idx}", f"số tb - ngày tb lần {r_idx}",
                f"sotbngaytb{r_idx}", f"so_tb_ngay_tb_{r_idx}"
            ]:
                norm_data[k_tb_r] = tb_r
                norm_data[normalize_key(k_tb_r)] = tb_r
                data_dict[k_tb_r] = tb_r

        if vv_r:
            for k_vv_r in [
                f"vv_thanh_toán{r_idx}", f"Vv_thanh_toán{r_idx}", f"vv_thanh_toan{r_idx}",
                f"vv_thanh_toán_{r_idx}", f"Vv_thanh_toán_{r_idx}",
                f"vv_thanh_toán_lần_{r_idx}", f"Vv_thanh_toán_lần_{r_idx}",
                f"v/v thanh toán{r_idx}", f"v/v thanh toán lần {r_idx}",
                f"vvthanhtoan{r_idx}", f"v_v_thanh_toan_{r_idx}"
            ]:
                norm_data[k_vv_r] = vv_r
                norm_data[normalize_key(k_vv_r)] = vv_r
                data_dict[k_vv_r] = vv_r

    # Đảm bảo Số tk và tên NH luôn luôn có giá trị đầy đủ trong norm_data và data_dict
    p_name_fill = str(norm_data.get("dự_án") or norm_data.get("project") or data_dict.get("project") or data_dict.get("Dự án") or "").strip().lower()
    def_fill_acc = "113002951528"
    def_fill_bank = "Ngân hàng TMCP Công thương Việt Nam - CN Bình Xuyên"
    if "nam long" in p_name_fill:
        def_fill_acc = "7063020417"
        def_fill_bank = "Phòng giao dịch Ngân hàng Chính sách xã hội Cái Răng"
    elif "hồng loan" in p_name_fill or "hong loan" in p_name_fill:
        def_fill_acc = "7063261897"
        def_fill_bank = "Phòng giao dịch Ngân hàng Chính sách xã hội Cái Răng"

    val_fill_acc = str(
        norm_data.get("số_tk") or norm_data.get("Số_tk") or norm_data.get("Số tk") or
        norm_data.get("so_tk") or norm_data.get("số_tài_khoản_trên_tbgn") or
        data_dict.get("Số tk") or data_dict.get("Số_tk") or data_dict.get("projectAccount") or
        data_dict.get("soTaiKhoan") or data_dict.get("soTaiKhoanTBGN") or def_fill_acc
    ).strip() or def_fill_acc

    val_fill_bank = str(
        norm_data.get("tên_nh") or norm_data.get("tên_NH") or norm_data.get("tên NH") or
        norm_data.get("ten_nh") or norm_data.get("tại_nh") or norm_data.get("Tại_NH") or
        data_dict.get("tên NH") or data_dict.get("tên_NH") or data_dict.get("projectBank") or
        data_dict.get("taiNH") or data_dict.get("tenNganHang") or def_fill_bank
    ).strip() or def_fill_bank

    for k_stk_fill in [
        "Số tk", "Số_tk", "số_tk", "so_tk", "Số_TK", "STK", "stk",
        "Số tài khoản", "số tài khoản", "số_tài_khoản", "Số_tài_khoản",
        "Số_tài_khoản_trên_TBGN", "số_tài_khoản_trên_tbgn"
    ]:
        norm_data[k_stk_fill] = val_fill_acc
        norm_data[normalize_key(k_stk_fill)] = val_fill_acc
        data_dict[k_stk_fill] = val_fill_acc

    for k_nh_fill in [
        "tên NH", "tên_NH", "Tên_NH", "ten_nh", "Tên NH",
        "Tại_NH", "tại_nh", "Tại NH", "tại NH",
        "tên ngân hàng", "Tên ngân hàng", "tên_ngân_hàng", "Tên_ngân_hàng"
    ]:
        norm_data[k_nh_fill] = val_fill_bank
        norm_data[normalize_key(k_nh_fill)] = val_fill_bank
        data_dict[k_nh_fill] = val_fill_bank

    if single_mode:
        # Nếu khách hàng Độc thân / Ly hôn: Xóa trắng toàn bộ biến vợ/chồng trong từ điển
        spouse_var_keys = [
            "DX vợ/chồng HOA", "dx vợ/chồng thường", "quan hệ với người vay",
            "tên CCCD/CC vợ/chồng viết đầy đủ", "tên CCCD/CC vợ/chồng", "Họ và tên Vợ/chồng",
            "Họ và tên vợ/chồng in hoa", "dân tộc vợ/chồng", "Giới tính vợ/chồng", "Ngày,tháng,năm sinh vợ/chồng",
            "CCCD Vợ/chồng", "Ngày cấp CCCD vợ/chồng", "Nơi cấp CCCD vợ/chồng", "Địa chỉ cư trú vợ/chồng",
            "nơi ở hiện tại vợ/chồng", "Thu nhập vợ/chồng", "nghề nghiệp vợ/chồng", "nơi làm việc vợ/chồng",
            "Số điện thoại vợ/chồng"
        ]
        for sk in spouse_var_keys:
            norm_data[normalize_key(sk)] = ""
            norm_data[sk] = ""
            sk_under = sk.replace(' ', '_').replace('/', '').replace('(', '').replace(')', '').replace('-', '').replace(',', '').replace('%', '')
            norm_data[sk_under] = ""
            norm_data[normalize_key(sk_under)] = ""

    norm_key_map = {}
    norm_no_acc_map = {}
    for d_source in (data_dict, norm_data):
        for k_src, v_src in d_source.items():
            if v_src is not None and str(v_src).strip() != "":
                nk_src = normalize_key(k_src)
                if nk_src:
                    norm_key_map[nk_src] = v_src
                    if 'strip_vn_accents' in globals():
                        nk_no_acc = normalize_key(strip_vn_accents(k_src))
                        if nk_no_acc:
                            norm_no_acc_map[nk_no_acc] = v_src

    def get_val_smart(key_name):
        if not key_name:
            return None
        k_clean = str(key_name).strip().strip('"').strip("'")
        if not k_clean:
            return None
        if k_clean in norm_data and norm_data[k_clean] != "":
            return norm_data[k_clean]
        if k_clean in data_dict and data_dict[k_clean] != "":
            return data_dict[k_clean]
        nk = normalize_key(k_clean)
        if nk in norm_data and norm_data[nk] != "":
            return norm_data[nk]
        if nk in norm_key_map:
            return norm_key_map[nk]
        k_sp = k_clean.replace('_', ' ')
        if k_sp in norm_data and norm_data[k_sp] != "":
            return norm_data[k_sp]
        if k_sp in data_dict and data_dict[k_sp] != "":
            return data_dict[k_sp]
        nk_sp = normalize_key(k_sp)
        if nk_sp in norm_key_map:
            return norm_key_map[nk_sp]
        k_un = k_clean.replace(' ', '_')
        if k_un in norm_data and norm_data[k_un] != "":
            return norm_data[k_un]
        if k_un in data_dict and data_dict[k_un] != "":
            return data_dict[k_un]
        nk_un = normalize_key(k_un)
        if nk_un in norm_key_map:
            return norm_key_map[nk_un]
        if 'strip_vn_accents' in globals():
            k_no_acc = strip_vn_accents(k_clean)
            nk_no_acc = normalize_key(k_no_acc)
            if nk_no_acc in norm_no_acc_map:
                return norm_no_acc_map[nk_no_acc]
        return None

    from io import BytesIO as _BytesIO
    _buf = _BytesIO()
    with zipfile.ZipFile(docx_path, 'r') as zin:
        with zipfile.ZipFile(_buf, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                content = zin.read(item.filename)
                is_content_xml = item.filename.startswith('word/') and item.filename.endswith('.xml') and any(
                    item.filename.startswith(p) for p in ['word/document', 'word/header', 'word/footer']
                )
                if not is_content_xml:
                    zout.writestr(item, content)
                    continue

                xml_str = content.decode('utf-8')
                if (item.filename == 'word/document.xml' or any(tag in xml_str for tag in ['MERGEFIELD', '«', '\xab', '<<', 'w:fldSimple'])):

                    # SMART DYNAMIC TABLE & PARAGRAPH: Tự động loại bỏ dữ liệu vợ/chồng nếu khách hàng Độc thân / Ly hôn
                    if item.filename == 'word/document.xml' and single_mode:
                        has_spouse_tags = bool(re.search(r'vợchồng|vợ/chồng|vochong|vợ_chồng', xml_str, re.IGNORECASE))
                        if has_spouse_tags:
                            # 1. Regex tìm và triệt tiêu hàng <w:tr> chỉ chứa thông tin của vợ/chồng trong bảng Table 2 (Mẫu 1)
                            # TUYỆT ĐỐI KHÔNG XÓA hàng nếu hàng đó có chứa thông tin của Khách hàng / Người vay (như Mẫu 11 Biên bản giao nhận hồ sơ)
                            def _filter_spouse_tr(m_tr):
                                tr_xml = m_tr.group(0)
                                if not re.search(r'Họ_và_tên_Vợchồng|CCCD_Vợchồng|Thu_nhập_vợchồng|nghề_nghiệp_vợchồng|quan_hệ_với_người_vay|tên_CCCDCC_vợchồng|Họ_và_tên_vợchồng_in_hoa|Họ và tên Vợ/chồng|CCCD Vợ/chồng', tr_xml, re.IGNORECASE):
                                    return tr_xml
                                
                                kh_tokens_tr = [
                                    'Họ_và_tên_KH', 'Họ và tên KH', 'KH_thường', 'khach_hang', 'khachhang',
                                    'Căn_cước_công_dân', 'Căn cước công dân',
                                    '«CCCD»', 'MERGEFIELD  CCCD ', 'MERGEFIELD CCCD ',
                                    'xác nhận lương', 'xac nhan luong', 'Thu_nhập_KH', 'Thu nhập KH'
                                ]
                                for kt in kh_tokens_tr:
                                    if kt.lower() in tr_xml.lower():
                                        return tr_xml
                                
                                return ''

                            xml_str = re.sub(r'<w:tr\b[^>]*>[\s\S]*?</w:tr>', _filter_spouse_tr, xml_str)

                            # 1b. Dọn dẹp đoạn văn bản ghép chứa cả Khách hàng và Vợ/chồng (như Mẫu 11: Biên bản giao nhận hồ sơ)
                            # Giữ lại trọn vẹn thông tin khách hàng đến dấu đóng ngoặc ')', cắt bỏ phần thông tin vợ/chồng phía sau
                            def _clean_joint_kh_spouse_p(p_xml):
                                has_kh = bool(re.search(r'Họ_và_tên_KH|Họ và tên KH', p_xml, re.IGNORECASE))
                                has_vc = bool(re.search(r'Họ_và_tên_Vợchồng|Họ và tên Vợ/chồng|vợchồng|vợ/chồng', p_xml, re.IGNORECASE))
                                if not (has_kh and has_vc):
                                    return p_xml

                                runs = list(re.finditer(r'<w:r\b[^>]*>[\s\S]*?</w:r>', p_xml))
                                kh_run_idx = -1
                                for i, r_m in enumerate(runs):
                                    r_str = r_m.group(0)
                                    if re.search(r'Họ_và_tên_KH|Họ và tên KH', r_str, re.IGNORECASE):
                                        kh_run_idx = i
                                        break
                                
                                if kh_run_idx == -1:
                                    return p_xml

                                close_idx = -1
                                has_extra_text = False
                                extra_run_str = ""
                                for j in range(kh_run_idx, len(runs)):
                                    r_str = runs[j].group(0)
                                    m_t = re.search(r'<w:t([^>]*)>([\s\S]*?\))([\s\S]*?)</w:t>', r_str)
                                    if m_t:
                                        close_idx = j
                                        t_attr = m_t.group(1)
                                        t_keep = m_t.group(2)
                                        t_tail = m_t.group(3)
                                        if t_tail.strip():
                                            has_extra_text = True
                                            extra_run_str = re.sub(r'<w:t[^>]*>[\s\S]*?</w:t>', f'<w:t{t_attr}>{t_keep}</w:t>', r_str)
                                        break

                                if close_idx == -1:
                                    return p_xml

                                if has_extra_text:
                                    return p_xml[:runs[close_idx].start()] + extra_run_str + '</w:p>'
                                else:
                                    return p_xml[:runs[close_idx].end()] + '</w:p>'

                            xml_str = re.sub(r'<w:p\b[^>]*>[\s\S]*?</w:p>', lambda m: _clean_joint_kh_spouse_p(m.group(0)), xml_str)

                            # 2. SMART DYNAMIC PARAGRAPH: Triệt tiêu hoàn toàn các đoạn <w:p> chỉ chứa thông tin vợ/chồng
                            # (Bao gồm BÊN THẾ CHẤP, BÊN VAY/BÊN C, Thu nhập Mẫu 3, Đăng ký thế chấp Mẫu 01a, Mẫu 05a...)
                            def _is_spouse_only_p(p_xml):
                                if not re.search(r'vợchồng|vợ/chồng|vochong', p_xml, re.IGNORECASE):
                                    return False
                                # ĐẶC QUYỀN MẪU 3 (BÁO CÁO THẨM ĐỊNH): Nếu đoạn chứa thu nhập, nghề nghiệp, nơi làm việc hoặc quan hệ vợ/chồng
                                # thì đây chính là đoạn mô tả vợ/chồng -> Triệt tiêu 100% khi khách hàng độc thân (kể cả khi đoạn có nhắc tên người vay)
                                if re.search(r'Thu_nhập_vợchồng|Thu_nhập_vợ_chồng|quan_hệ_với_người_vay|nơi_làm_việc_vợchồng|nghề_nghiệp_vợchồng', p_xml, re.IGNORECASE):
                                    return True

                                kh_tokens = [
                                    'Họ_và_tên_KH', 'Họ và tên KH', 'Số_điện_thoại_KH', 'Số điện thoại KH',
                                    'nơi_công_tác', 'mã_BHXH', 'mã_số_thuế', '«CCCD»', 'MERGEFIELD  CCCD ', 'MERGEFIELD CCCD ',
                                    '«Địa_chỉ_cư_trú»', 'MERGEFIELD  Địa_chỉ_cư_trú ', 'MERGEFIELD Địa_chỉ_cư_trú ',
                                    '«nơi_ở_hiện_tại»', 'MERGEFIELD  nơi_ở_hiện_tại ', 'MERGEFIELD nơi_ở_hiện_tại ',
                                    '«DX_HOA»', 'MERGEFIELD  DX_HOA ', 'MERGEFIELD DX_HOA ',
                                    '«dx_thường»', 'MERGEFIELD  dx_thường ', 'MERGEFIELD dx_thường '
                                ]
                                for kt in kh_tokens:
                                    if kt.lower() in p_xml.lower():
                                        return False
                                return True

                            def _filter_p_spouse(m):
                                p_xml = m.group(0)
                                if _is_spouse_only_p(p_xml):
                                    return ''
                                return p_xml

                            xml_str = re.sub(r'<w:p\b[^>]*>[\s\S]*?</w:p>', _filter_p_spouse, xml_str)

                        # 3. Định dạng lại đoạn Chi phí người phụ thuộc cho hồ sơ Độc thân:
                        # CHỈ ép thành "- Chi phí người phụ thuộc: 0 đồng." KHI VÀ CHỈ KHI:
                        # Khách hàng không có chi phí người phụ thuộc (<= 0) VÀ cũng không nhập thông tin người phụ thuộc nào.
                        # Nếu khách hàng có con nhỏ/người phụ thuộc (dù chi phí = 0 hay > 0), BẢO LƯU NGUYÊN VẸN để xuất thông tin người phụ thuộc!
                        raw_cppt_str = str(norm_data.get("chi_phí_người_phụ_thuộc") or norm_data.get("chiPhiNguoiPhuThuoc") or norm_data.get("Chi_phí_người_phụ_thuộc") or norm_data.get("Chi phí người phụ thuộc") or data_dict.get("chiPhiNguoiPhuThuoc") or "0").replace(".", "").replace(",", "").strip()
                        has_cppt = bool(raw_cppt_str and raw_cppt_str != "0" and raw_cppt_str.isdigit() and int(raw_cppt_str) > 0)
                        raw_info_pt = str(norm_data.get("thông_tin_phụ_thuộc") or norm_data.get("người_phụ_thuộc") or norm_data.get("thong_tin_phu_thuoc") or data_dict.get("thongTinPhuThuoc") or "").strip()
                        has_info_pt = bool(raw_info_pt and raw_info_pt != "—" and raw_info_pt != "0")

                        if not has_cppt and not has_info_pt and ('chi_phí_người_phụ_thuộc' in xml_str.lower() or 'chi phí người phụ thuộc' in xml_str.lower() or 'chiphinguoiphuthuoc' in xml_str.lower()):
                            pattern_p_cppt = re.compile(
                                r'<w:p\b[^>]*>(?:(?!<w:p\b)[\s\S])*?Chi_phí_người_phụ_thuộc(?:(?!<w:p\b)[\s\S])*?</w:p>',
                                re.IGNORECASE
                            )
                            def _rep_cppt_single(m):
                                p_xml = m.group(0)
                                m_ppr = re.search(r'<w:pPr>([\s\S]*?)</w:pPr>', p_xml)
                                ppr_inner = m_ppr.group(1) if m_ppr else ''
                                ppr_inner = re.sub(r'<w:color\s+w:val="[^"]*"\s*/>', '', ppr_inner)
                                ppr_inner = re.sub(r'<w:noProof\s*/>', '', ppr_inner)
                                base_ppr = f'<w:pPr>{ppr_inner}</w:pPr>' if ppr_inner else ''
                                
                                m_rpr = re.search(r'<w:rPr>([\s\S]*?)</w:rPr>', p_xml)
                                base_rpr = f'<w:rPr>{m_rpr.group(1)}</w:rPr>' if m_rpr else ''
                                base_rpr = re.sub(r'<w:color\s+w:val="[^"]*"\s*/>', '', base_rpr)
                                base_rpr = re.sub(r'<w:noProof\s*/>', '', base_rpr)
                                base_rpr = re.sub(r'<w:lang\s+[^>]*/>', '', base_rpr)
                                if base_rpr == '<w:rPr></w:rPr>':
                                    base_rpr = ''
                                return f'<w:p>{base_ppr}<w:r>{base_rpr}<w:t>- Chi phí người phụ thuộc: 0 đồng.</w:t></w:r></w:p>'
                            
                            xml_str = pattern_p_cppt.sub(_rep_cppt_single, xml_str)

                    # SMART GENDER CHECKBOX: Tự động đánh dấu ☑ cho Nam / Nữ theo Kiểu 1 (Mẫu 1)
                    raw_gt = str(norm_data.get('giới_tính') or norm_data.get('gioi_tinh') or 
                                 norm_data.get('Giới tính') or norm_data.get('gioiTinh') or '').strip().lower()
                    dx_norm = str(norm_data.get('dx_thường') or norm_data.get('dx_hoa') or '').strip().lower()
                    is_cust_male = (raw_gt == 'nam' or dx_norm in ['ông', 'ong'])
                    chk_nam_v = "☑" if is_cust_male else "☐"
                    chk_nu_v = "☐" if is_cust_male else "☑"

                    if 'giới tính' in xml_str.lower() and '☐' in xml_str:
                        pattern_p_gender = re.compile(r'<w:p\b[^>]*>(?:(?!<w:p\b)[\s\S])*?Giới tính:(?:(?!<w:p\b)[\s\S])*?</w:p>', re.IGNORECASE)
                        def _rep_gender_p(m_p):
                            p_txt = m_p.group(0)
                            p_sub = re.compile(
                                r'(<w:r\b[^>]*>(?:(?!<w:r\b)[\s\S])*?<w:t[^>]*>\s*\.?\s*Giới tính:\s*</w:t>[\s\S]*?</w:r>)'
                                r'[\s\S]*?'
                                r'(Nam[\s\S]*?☐[\s\S]*?Nữ[\s\S]*?☐[\s\S]*?</w:r>)',
                                re.IGNORECASE
                            )
                            return p_sub.sub(
                                f'<w:r><w:rPr><w:spacing w:val="-4"/><w:lang w:val="pt-BR"/></w:rPr>'
                                f'<w:t xml:space="preserve">. Giới tính: Nam  {chk_nam_v}\u00a0\u00a0\u00a0\u00a0\u00a0Nữ  {chk_nu_v}</w:t></w:r>',
                                p_txt
                            )
                        xml_str = pattern_p_gender.sub(_rep_gender_p, xml_str)

                        def _rep_plain_box(m_b):
                            return f"{m_b.group(1)}{chk_nam_v}{m_b.group(2)}{chk_nu_v}"
                        xml_str = re.sub(r'(Nam[^\w<]{1,10})☐([^\w<]{1,20}Nữ[^\w<]{1,10})☐', _rep_plain_box, xml_str)

                    # Track các field name đã được replace qua MERGEFIELD (bước 1 & 2)
                    _merged_keys = set()

                    def format_runs_with_italics(text, base_rpr=""):
                        clean_base_rpr = base_rpr.strip() if base_rpr else ""
                        clean_base_rpr = re.sub(r'<w:color\s+w:val="[^"]*"\s*/>', '', clean_base_rpr)
                        clean_base_rpr = re.sub(r'<w:noProof\s*/>', '', clean_base_rpr)
                        clean_base_rpr = re.sub(r'<w:lang\s+[^>]*/>', '', clean_base_rpr)
                        if clean_base_rpr == '<w:rPr></w:rPr>':
                            clean_base_rpr = ''

                        if clean_base_rpr:
                            if '<w:i/>' not in clean_base_rpr and '<w:i ' not in clean_base_rpr:
                                italic_rpr = re.sub(r'</w:rPr>', r'<w:i/><w:iCs/></w:rPr>', clean_base_rpr)
                            else:
                                italic_rpr = clean_base_rpr
                        else:
                            italic_rpr = '<w:rPr><w:i/><w:iCs/></w:rPr>'

                        pattern = re.compile(r'(\((?:Bằng chữ|bằng chữ):[^\)]+\))', re.IGNORECASE)
                        parts = pattern.split(str(text))
                        
                        out_runs = []
                        for part in parts:
                            if not part:
                                continue
                            is_italic = bool(pattern.fullmatch(part))
                            cur_rpr = italic_rpr if is_italic else clean_base_rpr
                            
                            escaped = part.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                            escaped_lines = escaped.replace('\r\n', '\n').replace('\r', '\n').split('\n')
                            
                            t_chunks = [f'<w:t xml:space="preserve">{line}</w:t>' for line in escaped_lines]
                            inner_t = '<w:br/>'.join(t_chunks)
                            out_runs.append(f'<w:r>{cur_rpr}{inner_t}</w:r>')
                            
                        return ''.join(out_runs)

                    # 0. Thay thế cấp đoạn <w:p> cho «tổng_hợp_phiếu_báo_có» / MERGEFIELD tổng_hợp_phiếu_báo_có
                    # Khi văn bản có nhiều dòng (4 đợt giải ngân), việc ngắt dòng bằng <w:br/> bên trong một <w:p>
                    # có thuộc tính căn đều (<w:jc w:val="both"/>) sẽ khiến Word kéo giãn khoảng cách chữ (giãn dòng).
                    # Khắc phục triệt để bằng cách biến mỗi đợt thành một thẻ <w:p> độc lập, giữ nguyên thụt đầu dòng
                    # và tự động in nghiêng các phần (Bằng chữ: ...).
                    th_bc_val = (norm_data.get(normalize_key("tổng_hợp_phiếu_báo_có")) or 
                                 norm_data.get("tổng_hợp_phiếu_báo_có") or norm_data.get("tong_hop_phieu_bao_co") or 
                                 norm_data.get("tonghopphieubaoco") or norm_data.get("tongHopPhieuBaoCo") or 
                                 norm_data.get("tổng hợp phiếu báo có"))
                    if th_bc_val and any(k in xml_str.lower() for k in ['báo_có', 'bao_co', 'baoco', 'báo có']):
                        pattern_p_thbc = re.compile(
                            r'<w:p\b[^>]*>(?:(?!<w:p\b)[\s\S])*?(?:MERGEFIELD\s+(?:"tổng_hợp_phiếu_báo_có"|tổng_hợp_phiếu_báo_có|tong_hop_phieu_bao_co|tonghopphieubaoco)|[«\xab]\s*(?:tổng_hợp_phiếu_báo_có|tong_hop_phieu_bao_co|tổng hợp phiếu báo có|tonghopphieubaoco)\s*[»\xbb]|(?:&lt;&lt;|<<)\s*(?:tổng_hợp_phiếu_báo_có|tong_hop_phieu_bao_co|tổng hợp phiếu báo có|tonghopphieubaoco)\s*(?:&gt;&gt;|>>))(?:(?!<w:p\b)[\s\S])*?</w:p>',
                            re.IGNORECASE
                        )
                        def _rep_p_thbc(match_p):
                            p_xml = match_p.group(0)
                            m_ppr = re.search(r'<w:pPr>([\s\S]*?)</w:pPr>', p_xml)
                            ppr_inner = m_ppr.group(1) if m_ppr else ''
                            ppr_inner = re.sub(r'<w:color\s+w:val="[^"]*"\s*/>', '', ppr_inner)
                            ppr_inner = re.sub(r'<w:noProof\s*/>', '', ppr_inner)
                            base_ppr = f'<w:pPr>{ppr_inner}</w:pPr>' if ppr_inner else ''
                            
                            m_rpr = re.search(r'<w:rPr>([\s\S]*?)</w:rPr>', p_xml)
                            base_rpr = f'<w:rPr>{m_rpr.group(1)}</w:rPr>' if m_rpr else ''
                            base_rpr = re.sub(r'<w:color\s+w:val="[^"]*"\s*/>', '', base_rpr)
                            base_rpr = re.sub(r'<w:noProof\s*/>', '', base_rpr)
                            base_rpr = re.sub(r'<w:lang\s+[^>]*/>', '', base_rpr)
                            if base_rpr == '<w:rPr></w:rPr>':
                                base_rpr = ''
                            
                            lines = [l.strip() for l in th_bc_val.replace('\r\n', '\n').replace('\r', '\n').split('\n') if l.strip()]
                            out_ps = [f'<w:p>{base_ppr}{format_runs_with_italics(line, base_rpr)}</w:p>' for line in lines]
                            return ''.join(out_ps)
                        
                        xml_str = pattern_p_thbc.sub(_rep_p_thbc, xml_str)
                        for k_bc in ["tổng_hợp_phiếu_báo_có", "tổng hợp phiếu báo có", "tong_hop_phieu_bao_co", "tonghopphieubaoco", "tongHopPhieuBaoCo"]:
                            _merged_keys.add(k_bc)
                            _merged_keys.add(normalize_key(k_bc))

                    # 0b. Thay thế cấp đoạn <w:p> cho «tổng_hợp_vốn_tự_có» / MERGEFIELD tổng_hợp_vốn_tự_có
                    # Khi người dùng nối chuỗi nhiều phiếu vốn tự có (1), (2)..., việc dùng <w:br/> trong một đoạn
                    # căn đều sẽ bị giãn dòng. Ta tách mỗi dòng thành một thẻ <w:p> độc lập chuẩn Word.
                    th_vtc_val = (norm_data.get(normalize_key("tổng_hợp_vốn_tự_có")) or 
                                  norm_data.get("tổng_hợp_vốn_tự_có") or norm_data.get("tong_hop_von_tu_co") or 
                                  norm_data.get("tonghopvontuco") or norm_data.get("tongHopVonTuCo") or 
                                  norm_data.get("tổng hợp vốn tự có") or norm_data.get("Tổng_hợp_vốn_tự_có"))
                    if th_vtc_val and any(k in xml_str.lower() for k in ['vốn_tự_có', 'von_tu_co', 'vontuco', 'vốn tự có']):
                        pattern_p_thvtc = re.compile(
                            r'<w:p\b[^>]*>(?:(?!<w:p\b)[\s\S])*?(?:MERGEFIELD\s+(?:"tổng_hợp_vốn_tự_có"|tổng_hợp_vốn_tự_có|tong_hop_von_tu_co|tonghopvontuco)|[«\xab]\s*(?:tổng_hợp_vốn_tự_có|tong_hop_von_tu_co|tổng hợp vốn tự có|tonghopvontuco)\s*[»\xbb]|(?:&lt;&lt;|<<)\s*(?:tổng_hợp_vốn_tự_có|tong_hop_von_tu_co|tổng hợp vốn tự có|tonghopvontuco)\s*(?:&gt;&gt;|>>))(?:(?!<w:p\b)[\s\S])*?</w:p>',
                            re.IGNORECASE
                        )
                        def _rep_p_thvtc(match_p):
                            p_xml = match_p.group(0)
                            m_ppr = re.search(r'<w:pPr>([\s\S]*?)</w:pPr>', p_xml)
                            ppr_inner = m_ppr.group(1) if m_ppr else ''
                            ppr_inner = re.sub(r'<w:color\s+w:val="[^"]*"\s*/>', '', ppr_inner)
                            ppr_inner = re.sub(r'<w:noProof\s*/>', '', ppr_inner)
                            base_ppr = f'<w:pPr>{ppr_inner}</w:pPr>' if ppr_inner else ''
                            
                            m_rpr = re.search(r'<w:rPr>([\s\S]*?)</w:rPr>', p_xml)
                            base_rpr = f'<w:rPr>{m_rpr.group(1)}</w:rPr>' if m_rpr else ''
                            base_rpr = re.sub(r'<w:color\s+w:val="[^"]*"\s*/>', '', base_rpr)
                            base_rpr = re.sub(r'<w:noProof\s*/>', '', base_rpr)
                            base_rpr = re.sub(r'<w:lang\s+[^>]*/>', '', base_rpr)
                            if base_rpr == '<w:rPr></w:rPr>':
                                base_rpr = ''
                            
                            lines = [l.strip() for l in th_vtc_val.replace('\r\n', '\n').replace('\r', '\n').split('\n') if l.strip()]
                            out_ps = [f'<w:p>{base_ppr}{format_runs_with_italics(line, base_rpr)}</w:p>' for line in lines]
                            return ''.join(out_ps)
                        
                        xml_str = pattern_p_thvtc.sub(_rep_p_thvtc, xml_str)
                        for k_vtc in ["tổng_hợp_vốn_tự_có", "tổng hợp vốn tự có", "tong_hop_von_tu_co", "tonghopvontuco", "tongHopVonTuCo"]:
                            _merged_keys.add(k_vtc)
                            _merged_keys.add(normalize_key(k_vtc))

                    # 0d. Thay thế cấp đoạn <w:p> cho «vốn_tự_có_n» / MERGEFIELD vốn_tự_có_n (Bỏ 2% KPBT)
                    th_vtc_n_val = (norm_data.get(normalize_key("vốn_tự_có_n")) or 
                                    norm_data.get("vốn_tự_có_n") or norm_data.get("von_tu_co_n") or 
                                    norm_data.get("vontucon") or norm_data.get("vốn tự có n") or 
                                    norm_data.get("Vốn_tự_có_n") or norm_data.get("vonTuCoN"))
                    if th_vtc_n_val and any(k in xml_str.lower() for k in ['vốn_tự_có_n', 'von_tu_co_n', 'vontucon', 'vốn tự có n']):
                        pattern_p_thvtc_n = re.compile(
                            r'<w:p\b[^>]*>(?:(?!<w:p\b)[\s\S])*?(?:MERGEFIELD\s+(?:"vốn_tự_có_n"|vốn_tự_có_n|von_tu_co_n|vontucon)|[«\xab]\s*(?:vốn_tự_có_n|von_tu_co_n|vốn tự có n|vontucon)\s*[»\xbb]|(?:&lt;&lt;|<<)\s*(?:vốn_tự_có_n|von_tu_co_n|vốn tự có n|vontucon)\s*(?:&gt;&gt;|>>))(?:(?!<w:p\b)[\s\S])*?<\/w:p>',
                            re.IGNORECASE
                        )
                        def _rep_p_thvtc_n(match_p):
                            p_xml = match_p.group(0)
                            m_ppr = re.search(r'<w:pPr>([\s\S]*?)</w:pPr>', p_xml)
                            ppr_inner = m_ppr.group(1) if m_ppr else ''
                            ppr_inner = re.sub(r'<w:color\s+w:val="[^"]*"\s*/>', '', ppr_inner)
                            ppr_inner = re.sub(r'<w:noProof\s*/>', '', ppr_inner)
                            base_ppr = f'<w:pPr>{ppr_inner}</w:pPr>' if ppr_inner else ''
                            
                            m_rpr = re.search(r'<w:rPr>([\s\S]*?)</w:rPr>', p_xml)
                            base_rpr = f'<w:rPr>{m_rpr.group(1)}</w:rPr>' if m_rpr else ''
                            base_rpr = re.sub(r'<w:color\s+w:val="[^"]*"\s*/>', '', base_rpr)
                            base_rpr = re.sub(r'<w:noProof\s*/>', '', base_rpr)
                            base_rpr = re.sub(r'<w:lang\s+[^>]*/>', '', base_rpr)
                            if base_rpr == '<w:rPr></w:rPr>':
                                base_rpr = ''
                            
                            lines = [l.strip() for l in th_vtc_n_val.replace('\r\n', '\n').replace('\r', '\n').split('\n') if l.strip()]
                            out_ps = [f'<w:p>{base_ppr}{format_runs_with_italics(line, base_rpr)}</w:p>' for line in lines]
                            return ''.join(out_ps)
                        
                        xml_str = pattern_p_thvtc_n.sub(_rep_p_thvtc_n, xml_str)
                        for k_vtcn in ["vốn_tự_có_n", "vốn tự có n", "von_tu_co_n", "vontucon", "vonTuCoN"]:
                            _merged_keys.add(k_vtcn)
                            _merged_keys.add(normalize_key(k_vtcn))

                    # 0c. Thay thế cấp đoạn <w:p> cho «kế_hoạch_giải_ngân» / MERGEFIELD kế_hoạch_giải_ngân
                    # Mỗi lần giải ngân (+ Lần 1: ..., + Lần 2: ...) được tách thành một thẻ <w:p> độc lập
                    # thay thế ký tự ngắt dòng \n, \r\n hoặc <br> bằng ngắt đoạn <w:p>, giữ nguyên thụt lề,
                    # căn dòng và tự động in nghiêng phần (Bằng chữ: ...) chống giãn chữ khi căn đều (justify),
                    # định dạng chuẩn mực y hệt như «tổng_hợp_phiếu_báo_có».
                    kh_gn_val = (norm_data.get(normalize_key("kế_hoạch_giải_ngân")) or 
                                 norm_data.get(normalize_key("Kế hoạch giải ngân")) or
                                 norm_data.get(normalize_key("ke_hoach_giai_ngan")) or
                                 norm_data.get("kế_hoạch_giải_ngân") or 
                                 norm_data.get("Kế_hoạch_giải_ngân") or
                                 norm_data.get("kế hoạch giải ngân") or 
                                 norm_data.get("Kế hoạch giải ngân") or
                                 norm_data.get("ke_hoach_giai_ngan") or 
                                 norm_data.get("kehoachgiaingan") or 
                                 norm_data.get("ke_hoạch_giải_ngân") or
                                 data_dict.get("kế_hoạch_giải_ngân") or
                                 data_dict.get("Kế hoạch giải ngân"))
                    if kh_gn_val and any(k in xml_str.lower() for k in ['giải_ngân', 'giai_ngan', 'giaingan', 'giải ngân']):
                        pattern_p_khgn = re.compile(
                            r'<w:p\b[^>]*>(?:(?!<w:p\b)[\s\S])*?(?:'
                            r'MERGEFIELD\s+(?:"\s*(?:kế[\s_]+hoạch[\s_]+giải[\s_]+ngân|ke[\s_]+hoach[\s_]+giai[\s_]+ngan)\s*"|(?:kế[\s_]+hoạch[\s_]+giải[\s_]+ngân|ke[\s_]+hoach[\s_]+giai[\s_]+ngan))|'
                            r'[«\xab]\s*(?:kế[\s_]+hoạch[\s_]+giải[\s_]+ngân|ke[\s_]+hoach[\s_]+giai[\s_]+ngan)\s*[»\xbb]|'
                            r'(?:&lt;&lt;|<<)\s*(?:kế[\s_]+hoạch[\s_]+giải[\s_]+ngân|ke[\s_]+hoach[\s_]+giai[\s_]+ngan)\s*(?:&gt;&gt;|>>)'
                            r')(?:(?!<w:p\b)[\s\S])*?</w:p>',
                            re.IGNORECASE
                        )
                        def _rep_p_khgn(match_p):
                            p_xml = match_p.group(0)
                            m_ppr = re.search(r'<w:pPr>([\s\S]*?)</w:pPr>', p_xml)
                            ppr_inner = m_ppr.group(1) if m_ppr else ''
                            ppr_inner = re.sub(r'<w:color\s+w:val="[^"]*"\s*/>', '', ppr_inner)
                            ppr_inner = re.sub(r'<w:noProof\s*/>', '', ppr_inner)
                            base_ppr = f'<w:pPr>{ppr_inner}</w:pPr>' if ppr_inner else ''
                            
                            m_rpr = re.search(r'<w:rPr>([\s\S]*?)</w:rPr>', p_xml)
                            base_rpr = f'<w:rPr>{m_rpr.group(1)}</w:rPr>' if m_rpr else ''
                            base_rpr = re.sub(r'<w:color\s+w:val="[^"]*"\s*/>', '', base_rpr)
                            base_rpr = re.sub(r'<w:noProof\s*/>', '', base_rpr)
                            base_rpr = re.sub(r'<w:lang\s+[^>]*/>', '', base_rpr)
                            if base_rpr == '<w:rPr></w:rPr>':
                                base_rpr = ''
                            
                            inner_p = re.sub(r'^<w:p\b[^>]*>', '', p_xml)
                            inner_p = re.sub(r'</w:p>$', '', inner_p)
                            inner_p = re.sub(r'^<w:pPr>[\s\S]*?</w:pPr>', '', inner_p)

                            pattern_field_khgn = re.compile(
                                r'(?:<w:fldSimple\b[^>]*?(?:kế[\s_]+hoạch[\s_]+giải[\s_]+ngân|ke[\s_]+hoach[\s_]+giai[\s_]+ngan)[\s\S]*?</w:fldSimple>|'
                                r'<w:r\b[^>]*>(?:(?!<w:r\b)[\s\S])*?<w:fldChar\s+w:fldCharType="begin"[\s\S]*?(?:kế[\s_]+hoạch[\s_]+giải[\s_]+ngân|ke[\s_]+hoach[\s_]+giai[\s_]+ngan)[\s\S]*?<w:fldChar\s+w:fldCharType="end"\s*/>[\s\S]*?</w:r>|'
                                r'<w:r\b[^>]*>(?:(?!<w:r\b)[\s\S])*?<w:fldChar\s+w:fldCharType="begin"[\s\S]*?</w:r>[\s\S]*?(?:kế[\s_]+hoạch[\s_]+giải[\s_]+ngân|ke[\s_]+hoach[\s_]+giai[\s_]+ngan)[\s\S]*?<w:r\b[^>]*>(?:(?!<w:r\b)[\s\S])*?<w:fldChar\s+w:fldCharType="end"\s*/>[\s\S]*?</w:r>|'
                                r'<w:r\b[^>]*>(?:(?!<w:r\b)[\s\S])*?<w:t[^>]*>(?:«|&lt;&lt;|<<)\s*(?:kế[\s_]+hoạch[\s_]+giải[\s_]+ngân|ke[\s_]+hoach[\s_]+giai[\s_]+ngan)\s*(?:»|&gt;&gt;|>>)</w:t>[\s\S]*?</w:r>)',
                                re.IGNORECASE
                            )

                            m_field = pattern_field_khgn.search(inner_p)
                            if not m_field:
                                prefix_runs = ""
                                suffix_runs = ""
                            else:
                                prefix_runs = inner_p[:m_field.start()]
                                suffix_runs = inner_p[m_field.end():]

                            raw_s = str(kh_gn_val)
                            # Chuẩn hóa các định dạng ngắt dòng \n, <br>, <br/>, <p> thành \n để ngắt đoạn <w:p>
                            raw_s = re.sub(r'<br\s*/?>', '\n', raw_s, flags=re.IGNORECASE)
                            raw_s = re.sub(r'</?p\b[^>]*>', '\n', raw_s, flags=re.IGNORECASE)
                            lines = [l.strip() for l in raw_s.replace('\r\n', '\n').replace('\r', '\n').split('\n') if l.strip()]

                            prefix_text = re.sub(r'<[^>]+>', '', prefix_runs).strip()
                            suffix_text = re.sub(r'<[^>]+>', '', suffix_runs).strip()

                            out_ps = []
                            if prefix_text:
                                out_ps.append(f'<w:p>{base_ppr}{prefix_runs}</w:p>')

                            for idx, line in enumerate(lines):
                                line_run = format_runs_with_italics(line, base_rpr)
                                if idx == len(lines) - 1 and suffix_text:
                                    out_ps.append(f'<w:p>{base_ppr}{line_run}{suffix_runs}</w:p>')
                                else:
                                    out_ps.append(f'<w:p>{base_ppr}{line_run}</w:p>')

                            return ''.join(out_ps)
                        
                        xml_str = pattern_p_khgn.sub(_rep_p_khgn, xml_str)
                        for k_kh in ["kế_hoạch_giải_ngân", "kế hoạch giải ngân", "Kế_hoạch_giải_ngân", "Kế hoạch giải ngân", "ke_hoach_giai_ngan", "kehoachgiaingan"]:
                            _merged_keys.add(k_kh)
                            _merged_keys.add(normalize_key(k_kh))

                    # 0e. Thay thế cấp đoạn <w:p> cho «Đối_tượng» / MERGEFIELD Đối_tượng
                    # Khi đối tượng vay vốn có nhiều dòng (ngắt dòng \n hoặc <br>),
                    # hệ thống tự động bóc tách từng dòng thành các thẻ đoạn <w:p> độc lập,
                    # bảo lưu chính xác phần tiền tố (Thuộc đối tượng: / - Thuộc đối tượng: )
                    # và hậu tố (dấu chấm kết thúc), khắc phục triệt để lỗi Word kéo giãn khoảng cách chữ khi căn đều (justify).
                    dt_raw = (norm_data.get("Đối_tượng") or norm_data.get("Đối tượng") or 
                              norm_data.get("đối_tượng") or norm_data.get("đối tượng") or 
                              norm_data.get("doi_tuong") or norm_data.get("doituong") or "")
                    if dt_raw and any(k in xml_str.lower() for k in ['đối_tượng', 'doi_tuong', 'doituong', 'đối tượng']):
                        clean_dt = str(dt_raw).replace('\r\n', '\n').replace('\r', '\n')
                        clean_dt = re.sub(r'<br\s*/?>', '\n', clean_dt, flags=re.IGNORECASE)
                        clean_dt = re.sub(r'</p>', '\n', clean_dt, flags=re.IGNORECASE)
                        clean_dt = re.sub(r'<p\b[^>]*>', '', clean_dt, flags=re.IGNORECASE)
                        dt_lines = [l.strip() for l in clean_dt.split('\n') if l.strip()]
                        
                        pattern_p_dt = re.compile(
                            r'<w:p\b[^>]*>(?:(?!<w:p\b)[\s\S])*?(?:MERGEFIELD\s+(?:"Đối_tượng"|Đối_tượng|"Đối tượng"|Đối tượng|doi_tuong|doituong|"đối_tượng"|đối_tượng)|[«\xab]\s*(?:Đối_tượng|Đối tượng|đối_tượng|đối tượng|doi_tuong|doituong)\s*[»\xbb]|(?:&lt;&lt;|<<)\s*(?:Đối_tượng|Đối tượng|đối_tượng|đối tượng|doi_tuong|doituong)\s*(?:&gt;&gt;|>>))(?:(?!<w:p\b)[\s\S])*?</w:p>',
                            re.IGNORECASE
                        )
                        
                        def _rep_p_dt(match_p):
                            p_xml = match_p.group(0)
                            m_ppr = re.search(r'<w:pPr>([\s\S]*?)</w:pPr>', p_xml)
                            ppr_inner = m_ppr.group(1) if m_ppr else ''
                            ppr_inner = re.sub(r'<w:color\s+w:val="[^"]*"\s*/>', '', ppr_inner)
                            ppr_inner = re.sub(r'<w:noProof\s*/>', '', ppr_inner)
                            base_ppr = f'<w:pPr>{ppr_inner}</w:pPr>' if ppr_inner else ''
                            
                            m_rpr = re.search(r'<w:rPr>([\s\S]*?)</w:rPr>', p_xml)
                            base_rpr = f'<w:rPr>{m_rpr.group(1)}</w:rPr>' if m_rpr else ''
                            base_rpr = re.sub(r'<w:color\s+w:val="[^"]*"\s*/>', '', base_rpr)
                            base_rpr = re.sub(r'<w:noProof\s*/>', '', base_rpr)
                            base_rpr = re.sub(r'<w:lang\s+[^>]*/>', '', base_rpr)
                            if base_rpr == '<w:rPr></w:rPr>':
                                base_rpr = ''
                                
                            pattern_field_dt = re.compile(
                                r'(?:<w:fldSimple\b[^>]*?(?:Đối_tượng|Đối tượng|doi_tuong|doituong|đối_tượng|đối tượng)[\s\S]*?</w:fldSimple>|'
                                r'<w:r\b[^>]*>(?:(?!<w:r\b)[\s\S])*?<w:fldChar\s+w:fldCharType="begin"[\s\S]*?(?:Đối_tượng|Đối tượng|doi_tuong|doituong|đối_tượng|đối tượng)[\s\S]*?<w:fldChar\s+w:fldCharType="end"/>[\s\S]*?</w:r>|'
                                r'<w:r\b[^>]*>(?:(?!<w:r\b)[\s\S])*?<w:fldChar\s+w:fldCharType="begin"[\s\S]*?</w:r>[\s\S]*?(?:Đối_tượng|Đối tượng|doi_tuong|doituong|đối_tượng|đối tượng)[\s\S]*?<w:r\b[^>]*>(?:(?!<w:r\b)[\s\S])*?<w:fldChar\s+w:fldCharType="end"/>[\s\S]*?</w:r>|'
                                r'<w:r\b[^>]*>(?:(?!<w:r\b)[\s\S])*?<w:t\b[^>]*>(?:«|&lt;&lt;|<<)\s*(?:Đối_tượng|Đối tượng|doi_tuong|doituong|đối_tượng|đối tượng)\s*(?:»|&gt;&gt;|>>)</w:t>[\s\S]*?</w:r>)',
                                re.IGNORECASE
                            )
                            
                            inner_p = re.sub(r'^<w:p\b[^>]*>', '', p_xml)
                            inner_p = re.sub(r'</w:p>$', '', inner_p)
                            inner_p = re.sub(r'^<w:pPr>[\s\S]*?</w:pPr>', '', inner_p)
                            
                            m_field = pattern_field_dt.search(inner_p)
                            if not m_field:
                                prefix_runs = ""
                                suffix_runs = ""
                            else:
                                prefix_runs = inner_p[:m_field.start()]
                                suffix_runs = inner_p[m_field.end():]
                                
                            def make_dt_run(text):
                                escaped = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                                return f'<w:r>{base_rpr}<w:t xml:space="preserve">{escaped}</w:t></w:r>'
                                
                            if not dt_lines:
                                return f'<w:p>{base_ppr}{prefix_runs}{suffix_runs}</w:p>'
                            if len(dt_lines) == 1:
                                return f'<w:p>{base_ppr}{prefix_runs}{make_dt_run(dt_lines[0])}{suffix_runs}</w:p>'
                                
                            out_ps = []
                            for idx, line in enumerate(dt_lines):
                                if idx == 0:
                                    out_ps.append(f'<w:p>{base_ppr}{prefix_runs}{make_dt_run(line)}</w:p>')
                                elif idx == len(dt_lines) - 1:
                                    out_ps.append(f'<w:p>{base_ppr}{make_dt_run(line)}{suffix_runs}</w:p>')
                                else:
                                    out_ps.append(f'<w:p>{base_ppr}{make_dt_run(line)}</w:p>')
                            return ''.join(out_ps)
                            
                        xml_str = pattern_p_dt.sub(_rep_p_dt, xml_str)
                        for tag_k in ["đối_tượng", "đối tượng", "Đối_tượng", "Đối tượng", "doi_tuong", "doituong"]:
                            _merged_keys.add(tag_k)
                            _merged_keys.add(normalize_key(tag_k))

                    # 1. w:fldSimple
                    def replace_fld_simple(m):
                        fld_name = m.group(1) or m.group(2)
                        val = get_val_smart(fld_name)
                        if val is None:
                            return m.group(0)
                        _merged_keys.add(normalize_key(fld_name))
                        _merged_keys.add(fld_name)
                        full_match = m.group(0)
                        m_rpr = re.search(r'<w:rPr>([\s\S]*?)</w:rPr>', full_match)
                        rpr_str = f'<w:rPr>{m_rpr.group(1)}</w:rPr>' if m_rpr else ''
                        rpr_str = re.sub(r'<w:noProof\s*/>', '', rpr_str)
                        if normalize_key(fld_name) in [normalize_key("kế_hoạch_giải_ngân"), normalize_key("Kế hoạch giải ngân"), normalize_key("ke_hoach_giai_ngan"), normalize_key("tổng_hợp_phiếu_báo_có"), normalize_key("tổng_hợp_vốn_tự_có"), normalize_key("vốn_tự_có_n")] or (isinstance(val, str) and '(bằng chữ:' in val.lower()):
                            return format_runs_with_italics(val, rpr_str)
                        escaped_val = escape_xml(val)
                        return f'<w:r>{rpr_str}<w:t xml:space="preserve">{escaped_val}</w:t></w:r>'

                    pattern_simple = re.compile(r'<w:fldSimple\s+w:instr="[^"]*MERGEFIELD\s+(?:"([^"]+)"|([^\s"\\<>]+))[^"]*">([\s\S]*?)</w:fldSimple>')
                    xml_str = pattern_simple.sub(replace_fld_simple, xml_str)

                    # 2. Complex fields: Chuyển đổi toàn bộ khối MERGEFIELD (cả đơn run lẫn đa run ECMA-376) thành run text chuẩn
                    def replace_complex_field(m):
                        fld_name = m.group(1) or m.group(2)
                        val = get_val_smart(fld_name)
                        if val is None:
                            return m.group(0)
                        _merged_keys.add(normalize_key(fld_name))
                        _merged_keys.add(fld_name)
                        full_match = m.group(0)
                        m_rpr = None
                        runs_in_match = re.findall(r'<w:r\b[^>]*>[\s\S]*?</w:r>', full_match)
                        for r_item in runs_in_match:
                            if '<w:t' in r_item and '<w:rPr>' in r_item:
                                m_rpr = re.search(r'<w:rPr>([\s\S]*?)</w:rPr>', r_item)
                                break
                        if not m_rpr:
                            for r_item in runs_in_match:
                                if '<w:rPr>' in r_item:
                                    m_rpr = re.search(r'<w:rPr>([\s\S]*?)</w:rPr>', r_item)
                                    break
                        rpr_str = f'<w:rPr>{m_rpr.group(1)}</w:rPr>' if m_rpr else ''
                        rpr_str = re.sub(r'<w:noProof\s*/>', '', rpr_str)
                        if normalize_key(fld_name) in [normalize_key("kế_hoạch_giải_ngân"), normalize_key("Kế hoạch giải ngân"), normalize_key("ke_hoach_giai_ngan"), normalize_key("tổng_hợp_phiếu_báo_có"), normalize_key("tổng_hợp_vốn_tự_có"), normalize_key("vốn_tự_có_n")] or (isinstance(val, str) and '(bằng chữ:' in val.lower()):
                            return format_runs_with_italics(val, rpr_str)
                        escaped_val = escape_xml(val)
                        return f'<w:r>{rpr_str}<w:t xml:space="preserve">{escaped_val}</w:t></w:r>'

                    # Multi-run complex field chuẩn ECMA-376 của Word
                    pattern_multi_complex = re.compile(
                        r'<w:r\b[^>]*>(?:(?!<w:r\b)[\s\S])*?<w:fldChar\s+[^>]*?w:fldCharType="begin"[\s\S]*?</w:r>'
                        r'[\s\S]*?'
                        r'<w:instrText[^>]*>\s*MERGEFIELD\s+(?:"([^"]+)"|([^\s"\\<>]+))[\s\S]*?</w:instrText>'
                        r'[\s\S]*?'
                        r'<w:r\b[^>]*>(?:(?!<w:r\b)[\s\S])*?<w:fldChar\s+[^>]*?w:fldCharType="end"[^>]*?>[\s\S]*?</w:r>',
                        re.IGNORECASE
                    )
                    xml_str = pattern_multi_complex.sub(replace_complex_field, xml_str)

                    # Single-run complex field dự phòng
                    pattern_single_complex = re.compile(
                        r'<w:r\b[^>]*>(?:(?!<w:r\b)[\s\S])*?<w:fldChar\s+[^>]*?w:fldCharType="begin"[\s\S]*?'
                        r'<w:instrText[^>]*>\s*MERGEFIELD\s+(?:"([^"]+)"|([^\s"\\<>]+))[\s\S]*?'
                        r'<w:fldChar\s+[^>]*?w:fldCharType="end"[^>]*?>(?:(?!<w:r\b)[\s\S])*?</w:r>',
                        re.IGNORECASE
                    )
                    xml_str = pattern_single_complex.sub(replace_complex_field, xml_str)

                    # 2.5. Thay thế riêng thẻ «tổng_hợp_phiếu_báo_có», «tổng_hợp_vốn_tự_có», «vốn_tự_có_n», «kế_hoạch_giải_ngân» dạng plain text
                    if th_bc_val and any(k in xml_str.lower() for k in ['báo_có', 'bao_co', 'baoco', 'báo có']):
                        pattern_plain_thbc = re.compile(
                            r'<w:r\b[^>]*>(?:(?!<w:r\b)[\s\S])*?<w:t[^>]*>[^<]*[«\xab]\s*(?:tổng_hợp_phiếu_báo_có|tong_hop_phieu_bao_co|tổng hợp phiếu báo có|tonghopphieubaoco)\s*[»\xbb][^<]*</w:t>(?:(?!<w:r\b)[\s\S])*?</w:r>',
                            re.IGNORECASE
                        )
                        def _rep_plain_thbc(m):
                            m_rpr = re.search(r'<w:rPr>([\s\S]*?)</w:rPr>', m.group(0))
                            rpr_str = f'<w:rPr>{m_rpr.group(1)}</w:rPr>' if m_rpr else ''
                            return format_runs_with_italics(th_bc_val, rpr_str)
                        xml_str = pattern_plain_thbc.sub(_rep_plain_thbc, xml_str)

                    if th_vtc_val and any(k in xml_str.lower() for k in ['vốn_tự_có', 'von_tu_co', 'vontuco', 'vốn tự có']):
                        pattern_plain_thvtc = re.compile(
                            r'<w:r\b[^>]*>(?:(?!<w:r\b)[\s\S])*?<w:t[^>]*>[^<]*[«\xab]\s*(?:tổng_hợp_vốn_tự_có|tong_hop_von_tu_co|tổng hợp vốn tự có|tonghopvontuco)\s*[»\xbb][^<]*</w:t>(?:(?!<w:r\b)[\s\S])*?</w:r>',
                            re.IGNORECASE
                        )
                        def _rep_plain_thvtc(m):
                            m_rpr = re.search(r'<w:rPr>([\s\S]*?)</w:rPr>', m.group(0))
                            rpr_str = f'<w:rPr>{m_rpr.group(1)}</w:rPr>' if m_rpr else ''
                            return format_runs_with_italics(th_vtc_val, rpr_str)
                        xml_str = pattern_plain_thvtc.sub(_rep_plain_thvtc, xml_str)

                    if th_vtc_n_val and any(k in xml_str.lower() for k in ['vốn_tự_có_n', 'von_tu_co_n', 'vontucon', 'vốn tự có n']):
                        pattern_plain_thvtcn = re.compile(
                            r'<w:r\b[^>]*>(?:(?!<w:r\b)[\s\S])*?<w:t[^>]*>[^<]*[«\xab]\s*(?:vốn_tự_có_n|von_tu_co_n|vốn tự có n|vontucon)\s*[»\xbb][^<]*</w:t>(?:(?!<w:r\b)[\s\S])*?</w:r>',
                            re.IGNORECASE
                        )
                        def _rep_plain_thvtcn(m):
                            m_rpr = re.search(r'<w:rPr>([\s\S]*?)</w:rPr>', m.group(0))
                            rpr_str = f'<w:rPr>{m_rpr.group(1)}</w:rPr>' if m_rpr else ''
                            return format_runs_with_italics(th_vtc_n_val, rpr_str)
                        xml_str = pattern_plain_thvtcn.sub(_rep_plain_thvtcn, xml_str)

                    if kh_gn_val and any(k in xml_str.lower() for k in ['giải_ngân', 'giai_ngan', 'giaingan', 'giải ngân']):
                        pattern_plain_khgn = re.compile(
                            r'<w:r\b[^>]*>(?:(?!<w:r\b)[\s\S])*?<w:t[^>]*>[^<]*[«\xab]\s*(?:kế[\s_]+hoạch[\s_]+giải[\s_]+ngân|ke[\s_]+hoach[\s_]+giai[\s_]+ngan)\s*[»\xbb][^<]*</w:t>(?:(?!<w:r\b)[\s\S])*?</w:r>',
                            re.IGNORECASE
                        )
                        def _rep_plain_khgn(m):
                            m_rpr = re.search(r'<w:rPr>([\s\S]*?)</w:rPr>', m.group(0))
                            rpr_str = f'<w:rPr>{m_rpr.group(1)}</w:rPr>' if m_rpr else ''
                            return format_runs_with_italics(kh_gn_val, rpr_str)
                        xml_str = pattern_plain_khgn.sub(_rep_plain_khgn, xml_str)


                    # 3. Placeholder «...» và <<...>>: Quét và thay thế toàn diện theo từ điển chuẩn hóa
                    # Loại bỏ trước toàn bộ thẻ kiểm tra chính tả <w:proofErr> của Word gây chia cắt run
                    xml_str = re.sub(r'<w:proofErr\b[^>]*/>', '', xml_str)

                    # Tránh lặp thông tin phụ thuộc nếu trong cùng 1 đoạn văn bản có cả 2 thẻ «người_phụ_thuộc» và «thông_tin_phụ_thuộc»
                    if any(k in xml_str.lower() for k in ['phụ_thuộc', 'phu_thuoc', 'phụ thuộc']):
                        def _dedup_pt_p(m_p):
                            p_xml = m_p.group(0)
                            plain_p = re.sub(r'<[^>]+>', '', p_xml)
                            if ('thông_tin_phụ_thuộc' in plain_p or 'thong_tin_phu_thuoc' in plain_p) and ('người_phụ_thuộc' in plain_p or 'nguoi_phu_thuoc' in plain_p):
                                p_xml = re.sub(r'[«\xab](?:<[^>]+>)*\s*người(?:<[^>]+>)*[_\s]*(?:<[^>]+>)*phụ(?:<[^>]+>)*[_\s]*(?:<[^>]+>)*thuộc(?:<[^>]+>)*\s*[»\xbb]\s*', '', p_xml, flags=re.IGNORECASE)
                            return p_xml
                        xml_str = re.sub(r'<w:p\b[^>]*>[\s\S]*?</w:p>', _dedup_pt_p, xml_str)

                    def _rep_any_placeholder(m_pl):
                        raw_tag = m_pl.group(1) or m_pl.group(2)
                        if not raw_tag:
                            return m_pl.group(0)
                        clean_tag = re.sub(r'<[^>]+>', '', raw_tag).strip().strip('"').strip("'")
                        val = get_val_smart(clean_tag)
                        if val is None:
                            return m_pl.group(0)
                        _merged_keys.add(clean_tag)
                        _merged_keys.add(normalize_key(clean_tag))
                        return escape_xml(str(val))

                    # Regex bắt trọn cả thẻ nằm gọn trong 1 run lẫn thẻ bị phân mảnh nhiều run trong cùng đoạn <w:p>
                    pattern_any_pl = re.compile(
                        r'[«\xab]((?:(?!<w:p\b)[^»\xbb])+?)[»\xbb]|(?:&lt;&lt;|<<)((?:(?!<w:p\b)[^>])+?)(?:&gt;&gt;|>>)'
                    )
                    xml_str = pattern_any_pl.sub(_rep_any_placeholder, xml_str)

                    # Quét dự phòng siêu tốc CHỈ KHI trong xml_str vẫn còn sót lại ký tự « hoặc <<
                    if '«' in xml_str or '\xab' in xml_str or '<<' in xml_str or '&lt;&lt;' in xml_str:
                        rem_matches = re.finditer(
                            r'[«\xab]((?:(?!<w:p\b)[^»\xbb])+?)[»\xbb]|(?:&lt;&lt;|<<)((?:(?!<w:p\b)[^>])+?)(?:&gt;&gt;|>>)',
                            xml_str
                        )
                        for r_m in list(rem_matches):
                            raw_k = r_m.group(1) or r_m.group(2)
                            if not raw_k: continue
                            clean_k = re.sub(r'<[^>]+>', '', raw_k).strip().strip('"').strip("'")
                            val = get_val_smart(clean_k)
                            if val is not None:
                                escaped_v = escape_xml(str(val))
                                xml_str = xml_str.replace(r_m.group(0), escaped_v)

                    # Loại bỏ khoảng trắng dư thừa phát sinh từ việc ghép các thẻ (ví dụ "  ngày", "  ký giữa", tabs thừa)
                    def clean_wt_spaces(m):
                        t_attrs = m.group(1)
                        text_content = m.group(2)
                        cleaned = text_content.replace('\t', '')
                        cleaned = re.sub(r'[ ]{2,}', ' ', cleaned)
                        return f'<w:t{t_attrs}>{cleaned}</w:t>'

                    xml_str = re.sub(r'<w:t([^>]*)>(.*?)</w:t>', clean_wt_spaces, xml_str, flags=re.DOTALL)

                    content = xml_str.encode('utf-8')

                zout.writestr(item, content)
    # Ghi buffer hoàn chỉnh ra file thực (chỉ sau khi toàn bộ zip entries đã được ghi xong)
    try:
        with open(out_path, 'wb') as _f:
            _f.write(_buf.getvalue())
    except PermissionError:
        # Nếu file Word đang mở trên máy người dùng (bị Word khóa), ghi tạm ra file _moi.docx để tránh bị gián đoạn
        alt_path = (out_path[:-5] + "_moi.docx") if out_path.lower().endswith(".docx") else (out_path + "_moi.docx")
        with open(alt_path, 'wb') as _f:
            _f.write(_buf.getvalue())

def find_best_excel_file():
    """Tự động quét tìm file Excel dữ liệu mới nhất trong thư mục ứng dụng hoặc thư mục Downloads."""
    candidates = []
    # 1. Thư mục chương trình
    base_files = glob.glob(os.path.join(BASE_DIR, "DATA_MAIL_MERGE*.xlsx")) + glob.glob(os.path.join(BASE_DIR, "*.xlsx"))
    for f in base_files:
        if not os.path.basename(f).startswith("~$"):
            candidates.append((os.path.getmtime(f), f))

    # 2. Thư mục Downloads của người dùng
    downloads = os.path.expanduser('~/Downloads')
    if os.path.exists(downloads):
        dl_files = (
            glob.glob(os.path.join(downloads, "DATA_MAIL_MERGE*.xlsx")) +
            glob.glob(os.path.join(downloads, "*206cot*.xlsx")) +
            glob.glob(os.path.join(downloads, "*207cot*.xlsx")) +
            glob.glob(os.path.join(downloads, "*262cot*.xlsx")) +
            glob.glob(os.path.join(downloads, "*NHCSXH*.xlsx"))
        )
        for f in dl_files:
            if not os.path.basename(f).startswith("~$"):
                candidates.append((os.path.getmtime(f), f))

    if not candidates:
        return None
    candidates.sort(key=lambda x: x[0], reverse=True)
    return candidates[0][1]

def run_mail_merge(excel_path=None, selected_indices=None, folder_mapping=None, selected_templates=None, log_callback=print, marital_mapping=None):
    log_callback("📄 Bắt đầu tiến trình xuất file hàng loạt...")
    
    if not excel_path:
        excel_path = find_best_excel_file()
        if not excel_path:
            log_callback("❌ Không tìm thấy file Excel nào trong thư mục ứng dụng hoặc Downloads!")
            log_callback("👉 Vui lòng mở Web và bấm 'Xuất File Excel 262 Cột' hoặc bấm chọn file Excel trên máy.")
            return False, "Không tìm thấy file Excel nguồn."

    log_callback(f"📂 Đọc file Excel: {os.path.basename(excel_path)}")
    
    try:
        import openpyxl
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        ws = wb['data'] if 'data' in wb.sheetnames else wb.active
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        
        db_records_map = {}
        if os.path.exists(DB_FILE):
            try:
                with open(DB_FILE, 'r', encoding='utf-8') as f:
                    db_data = json.load(f)
                    if isinstance(db_data, list):
                        for item in db_data:
                            if not isinstance(item, dict): continue
                            c_id = str(item.get("cccd") or item.get("soCCCD") or item.get("CCCD") or "").strip().replace(' ', '').replace('.', '').replace('-', '').replace('\t', '')
                            n_id = str(item.get("hoTen") or item.get("ho_ten") or item.get("họ và tên") or "").strip().lower()
                            rec_id = str(item.get("_id") or item.get("id") or "").strip()
                            if c_id: db_records_map[c_id] = item
                            if n_id: db_records_map[n_id] = item
                            if rec_id: db_records_map[rec_id] = item
            except Exception:
                pass

        records = []
        for r_idx in range(2, ws.max_row + 1):
            row_vals = [ws.cell(r_idx, c).value for c in range(1, len(headers) + 1)]
            if any(row_vals):
                rec = {}
                for idx, h in enumerate(headers):
                    if h:
                        val = row_vals[idx] if idx < len(row_vals) else ""
                        # Nếu số tiền = 0 hoặc bằng chữ là "không đồng" thì chuyển thành rỗng (null)
                        if val == 0 or str(val).strip() == "0" or str(val).strip().lower() in ["không đồng", "không", "null", "none"]:
                            val = ""
                        # Chuẩn hóa đọc số tiền: số 4 luôn đọc là 'bốn', không đọc là 'tư'
                        if isinstance(val, str) and "tư" in val.lower():
                            val = re.sub(r'\b(mươi|mười|lẻ)\s+tư\b', r'\1 bốn', val, flags=re.IGNORECASE)
                            val = re.sub(r'\bTư\s+(nghìn|ngàn|triệu|tỷ|đồng)\b', r'Bốn \1', val)
                            val = re.sub(r'\btư\s+(nghìn|ngàn|triệu|tỷ|đồng)\b', r'bốn \1', val)
                        # Chuẩn hóa đọc số tiền: thêm dấu phẩy sau các bậc triệu, ngàn, nghìn, tỷ nếu chưa có
                        if isinstance(val, str) and ("đồng" in val.lower() or "ngàn" in val.lower() or "nghìn" in val.lower() or "triệu" in val.lower() or "tỷ" in val.lower()):
                            val = re.sub(r'\b(triệu|ngàn|nghìn|tỷ)(?!\s*,)\s+(?=(?:không|một|hai|ba|bốn|năm|sáu|bảy|tám|chín|mười|lẻ)\b)', r'\1, ', val, flags=re.IGNORECASE)
                        if isinstance(val, str):
                            val = val.replace('\t', '')
                            if '\n' not in val:
                                val = re.sub(r' {2,}', ' ', val).strip()
                            else:
                                val = val.strip()
                        rec[str(h).strip()] = val if val is not None else ""

                # Đồng bộ tìm matched_db từ database_records.json cho dòng hiện tại (tránh rò rỉ biến từ dòng trước)
                c_key = str(rec.get("Số CCCD") or rec.get("cccd") or rec.get("CCCD") or rec.get("Số_CCCD") or "").strip().replace(' ', '').replace('.', '').replace('-', '').replace('\t', '')
                n_key = str(rec.get("Họ và tên") or rec.get("họ và tên") or rec.get("hoTen") or rec.get("Họ tên") or rec.get("Họ_và_tên") or "").strip().lower()
                id_key = str(rec.get("_id") or rec.get("id") or "").strip()
                matched_db = db_records_map.get(c_key) or db_records_map.get(n_key) or (db_records_map.get(id_key) if id_key else None)

                if matched_db:
                    # Bù đắp toàn bộ 206 cột dữ liệu đã được tính toán chuẩn xác từ database_records.json
                    db_row_full = build_row206_from_record(matched_db)
                    for k_full, v_full in db_row_full.items():
                        if v_full is not None and str(v_full).strip() != "":
                            if rec.get(k_full) is None or str(rec.get(k_full)).strip() == "":
                                rec[k_full] = v_full

                    # Bù đắp thêm dữ liệu đặc thù nếu record trong Excel bị thiếu/rỗng
                    if matched_db.get("tinhTrangHonNhan") and not rec.get("Tình trạng hôn nhân"):
                        rec["tinhTrangHonNhan"] = matched_db["tinhTrangHonNhan"]
                        rec["Tình trạng hôn nhân"] = matched_db["tinhTrangHonNhan"]
                    if matched_db.get("tongHopPhieuBaoCo") and not rec.get("tổng hợp phiếu báo có"):
                        rec["tongHopPhieuBaoCo"] = matched_db["tongHopPhieuBaoCo"]
                        rec["tổng_hợp_phiếu_báo_có"] = matched_db["tongHopPhieuBaoCo"]
                        rec["tổng hợp phiếu báo có"] = matched_db["tongHopPhieuBaoCo"]
                    if matched_db.get("tongHopVonTuCo") and not rec.get("tổng hợp vốn tự có"):
                        rec["tongHopVonTuCo"] = matched_db["tongHopVonTuCo"]
                        rec["tổng_hợp_vốn_tự_có"] = matched_db["tongHopVonTuCo"]
                        rec["tổng hợp vốn tự có"] = matched_db["tongHopVonTuCo"]
                    if matched_db.get("noiDungChungTuDot1") and not rec.get("chứng_từ_vốn_tự_có_đợt_1"):
                        rec["noiDungChungTuDot1"] = matched_db["noiDungChungTuDot1"]
                        rec["chứng_từ_vốn_tự_có_đợt_1"] = matched_db["noiDungChungTuDot1"]
                    if matched_db.get("noiDungChungTuDot2") and not rec.get("chứng_từ_vốn_tự_có_đợt_2"):
                        rec["noiDungChungTuDot2"] = matched_db["noiDungChungTuDot2"]
                        rec["chứng_từ_vốn_tự_có_đợt_2"] = matched_db["noiDungChungTuDot2"]
                    if matched_db.get("giaiChapAnPhu") and not rec.get("Giải chấp"):
                        rec["giaiChapAnPhu"] = matched_db["giaiChapAnPhu"]
                        rec["Giải chấp"] = matched_db["giaiChapAnPhu"]
                    if matched_db.get("giaiChapHongLoan") and not rec.get("Giải chấp Hồng Loan"):
                        rec["giaiChapHongLoan"] = matched_db["giaiChapHongLoan"]
                        rec["Giải chấp Hồng Loan"] = matched_db["giaiChapHongLoan"]
                    if matched_db.get("thongTinPhuThuoc") and not rec.get("thông tin phụ thuộc"):
                        rec["thongTinPhuThuoc"] = matched_db["thongTinPhuThuoc"]
                        rec["thông tin phụ thuộc"] = matched_db["thongTinPhuThuoc"]
                        rec["thông_tin_phụ_thuộc"] = matched_db["thongTinPhuThuoc"]
                        rec["người_phụ_thuộc"] = matched_db["thongTinPhuThuoc"]
                    if matched_db.get("xacNhanLuong") and not rec.get("xác nhận lương"):
                        rec["xacNhanLuong"] = matched_db["xacNhanLuong"]
                        rec["xác nhận lương"] = matched_db["xacNhanLuong"]
                        rec["xác_nhận_lương"] = matched_db["xacNhanLuong"]
                    if matched_db.get("maKH") and not rec.get("mã khách hàng"):
                        rec["maKH"] = matched_db["maKH"]
                        rec["mã khách hàng"] = matched_db["maKH"]
                        rec["mã_khách_hàng"] = matched_db["maKH"]
                    if matched_db.get("soTK102") and not rec.get("số tài khoản 102"):
                        rec["soTK102"] = matched_db["soTK102"]
                        rec["số tài khoản 102"] = matched_db["soTK102"]
                        rec["số_tài_khoản_102"] = matched_db["soTK102"]
                    if matched_db.get("canBo") and not rec.get("cán bộ TD"):
                        rec["canBo"] = matched_db["canBo"]
                        rec["cán bộ TD"] = matched_db["canBo"]
                    if matched_db.get("dxCanBo") and not rec.get("danh xưng cán bộ"):
                        rec["dxCanBo"] = matched_db["dxCanBo"]
                        rec["danh xưng cán bộ"] = matched_db["dxCanBo"]
                    if matched_db.get("sdtCanBo") and not rec.get("sđt cán bộ TD"):
                        rec["sdtCanBo"] = matched_db["sdtCanBo"]
                        rec["sđt cán bộ TD"] = matched_db["sdtCanBo"]

                # Tự động tính lại Tổng chi phí sinh hoạt = chi phí sinh hoạt + chi phí người phụ thuộc + chi phí khác nếu có dữ liệu
                try:
                    def _parse_num(val):
                        if not val: return 0.0
                        cleaned = str(val).replace(".", "").replace(",", "").strip()
                        return float(cleaned) if cleaned else 0.0

                    cp_sh = _parse_num(rec.get("Chi phí sinh hoạt"))
                    cp_pt = _parse_num(rec.get("Chi phí người phụ thuộc"))
                    cp_kh = _parse_num(rec.get("Chi phí khác"))
                    if cp_sh > 0 or cp_pt > 0 or cp_kh > 0:
                        tot_cp = int(round(cp_sh + cp_pt + cp_kh))
                        rec["Tổng chi phí hàng tháng"] = f"{tot_cp:,}".replace(",", ".")

                    # N112: Đóng số tiền lẻ = Tổng vốn tự có - Số tiền đợt 1
                    vtc_val = _parse_num(rec.get("Vốn tự có") or rec.get("vốn tự có") or rec.get("tongVTC") or 0)
                    amt1_val = _parse_num(rec.get("số tiền đợt 1") or rec.get("số_tiền_đợt_1") or rec.get("amt1") or 0)
                    if vtc_val > 0 and amt1_val > 0:
                        calc_le = max(0, int(round(vtc_val - amt1_val)))
                        if not rec.get("đóng số tiền lẻ") or not rec.get("bằng chữ số tiền lẻ"):
                            rec["đóng số tiền lẻ"] = f"{calc_le:,}".replace(",", ".")
                            rec["đóng_số_tiền_lẻ"] = rec["đóng số tiền lẻ"]

                    # Tự động tính Số tiền trả gốc hàng tháng, tháng cuối trả nếu chưa có
                    loan_v = _parse_num(rec.get("Vốn vay") or rec.get("vonVay") or rec.get("Vốn_vay"))
                    thoi_han_n = int(_parse_num(rec.get("Thời hạn vay") or rec.get("thoiHanVay") or 300)) or 300

                    goc_ht_db = _parse_num(matched_db.get("traGocHangThang")) if matched_db else 0
                    thang_cuoi_db = _parse_num(matched_db.get("thangCuoiTra")) if matched_db else 0
                    goc_kh_db = _parse_num(matched_db.get("traGocKyHan")) if matched_db else 0
                    goc_kh_cuoi_db = _parse_num(matched_db.get("traGocKyHanCuoi")) if matched_db else 0

                    g_ht = _parse_num(rec.get("số tiền trả gốc hàng tháng") or rec.get("số_tiền_trả_gốc_hàng_tháng")) or goc_ht_db
                    if g_ht <= 0 and loan_v > 0 and thoi_han_n > 0:
                        g_ht = math.ceil(loan_v / thoi_han_n / 1000) * 1000

                    t_cuoi = _parse_num(rec.get("tháng cuối trả") or rec.get("tháng_cuối_trả")) or thang_cuoi_db
                    if t_cuoi <= 0 and loan_v > 0 and thoi_han_n > 0:
                        t_cuoi = loan_v - g_ht * (thoi_han_n - 1)

                    g_kh = _parse_num(rec.get("số tiền trả gốc cho mỗi kỳ hạn") or rec.get("số_tiền_trả_gốc_cho_mỗi_kỳ_hạn")) or goc_kh_db
                    if g_kh <= 0 and loan_v > 0:
                        g_kh = round(loan_v / 50)

                    g_kh_cuoi = _parse_num(rec.get("số tiền trả gốc cho kỳ hạn cuối cùng") or rec.get("số_tiền_trả_gốc_cho_kỳ_hạn_cuối_cùng")) or goc_kh_cuoi_db
                    if g_kh_cuoi <= 0 and loan_v > 0:
                        g_kh_cuoi = loan_v - g_kh * 49

                    s_g_ht = f"{int(g_ht):,}".replace(",", ".") if g_ht > 0 else ""
                    s_t_cuoi = f"{int(t_cuoi):,}".replace(",", ".") if t_cuoi > 0 else ""
                    s_g_kh = f"{int(g_kh):,}".replace(",", ".") if g_kh > 0 else ""
                    s_g_kh_cuoi = f"{int(g_kh_cuoi):,}".replace(",", ".") if g_kh_cuoi > 0 else ""

                    rec["số tiền trả gốc hàng tháng"] = s_g_ht
                    rec["số_tiền_trả_gốc_hàng_tháng"] = s_g_ht
                    rec["tháng cuối trả"] = s_t_cuoi
                    rec["tháng_cuối_trả"] = s_t_cuoi
                    rec["số tiền trả gốc cho mỗi kỳ hạn"] = s_g_kh
                    rec["số_tiền_trả_gốc_cho_mỗi_kỳ_hạn"] = s_g_kh
                    rec["số tiền trả gốc cho kỳ hạn cuối cùng"] = s_g_kh_cuoi
                    rec["số_tiền_trả_gốc_cho_kỳ_hạn_cuối_cùng"] = s_g_kh_cuoi

                    # - Số tiền còn phải giải ngân đợt 3 = số tiền giải ngân đợt 3 = numAmt3
                    rem3_val = rec.get("Số tiền còn phải GN đợt 3") or rec.get("Số_tiền_còn_phải_GN_đợt_3") or rec.get("số tiền đợt 3") or ""
                    rem3_words = rec.get("Bằng chữ số tiền còn phải GN đợt 3") or rec.get("Bằng_chữ_số_tiền_còn_phải_GN_đợt_3") or rec.get("bằng chữ số tiền đợt 3") or ""
                    rec["Số tiền còn phải GN đợt 3"] = rem3_val
                    rec["Bằng chữ số tiền còn phải GN đợt 3"] = rem3_words
                    rec["Số_tiền_còn_phải_GN_đợt_3"] = rem3_val
                    rec["Bằng_chữ_số_tiền_còn_phải_GN_đợt_3"] = rem3_words
                    rec["Còn_phải_GN_đợt_3"] = rem3_val
                    rec["Bằng_chữ_còn_phải_GN_đợt_3"] = rem3_words

                    # Số tiền còn phải giải ngân đợt 4, 5, 6, 7: Mặc định "" (rỗng)
                    for k_dot in [4, 5, 6, 7]:
                        rem_k_val = rec.get(f"remain{k_dot}") or rec.get(f"Số tiền còn phải GN đợt {k_dot}") or ""
                        rec[f"Số tiền còn phải GN đợt {k_dot}"] = str(rem_k_val).strip() if rem_k_val else ""
                        rec[f"Số_tiền_còn_phải_GN_đợt_{k_dot}"] = rec[f"Số tiền còn phải GN đợt {k_dot}"]
                        rec[f"Còn_phải_GN_đợt_{k_dot}"] = rec[f"Số tiền còn phải GN đợt {k_dot}"]
                        rem_k_w = rec.get(f"Bằng chữ số tiền còn phải GN đợt {k_dot}") or (_word_m(rem_k_val) if rem_k_val else "")
                        rec[f"Bằng chữ số tiền còn phải GN đợt {k_dot}"] = rem_k_w
                        rec[f"Bằng_chữ_số_tiền_còn_phải_GN_đợt_{k_dot}"] = rem_k_w
                        rec[f"Bằng_chữ_còn_phải_GN_đợt_{k_dot}"] = rem_k_w

                    # - Số tiền còn giải ngân đợt 8 = Số tiền còn phải GN đợt 8 (số tiền đợt 8)
                    auto_rem8_doc = rec.get("số tiền đợt 8") or rec.get("số_tiền_đợt_8") or ""
                    auto_rem8_words = rec.get("bằng chữ số tiền đợt 8") or rec.get("bằng_chữ_số_tiền_đợt_8") or ""
                    if not rec.get("Còn_phải_GN_đợt_8"):
                        rec["Còn_phải_GN_đợt_8"] = auto_rem8_doc
                        rec["Bằng_chữ_còn_phải_GN_đợt_8"] = auto_rem8_words
                    rec["Số_tiền_còn_phải_GN_đợt_8"] = rec.get("Số_tiền_còn_phải_GN_đợt_8") or rec["Còn_phải_GN_đợt_8"]
                    rec["Bằng_chữ_số_tiền_còn_phải_GN_đợt_8"] = rec.get("Bằng_chữ_số_tiền_còn_phải_GN_đợt_8") or rec["Bằng_chữ_còn_phải_GN_đợt_8"]

                    # Chuẩn hóa ngày đợt 1 nếu có tiền tố "Ngày " để tránh trùng lặp "Ngày Ngày..." trong mẫu Word
                    if rec.get("ngày_đợt_1"):
                        raw_d1 = str(rec["ngày_đợt_1"]).strip()
                        if raw_d1.lower().startswith("ngày "):
                            clean_d1 = raw_d1[5:].strip()
                            rec["ngày_đợt_1"] = clean_d1
                            rec["ngay_dot_1"] = clean_d1

                    # Phụ lục, HĐMB & Giá bán đã bao gồm thuế GTGT
                    if rec.get("phụ lục"):
                        clean_pl = str(rec["phụ lục"]).replace('\t', '').strip()
                        rec["phụ_lục"] = clean_pl
                        rec["phu_luc"] = clean_pl
                        rec["phụ lục"] = clean_pl
                    if rec.get("ngày phụ lục"):
                        clean_npl = str(rec["ngày phụ lục"]).replace('\t', '').strip()
                        rec["ngày_phụ_lục"] = clean_npl
                        rec["ngay_phu_luc"] = clean_npl
                        rec["ngàythángnăm_phụ_lục"] = clean_npl
                    if rec.get("ngày,tháng,năm Phụ lục 4"):
                        clean_npl4 = str(rec["ngày,tháng,năm Phụ lục 4"]).replace('\t', '').strip()
                        rec["ngàythángnăm_Phụ_lục_4"] = clean_npl4
                    if rec.get("số hđmb 1") or rec.get("số hđmb1"):
                        clean_hdmb1 = str(rec.get("số hđmb 1") or rec.get("số hđmb1")).replace('\t', '').strip()
                        rec["số_hđmb_1"] = clean_hdmb1
                        rec["số_hđmb1"] = clean_hdmb1
                    if rec.get("ngày hđmb 1") or rec.get("ngày hđmb1"):
                        clean_ngay_hdmb1 = str(rec.get("ngày hđmb 1") or rec.get("ngày hđmb1")).replace('\t', '').strip()
                        rec["ngày_hđmb_1"] = clean_ngay_hdmb1
                        rec["ngày_hđmb1"] = clean_ngay_hdmb1
                    for vv_key in ["v/v thanh toán", "v/v thanh toán1", "v/v thanh toán2", "v/v thanh toán3"]:
                        if rec.get(vv_key):
                            rec[vv_key] = str(rec[vv_key]).replace('\t', '').strip()
                    if rec.get("Giá bán đã bao gồm thuế GTGT") or rec.get("Giá bán da bao gồm thuế GTGT"):
                        vat_price = rec.get("Giá bán đã bao gồm thuế GTGT") or rec.get("Giá bán da bao gồm thuế GTGT")
                        rec["Giá_bán_đã_bao_gồm_thuế_GTGT"] = vat_price
                        rec["Giá_bán_da_bao_gồm_thuế_GTGT"] = vat_price

                    # Chuẩn hóa ngày tháng năm TB phê duyệt cho vay
                    raw_tb_pd_m = str(
                        rec.get("ngàythángnăm_tb_phê_duyệt_cho_vay") or rec.get("ngày,tháng,năm tb phê duyệt cho vay") or
                        rec.get("ngày_tháng_năm_tb_phê_duyệt_cho_vay") or rec.get("ngày tb phê duyệt cho vay") or
                        (matched_db.get("ngayThangNamThongBao") if matched_db else "") or
                        (matched_db.get("ngayThongBao") if matched_db else "") or
                        (matched_db.get("thamDinhDen") if matched_db else "") or
                        rec.get("thamDinhDen") or rec.get("thẩm định đến ngày") or ""
                    ).strip()
                    full_tb_m, d_tb_m, m_tb_m, y_tb_m = format_admin_date(raw_tb_pd_m)
                    fmt_tb_m = _fmt_date(raw_tb_pd_m)
                    if not full_tb_m and fmt_tb_m:
                        m_parts = fmt_tb_m.split('/')
                        if len(m_parts) == 3:
                            full_tb_m = f"ngày {m_parts[0]} tháng {m_parts[1]} năm {m_parts[2]}"

                    for k_tb_m in [
                        "ngàythángnăm_tb_phê_duyệt_cho_vay", "ngày,tháng,năm tb phê duyệt cho vay",
                        "ngày_tháng_năm_tb_phê_duyệt_cho_vay", "ngaythangnam_tb_phe_duyet_cho_vay",
                        "Ngày__tháng__năm_TB_phê_duyệt_cho_vay", "Ngày_tháng_năm_tb_phê_duyệt_cho_vay",
                        "ngày tháng năm tb phê duyệt cho vay"
                    ]:
                        rec[k_tb_m] = full_tb_m or fmt_tb_m

                    for k_tb_m_s in ["ngày tb phê duyệt cho vay", "ngày_tb_phê_duyệt_cho_vay", "ngay_tb_phe_duyet_cho_vay", "ngayThongBao"]:
                        rec[k_tb_m_s] = fmt_tb_m or full_tb_m

                    # Chuẩn hóa toàn diện ngày tháng năm HĐTD, HĐTC, HĐ3B
                    ngay_vay_raw_m = str(rec.get("ngày vay") or rec.get("ngày_vay") or rec.get("ngayVay") or (matched_db.get("ngayVay") if matched_db else "") or "").strip()
                    ngay_hdtd_raw_m = str(rec.get("ngayHDTD") or rec.get("ngày hđtd") or rec.get("ngày_hđtd") or rec.get("ngày/tháng/năm hđtd") or rec.get("ngàythángnăm_hđtd") or (matched_db.get("ngayHDTD") if matched_db else "") or ngay_vay_raw_m or "").strip()
                    full_hdtd_m, d_hdtd_m, m_hdtd_m, y_hdtd_m = format_admin_date(ngay_hdtd_raw_m)
                    fmt_hdtd_m = _fmt_date(ngay_hdtd_raw_m)

                    rec["ngày hđtd"] = fmt_hdtd_m
                    rec["ngày_hđtd"] = fmt_hdtd_m
                    rec["ngay_hdtd"] = fmt_hdtd_m
                    rec["ngayHDTD"] = fmt_hdtd_m
                    rec["ngày/tháng/năm hđtd"] = full_hdtd_m or fmt_hdtd_m
                    rec["ngàythángnăm_hđtd"] = full_hdtd_m or fmt_hdtd_m
                    rec["ngày_tháng_năm_hđtd"] = full_hdtd_m or fmt_hdtd_m
                    rec["ngaythangnam_hdtd"] = full_hdtd_m or fmt_hdtd_m
                    rec["Ngày__tháng__năm_HĐTD"] = full_hdtd_m or fmt_hdtd_m
                    rec["Ngày_tháng_năm_HĐTD"] = full_hdtd_m or fmt_hdtd_m

                    y_td = y_hdtd_m or (d_hdtd_m if len(d_hdtd_m) == 4 else "") or ""
                    if not y_td:
                        m_y = re.search(r'\b(20\d\d)\b', str(ngay_hdtd_raw_m or ngay_vay_raw_m))
                        if m_y: y_td = m_y.group(1)
                    if y_td:
                        rec["Năm_HĐTD"] = y_td
                        rec["Năm HĐTD"] = y_td
                        rec["năm_hđtd"] = y_td

                    # HĐTC
                    ngay_hdtc_raw_m = str(rec.get("ngayHDTC") or rec.get("ngày hđtc") or rec.get("ngày_hđtc") or rec.get("ngày,tháng,năm hđtc") or rec.get("ngàythángnăm_hđtc") or (matched_db.get("ngayHDTC") if matched_db else "") or ngay_hdtd_raw_m or "").strip()
                    full_hdtc_m, d_hdtc_m, m_hdtc_m, y_hdtc_m = format_admin_date(ngay_hdtc_raw_m)
                    fmt_hdtc_m = _fmt_date(ngay_hdtc_raw_m)

                    rec["ngày hđtc"] = fmt_hdtc_m
                    rec["ngày_hđtc"] = fmt_hdtc_m
                    rec["ngay_hdtc"] = fmt_hdtc_m
                    rec["ngayHDTC"] = fmt_hdtc_m
                    rec["ngày,tháng,năm hđtc"] = full_hdtc_m or fmt_hdtc_m
                    rec["ngàythángnăm_hđtc"] = full_hdtc_m or fmt_hdtc_m
                    rec["ngày_tháng_năm_hđtc"] = full_hdtc_m or fmt_hdtc_m
                    rec["ngaythangnam_hdtc"] = full_hdtc_m or fmt_hdtc_m
                    rec["Ngày__tháng__năm_HĐTC"] = full_hdtc_m or fmt_hdtc_m
                    rec["Ngày_tháng_năm_HĐTC"] = full_hdtc_m or fmt_hdtc_m

                    y_tc = y_hdtc_m or y_td
                    if y_tc:
                        rec["Năm_HĐTC"] = y_tc
                        rec["Năm HĐTC"] = y_tc
                        rec["năm_hđtc"] = y_tc

                    # HĐ3B
                    ngay_hd3b_raw_m = str(rec.get("ngayHD3B") or rec.get("ngày hđ3b") or rec.get("ngày_hđ3b") or rec.get("ngày,tháng,năm hđ3b") or rec.get("ngàythángnăm_hđ3b") or (matched_db.get("ngayHD3B") if matched_db else "") or ngay_hdtd_raw_m or "").strip()
                    full_hd3b_m, d_hd3b_m, m_hd3b_m, y_hd3b_m = format_admin_date(ngay_hd3b_raw_m)
                    fmt_hd3b_m = _fmt_date(ngay_hd3b_raw_m)

                    rec["ngày hđ3b"] = fmt_hd3b_m
                    rec["ngày_hđ3b"] = fmt_hd3b_m
                    rec["ngay_hd3b"] = fmt_hd3b_m
                    rec["ngayHD3B"] = fmt_hd3b_m
                    rec["ngày,tháng,năm hđ3b"] = full_hd3b_m or fmt_hd3b_m
                    rec["ngàythángnăm_hđ3b"] = full_hd3b_m or fmt_hd3b_m
                    rec["ngày_tháng_năm_hđ3b"] = full_hd3b_m or fmt_hd3b_m
                    rec["ngaythangnam_hd3b"] = full_hd3b_m or fmt_hd3b_m

                    # Mã KH, Số TK 102, Thông tin phụ thuộc
                    makh_val = rec.get("mã khách hàng") or rec.get("Mã khách hàng") or rec.get("maKH") or ""
                    rec["mã_khách_hàng"] = makh_val
                    rec["Mã_khách_hàng"] = makh_val
                    rec["mã_kh"] = makh_val
                    rec["Mã_KH"] = makh_val
                    rec["ma_kh"] = makh_val
                    rec["makh"] = makh_val
                    rec["maKH"] = makh_val

                    stk102_val = rec.get("số tài khoản 102") or rec.get("Số tài khoản 102") or rec.get("soTK102") or ""
                    rec["số_tài_khoản_102"] = stk102_val
                    rec["Số_tài_khoản_102"] = stk102_val
                    rec["số_tk_102"] = stk102_val
                    rec["Số_TK_102"] = stk102_val
                    rec["stk_102"] = stk102_val
                    rec["STK_102"] = stk102_val
                    rec["so_tk_102"] = stk102_val
                    rec["soTK102"] = stk102_val

                    ttpt_val = rec.get("thông tin phụ thuộc") or rec.get("Thông tin phụ thuộc") or rec.get("thongTinPhuThuoc") or ""
                    rec["thông_tin_phụ_thuộc"] = ttpt_val
                    rec["Thông_tin_phụ_thuộc"] = ttpt_val
                    rec["thong_tin_phu_thuoc"] = ttpt_val
                    rec["người_phụ_thuộc"] = ttpt_val
                    rec["Người_phụ_thuộc"] = ttpt_val
                    rec["thongTinPhuThuoc"] = ttpt_val

                    # Danh xưng cán bộ tín dụng
                    dxcb_val = rec.get("danh xưng cán bộ tín dụng") or rec.get("danh xưng cán bộ") or rec.get("dxCanBo") or rec.get("gt_cbtd") or "Ông"
                    rec["danh_xưng_cán_bộ_tín_dụng"] = dxcb_val
                    rec["Danh_xưng_cán_bộ_tín_dụng"] = dxcb_val
                    rec["danh xưng cán bộ tín dụng"] = dxcb_val
                    rec["danh_xưng_cán_bộ"] = dxcb_val
                    rec["Danh_xưng_cán_bộ"] = dxcb_val
                    rec["danh xưng cán bộ"] = dxcb_val
                    rec["danh_xung_can_bo"] = dxcb_val
                    rec["dx_cán_bộ"] = dxcb_val
                    rec["dx_can_bo"] = dxcb_val
                    rec["dxCanBo"] = dxcb_val
                    rec["gt_cbtd"] = dxcb_val
                    rec["GT_CBTD"] = dxcb_val
                    rec["Danh_Xưng_CBTD_BBKT"] = dxcb_val
                    rec["danh_xưng_cbtd_bbkt"] = dxcb_val
                    if not rec.get("Tên_CBTD_BBKT"):
                        rec["Tên_CBTD_BBKT"] = rec.get("cán bộ TD") or rec.get("canBo") or ""

                    gc_ap_merge = str(rec.get("Giải chấp") or rec.get("giaiChapAnPhu") or (matched_db.get("giaiChapAnPhu") if matched_db else "") or "").strip()
                    gc_hl_merge = str(rec.get("Giải chấp Hồng Loan") or rec.get("giaiChapHongLoan") or (matched_db.get("giaiChapHongLoan") if matched_db else "") or "").strip()
                    rec["Giải chấp"] = gc_ap_merge
                    rec["Giải_chấp"] = gc_ap_merge
                    rec["giải_chấp"] = gc_ap_merge
                    rec["giải chấp"] = gc_ap_merge
                    rec["Giải chấp Hồng Loan"] = gc_hl_merge
                    rec["Giải_chấp_Hồng_Loan"] = gc_hl_merge
                    rec["giải_chấp_hồng_loan"] = gc_hl_merge
                    rec["giải chấp hồng loan"] = gc_hl_merge

                    # Nguồn vốn vay
                    nv_merge = str(rec.get("nguonVon") or rec.get("Nguồn vốn") or rec.get("Nguồn_vốn") or rec.get("nguồn_vốn") or (matched_db.get("nguonVon") if matched_db else "") or "Nguồn vốn ngân sách thành phố").strip() or "Nguồn vốn ngân sách thành phố"
                    rec["Nguồn vốn"] = nv_merge
                    rec["Nguồn_vốn"] = nv_merge
                    rec["nguồn_vốn"] = nv_merge.lower()
                    rec["NGUỒN_VỐN"] = nv_merge.upper()
                    rec["nguonVon"] = nv_merge
                    rec["nguon_von"] = remove_accents(nv_merge).lower() if 'remove_accents' in globals() else nv_merge.lower()

                    # Tài khoản và ngân hàng thụ hưởng dự án (Số tk, tên NH)
                    p_merge_cur = str(rec.get("dự_án") or rec.get("Dự án") or rec.get("project") or (matched_db.get("project") if matched_db else "") or "").strip().lower()
                    def_merge_acc = "113002951528"
                    def_merge_bank = "Ngân hàng TMCP Công thương Việt Nam - CN Bình Xuyên"
                    if "nam long" in p_merge_cur:
                        def_merge_acc = "7063020417"
                        def_merge_bank = "Phòng giao dịch Ngân hàng Chính sách xã hội Cái Răng"
                    elif "hồng loan" in p_merge_cur or "hong loan" in p_merge_cur:
                        def_merge_acc = "7063261897"
                        def_merge_bank = "Phòng giao dịch Ngân hàng Chính sách xã hội Cái Răng"

                    acc_merge_val = str(
                        rec.get("Số tk") or rec.get("Số_tk") or rec.get("số_tk") or rec.get("so_tk") or
                        rec.get("projectAccount") or rec.get("soTaiKhoan") or rec.get("soTaiKhoanTBGN") or
                        (matched_db.get("projectAccount") if matched_db else "") or
                        (matched_db.get("soTaiKhoan") if matched_db else "") or
                        (matched_db.get("soTaiKhoanTBGN") if matched_db else "") or def_merge_acc
                    ).strip() or def_merge_acc

                    bank_merge_val = str(
                        rec.get("tên NH") or rec.get("tên_NH") or rec.get("ten_nh") or rec.get("Tại_NH") or rec.get("tại_nh") or
                        rec.get("projectBank") or rec.get("taiNH") or rec.get("tenNganHang") or
                        (matched_db.get("projectBank") if matched_db else "") or
                        (matched_db.get("taiNH") if matched_db else "") or
                        (matched_db.get("tenNganHang") if matched_db else "") or def_merge_bank
                    ).strip() or def_merge_bank

                    for k_stk_m in [
                        "Số tk", "Số_tk", "số_tk", "so_tk", "Số_TK", "STK", "stk",
                        "Số tài khoản", "số tài khoản", "số_tài_khoản", "Số_tài_khoản",
                        "Số_tài_khoản_trên_TBGN", "số_tài_khoản_trên_tbgn"
                    ]:
                        rec[k_stk_m] = acc_merge_val

                    for k_nh_m in [
                        "tên NH", "tên_NH", "Tên_NH", "ten_nh", "Tên NH",
                        "Tại_NH", "tại_nh", "Tại NH", "tại NH",
                        "tên ngân hàng", "Tên ngân hàng", "tên_ngân_hàng", "Tên_ngân_hàng"
                    ]:
                        rec[k_nh_m] = bank_merge_val

                    def _fmt_money(val):
                        if not val: return ""
                        vn = _parse_num(val)
                        if vn > 0:
                            return f"{int(round(vn)):,}".replace(",", ".")
                        return str(val).strip()

                    # Chứng từ vốn tự có đợt 1, đợt 2 và tổng hợp phiếu báo có
                    vtc_dot1 = rec.get("chứng_từ_vốn_tự_có_đợt_1") or rec.get("chứng từ vốn tự có đợt 1") or rec.get("noiDungChungTuDot1") or ""
                    vtc_dot2 = rec.get("chứng_từ_vốn_tự_có_đợt_2") or rec.get("chứng từ vốn tự có đợt 2") or rec.get("noiDungChungTuDot2") or ""
                    th_bc = rec.get("tổng_hợp_phiếu_báo_có") or rec.get("tổng hợp phiếu báo có") or rec.get("tongHopPhieuBaoCo") or ""

                    # Nếu chưa có trong record Excel, tự động suy luận nhanh từ các cột vốn tự có 1..8
                    if not vtc_dot1 or not th_bc:
                        bc_list = []
                        for bi in range(1, 9):
                            bc_raw = str(rec.get(f"vốn tự có {bi}") or rec.get(f"vốn_tự_có_{bi}") or rec.get(f"baoCo{bi}") or "").strip()
                            bc_val = re.sub(r'(số giao dịch|chứng từ:|toán số:)\s*([A-Za-z0-9])', r'\1 \2', bc_raw, flags=re.IGNORECASE)
                            tien_val = _parse_num(rec.get(f"số tiền vốn tự có {bi}") or rec.get(f"số_tiền_vốn_tự_có_{bi}") or rec.get(f"vonTuCo{bi}") or 0)
                            if bc_val or tien_val > 0:
                                tien_str = f"{int(tien_val):,}".replace(",", ".") if tien_val > 0 else ""
                                desc = bc_val
                                if desc and tien_str:
                                    desc += f", số tiền: {tien_str} đồng"
                                elif tien_str:
                                    desc = f"Phiếu nộp tiền số tiền: {tien_str} đồng"
                                if desc:
                                    bc_list.append((tien_val, desc))

                        if not th_bc:
                            # Tự động tạo mẫu chuẩn 4 đợt hoặc 5 đợt nếu chưa có trong record
                            amt1_str = _fmt_money(rec.get("số tiền đợt 1")) or "«số_tiền_đợt_1»"
                            amt1_words = rec.get("bằng chữ số tiền đợt 1") or "«bằng_chữ_số_tiền_đợt_1»"
                            amt2_str = _fmt_money(rec.get("số tiền đợt 2")) or "«số_tiền_đợt_2»"
                            amt2_words = rec.get("bằng chữ số tiền đợt 2") or "«bằng_chữ_số_tiền_đợt_2»"
                            amt7_str = _fmt_money(rec.get("số tiền đợt 7")) or "«số_tiền_đợt_7»"
                            amt7_words = rec.get("bằng chữ số tiền đợt 7") or "«bằng_chữ_số_tiền_đợt_7»"
                            amt8_str = _fmt_money(rec.get("số tiền đợt 8")) or "«số_tiền_đợt_8»"
                            amt8_words = rec.get("bằng chữ số tiền đợt 8") or rec.get("bằng_chữ_số_tiền_đợt_8") or "«bằng_chữ_số_tiền_đợt_8»"

                            rem2_str = _fmt_money(rec.get("Số tiền còn phải GN đợt 2") or rec.get("Số_tiền_còn_phải_GN_đợt_2")) or "«Số_tiền_còn_phải_GN_đợt_2»"
                            rem2_words = rec.get("Bằng chữ số tiền còn phải GN đợt 2") or rec.get("Bằng_chữ_số_tiền_còn_phải_GN_đợt_2") or "«Bằng_chữ_số_tiền_còn_phải_GN_đợt_2»"

                            def format_schedule_cond(val, fallback):
                                if not val or not str(val).strip():
                                    return fallback
                                s = str(val).strip()
                                parts = s.split("-")
                                if len(parts) == 3 and len(parts[0]) == 4:
                                    s = f"{parts[2]}/{parts[1]}/{parts[0]}"
                                if re.match(r'^\d{1,2}/\d{1,2}/\d{2,4}$', s):
                                    return f"Ngày {s}"
                                return s

                            raw_date2 = rec.get("ngày đợt 2") or rec.get("ngày_đợt_2") or rec.get("desc2") or rec.get("date2") or ""
                            date2_str = format_schedule_cond(raw_date2, "«ngày đợt 2»")

                            raw_date7 = rec.get("ngày đợt 7") or rec.get("ngày_đợt_7") or rec.get("desc7") or rec.get("date7") or ""
                            date7_str = format_schedule_cond(raw_date7, "«ngày đợt 7»")

                            raw_date8 = rec.get("ngày đợt 8") or rec.get("ngày_đợt_8") or rec.get("desc8") or rec.get("date8") or ""
                            date8_str = format_schedule_cond(raw_date8, "«ngày đợt 8»")

                            kpbt_str = _fmt_money(rec.get("kinh phí bảo trì 2%") or rec.get("kinh_phí_bảo_trì_2%") or rec.get("kinh_phí_bảo_trì_2")) or "«kinh_phí_bảo_trì_2»"

                            def _get_bc_clean(key):
                                v = str(rec.get(key) or "").strip()
                                if v:
                                    v = re.sub(r'(số giao dịch|chứng từ:|toán số:)\s*([A-Za-z0-9])', r'\1 \2', v, flags=re.IGNORECASE)
                                    v = re.sub(r'[;,\s]+$', '', v)
                                return v

                            bc1_val = _get_bc_clean("vốn tự có 1") or _get_bc_clean("vốn_tự_có_1") or "«vốn_tự_có_1»"
                            bc2_val = _get_bc_clean("vốn tự có 2") or _get_bc_clean("vốn_tự_có_2") or "«vốn_tự_có_2»"
                            bc3_val = _get_bc_clean("vốn tự có 3") or _get_bc_clean("vốn_tự_có_3") or "«vốn_tự_có_3»"
                            bc4_val = _get_bc_clean("vốn tự có 4") or _get_bc_clean("vốn_tự_có_4") or "«vốn_tự_có_4»"
                            bc5_val = _get_bc_clean("vốn tự có 5") or _get_bc_clean("vốn_tự_có_5") or "«vốn_tự_có_5»"

                            vtc1_str = _fmt_money(rec.get("số tiền vốn tự có 1") or rec.get("số_tiền_vốn_tự_có_1")) or "«số_tiền_vốn_tự_có_1»"
                            vtc1_words = rec.get("bằng chữ số tiền vốn tự có 1") or rec.get("bằng_chữ_số_tiền_vốn_tự_có_1") or "«bằng_chữ_số_tiền_vốn_tự_có_1»"

                            vtc2_str = _fmt_money(rec.get("số tiền vốn tự có 2") or rec.get("số_tiền_vốn_tự_có_2")) or "«số_tiền_vốn_tự_có_2»"
                            vtc2_words = rec.get("bằng chữ số tiền vốn tự có 2") or rec.get("bằng_chữ_số_tiền_vốn_tự_có_2") or "«bằng_chữ_số_tiền_vốn_tự_có_2»"

                            vtc3_str = _fmt_money(rec.get("số tiền vốn tự có 3") or rec.get("số_tiền_vốn_tự_có_3")) or "«số_tiền_vốn_tự_có_3»"
                            vtc3_words = rec.get("bằng chữ số tiền vốn tự có 3") or rec.get("bằng_chữ_số_tiền_vốn_tự_có_3") or "«bằng_chữ_số_tiền_vốn_tự_có_3»"

                            vtc4_str = _fmt_money(rec.get("số tiền vốn tự có 4") or rec.get("số_tiền_vốn_tự_có_4")) or "«số_tiền_vốn_tự_có_4»"
                            vtc4_words = rec.get("bằng chữ số tiền vốn tự có 4") or rec.get("bằng_chữ_số_tiền_vốn_tự_có_4") or "«bằng_chữ_số_tiền_vốn_tự_có_4»"

                            vtc5_str = _fmt_money(rec.get("số tiền vốn tự có 5") or rec.get("số_tiền_vốn_tự_có_5")) or "«số_tiền_vốn_tự_có_5»"
                            vtc5_words = rec.get("bằng chữ số tiền vốn tự có 5") or rec.get("bằng_chữ_số_tiền_vốn_tự_có_5") or "«bằng_chữ_số_tiền_vốn_tự_có_5»"

                            # Tự động nhận diện các đợt giải ngân theo tỷ lệ % và số tiền
                            active_cols = []
                            pcts_map = {}
                            amts_map = {}
                            for col in range(1, 9):
                                p_val = _parse_num(rec.get(f"pct{col}") or rec.get(f"tỷ lệ đợt {col}") or rec.get(f"tỷ_lệ_đợt_{col}") or 0)
                                a_val = _parse_num(rec.get(f"số tiền đợt {col}") or rec.get(f"số_tiền_đợt_{col}") or 0)
                                raw_d = str(rec.get(f"ngày đợt {col}") or rec.get(f"ngày_đợt_{col}") or rec.get(f"desc{col}") or rec.get(f"date{col}") or "").strip()
                                pcts_map[col] = p_val
                                amts_map[col] = a_val
                                is_active = (p_val > 0 or a_val > 0 or (col == 1 and (pcts_map.get(1, 0) > 0 or a_val > 0 or bool(raw_d))))
                                if is_active:
                                    active_cols.append(col)

                            # Nếu không có cột nào có % > 0: mặc định mẫu 4 đợt (1, 2, 7, 8)
                            if not any(pcts_map[c] > 0 for c in range(1, 9)):
                                active_cols = [1, 2, 7, 8]
                                pcts_map[1] = 0.2
                                pcts_map[2] = 0.5
                                pcts_map[7] = 0.25
                                pcts_map[8] = 0.05

                            # Hàm đọc số thành chữ tiếng Việt dùng cho mail merge
                            def _read_num_vn(val):
                                if not val:
                                    return ""
                                try:
                                    n = int(round(float(str(val).replace(",", "").replace(".", "").strip())))
                                except:
                                    return ""
                                if n == 0:
                                    return ""
                                neg = n < 0
                                abs_v = abs(n)
                                d_arr = ["không", "một", "hai", "ba", "bốn", "năm", "sáu", "bảy", "tám", "chín"]
                                sc_arr = ["", "ngàn", "triệu", "tỷ", "ngàn tỷ", "triệu tỷ", "tỷ tỷ"]

                                def r3(num, f_hundred):
                                    a = num // 100
                                    b = (num % 100) // 10
                                    c = num % 10
                                    s = []
                                    if a > 0:
                                        s.append(d_arr[a] + " trăm")
                                    elif f_hundred:
                                        s.append("không trăm")
                                    if b == 0:
                                        if c > 0 and (a > 0 or f_hundred):
                                            s.append("lẻ")
                                    elif b == 1:
                                        s.append("mười")
                                    else:
                                        s.append(d_arr[b] + " mươi")
                                    if c > 0:
                                        w = d_arr[c]
                                        if b >= 2 and c == 1:
                                            w = "mốt"
                                        elif b >= 1 and c == 5:
                                            w = "lăm"
                                        elif c == 4:
                                            w = "bốn"
                                        s.append(w)
                                    return " ".join(s)

                                grps = []
                                tmp = abs_v
                                while tmp > 0:
                                    grps.append(tmp % 1000)
                                    tmp //= 1000

                                prts = []
                                for i in range(len(grps) - 1, -1, -1):
                                    g = grps[i]
                                    if not g:
                                        continue
                                    f_h = (i < len(grps) - 1 and g < 100)
                                    t = r3(g, f_h)
                                    if sc_arr[i]:
                                        t += " " + sc_arr[i]
                                    prts.append(t)

                                res = ", ".join(prts) + " đồng"
                                res = re.sub(r'\b(mươi|mười|lẻ)\s+tư\b', r'\1 bốn', res, flags=re.IGNORECASE)
                                res = re.sub(r'\b(tư)\s+(nghìn|ngàn|triệu|tỷ|đồng)\b', r'bốn \2', res, flags=re.IGNORECASE)
                                res = re.sub(r'\b(Tư)\s+(nghìn|ngàn|triệu|tỷ|đồng)\b', r'Bốn \2', res, flags=re.IGNORECASE)
                                if neg:
                                    res = "Âm " + res
                                return res[0].upper() + res[1:]

                            # Thu thập và phân bổ động tất cả các Vouchers cho Đợt 1 và Đợt 2 theo nghiệp vụ ngân hàng
                            vouchers = []
                            for i in range(1, 9):
                                n_vtc = _parse_num(rec.get(f"số tiền vốn tự có {i}") or rec.get(f"số_tiền_vốn_tự_có_{i}") or 0)
                                b_co = _get_bc_clean(f"vốn tự có {i}") or _get_bc_clean(f"vốn_tự_có_{i}") or ""
                                if n_vtc > 0 or b_co:
                                    v_str = _fmt_money(n_vtc) if n_vtc > 0 else f"«số_tiền_vốn_tự_có_{i}»"
                                    v_words = rec.get(f"bằng chữ số tiền vốn tự có {i}") or rec.get(f"bằng_chữ_số_tiền_vốn_tự_có_{i}") or (_read_num_vn(n_vtc) if n_vtc > 0 else f"«bằng_chữ_số_tiền_vốn_tự_có_{i}»")
                                    b_val = b_co or f"«vốn_tự_có_{i}»"
                                    vouchers.append({
                                        'index': i,
                                        'amount': n_vtc,
                                        'bc_val': b_val,
                                        'v_str': v_str,
                                        'v_words': v_words
                                    })

                            amt1_n = amts_map.get(1, 0) or _parse_num(rec.get("số tiền đợt 1") or 0)
                            so_du_dot1 = 0
                            phieu_dot1_text = ""
                            p1_vouchers = []
                            next_voucher_idx = 0
                            sum_p1 = 0

                            if vouchers:
                                if amt1_n > 0:
                                    for k_v, v_item in enumerate(vouchers):
                                        p1_vouchers.append(v_item)
                                        sum_p1 += v_item['amount']
                                        next_voucher_idx = k_v + 1
                                        if sum_p1 >= amt1_n:
                                            break
                                else:
                                    bc2_val = _get_bc_clean("vốn tự có 2") or _get_bc_clean("vốn_tự_có_2") or ""
                                    vtc2_n = _parse_num(rec.get("số tiền vốn tự có 2") or rec.get("số_tiền_vốn_tự_có_2") or 0)
                                    has_p2 = bool(vtc2_n > 0 or (bc2_val and bc2_val != "«vốn_tự_có_2»"))
                                    p1_vouchers = vouchers[:2] if (has_p2 and len(vouchers) > 1) else [vouchers[0]]
                                    next_voucher_idx = len(p1_vouchers)
                                    sum_p1 = sum(v['amount'] for v in p1_vouchers)

                            if p1_vouchers:
                                list_p1_items = [f"số tiền {v['v_str']} đồng (Bằng chữ: {v['v_words']}) theo: {v['bc_val']}" for v in p1_vouchers]
                                if len(list_p1_items) == 1:
                                    base_p1_list = list_p1_items[0]
                                elif len(list_p1_items) == 2:
                                    base_p1_list = f"{list_p1_items[0]} và {list_p1_items[1]}"
                                else:
                                    base_p1_list = f"{', '.join(list_p1_items[:-1])} và {list_p1_items[-1]}"

                                if amt1_n > 0 and sum_p1 > amt1_n:
                                    so_du_dot1 = sum_p1 - amt1_n
                                    amt1_s = _fmt_money(amt1_n)
                                    amt1_w = _read_num_vn(amt1_n)
                                    sd_s = _fmt_money(so_du_dot1)
                                    sd_w = _read_num_vn(so_du_dot1)
                                    phieu_dot1_text = f"Bên mua đã nộp {base_p1_list}, trong đó số tiền của đợt 1 là {amt1_s} đồng (Bằng chữ: {amt1_w}) và số tiền của đợt 2 là {sd_s} đồng (Bằng chữ: {sd_w})."
                                else:
                                    phieu_dot1_text = f"Bên mua đã nộp {base_p1_list}."

                            # Thu thập vouchers Đợt 2 (số dư Đợt 1 + các phiếu còn lại từ next_voucher_idx)
                            list_p2 = []
                            if so_du_dot1 > 0:
                                list_p2.append(f"số tiền {_fmt_money(so_du_dot1)} đồng (Bằng chữ: {_read_num_vn(so_du_dot1)})")

                            for k_v in range(next_voucher_idx, len(vouchers)):
                                v_item = vouchers[k_v]
                                list_p2.append(f"số tiền {v_item['v_str']} đồng (Bằng chữ: {v_item['v_words']}) theo: {v_item['bc_val']}")

                            phieu_dot2_text = ""
                            if len(list_p2) == 1:
                                phieu_dot2_text = f" Bên mua đã nộp {list_p2[0]}."
                            elif len(list_p2) == 2:
                                phieu_dot2_text = f" Bên mua đã nộp {list_p2[0]} và {list_p2[1]}."
                            elif len(list_p2) > 2:
                                phieu_dot2_text = f" Bên mua đã nộp {', '.join(list_p2[:-1])} và {list_p2[-1]}."

                            out_dots = []
                            cur_stt = 1
                            last_active_col = active_cols[-1] if active_cols else -1
                            for col in active_cols:
                                p_val = pcts_map[col]
                                a_val = amts_map[col]
                                p_str = f"{int(round(p_val * 100))}%" if (0 < p_val <= 1) else (f"{int(round(p_val))}%" if p_val > 1 else f"«%_đợt_{col}»")
                                a_str = _fmt_money(a_val) if a_val > 0 else (_fmt_money(rec.get(f"số tiền đợt {col}")) or f"«số_tiền_đợt_{col}»")
                                a_words = rec.get(f"bằng chữ số tiền đợt {col}") or rec.get(f"bằng_chữ_số_tiền_đợt_{col}") or f"«bằng_chữ_số_tiền_đợt_{col}»"
                                raw_d = str(rec.get(f"desc{col}") or rec.get(f"nội dung đợt {col}") or rec.get(f"nội_dung_đợt_{col}") or rec.get(f"desc{cur_stt}") or rec.get(f"nội dung đợt {cur_stt}") or rec.get(f"nội_dung_đợt_{cur_stt}") or rec.get(f"ngày đợt {col}") or rec.get(f"ngày_đợt_{col}") or rec.get(f"date{col}") or rec.get(f"ngày đợt {cur_stt}") or rec.get(f"ngày_đợt_{cur_stt}") or rec.get(f"date{cur_stt}") or "").strip()

                                clean_c = raw_d.strip()
                                conn = " với số tiền "

                                if not clean_c:
                                    clean_c = f"«nội_dung_đợt_{cur_stt}»"
                                    conn = " với số tiền "
                                else:
                                    if re.search(r'(?:với\s+)?số\s+tiền(?:\s+tương\s+đương)?(?:\s+là)?\s*[:]?\s*$', clean_c, flags=re.IGNORECASE):
                                        conn = "" if clean_c.endswith(" ") else " "
                                    elif re.search(r'[,]\s*$', clean_c):
                                        conn = "số tiền là " if clean_c.endswith(" ") else " số tiền là "
                                    elif re.search(r'[:]\s*$', clean_c):
                                        conn = "" if clean_c.endswith(" ") else " "
                                    else:
                                        conn = " với số tiền "

                                line_c = f"+ Đợt {cur_stt}: {clean_c}{conn}{a_str} đồng (Bằng chữ: {a_words})."

                                if col == 1 and phieu_dot1_text:
                                    line_c += f" {phieu_dot1_text}"
                                elif col == 2 and phieu_dot2_text:
                                    line_c += f" {phieu_dot2_text.strip()}"

                                if col > 1 and col != last_active_col:
                                    rem_num = _parse_num(rec.get(f"Số tiền còn phải GN đợt {col}") or rec.get(f"Số_tiền_còn_phải_GN_đợt_{col}") or 0)
                                    if rem_num > 0:
                                        rem_s = f"{int(rem_num):,}".replace(",", ".")
                                        rem_w = rec.get(f"Bằng chữ số tiền còn phải GN đợt {col}") or rec.get(f"Bằng_chữ_số_tiền_còn_phải_GN_đợt_{col}") or _read_num_vn(rem_num)
                                        line_c += f" Số tiền còn phải thanh toán là {rem_s} đồng" + (f" (Bằng chữ: {rem_w})" if rem_w else "") + "."

                                # Đối với An Phú: Đợt kế cuối (tức Đợt 3 / Cột 7) nối thêm câu KPBT 2%
                                p_name_cur = str(rec.get("project") or rec.get("dự án") or rec.get("duAn") or rec.get("du_an") or rec.get("Dự án") or rec.get("tenDuAn") or "").strip().lower()
                                is_an_phu_2026 = "an phú" in p_name_cur or "an phu" in p_name_cur or (len(active_cols) == 4 and active_cols[0] == 1 and active_cols[1] == 2 and active_cols[2] == 7 and active_cols[3] == 8)
                                penultimate_active = active_cols[-2] if len(active_cols) >= 2 else -1
                                if is_an_phu_2026 and (col == penultimate_active or col == 7):
                                    kpbt_v = _parse_num(rec.get("kinh phí bảo trì 2%") or rec.get("kinh_phí_bảo_trì_2%") or rec.get("kinh_phí_bảo_trì_2") or rec.get("kinhPhiBaoTri") or 0)
                                    kpbt_s = _fmt_money(kpbt_v) if kpbt_v > 0 else (str(rec.get("kinh phí bảo trì 2%") or rec.get("kinh_phí_bảo_trì_2%") or rec.get("kinh_phí_bảo_trì_2") or rec.get("kinhPhiBaoTri") or "").strip() or "«kinh phí bảo trì 2%»")
                                    line_c += f" Trong đó, Bên mua tự thanh toán bằng vốn tự có kinh phí bảo trì 2% số tiền là {kpbt_s} đồng cho Bên bán khi có thông báo bàn giao căn hộ và trước khi nhận bàn giao Căn hộ."

                                out_dots.append(re.sub(r'\s+', ' ', line_c).strip())
                                cur_stt += 1

                            # Chuẩn hóa dấu ngắt câu: các đợt trước đợt cuối kết thúc bằng (;), riêng đợt cuối kết thúc bằng (.)
                            for idx in range(len(out_dots)):
                                out_dots[idx] = re.sub(r'[.;\s]+$', '', out_dots[idx])
                                if idx < len(out_dots) - 1:
                                    out_dots[idx] += ";"
                                else:
                                    out_dots[idx] += "."

                            th_bc = "\n".join(out_dots)

                        if not vtc_dot1 and bc_list:
                            gn1_val = _parse_num(rec.get("số tiền GN đầu tiên") or rec.get("số tiền đợt 1") or 0)
                            p1_tien, p1_desc = bc_list[0]
                            if gn1_val > 0:
                                if p1_tien == gn1_val:
                                    vtc_dot1 = p1_desc
                                    if len(bc_list) > 1 and not vtc_dot2:
                                        vtc_dot2 = "; ".join([item[1] for item in bc_list[1:]])
                                elif p1_tien < gn1_val:
                                    if len(bc_list) > 1:
                                        p2_tien, p2_desc = bc_list[1]
                                        p12_tien = p1_tien + p2_tien
                                        if p12_tien >= gn1_val:
                                            so_du = p12_tien - gn1_val
                                            if so_du > 0:
                                                so_du_str = f"{int(so_du):,}".replace(",", ".")
                                                vtc_dot1 = f"{p1_desc}; {p2_desc} (trong đó số dư {so_du_str} đồng chuyển thanh toán đợt 2)"
                                                if not vtc_dot2:
                                                    rem_v = [f"Số dư từ chứng từ vốn tự có đợt 1 chuyển sang: {so_du_str} đồng"] + [item[1] for item in bc_list[2:]]
                                                    vtc_dot2 = "; ".join(rem_v)
                                            else:
                                                vtc_dot1 = f"{p1_desc}; {p2_desc}"
                                        else:
                                            vtc_dot1 = f"{p1_desc}; {p2_desc}"
                                    else:
                                        vtc_dot1 = p1_desc
                                elif p1_tien > gn1_val:
                                    so_du = p1_tien - gn1_val
                                    so_du_str = f"{int(so_du):,}".replace(",", ".")
                                    vtc_dot1 = f"{p1_desc} (trong đó số dư {so_du_str} đồng chuyển thanh toán đợt 2)"
                                    if not vtc_dot2:
                                        rem_v = [f"Số dư từ chứng từ vốn tự có đợt 1 chuyển sang: {so_du_str} đồng"] + [item[1] for item in bc_list[1:]]
                                        vtc_dot2 = "; ".join(rem_v)
                            else:
                                vtc_dot1 = p1_desc

                    rec["chứng_từ_vốn_tự_có_đợt_1"] = vtc_dot1
                    rec["chứng từ vốn tự có đợt 1"] = vtc_dot1
                    rec["chứng_từ_đợt_1"] = vtc_dot1
                    rec["chung_tu_vtc_dot_1"] = vtc_dot1
                    rec["chung_tu_dot_1"] = vtc_dot1

                    rec["chứng_từ_vốn_tự_có_đợt_2"] = vtc_dot2
                    rec["chứng từ vốn tự có đợt 2"] = vtc_dot2
                    rec["chứng_từ_đợt_2"] = vtc_dot2
                    rec["chung_tu_vtc_dot_2"] = vtc_dot2
                    rec["chung_tu_dot_2"] = vtc_dot2

                    # Đảm bảo các tag placeholder ngày đợt 2, 7, 8 trong th_bc được thay thế nếu còn sót
                    if th_bc and "«ngày" in th_bc:
                        def _get_cond(key_list, fallback):
                            for k in key_list:
                                v = rec.get(k)
                                if v and str(v).strip():
                                    s = str(v).strip()
                                    parts = s.split("-")
                                    if len(parts) == 3 and len(parts[0]) == 4:
                                        s = f"{parts[2]}/{parts[1]}/{parts[0]}"
                                    if re.match(r'^\d{1,2}/\d{1,2}/\d{2,4}$', s):
                                        return f"Ngày {s}"
                                    return s
                            return fallback
                        for k in range(1, 9):
                            ck = _get_cond([f"ngày đợt {k}", f"ngày_đợt_{k}", f"desc{k}", f"date{k}"], f"«ngày đợt {k}»")
                            th_bc = th_bc.replace(f"««ngày đợt {k}»»", ck).replace(f"«ngày_đợt_{k}»", ck).replace(f"«ngày đợt {k}»", ck)

                    # Đảm bảo đối với An Phú: nếu th_bc đã có từ trước mà chưa nối câu KPBT 2% thì nối vào đợt kế cuối
                    p_name_cur = str(rec.get("project") or rec.get("dự án") or rec.get("duAn") or rec.get("du_an") or rec.get("Dự án") or rec.get("tenDuAn") or "").strip().lower()
                    is_an_phu_2026 = "an phú" in p_name_cur or "an phu" in p_name_cur or "bình xuyên" in str(rec).lower() or "113002951528" in str(rec)
                    if is_an_phu_2026 and th_bc and "kinh phí bảo trì 2%" not in th_bc:
                        kpbt_v = _parse_num(rec.get("kinh phí bảo trì 2%") or rec.get("kinh_phí_bảo_trì_2%") or rec.get("kinh_phí_bảo_trì_2") or rec.get("kinhPhiBaoTri") or 0)
                        kpbt_s = _fmt_money(kpbt_v) if kpbt_v > 0 else (str(rec.get("kinh phí bảo trì 2%") or rec.get("kinh_phí_bảo_trì_2%") or rec.get("kinh_phí_bảo_trì_2") or rec.get("kinhPhiBaoTri") or "").strip() or "«kinh phí bảo trì 2%»")
                        th_lines = th_bc.split("\n")
                        if len(th_lines) >= 2:
                            target_idx = len(th_lines) - 2
                            line_target = re.sub(r'[.;\s]+$', '', th_lines[target_idx])
                            th_lines[target_idx] = line_target + f" Trong đó, Bên mua tự thanh toán bằng vốn tự có kinh phí bảo trì 2% số tiền là {kpbt_s} đồng cho Bên bán khi có thông báo bàn giao căn hộ và trước khi nhận bàn giao Căn hộ.;"
                            th_bc = "\n".join(th_lines)

                    rec["tổng_hợp_phiếu_báo_có"] = th_bc
                    rec["tổng hợp phiếu báo có"] = th_bc
                    rec["tong_hop_phieu_bao_co"] = th_bc
                    rec["tongHopPhieuBaoCo"] = th_bc
                    rec["TỔNG_HỢP_PHIẾU_BÁO_CÓ"] = th_bc
                    rec["Tổng_hợp_phiếu_báo_có"] = th_bc
                    rec["tong_hop_bao_co"] = th_bc

                    # Tổng hợp vốn tự có nối chuỗi tự sinh thứ tự (1), (2)... và 2% KPBT
                    th_vtc = rec.get("tổng_hợp_vốn_tự_có") or rec.get("tổng hợp vốn tự có") or rec.get("tongHopVonTuCo") or ""
                    if not th_vtc:
                        lines_vtc = []
                        order_idx = 1
                        for bi in range(1, 9):
                            bc_raw = str(rec.get(f"vốn tự có {bi}") or rec.get(f"vốn_tự_có_{bi}") or rec.get(f"baoCo{bi}") or "").strip()
                            bc_val = re.sub(r'(số giao dịch|chứng từ:|toán số:)\s*([A-Za-z0-9])', r'\1 \2', bc_raw, flags=re.IGNORECASE)
                            bc_val = re.sub(r'[;,\s]+$', '', bc_val)
                            tien_val = _parse_num(rec.get(f"số tiền vốn tự có {bi}") or rec.get(f"số_tiền_vốn_tự_có_{bi}") or rec.get(f"vonTuCo{bi}") or 0)
                            if bc_val or tien_val > 0:
                                tien_str = f"{int(tien_val):,}".replace(",", ".") if tien_val > 0 else f"«số_tiền_vốn_tự_có_{bi}»"
                                bc_str = bc_val if bc_val else f"«vốn_tự_có_{bi}»"
                                lines_vtc.append(f"({order_idx}) {bc_str}, số tiền: {tien_str} đồng.")
                                order_idx += 1
                        if not lines_vtc:
                            lines_vtc.append("(1) «vốn_tự_có_1», số tiền: «số_tiền_vốn_tự_có_1» đồng.")
                            order_idx = 2
                        lines_vtc_n = list(lines_vtc)
                        kpbt_val = _parse_num(rec.get("kinh phí bảo trì 2%") or rec.get("kinh_phí_bảo_trì_2") or rec.get("kinhPhiBaoTri") or 0)
                        kpbt_str = f"{int(kpbt_val):,}".replace(",", ".") if kpbt_val > 0 else "«kinh_phí_bảo_trì_2»"
                        lines_vtc.append(f"({order_idx}) Số tiền thanh toán 2% kinh phí bảo trì {kpbt_str} đồng. Bên mua tự thanh toán bằng vốn tự có cho Bên Bán trong thời hạn từ ngày Bên Bán gửi thông báo bàn giao Căn hộ đến trước thời điểm ký Biên bản bàn giao căn hộ.")
                        th_vtc = "\n".join(lines_vtc)
                        th_vtc_n = "\n".join(lines_vtc_n)

                    rec["tổng_hợp_vốn_tự_có"] = th_vtc
                    rec["tổng hợp vốn tự có"] = th_vtc
                    rec["tong_hop_von_tu_co"] = th_vtc
                    rec["tongHopVonTuCo"] = th_vtc

                    rec["vốn_tự_có_n"] = th_vtc_n
                    rec["vốn tự có n"] = th_vtc_n
                    rec["von_tu_co_n"] = th_vtc_n
                    rec["vonTuCoN"] = th_vtc_n

                    # Kế hoạch giải ngân (Mục 3 HĐTD)
                    kh_gn_docx = rec.get("Kế hoạch giải ngân") or rec.get("kế_hoạch_giải_ngân") or rec.get("keHoachGiaiNgan")
                    if not kh_gn_docx:
                        act_cols = []
                        for col in range(2, 9):
                            p_v = _parse_num(rec.get(f"pct{col}") or rec.get(f"tỷ lệ đợt {col}"))
                            a_v = _parse_num(rec.get(f"số tiền đợt {col}"))
                            d_v = str(rec.get(f"ngày đợt {col}") or rec.get(f"date{col}") or rec.get(f"desc{col}") or "").strip()
                            if p_v > 0 or a_v > 0:
                                act_cols.append(col)
                        c_list = act_cols if act_cols else [2, 3, 4, 8]
                        m_len = len(c_list)
                        lines_kh = []
                        is_hl_nl = ("hồng loan" in str(rec.get("project") or rec.get("Dự án") or "").lower() or 
                                    "nam long" in str(rec.get("project") or rec.get("Dự án") or "").lower())
                        maint_num = _parse_num(rec.get("kinh phí bảo trì 2%"))

                        def _get_money_connector(content):
                            if not content: return ", số tiền "
                            if re.search(r'(?:với\s+)?số\s+tiền(?:\s+tương\s+đương)?(?:\s+là)?\s*[:]?\s*$', content, flags=re.IGNORECASE):
                                return "" if content.endswith(" ") else " "
                            if re.search(r'[,]\s*$', content):
                                return "số tiền là " if content.endswith(" ") else " số tiền là "
                            if re.search(r'[:]\s*$', content):
                                return "" if content.endswith(" ") else " "
                            return ", số tiền "

                        def _fmt_sc(val, fallback):
                            s = str(val or "").strip()
                            if not s: return fallback
                            if re.match(r'^\d{4}-\d{2}-\d{2}$', s):
                                parts = s.split('-')
                                return f"Ngày {parts[2]}/{parts[1]}/{parts[0]}"
                            if re.match(r'^\d{1,2}/\d{1,2}/\d{2,4}$', s):
                                return f"Ngày {s}"
                            return s

                        for idx, col in enumerate(c_list):
                            lan = idx + 1
                            is_last = (idx == m_len - 1)
                            label_prefix = f"+ Lần {lan}:"

                            if idx == 0:
                                raw_d1 = rec.get("ngayGNDauTien") or rec.get("ngayVay") or rec.get("ngày GN đầu tiên") or rec.get("ngày vay") or rec.get("date2")
                                d1_s = _fmt_sc(raw_d1, "«ngày_GN_đầu_tiên»")
                                if not re.match(r'^ngày\s+', d1_s, flags=re.IGNORECASE):
                                    d1_s = f"Ngày {d1_s}"
                                rem2_n = _parse_num(rec.get("Số tiền còn phải GN đợt 2") or rec.get("số tiền GN đầu tiên") or rec.get("soTienGNDauTien") or rec.get("số tiền đợt 2"))
                                rem2_str = _fmt_money(rem2_n) if rem2_n else "«số_tiền_GN_đầu_tiên»"
                                rem2_w = _read_num_vn(rem2_n) if rem2_n else "«bằng_chữ_số_tiền_GN_đầu_tiên»"
                                lines_kh.append(f"{label_prefix} {d1_s}, số tiền {rem2_str} đồng (Bằng chữ: {rem2_w})")
                            elif col == 7:
                                raw_d7 = rec.get("desc7") or rec.get("date7") or rec.get("desc5") or rec.get("date5") or rec.get("nội dung đợt 7") or rec.get("ngày đợt 7")
                                fallback_tag7 = "«nội dung đợt 5»" if (lan == 4 and m_len == 5) else ("«nội dung đợt 4»" if (lan == 3 and m_len == 4) else "«nội dung đợt 7»")
                                d7_s = _fmt_sc(raw_d7, fallback_tag7)
                                a7_n = _parse_num(rec.get("số tiền đợt 7") or rec.get("số tiền đợt 5"))
                                a7_no_kp = a7_n if is_hl_nl else (a7_n - maint_num if a7_n > maint_num else a7_n)
                                amt_tag7 = "«số_tiền_đợt_5»" if (lan == 4 and m_len == 5) else ("«số_tiền_đợt_4»" if (lan == 3 and m_len == 4) else "«số_tiền_đợt_7_KPBT»")
                                amt_words_tag7 = "«bằng_chữ_số_tiền_đợt_5»" if (lan == 4 and m_len == 5) else ("«bằng_chữ_số_tiền_đợt_4»" if (lan == 3 and m_len == 4) else "«bằng_chữ_số_tiền_đợt_7_KPBT»")
                                a7_str = _fmt_money(a7_no_kp) if a7_no_kp else amt_tag7
                                a7_w = _read_num_vn(a7_no_kp) if a7_no_kp else amt_words_tag7
                                conn7 = _get_money_connector(d7_s)
                                lines_kh.append(f"{label_prefix} {d7_s}{conn7}{a7_str} đồng (Bằng chữ: {a7_w})")
                            else:
                                raw_d = rec.get(f"desc{col}") or rec.get(f"date{col}") or rec.get(f"nội dung đợt {col}") or rec.get(f"ngày đợt {col}")
                                fallback_tag = f"«nội dung đợt {col}»"
                                d_s = _fmt_sc(raw_d, fallback_tag)
                                a_n = _parse_num(rec.get(f"số tiền đợt {col}"))
                                a_str = _fmt_money(a_n) if a_n else f"«số_tiền_đợt_{col}»"
                                a_w = _read_num_vn(a_n) if a_n else f"«bằng_chữ_số_tiền_đợt_{col}»"
                                conn = _get_money_connector(d_s)
                                lines_kh.append(f"{label_prefix} {d_s}{conn}{a_str} đồng (Bằng chữ: {a_w})")

                        for idx in range(len(lines_kh)):
                            lines_kh[idx] = re.sub(r'[.;\s]+$', '', lines_kh[idx])
                            if idx < len(lines_kh) - 1:
                                lines_kh[idx] += ";"
                            else:
                                lines_kh[idx] += "."

                        kh_gn_docx = "\n".join(lines_kh)

                    rec["kế_hoạch_giải_ngân"] = kh_gn_docx
                    rec["kế hoạch giải ngân"] = kh_gn_docx
                    rec["Kế_hoạch_giải_ngân"] = kh_gn_docx
                    rec["Kế hoạch giải ngân"] = kh_gn_docx
                    rec["ke_hoach_giai_ngan"] = kh_gn_docx
                    rec["TỔNG_HỢP_VỐN_TỰ_CÓ"] = th_vtc
                    rec["tonghopvontuco"] = th_vtc

                    # Các đợt thông báo giải ngân (cả đợt 1 và các đợt 2..10)
                    can_ho_r = str(rec.get("Căn hộ số") or rec.get("canHoSo") or rec.get("căn hộ số") or "").strip()
                    tb_list_r = rec.get("thongBaoGiaiNgan") if isinstance(rec.get("thongBaoGiaiNgan"), list) else []
                    tb_map_r = {}
                    vv_map_r = {}
                    for it_r in tb_list_r:
                        if isinstance(it_r, dict):
                            d_idx_r = it_r.get("dot")
                            if d_idx_r is not None:
                                tb_map_r[d_idx_r] = str(it_r.get("soTB") or "").strip()
                                vv_map_r[d_idx_r] = str(it_r.get("vv") or "").strip()

                    tb_1_r = str(
                        rec.get("số_tb__ngày_tb") or rec.get("Số_tb__ngày_tb") or rec.get("số tb - ngày tb") or
                        rec.get("soTBNgayTB") or rec.get("soThongBao") or rec.get("số_tb__ngày_tb1") or
                        rec.get("số tb - ngày tb1") or rec.get("soTBNgayTB1") or tb_map_r.get(1) or ""
                    ).strip()
                    vv_1_r = str(
                        rec.get("vv_thanh_toán") or rec.get("Vv_thanh_toán") or rec.get("v/v thanh toán") or
                        rec.get("vvThanhToan") or rec.get("vv_thanh_toán1") or rec.get("v/v thanh toán1") or
                        rec.get("vvThanhToan1") or vv_map_r.get(1) or ""
                    ).strip()
                    if not vv_1_r and can_ho_r:
                        vv_1_r = f"Thanh toán tiền mua căn hộ số {can_ho_r}"

                    rec["số_tb__ngày_tb"] = tb_1_r
                    rec["Số_tb__ngày_tb"] = tb_1_r
                    rec["so_tb__ngay_tb"] = tb_1_r
                    rec["số tb - ngày tb"] = tb_1_r
                    rec["vv_thanh_toán"] = vv_1_r
                    rec["Vv_thanh_toán"] = vv_1_r
                    rec["vv_thanh_toan"] = vv_1_r
                    rec["v/v thanh toán"] = vv_1_r
                    rec["số_tb__ngày_tb1"] = tb_1_r
                    rec["Số_tb__ngày_tb1"] = tb_1_r
                    rec["số_tb__ngày_tb_1"] = tb_1_r
                    rec["vv_thanh_toán1"] = vv_1_r
                    rec["Vv_thanh_toán1"] = vv_1_r
                    rec["vv_thanh_toán_1"] = vv_1_r

                    for k in range(2, 11):
                        tb_k = str(
                            rec.get(f"số tb - ngày tb{k}") or rec.get(f"soTBNgayTB{k}") or
                            rec.get(f"số tb - ngày tb lần {k}") or rec.get(f"số_tb__ngày_tb{k}") or
                            rec.get(f"Số_tb__ngày_tb{k}") or rec.get(f"số_tb__ngày_tb_lần_{k}") or
                            rec.get(f"Số_tb__ngày_tb_lần_{k}") or tb_map_r.get(k) or ""
                        ).strip()
                        vv_k = str(
                            rec.get(f"v/v thanh toán{k}") or rec.get(f"vvThanhToan{k}") or
                            rec.get(f"v/v thanh toán lần {k}") or rec.get(f"vv_thanh_toán{k}") or
                            rec.get(f"Vv_thanh_toán{k}") or rec.get(f"vv_thanh_toán_lần_{k}") or
                            rec.get(f"Vv_thanh_toán_lần_{k}") or vv_map_r.get(k) or ""
                        ).strip()
                        if not vv_k and tb_k and can_ho_r:
                            vv_k = f"Thanh toán tiền mua căn hộ số {can_ho_r} đợt {k}"

                        rec[f"số_tb__ngày_tb{k}"] = tb_k
                        rec[f"Số_tb__ngày_tb{k}"] = tb_k
                        rec[f"số_tb__ngày_tb_{k}"] = tb_k
                        rec[f"số_tb__ngày_tb_lần_{k}"] = tb_k
                        rec[f"Số_tb__ngày_tb_lần_{k}"] = tb_k
                        rec[f"so_tb__ngay_tb{k}"] = tb_k
                        rec[f"vv_thanh_toán{k}"] = vv_k
                        rec[f"Vv_thanh_toán{k}"] = vv_k
                        rec[f"vv_thanh_toán_{k}"] = vv_k
                        rec[f"vv_thanh_toán_lần_{k}"] = vv_k
                        rec[f"Vv_thanh_toán_lần_{k}"] = vv_k
                        rec[f"vv_thanh_toan{k}"] = vv_k

                    # Ngày làm đơn vay vốn & ngày/tháng/năm nhận hs vay vốn
                    raw_nl_rec = str(rec.get("Ngày làm đơn vay vốn") or rec.get("Ngày_làm_đơn_vay_vốn") or rec.get("ngayTiepNhan") or rec.get("ngayLamDon") or rec.get("ngayLamDonVay") or rec.get("ngayNhanHs") or rec.get("Ngày tiếp nhận") or "").strip()
                    full_nl_r, d_nl_r, m_nl_r, y_nl_r = format_admin_date(raw_nl_rec)
                    if not d_nl_r and rec.get("ngày nhận hs vay vốn"):
                        d_nl_r = str(rec.get("ngày nhận hs vay vốn")).strip()
                    if not m_nl_r and rec.get("tháng nhận hs vay vốn"):
                        m_nl_r = str(rec.get("tháng nhận hs vay vốn")).strip()
                    if not y_nl_r and rec.get("năm nhận hs vay vốn"):
                        y_nl_r = str(rec.get("năm nhận hs vay vốn")).strip()
                    if not full_nl_r and d_nl_r and m_nl_r and y_nl_r:
                        full_nl_r = f"ngày {d_nl_r.zfill(2)} tháng {m_nl_r.zfill(2)} năm {y_nl_r}"

                    if full_nl_r:
                        rec["Ngày làm đơn vay vốn"] = full_nl_r
                        rec["Ngày_làm_đơn_vay_vốn"] = full_nl_r
                        rec["ngay_lam_don_vay_von"] = full_nl_r
                        rec["ngày_làm_đơn_vay_vốn"] = full_nl_r
                        rec["ngày làm đơn vay vốn"] = full_nl_r
                        rec["NGÀY_LÀM_ĐƠN_VAY_VỐN"] = full_nl_r
                    if d_nl_r:
                        rec["ngày nhận hs vay vốn"] = d_nl_r
                        rec["ngày_nhận_hs_vay_vốn"] = d_nl_r
                        rec["ngay_nhan_hs_vay_von"] = d_nl_r
                    if m_nl_r:
                        rec["tháng nhận hs vay vốn"] = m_nl_r
                        rec["tháng_nhận_hs_vay_vốn"] = m_nl_r
                        rec["thang_nhan_hs_vay_von"] = m_nl_r
                    if y_nl_r:
                        rec["năm nhận hs vay vốn"] = y_nl_r
                        rec["năm_nhận_hs_vay_vốn"] = y_nl_r
                        rec["nam_nhan_hs_vay_von"] = y_nl_r

                    # 3 trường Số HĐTD, Số HĐTC, Số BB xác định đấu giá
                    cccd_clean_rec = str(rec.get("CCCD") or rec.get("cccd") or "").replace(" ", "").strip()
                    nam_hs_rec = str(rec.get("namNhanHs") or rec.get("năm nhận hs vay vốn") or "").strip()
                    if not nam_hs_rec:
                        raw_dt_rec = str(rec.get("ngayTiepNhan") or rec.get("ngayNhanHs") or rec.get("ngayVay") or rec.get("ngày vay") or "").strip()
                        m_y_rec = re.search(r'\b(20\d\d)\b', raw_dt_rec)
                        nam_hs_rec = m_y_rec.group(1) if m_y_rec else str(datetime.now().year)

                    so_hdtd_rec = str(rec.get("soHDTD") or rec.get("Số hợp đồng tín dụng") or rec.get("số_hđtd") or "").strip()
                    if not so_hdtd_rec and cccd_clean_rec:
                        so_hdtd_rec = f"{cccd_clean_rec}/{nam_hs_rec}/HĐTD"

                    so_hdtc_rec = str(rec.get("soHDTC") or rec.get("Số hợp đồng thế chấp") or rec.get("số_hđtc") or "").strip()
                    if not so_hdtc_rec and cccd_clean_rec:
                        so_hdtc_rec = f"{cccd_clean_rec}/{nam_hs_rec}/HĐTTC"

                    so_bbxd_rec = str(rec.get("soBienBanXacDinh") or rec.get("soBBXD") or rec.get("soBBDH") or rec.get("Số biên bản xác định đấu giá") or rec.get("số_biên_bản_xác_định") or rec.get("so_bbdh") or "").strip()
                    if not so_bbxd_rec and cccd_clean_rec:
                        so_bbxd_rec = f"{cccd_clean_rec}/{nam_hs_rec}/BBĐH"

                    rec["Số hợp đồng tín dụng"] = so_hdtd_rec
                    rec["số_hđtd"] = so_hdtd_rec
                    rec["so_hdtd"] = so_hdtd_rec
                    rec["số hợp đồng tín dụng"] = so_hdtd_rec
                    rec["Số_HĐTD"] = so_hdtd_rec
                    rec["số_hợp_đồng_tín_dụng"] = so_hdtd_rec

                    rec["Số hợp đồng thế chấp"] = so_hdtc_rec
                    rec["số_hđtc"] = so_hdtc_rec
                    rec["so_hdtc"] = so_hdtc_rec
                    rec["số hợp đồng thế chấp"] = so_hdtc_rec
                    rec["Số_HĐTC"] = so_hdtc_rec
                    rec["số_hợp_đồng_thế_chấp"] = so_hdtc_rec

                    rec["Số biên bản xác định đấu giá"] = so_bbxd_rec
                    rec["số_biên_bản_xác_định"] = so_bbxd_rec
                    rec["so_bbdh"] = so_bbxd_rec
                    rec["so_bbxd"] = so_bbxd_rec
                    rec["số_biên_bản_đấu_giá"] = so_bbxd_rec
                    rec["số biên bản xác định"] = so_bbxd_rec
                    rec["Số_biên_bản_xác_định"] = so_bbxd_rec
                    rec["số_biên_bản_xác_định_đấu_giá"] = so_bbxd_rec

                    # 15. BIÊN BẢN KIỂM TRA (MẪU 11/NƠXH - MẪU 17)
                    bbkt_ngay_rec = str(rec.get("Ngày bb kiểm tra") or rec.get("ngayBBKT") or rec.get("Ngày_tháng_năm_BBKT") or "").strip()
                    bbkt_nam_rec = str(rec.get("Năm tài chính") or rec.get("namTaiChinh") or rec.get("Năm_HĐTD") or "").strip() or nam_hs_rec

                    bbkt_tong_gn_rec = _parse_m(rec.get("Tổng số tiền đã giải ngân") or rec.get("tongTienDaGiaiNganBBKT") or rec.get("Vốn vay") or rec.get("vonVay") or 0)
                    if bbkt_tong_gn_rec <= 0 and rec.get("amt1"):
                        bbkt_tong_gn_rec = _parse_m(rec.get("amt1"))

                    bbkt_goc_tra_rec = _parse_m(rec.get("Lũy kế số nợ gốc đã trả") or rec.get("luyKeGocDaTraBBKT") or 0)

                    if rec.get("Dư nợ vay") is not None and str(rec.get("Dư nợ vay")).strip() != "":
                        bbkt_du_no_rec = _parse_m(rec.get("Dư nợ vay"))
                    elif rec.get("duNoVayBBKT") is not None and str(rec.get("duNoVayBBKT")).strip() != "":
                        bbkt_du_no_rec = _parse_m(rec.get("duNoVayBBKT"))
                    else:
                        bbkt_du_no_rec = max(0.0, bbkt_tong_gn_rec - bbkt_goc_tra_rec)

                    bbkt_chu_du_no_rec = str(rec.get("bằng chữ dư nợ vay") or rec.get("bangChuDuNoVayBBKT") or rec.get("Bằng_chữ_Lũy_kế_số_nợ_gốc_đã_trả_BBKT1") or "").strip()
                    if not bbkt_chu_du_no_rec and bbkt_du_no_rec > 0:
                        bbkt_chu_du_no_rec = _word_m(bbkt_du_no_rec)

                    bbkt_gttt_rec = str(rec.get("giá trị thị trường") or rec.get("giaTriThiTruongBBKT") or rec.get("Giá_trị_hiện_tại_của_TSTC_BBKT") or "").strip()

                    rec["Ngày bb kiểm tra"] = bbkt_ngay_rec
                    rec["Ngày_tháng_năm_BBKT"] = bbkt_ngay_rec
                    rec["ngày_tháng_năm_bbkt"] = bbkt_ngay_rec
                    rec["ngày_bb_kiểm_tra"] = bbkt_ngay_rec
                    rec["ngay_bb_kiem_tra"] = bbkt_ngay_rec

                    rec["Năm tài chính"] = bbkt_nam_rec
                    rec["Năm_HĐTD"] = bbkt_nam_rec
                    rec["năm_hđtd"] = bbkt_nam_rec
                    rec["nam_hdtd"] = bbkt_nam_rec
                    rec["năm_tài_chính"] = bbkt_nam_rec
                    rec["nam_tai_chinh"] = bbkt_nam_rec

                    str_tong_gn = f"{int(bbkt_tong_gn_rec):,}".replace(",", ".") if bbkt_tong_gn_rec > 0 else ""
                    rec["Tổng số tiền đã giải ngân"] = str_tong_gn
                    rec["tổng_số_tiền_đã_giải_ngân"] = str_tong_gn
                    rec["tong_so_tien_da_giai_ngan"] = str_tong_gn

                    str_goc_tra = f"{int(bbkt_goc_tra_rec):,}".replace(",", ".") if bbkt_goc_tra_rec > 0 else "0"
                    rec["Lũy kế số nợ gốc đã trả"] = str_goc_tra
                    rec["Lũy_kế_số_nợ_gốc_đã_trả_BBKT"] = str_goc_tra
                    rec["lũy_kế_số_nợ_gốc_đã_trả_bbkt"] = str_goc_tra
                    rec["lũy_kế_số_nợ_gốc_đã_trả"] = str_goc_tra
                    rec["luy_ke_so_no_goc_da_tra"] = str_goc_tra

                    str_du_no = f"{int(bbkt_du_no_rec):,}".replace(",", ".") if bbkt_du_no_rec > 0 else "0"
                    rec["Dư nợ vay"] = str_du_no
                    rec["Dư_nợ_vay_hiện_tại_BBKT"] = str_du_no
                    rec["dư_nợ_vay_hiện_tại_bbkt"] = str_du_no
                    rec["dư_nợ_vay"] = str_du_no
                    rec["du_no_vay"] = str_du_no

                    rec["bằng chữ dư nợ vay"] = bbkt_chu_du_no_rec
                    rec["Bằng_chữ_Lũy_kế_số_nợ_gốc_đã_trả_BBKT1"] = bbkt_chu_du_no_rec
                    rec["bằng_chữ_lũy_kế_số_nợ_gốc_đã_trả_bbkt1"] = bbkt_chu_du_no_rec
                    rec["bang_chu_du_no_vay"] = bbkt_chu_du_no_rec

                    rec["giá trị thị trường"] = bbkt_gttt_rec
                    rec["Giá_trị_hiện_tại_của_TSTC_BBKT"] = bbkt_gttt_rec
                    rec["giá_trị_hiện_tại_của_tstc_bbkt"] = bbkt_gttt_rec
                    rec["gia_tri_thi_truong"] = bbkt_gttt_rec

                    if not rec.get("Số_Căn_hộ_BBKT"):
                        rec["Số_Căn_hộ_BBKT"] = rec.get("Căn hộ số") or rec.get("căn hộ số") or rec.get("canHoSo") or ""

                    # 16. TỜ TRÌNH ĐỀ NGHỊ GIẢI NGÂN (MẪU 16/NƠXH - MẪU 18)
                    raw_ttr_date_rec = str(rec.get("Ngày tờ trình") or rec.get("ngày tờ trình") or rec.get("ngayToTrinh") or "").strip()
                    ttr_hoa_rec = ""
                    ttr_thuong_rec = ""
                    if raw_ttr_date_rec:
                        m_full = re.search(r'(?:[Nn]gày\s*)?(\d{1,2})\s*tháng\s*(\d{1,2})\s*năm\s*(\d{4})', raw_ttr_date_rec, re.IGNORECASE)
                        if m_full:
                            d, m, y = m_full.group(1), m_full.group(2), m_full.group(3)
                            ttr_hoa_rec = f"Ngày {d} tháng {m} năm {y}"
                            ttr_thuong_rec = f"ngày {d} tháng {m} năm {y}"
                        elif re.match(r'^\d{4}-\d{2}-\d{2}$', raw_ttr_date_rec):
                            parts = raw_ttr_date_rec.split('-')
                            d, m, y = str(int(parts[2])), str(int(parts[1])), parts[0]
                            ttr_hoa_rec = f"Ngày {d} tháng {m} năm {y}"
                            ttr_thuong_rec = f"ngày {d} tháng {m} năm {y}"
                        elif re.match(r'^\d{1,2}/\d{1,2}/\d{4}$', raw_ttr_date_rec):
                            parts = raw_ttr_date_rec.split('/')
                            d, m, y = str(int(parts[0])), str(int(parts[1])), parts[2]
                            ttr_hoa_rec = f"Ngày {d} tháng {m} năm {y}"
                            ttr_thuong_rec = f"ngày {d} tháng {m} năm {y}"
                        else:
                            ttr_hoa_rec = raw_ttr_date_rec
                            ttr_thuong_rec = raw_ttr_date_rec[:1].lower() + raw_ttr_date_rec[1:] if raw_ttr_date_rec else ""

                    # Tổng giải ngân TTr
                    ttr_tong_da_gn_rec = 0
                    if rec.get("Tổng số tiền đã giải ngân TTr") is not None and str(rec.get("Tổng số tiền đã giải ngân TTr")).strip() != "":
                        ttr_tong_da_gn_rec = _parse_m(rec.get("Tổng số tiền đã giải ngân TTr"))
                    elif rec.get("tongTienDaGiaiNganTTr") is not None and str(rec.get("tongTienDaGiaiNganTTr")).strip() != "":
                        ttr_tong_da_gn_rec = _parse_m(rec.get("tongTienDaGiaiNganTTr"))
                    else:
                        amt1_val_rec = _parse_m(rec.get("số tiền GN đầu tiên") or rec.get("số tiền đợt 1") or rec.get("amt1") or 0)
                        running_tot = amt1_val_rec
                        for bi in range(2, 9):
                            has_dot = bool(rec.get(f"tb{bi}") or rec.get(f"date{bi}") or rec.get(f"số thông báo đợt {bi}") or rec.get(f"số tb - ngày tb{bi}"))
                            dot_amt = _parse_m(rec.get(f"số tiền đợt {bi}") or rec.get(f"amt{bi}") or rec.get(f"soTienTB{bi}"))
                            if has_dot and dot_amt > 0:
                                running_tot += dot_amt
                        ttr_tong_da_gn_rec = running_tot

                    # Tổng còn phải giải ngân
                    von_vay_rec = _parse_m(rec.get("Vốn vay") or rec.get("Vốn_vay") or rec.get("vonVay") or 0)
                    if rec.get("Tổng số tiền còn phải giải ngân") is not None and str(rec.get("Tổng số tiền còn phải giải ngân")).strip() != "":
                        ttr_con_phai_gn_rec = _parse_m(rec.get("Tổng số tiền còn phải giải ngân"))
                    elif rec.get("tongTienConPhaiGNTTr") is not None and str(rec.get("tongTienConPhaiGNTTr")).strip() != "":
                        ttr_con_phai_gn_rec = _parse_m(rec.get("tongTienConPhaiGNTTr"))
                    else:
                        ttr_con_phai_gn_rec = max(0, von_vay_rec - ttr_tong_da_gn_rec) if von_vay_rec > 0 else 0

                    s_ttr_tong_da_gn_rec = f"{int(ttr_tong_da_gn_rec):,}".replace(",", ".") if ttr_tong_da_gn_rec > 0 else "0"
                    s_ttr_con_phai_gn_rec = f"{int(ttr_con_phai_gn_rec):,}".replace(",", ".") if ttr_con_phai_gn_rec > 0 else "0"

                    rec["Ngày tờ trình"] = ttr_hoa_rec
                    rec["ngày tờ trình"] = ttr_thuong_rec
                    rec["Tổng số tiền đã giải ngân TTr"] = s_ttr_tong_da_gn_rec
                    rec["Tổng số tiền còn phải giải ngân"] = s_ttr_con_phai_gn_rec

                    rec["Ngày_tờ_trình"] = ttr_hoa_rec
                    rec["ngày_tờ_trình"] = ttr_thuong_rec
                    rec["Ngay_to_trinh"] = ttr_hoa_rec
                    rec["ngay_to_trinh"] = ttr_thuong_rec
                    rec["Dư_nợ_trước_ngày_nhận_nợ_L2"] = s_ttr_tong_da_gn_rec
                    rec["dư_nợ_trước_ngày_nhận_nợ_l2"] = s_ttr_tong_da_gn_rec
                    rec["Tổng_số_tiền_còn_phải_giải_ngân"] = s_ttr_con_phai_gn_rec
                    rec["tổng_số_tiền_còn_phải_giải_ngân"] = s_ttr_con_phai_gn_rec
                    rec["tong_so_tien_con_phai_giai_ngan"] = s_ttr_con_phai_gn_rec
                except Exception:
                    pass

                records.append(rec)
    except Exception as e:
        log_callback(f"❌ Lỗi khi đọc file Excel: {e}")
        return False, f"Lỗi đọc Excel: {e}"

    if not records:
        log_callback("⚠ Không có dòng dữ liệu nào trong file Excel.")
        return False, "File Excel không có dữ liệu."

    log_callback(f"👥 Tìm thấy {len(records)} hồ sơ khách hàng cần xuất biểu mẫu.")

    templates = sorted(glob.glob(os.path.join(TEMPLATE_DIR, "*.docx")))
    templates = [t for t in templates if not os.path.basename(t).startswith("~$")]

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    total_files_exported = 0

    # Lọc danh sách hồ sơ nếu người dùng chỉ định (chọn 1 hoặc nhiều hồ sơ)
    if selected_indices is not None and len(selected_indices) > 0:
        records_to_process = [(i, records[i]) for i in selected_indices if 0 <= i < len(records)]
    else:
        records_to_process = [(i, r) for i, r in enumerate(records)]

    if not records_to_process:
        log_callback("⚠ Không có hồ sơ nào được chọn để xuất.")
        return False, "Không có hồ sơ nào được chọn."

    log_callback(f"👥 Đang tiến hành xuất biểu mẫu cho {len(records_to_process)} hồ sơ được chọn...")

    for order_idx, (real_idx, rec) in enumerate(records_to_process, 1):
        ho_ten = rec.get("Họ và tên KH hoa") or rec.get("Họ và tên KH thường") or f"KhachHang_{order_idx}"
        ho_ten_clean = re.sub(r'[\s/\\?%*:|"<>]+', '_', str(ho_ten).strip())
        can_ho = rec.get("Căn hộ số") or ""
        can_ho_clean = re.sub(r'[\s/\\?%*:|"<>]+', '_', str(can_ho).strip())
        
        folder_name = f"{order_idx:02d}_{ho_ten_clean}" + (f"_Can_{can_ho_clean}" if can_ho_clean else "")
        cust_dir = os.path.join(OUTPUT_DIR, folder_name)
        os.makedirs(cust_dir, exist_ok=True)
        # Dọn dẹp file Word cũ trong thư mục hồ sơ này để không bị lẫn mẫu từ các lần xuất trước
        for existing_f in glob.glob(os.path.join(cust_dir, "*.docx")):
            try:
                os.remove(existing_f)
            except Exception:
                pass

        # 1. Xác định thư mục mẫu cụ thể từ Template/
        case_folder = ""
        if folder_mapping and real_idx in folder_mapping:
            case_folder = folder_mapping[real_idx]
        if not case_folder:
            case_folder = rec.get("_resolved_folder") or resolve_template_folder(rec)

        if marital_mapping and real_idx in marital_mapping:
            chosen_marital = str(marital_mapping[real_idx] or "").strip()
            if chosen_marital:
                rec["tinhTrangHonNhan"] = chosen_marital
                rec["Tình trạng hôn nhân"] = chosen_marital

        rec["_resolved_folder"] = case_folder
        rec["case_folder"] = case_folder

        if is_single_profile(rec):
            rec["tinhTrangHonNhan"] = "Độc thân"
            rec["Tình trạng hôn nhân"] = "Độc thân"
            for k_sp in [
                "Họ và tên Vợ/chồng", "Họ và tên vợ/chồng in hoa", "CCCD Vợ/chồng",
                "Thu nhập vợ/chồng", "nghề nghiệp vợ/chồng", "nơi làm việc vợ/chồng",
                "quan hệ với người vay", "tên CCCD/CC vợ/chồng", "tên CCCD/CC vợ/chồng viết đầy đủ",
                "Ngày,tháng,năm sinh vợ/chồng", "Ngày cấp CCCD vợ/chồng", "Nơi cấp CCCD vợ/chồng",
                "Địa chỉ cư trú vợ/chồng", "nơi ở hiện tại vợ/chồng", "dân tộc vợ/chồng",
                "Giới tính vợ/chồng", "DX vợ/chồng HOA", "dx vợ/chồng thường", "Số điện thoại vợ/chồng",
                "vcHoTen", "vcCCCD", "vcDanhXung", "vcThuNhap", "vcNgheNghiep", "vcNoiCongTac", "vcSDT"
            ]:
                rec[k_sp] = ""
            rec["Thu nhập vợ/chồng"] = ""
            if rec.get("Thu nhập người vay"):
                rec["Tổng thu nhập"] = rec.get("Thu nhập người vay")

        # Tìm đường dẫn thư mục mẫu case
        case_dir = get_case_folder_path(case_folder)
        case_docx = []
        if case_dir:
            all_case_files = [f for f in os.listdir(case_dir) if f.lower().endswith('.docx') and not f.startswith('~$')]
            all_case_files.sort(key=natural_sort_key)
            
            cust_is_single = is_single_profile(rec)
            for f in all_case_files:
                f_norm = strip_vn_accents(f.lower())
                if cust_is_single:
                    # Nếu độc thân: nếu có file độc thân riêng (như Mẫu 9 01a độc thân) thì loại bỏ file có gia đình
                    if "co gia dinh" in f_norm and any(
                        "doc than" in strip_vn_accents(other_f.lower()) and get_form_key(other_f) == get_form_key(f)
                        for other_f in all_case_files
                    ):
                        continue
                    if "mau so 05a" in f_norm and "co gia dinh" in f_norm:
                        continue
                else:
                    if "doc than" in f_norm:
                        continue
                case_docx.append(f)

        if case_docx:
            log_callback(f"\n[{order_idx}/{len(records_to_process)}] Hồ sơ: {ho_ten}")
            log_callback(f"   📁 Thư mục mẫu: {case_folder} (Có {len(case_docx)} biểu mẫu trong nhóm):")
            
            # Lọc biểu mẫu tùy chọn nếu người dùng đã chỉ định danh sách mẫu cần xuất
            if selected_templates is not None:
                sel_set = set(selected_templates)
                selected_keys = {get_form_key(s) for s in selected_templates}
                filtered_docx = [f for f in case_docx if (f in sel_set or get_form_key(f) in selected_keys)]
                case_docx = filtered_docx
                log_callback(f"   🎯 Tùy chọn xuất: Xuất {len(case_docx)} biểu mẫu đã chọn: {[f[:35] for f in case_docx]}")

            templates_to_use = [(os.path.join(case_dir, f), f) for f in case_docx]
        else:
            cust_proj, cust_marital = determine_customer_marital_and_project(rec)
            marital_label = "Độc thân / Ly hôn" if cust_marital == "single" else "Có gia đình"
            proj_label = cust_proj if cust_proj != "all" else "Tất cả dự án"
            log_callback(f"\n[{order_idx}/{len(records_to_process)}] Hồ sơ: {ho_ten} -> Nhóm: [{proj_label} • {marital_label}]")
            
            # Fallback: Lọc từ thư mục template gốc
            matched_templates = []
            for t in templates:
                fname = os.path.basename(t)
                code, t_proj, t_marital, title = get_template_info(fname)
                if t_proj != "all" and cust_proj != "all":
                    if t_proj.lower().strip() != cust_proj.lower().strip():
                        continue
                if t_marital != "all":
                    if t_marital != cust_marital:
                        continue
                matched_templates.append((t, fname))
            matched_templates.sort(key=lambda x: natural_sort_key(x[1]))
            if selected_templates is not None:
                sel_set = set(selected_templates)
                selected_keys = {get_form_key(s) for s in selected_templates}
                matched_templates = [t for t in matched_templates if (t[1] in sel_set or get_form_key(t[1]) in selected_keys)]
            templates_to_use = matched_templates
            log_callback(f"   📋 Đã khớp {len(templates_to_use)}/{len(templates)} mẫu biểu:")

        if not templates_to_use:
            log_callback(f"   ⚠ Không tìm thấy biểu mẫu Word nào cho hồ sơ {ho_ten}!")
            continue

        for item_idx, (t, fname) in enumerate(templates_to_use, 1):
            t_base = os.path.splitext(fname)[0]
            out_file = os.path.join(cust_dir, f"{t_base}_{ho_ten_clean}.docx")
            try:
                fill_docx_template(t, out_file, rec)
                log_callback(f"   [{item_idx:02d}] ✓ {os.path.basename(out_file)}")
                total_files_exported += 1
            except Exception as e:
                log_callback(f"   [{item_idx:02d}] ✗ Lỗi {t_base}: {e}")

    log_callback("\n" + "="*50)
    log_callback("✅ ĐÃ HOÀN THÀNH XUẤT FILE CHO TẤT CẢ HỒ SƠ!")
    log_callback(f"📂 Thư mục kết quả: {OUTPUT_DIR}")
    return True, f"Đã xuất thành công {total_files_exported} file Word cho {len(records_to_process)} bộ hồ sơ vào thư mục output_word/"

# ==============================================================================
# PHẦN 3: LOCAL HTTP SERVER & API
# ==============================================================================
httpd_server = None
server_port = 8765

class QuietHandler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=BASE_DIR, **kwargs)

    def log_message(self, format, *args):
        pass

    def end_headers(self):
        self.send_header('Access-Control-Allow-Origin', '*')
        self.send_header('Access-Control-Allow-Methods', 'GET, POST, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type')
        self.send_header('Cache-Control', 'no-cache, no-store, must-revalidate, max-age=0')
        self.send_header('Pragma', 'no-cache')
        self.send_header('Expires', '0')
        super().end_headers()

    def do_OPTIONS(self):
        self.send_response(200)
        self.end_headers()

    def do_GET(self):
        if self.path == "/api/records":
            self.send_response(200)
            self.send_header('Content-Type', 'application/json; charset=utf-8')
            self.end_headers()
            if os.path.exists(DB_FILE):
                try:
                    with open(DB_FILE, 'r', encoding='utf-8') as f:
                        data = json.load(f)
                except Exception:
                    data = []
                resp = {"fileFound": True, "records": data, "baseDir": BASE_DIR}
            else:
                resp = {"fileFound": False, "records": [], "baseDir": BASE_DIR}
            self.wfile.write(json.dumps(resp, ensure_ascii=False).encode('utf-8'))
            return

        if self.path == "/api/check_update":
            self.send_response(200)
            self.send_header('Content-Type', 'application/json; charset=utf-8')
            self.end_headers()
            res_update = check_github_release()
            self.wfile.write(json.dumps(res_update, ensure_ascii=False).encode('utf-8'))
            return

        if self.path == "/api/update_progress":
            self.send_response(200)
            self.send_header('Content-Type', 'application/json; charset=utf-8')
            self.end_headers()
            self.wfile.write(json.dumps(update_status, ensure_ascii=False).encode('utf-8'))
            return

        super().do_GET()

    def do_POST(self):
        if self.path == "/api/records":
            length = int(self.headers.get('Content-Length', 0))
            body = self.rfile.read(length)
            try:
                payload = json.loads(body.decode('utf-8'))
                if isinstance(payload, dict) and "records" in payload:
                    data = payload["records"]
                    rows206 = payload.get("rows206", [])
                else:
                    data = payload
                    rows206 = []

                # Lưu vào database_records.json
                safe_unhide_file(DB_FILE)
                with open(DB_FILE, 'w', encoding='utf-8') as f:
                    json.dump(data, f, ensure_ascii=False, indent=2)
                safe_rehide_file(DB_FILE)

                # Lưu rows206 cache
                rows206_file = os.path.join(BASE_DIR, "database_records_rows206.json")
                if rows206 and len(rows206) > 0:
                    try:
                        safe_unhide_file(rows206_file)
                        with open(rows206_file, 'w', encoding='utf-8') as rf:
                            json.dump(rows206, rf, ensure_ascii=False, indent=2)
                        safe_rehide_file(rows206_file)
                    except Exception:
                        pass
                else:
                    if os.path.exists(rows206_file):
                        try:
                            safe_unhide_file(rows206_file)
                            with open(rows206_file, 'w', encoding='utf-8') as rf:
                                json.dump([], rf, ensure_ascii=False)
                            safe_rehide_file(rows206_file)
                        except Exception:
                            pass

                # Tự động xuất file Excel đồng bộ DATA_MAIL_MERGE_SYNC.xlsx và cache rows206 trong luồng nền
                sync_file = os.path.join(BASE_DIR, "DATA_MAIL_MERGE_SYNC.xlsx")
                if data and len(data) > 0:
                    threading.Thread(target=sync_all_database_records, args=(data,), daemon=True).start()
                elif len(data) == 0 and os.path.exists(sync_file):
                    try:
                        safe_unhide_file(sync_file)
                        os.remove(sync_file)
                    except Exception:
                        pass

                self.send_response(200)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.end_headers()
                self.wfile.write(json.dumps({"success": True}).encode('utf-8'))
            except Exception as e:
                self.send_response(500)
                self.end_headers()
                self.wfile.write(str(e).encode('utf-8'))
            return

        if self.path == "/api/save_template":
            length = int(self.headers.get('Content-Length', 0))
            body = self.rfile.read(length)
            try:
                req = json.loads(body.decode('utf-8'))
                fname = req.get('filename')
                b64_str = req.get('b64')
                project = req.get('project', 'all')
                marital = req.get('maritalTarget', 'all')
                code = req.get('code', 'Mẫu Word')
                title = req.get('title', '')
                if fname and b64_str:
                    os.makedirs(TEMPLATE_DIR, exist_ok=True)
                    out_path = os.path.join(TEMPLATE_DIR, fname)
                    with open(out_path, 'wb') as f:
                        f.write(base64.b64decode(b64_str))

                    # Lưu metadata
                    meta = load_template_metadata()
                    meta[fname] = {
                        "code": code or "Mẫu Word",
                        "title": title or fname.replace('.docx', ''),
                        "project": project or "all",
                        "maritalTarget": marital or "all"
                    }
                    save_template_metadata(meta)

                    # Tự động đồng bộ ngay lập tức
                    run_sync_templates()
                    self.send_response(200)
                    self.send_header('Content-Type', 'application/json; charset=utf-8')
                    self.end_headers()
                    self.wfile.write(json.dumps({"success": True, "message": f"Đã lưu {fname} vào template/ và đồng bộ xong!"}).encode('utf-8'))
                    return
            except Exception as e:
                self.send_response(500)
                self.end_headers()
                self.wfile.write(str(e).encode('utf-8'))
                return

        if self.path == "/api/save_excel":
            length = int(self.headers.get('Content-Length', 0))
            body = self.rfile.read(length)
            try:
                req = json.loads(body.decode('utf-8'))
                fname = req.get('filename')
                b64_str = req.get('b64')
                if fname and b64_str:
                    out_path = os.path.join(BASE_DIR, fname)
                    with open(out_path, 'wb') as f:
                        f.write(base64.b64decode(b64_str))
                    # Nếu lưu file DATA_MAIL_MERGE, cập nhật luôn DATA_MAIL_MERGE_SYNC.xlsx
                    if "DATA_MAIL_MERGE" in fname.upper() and fname != "DATA_MAIL_MERGE_SYNC.xlsx":
                        sync_file = os.path.join(BASE_DIR, "DATA_MAIL_MERGE_SYNC.xlsx")
                        shutil.copy2(out_path, sync_file)
                    self.send_response(200)
                    self.send_header('Content-Type', 'application/json; charset=utf-8')
                    self.end_headers()
                    self.wfile.write(json.dumps({"success": True, "message": f"Đã lưu {fname} vào thư mục ứng dụng!"}).encode('utf-8'))
                    return
            except Exception as e:
                self.send_response(500)
                self.end_headers()
                self.wfile.write(str(e).encode('utf-8'))
                return

        if self.path == "/api/perform_update":
            length = int(self.headers.get('Content-Length', 0))
            body = self.rfile.read(length) if length > 0 else b'{}'
            try:
                req = json.loads(body.decode('utf-8')) if body else {}
                files_list = req.get('files')
                
                # Chạy luồng nền tải và cập nhật
                threading.Thread(target=download_and_apply_update, args=(files_list,), daemon=True).start()
                self.send_response(200)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.end_headers()
                self.wfile.write(json.dumps({"status": "started", "message": "Bắt đầu cập nhật các tệp từ GitHub..."}).encode('utf-8'))
                return
            except Exception as e:
                self.send_response(500)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.end_headers()
                self.wfile.write(json.dumps({"success": False, "error": str(e)}).encode('utf-8'))
                return

        super().do_POST()

def start_local_server():
    global httpd_server, server_port
    if httpd_server is not None:
        return f"http://127.0.0.1:{server_port}/he-thong-quan-ly-ho-so-demo.html"

    for port in range(8765, 8780):
        try:
            httpd_server = HTTPServer(("127.0.0.1", port), QuietHandler)
            server_port = port
            t = threading.Thread(target=httpd_server.serve_forever, daemon=True)
            t.start()
            return f"http://127.0.0.1:{server_port}/he-thong-quan-ly-ho-so-demo.html"
        except OSError:
            continue
    return f"file:///{HTML_FILE.replace(os.sep, '/')}"

def open_web_app():
    url = start_local_server()
    cache_url = f"{url}?v={int(time.time())}"
    webbrowser.open(cache_url)

# ==============================================================================
# PHẦN 4: GIAO DIỆN ĐỒ HỌA NGƯỜI DÙNG (TKINTER GUI)
# ==============================================================================
# VBSP Official Logo (Base64 PNG 54x56)
VBSP_LOGO_B64 = "iVBORw0KGgoAAAANSUhEUgAAADYAAAA4CAIAAAA5LZz6AAAUBUlEQVR4nMVZCZRdRbWtc6ruffN7PaQzdGfozkgCAgEkyqQhBFAgXyYFFKMCMsR8viB8/YGABD4qKl8Gw4wIhEAIEgQCRsIYCKNISMg8ddJJp+c3v3ur6rhOvQ4iuAx8Fsu7sjL1vVWnzrDP3qcEferHWrLWhloTUeH1be/NuNf2GksUGmuIf/Yp11fiUz8AwhiBUgabejZf/JBZtqU19IZffyrGBVgQCJ9yffxUX5MgEsZYIYHa8pvPvk+9trOmri77+9c3X7wAKyCAnfjvNJEttBYkiq7yuvPniZe3eCk/IB1PxYr3vLLl0kfAoAAO9L/NRGFJINi2/Oaz7qUla/10lIwRnD86nkhlb39566V/QCP5LJ/Cl5/MRHK7CUGWSJMl4DzLvrIp/6dV8XiE60VIEFYaZa1IJRK5W57fNvtxIAyFNdbuTg36bL0IAgSBtYJAYJnWTL+z/O7OEVd9I1fOewYABRAYdGchm0iku29cuu3qxR5KAiOsABKftHw+sYnkkktK8kpi838+APNXd/9miYnBgIuPLxQKkvgI/Bqwu4Bs2k/2XPfnHT/7kwKPhDVuhc/YRGssCiiLDRc9mJ/3VrwmkSC/bfaCaHND7Q8mF/I5CVJaAhIWhCEhwGa8WOc1T+381XOASFYT2c/QRF4dQZVF648eLt7zRioR1zYUClM6um3W/OQBo9PTD8319YFS1XUJhBVkkJJ+qmPO4523vqKk5wz/bEysYgdqbP3Jwuxdr2XiCcudQ1kSNiJiebH1x/Mzxx6QOPmAYk8WJCJHFAQhwyaGCRlrm/WHnvvfRCWtNkT2Y9bNxzGRF+IeJwgJt/xkUe9ty9LxmHaQg+wpIwyIWCzWVdpyxX2DZ3w1ctyEYjaLUqEFwV9ZQ0BSZLTccsmD3Y+9h0oaE1bX3qOhH8NEApc9JAlbZ/0xe/NzmViSnVCtbUFAnP6oQ5mIxtb2bLr8vqGzTlFHTsjn8iAtOAv4T2utJ1MlbJ3x++wz65WKMG59jOrZs4nunBZQ7pjzVN8Nz9XGUtz0QKDltQ2vD4xCYElbUZMRy1tb5zzYcs0p/hebg1wZFDdqEugb5Ez2vGSfbT1/Xvn1HaCQjIE9RXtPJjLUGkC189dLe361NBGPBWhCZL/JaoSAbZRcucICYGDjyXT49Ootv3ls5PXfCicO1NmKQgyl0MDuNkQy4vttxQ0X3KPXdqHCPTbxPZhI1iCq9rnLun+6OOZF2DaLQP2YYlFiYAUDOaCVQggjraEwlqkN5v217Y6nR996XnlMShfKLmctEPveWqHiUW9l98YZ95mOQCJX4ic00X3AfcpYkLLjvrd2/eSRuIwKALLAQSP+ClDqsBI2xirGKBIGCQg99glpa+KZdN8tyzoeemHs7y4oNEWxECiQLm35fIZ0NBmzL27d+MMHROAyiX1JH9+LnB5VCtP3+OqdF81PC08r4BREy6EFKxSU80X55ZaRT86MnXFQtpj3BJMaA0BCCmE06XhNXfcvn+x89p0xd5yXq0UdhggoCJFAWQqtiaUz5uG3tl31GAAyBpEwH9NETjTDMSy8un3rhfMSAQrP8zSXLvdeIRAxzJXk4c3Nt50ZpjF11PjMaQeX+wpCcd0rrvX+OKQTDd2zFxXXto367VlFFWBoiTuPILcSOztZ23vT87vuelVKSdZIsh8N+j8z0RJIDDd2bj/3/nh7hSIRY0EjbwtCWIVBPhCfH9pyx5lqYBw7KpuvWlg7Zf/ISfuVeksKhXMk55ckbcAmI+ntlzwYlotDf/XNHipKEkicl+gapBEiBfHO2Y/mXtqMfLpqsOlfmUic+sJ2VzbPvB/XdGLCR83ghWSRCBFNsUz71DffNV0NTxMRxr1Yh956xcK6Uw9RXxlbzOUUorQg2U8cOitNEuT2Gfd5g1KD/+dr+WzAUgGEcqhuhbE+Rnpt60Xz9dackWg+0sH/bqKLjbBEaKHtsseCpa0qHRGaywad91FiWA7DlkTLnd/xR9VqbVxWkkii2tHXNuehpm9PVoeN1L0564MRUjAfkowAUiZL3qbz7kofNDJ13qRKX5+QnmG2gQBExsp43Huns/XyR1TIjnCp9gETSTBHZQLNRzZSyq7bXu67d3kslbBGs/+E0CAA0QS2WIfNt0z3PzeQtJWMFvy91eTF4mptz/afPdo449jwwCFhX94hNmMTYzwBxUS022ye+buGE74gp47T2SxJ3ho5cQWZ0EvFCwtX7LrjFQ43OcfsptDomhh3ITShlDL3xvb2axYlZZw4aSQQty5AQqOLfjjsxjMTh4ywoQHZf1ASQnIlW0jH4J22rf/32NAfnyLGDcRsQShlkAg0o6khmYjh6vbOx18ZcfOZwfC0qYTsaF4fCECDTqtI+/8+mX97B6C0Vu9GFxdoLkALVoLJhq2XLlA9ZH2J1lgBoWQo8yxktRk8Z1rNtL20NkJhVRLsthK5loyWNTXi1e1tcxcPvfzk0siUKBSVkEy0QUgCMOQpH3IVrzk1/Jenaz9EAwK4kgUJadBEMN4RtF32iC2RcxuHwPw9F8kiqPablrKKS8bJaKeBrbJGCigVKrXfP6Lh3EPJaCW5i7xvoOuwZND1wNAkUmlaur5t3vPDZp9aHuJTKUCUQMSAAKAFaV8JovQJ4+ovnlopFQGYTXLvtF5orUz74bPrd/3uZYXKMXTeCHkXhzKlN3f2zn0hHkty0+N9OcogoVIs4dRRTZcfrZlSSfa8azH9VoKwkjOSEw9ESEEskzaPr+x64o3hs7+ZrRNQ1k7tc1q510gAaK0bLpwijxsbFAuAnuOhoSJJRHEV67vh2XBz1vnXIvdM69iKprYbnsGuAJSUxrJAAWGUsRXSQ1PDfnGqyHiOcrnNqrDlfreWoEIAHLSqYgkpjKXT5Qf+0rtsRePVZ+QiAYRMMISwQBLd0Xj1GDbNOdE0xqFsCdGCASJWtBHPbu7uuPl5R6B4LxTWImDuxXWlJ1d48YjlzAf2PFnfihKVBl05LbpXPVmN6BRx1fv8i4MsMzEakQorFVZPHEklrdRkVF2i7+6Xym9vabr6tDyWQbN65ZRwmcHBtyY6vqHhkuOKpij5AOgYHdPKaCze/fBr5RW7ECQxHiOzj+67X4nmNZN5XgIsgFV+OV+JnbTvgNMmahMAsEVO3u3OQ241VtVFR914hh6dqZQDp/H6+6QMRSaV7p67uNLaNeSKr/dRWXPbE+AIkSNxaK0dMP0L3lGjw3wJgcvbgmG3SeXtKHbcu5ypE1hExPLK9sqzGyEescxCnHcBIAiDIYmmHx8vQEghgXmK0/kf7KGIZHR03yHNt0+vNHq2xGAEwlhAC6iFTaZquq9frAuFpv+eFhQrXDNMRKrY7M4cEUMu+WolzayAuZMj84Yo4kfzT6zQrVlEycmVXbom7Myh5GwlJlzCI2lKhfozJ0U4xFa4KnZRhiqHfd+PIBVpHZ80dMQt07MDhAgCgb4FY2UgCI2QmUi869pF4Pm1Fx2XC7NVBHGx4H5K1qQOHZGctl+llAdUPCtwKsTEpNnSlV26FhmUNeVeWBNh11iGaeZawmgTNqdrv/2FquqrIguLuX9sTf01I5UNw9TkkSNv+nY+gRRqnxAZSwyQMaiiXqb12odlY7zunGOLuUI/GWW10d/rBnzvCF0TgZCLXXCekCQbCW32+bUc98q6rtKK7SoS42EmEFoJAOWgmDpsbGR0Hc9FeAiyG6E+ooSqowdUnjUmfdxejdefllMBGWslVsUND36Q0jay88qFqfHD66YfylAiEQAVJyT35OSBTfEvtuhKuZrwXBBEGIkUVrSZXSUsvLlJtufJ84BXJM0aiISk1OQJDh32INB2V7hABKsr9V/ft/EXJ+dE6IWS0JUeaG7SCqMVaPvFg5EAgJUFfbDmhIeJI/fRpLkVcryYuvNQoLWnsqoNs2+2eqErwqopQCI0ODiVmNTMwvkDmfdPn93gw0KBe6sJB3xv0qCrpuVMFkghTwYcFbfkeVHVGXYuXF7Nww8uYoVIHtRMmQjnPU+FXAIwYwgL727D4roOFxRrXCg9AaKi5fgh3ogapnOMWHt4iEG5ykQUtwpjB848rH7WVwpBQQqQ1qv2IA45etKPfviQznGxlnoYXqONBgC2hBOZuVp23S6Mtue4jfJJ+vuDJoo0N4AHTEr3aKBwNIEN7H8bBBirh1w6Nf1fX86Vskq4kQSwHnC07yNL8rYW6mNeU62saLccAkOoJqRgax+qXFDF/d1bMOhE6tPO/+6Ee7ZRvD9k4lUkOVJBw2YfH//+4blC1od+te/ExYdJtZNyQEpgJsa+diXD+oKzTPodeTSOhlS/q84GmWnuruF/rXD/0RVQrXlGDmbTgqRuufbk6BkH5nNZRB9Yx7pO/uFPwSkFhiGenfYLWVbb3KQ9QEhGLKvw6t0DcV8RotTZU/3XJ5+pCjfqcYSDyEZNy42n+ydODPvyQlXJ84eTmydGYEWFbFeBJKOAqxb+XhgTDkxiMDCJ2iDxfzGVZwSQxTU7qGgcavw/xuhUNZJrzwqM4+jfnu4fM7aUD4VS1TbwodcFoe4o2E2dilWBm3O4gACZ2PA6VHsPJKv7R9husi99j97ZUXmrjUHbfLKJqugvif7ZAoK01tpaNXzumfLQZl0uCiU//D6RBCy8vFG39iqPT1WdtQGJQFFy70asnzTa+GiZalYrUwiFmNXbb1kCoRAKrHE3Af3n/ZfOc49jecgy2XVid04BjfGWW88IPtcY5LkBumf37RwKEYrO+19UmkvA8QDutDzQrI0l9h2ByUNG2ZZ6LGqGMK48Ya3x49HyH1dtufAhyBqUyvCchi8xqvPA9xVtP7slBlz3V97UJTSLsyrIcHAQjDbe8NSou88S+wxyYxEGc6utYMDDtp8vLj+/IRKLWWJM5DsHCbZUiewzJDphEKqmROar+5ZCrVg0SrTMco2wMT9ZvGf5ulN/W1i6QSGAUgLRorVgDKeDy3FHBrTg6yCGPoZnl+2WhOYIk9G8K3dU5CnPgY3NFxzNJJAnByg9ST1B6yWLun+5JCZTho/pQkCoCAJrMsdNhBjwDUWwunPTsTdE+qxRzIar6eRZlCDz5WI4IKomNdZO3bf2kFFyYAZTEfIdt3i/PcMH5JYQhkeI8MHrT8NYI13DdsikbZALMBf0vbCm57Zl4ZutkViC6YrQzv3SJwy0rrQkxz5xoWqMKaspsteA2nMP2TXnzykvxm5w5FgjX14giMzBYxIn7bfjgWUdP3/aRGRmxuHN508WINqu/1PHS6vHXXNG27yX9Ds7mfIqrPvawXVnTLRFvfOuZfln1sgBiYHnHpE4aNj2W5/L/3Gl70eCoAzjB4e1ifIDr0c2dEdVJB5PWIeXLG34/khYpHJYajjraNUUp5B1C2dQw8wjI1NG6mwR+DqCuT1P+jk+srChte4/9tn70fOjYwb5q3bYNe1Vn9k3d9gn14ueoPTUWrG2SyV8s72waca80oqOznte2Tb7YQQovbBu5Yk36bYCre8Nn1ihV+7UG9or77W3HDZhzJUnq8EZT0W0owKuuHhqBApKxWL0qLEN3zvU8uQZ2UQGsqTf8uvT9IQGmy2BG/RXtTb6UVrb1fXAGwJE+th9LEhdqcIQaalUIiElGhWmfnhky71njVtykReD0rvtfS+sbjjmwJGPnDP26R/5Y+pK23v9RExIb/DPTxn/lytGXHjMezNu8YbVJ04/uJwrIJLieQA7xXjSlCs0Ij702pMhoVjloZvpVC+9vTF1zXd+Jxhbp/uKiq/s3JgDrEde170vU9Fkpn1ODIqbzrJDC4BKADYwIDzpF5e8t2Pe661X/kE21qYPHxUd3dD3+FttMxbmV23b/6Efpj7fFBbzGoRqiGBUpaeOTtia7mXvec21rLgImaQC5z4Ui8WMGnbT9MjeDVazBmLbmevz3AKtCaITB7csONdOHl7qy5H1JCgrKjLu02tbO25fFmmpjZ20v+nOcaMH1id8b2pI+aq4dPXOnz7dO/eFCCHJsHHmlOiJE9vvfWnLCde/e85dujeUIqqsqmzsKW3v6Zr/diFtB03dz7yzy0OeKwDIKJDJloqDky13fzd55CirLSlkeUrVm88qqZEeGRMdUztmwQWpWUcHCV3JFYWWJCnqx3b9eknxjbZhl00LsBKu62NlmwuZxYIol8r1M6dMXH7JhNcvL9tw49WLwYuOvnn6XssvHXT1ycFTa7KPvE1JKQG3zVrw9hfnbL3m0QnXfavn1Q25Ba+pdBwt2ED3FUpi8vBRi85PTRltjWZ5Up0bAc9+XcvhhggkJSvMJDbNPqbu+L07bn+xsHiN2NnLRLSdtpx999jHLhwz65S/fvfmmmQDrWz3ozygAUBRExW1EVkTiQ4aaDb1rT3tlsThY4ZddmxswpCe+5ab1h5QVkN56FlfUQcP6/3z2i3n3CN35qNalm3FGiPG1NSfPXXg2YdDQvFoVyoEl4Ku3TH9cBqDt+IyQMeOrYkeMHTo3NP1qo7uJasqL64P1rfbbT1vnXj9qBu+OfLi43Zcvkh054MM9zY0Nn/T0k3PrQ135XNvbxx551nFd9t2XbfY7wt0Wza/rWPYEePyz6ywOohMHhk/YmTfxu7e1VszDbVU4+M+jbXH7lt/wv5yeIqN0IxcbItrMg54Qb3PDHePkXgsIliFWAvkT2gYPOFL4gdf0jtyldbe/IZdhY1tTSdOqjl8XOfCN60kf3R9fMp4Wt1llYfN3tBvfaP2G/unj9rLtHf3zH9def7Qy05MTh6Re2NtdO8mBjhjU/sNabru1GjLwNTYwdFxg4TviK7hqYNQzor3J1ru+RvOteaKW6sc9AAAAABJRU5ErkJggg=="

# ==============================================================================
# HỘP THOẠI LỰA CHỌN HỒ SƠ ĐỂ XUẤT BIỂU MẪU WORD (CHỌN 1 HOẶC TẤT CẢ)
# ==============================================================================
class AddTemplateConfigDialog(tk.Toplevel):
    def __init__(self, parent, filepaths, on_success):
        super().__init__(parent)
        self.parent = parent
        self.filepaths = filepaths
        self.on_success = on_success

        self.title("PHÂN LOẠI & THÊM BIỂU MẪU WORD (.DOCX)")
        win_w, win_h = 760, 540
        self.update_idletasks()
        p_x = parent.winfo_rootx() + (parent.winfo_width() - win_w) // 2
        p_y = parent.winfo_rooty() + (parent.winfo_height() - win_h) // 2
        self.geometry(f"{win_w}x{win_h}+{max(0, p_x)}+{max(0, p_y)}")
        self.minsize(680, 460)
        self.configure(bg="#f8fafc")
        self.transient(parent)
        self.grab_set()

        # Header
        header = tk.Frame(self, bg="#ffffff", padx=18, pady=12, highlightthickness=1, highlightbackground="#e2e8f0")
        header.pack(fill="x")
        tk.Label(
            header, text=f"📋 THÊM & PHÂN LOẠI {len(filepaths)} BIỂU MẪU WORD (.DOCX)",
            font=("Segoe UI", 11, "bold"), fg="#1e293b", bg="#ffffff"
        ).pack(anchor="w")
        tk.Label(
            header, text="Thiết lập Dự án và Tình trạng hôn nhân áp dụng để hệ thống tự động xuất đúng mẫu khi tạo hồ sơ.",
            font=("Segoe UI", 8), fg="#64748b", bg="#ffffff"
        ).pack(anchor="w", pady=(2, 0))

        # Scrollable list of files
        body_frame = tk.Frame(self, bg="#f8fafc")
        body_frame.pack(fill="both", expand=True, padx=16, pady=10)

        canvas = tk.Canvas(body_frame, bg="#f8fafc", highlightthickness=0)
        scrollbar = ttk.Scrollbar(body_frame, orient="vertical", command=canvas.yview)
        scrollable_frame = tk.Frame(canvas, bg="#f8fafc")

        scrollable_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )
        canvas.create_window((0, 0), window=scrollable_frame, anchor="nw", width=710)
        canvas.configure(yscrollcommand=scrollbar.set)

        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        self.items_data = []

        project_options = ["Tất cả dự án", "An Phú", "Hồng Loan", "Nam Long"]
        marital_options = [
            ("Tất cả đối tượng", "all"),
            ("Độc thân / Ly hôn (Chưa kết hôn)", "single"),
            ("Có gia đình / Đã kết hôn (Vợ chồng)", "married")
        ]

        for idx, fpath in enumerate(filepaths, 1):
            fname = os.path.basename(fpath)
            detected_code, detected_proj, detected_marital, detected_title = detect_metadata(fname)

            card = tk.LabelFrame(
                scrollable_frame, text=f" [{idx:02d}] {fname} ",
                font=("Segoe UI", 9, "bold"), fg="#334155", bg="#ffffff",
                padx=12, pady=8, highlightthickness=1, highlightbackground="#cbd5e1"
            )
            card.pack(fill="x", pady=5)

            # Row 1: Code / Title
            row1 = tk.Frame(card, bg="#ffffff")
            row1.pack(fill="x", pady=2)
            tk.Label(row1, text="Mã mẫu:", font=("Segoe UI", 8, "bold"), fg="#475569", bg="#ffffff", width=8, anchor="w").pack(side="left")
            ent_code = ttk.Entry(row1, width=15)
            ent_code.insert(0, detected_code)
            ent_code.pack(side="left", padx=(0, 15))

            tk.Label(row1, text="Tiêu đề mẫu:", font=("Segoe UI", 8, "bold"), fg="#475569", bg="#ffffff", width=11, anchor="w").pack(side="left")
            ent_title = ttk.Entry(row1)
            ent_title.insert(0, detected_title)
            ent_title.pack(side="left", fill="x", expand=True)

            # Row 2: Project & Marital
            row2 = tk.Frame(card, bg="#ffffff")
            row2.pack(fill="x", pady=(6, 2))

            tk.Label(row2, text="Dự án:", font=("Segoe UI", 8, "bold"), fg="#475569", bg="#ffffff", width=8, anchor="w").pack(side="left")
            cb_proj = ttk.Combobox(row2, values=project_options, state="normal", width=18)
            cur_proj_str = "Tất cả dự án" if detected_proj == "all" else detected_proj
            cb_proj.set(cur_proj_str)
            cb_proj.pack(side="left", padx=(0, 15))

            tk.Label(row2, text="Đối tượng:", font=("Segoe UI", 8, "bold"), fg="#475569", bg="#ffffff", width=11, anchor="w").pack(side="left")
            cb_marital = ttk.Combobox(row2, values=[m[0] for m in marital_options], state="readonly", width=32)
            if detected_marital == "single":
                cb_marital.current(1)
            elif detected_marital == "married":
                cb_marital.current(2)
            else:
                cb_marital.current(0)
            cb_marital.pack(side="left", fill="x", expand=True)

            self.items_data.append({
                "fpath": fpath,
                "fname": fname,
                "ent_code": ent_code,
                "ent_title": ent_title,
                "cb_proj": cb_proj,
                "cb_marital": cb_marital,
                "marital_options": marital_options
            })

        # Footer
        footer = tk.Frame(self, bg="#ffffff", padx=16, pady=10, highlightthickness=1, highlightbackground="#e2e8f0")
        footer.pack(fill="x", side="bottom")

        btn_cancel = tk.Button(
            footer, text="Hủy", font=("Segoe UI", 9), bg="#e2e8f0", fg="#475569",
            relief="flat", padx=14, pady=6, cursor="hand2", command=self.destroy
        )
        btn_cancel.pack(side="right", padx=(6, 0))

        btn_confirm = tk.Button(
            footer, text="✅ Xác Nhận Thêm & Đồng Bộ Vào Hệ Thống",
            font=("Segoe UI", 9, "bold"), bg="#059669", fg="white",
            activebackground="#047857", activeforeground="white",
            relief="flat", padx=18, pady=6, cursor="hand2", command=self.on_confirm
        )
        btn_confirm.pack(side="right")

    def on_confirm(self):
        meta = load_template_metadata()
        os.makedirs(TEMPLATE_DIR, exist_ok=True)
        count = 0

        for item in self.items_data:
            fpath = item["fpath"]
            fname = item["fname"]
            dest = os.path.join(TEMPLATE_DIR, fname)
            if os.path.abspath(fpath) != os.path.abspath(dest):
                shutil.copy2(fpath, dest)
            count += 1

            code = item["ent_code"].get().strip() or "Mẫu Word"
            title = item["ent_title"].get().strip() or fname
            proj_val = item["cb_proj"].get().strip()
            proj = "all" if proj_val in ["Tất cả dự án", "all", ""] else proj_val

            marital_idx = item["cb_marital"].current()
            if marital_idx >= 0 and marital_idx < len(item["marital_options"]):
                marital_val = item["marital_options"][marital_idx][1]
            else:
                marital_val = "all"

            meta[fname] = {
                "code": code,
                "title": title,
                "project": proj,
                "maritalTarget": marital_val
            }

        save_template_metadata(meta)
        self.destroy()
        self.on_success(count)


class SelectRecordsDialog(tk.Toplevel):
    def __init__(self, parent, excel_path, on_submit, on_browse_another):
        super().__init__(parent)
        self.parent = parent
        self.excel_path = excel_path
        self.on_submit = on_submit
        self.on_browse_another = on_browse_another
        
        self.title("CHỌN HỒ SƠ XUẤT FILE HÀNG LOẠT (BIỂU MẪU WORD)")
        win_w, win_h = 1000, 750
        self.update_idletasks()
        screen_w = self.winfo_screenwidth()
        screen_h = self.winfo_screenheight()
        p_x = max(0, (screen_w - win_w) // 2)
        p_y = max(0, (screen_h - win_h) // 2)
        self.geometry(f"{win_w}x{win_h}+{p_x}+{p_y}")
        self.minsize(920, 620)
        self.configure(bg="#f8fafc")
        self.transient(parent)

        # Ưu tiên hiển thị nổi bật lên TOP màn hình
        self.lift()
        self.attributes("-topmost", True)
        self.focus_force()
        self.grab_set()

        self.available_folders = get_case_template_folders()
        self.records = self.load_records(excel_path)

        # Rút trích danh sách các dự án duy nhất có trong dữ liệu
        all_projs = set()
        for r in self.records:
            p = str(r.get("Dự án") or r.get("project") or r.get("duAn") or "").strip()
            if p and p.lower() not in ["none", "null", "all", "tất cả dự án", ""]:
                if "an phú" in p.lower() or "an phu" in p.lower():
                    all_projs.add("An Phú")
                else:
                    all_projs.add(p)

        # Nếu all_projs chưa có hoặc còn thiếu, bổ sung các dự án chính dựa trên biểu mẫu & mặc định
        default_candidates = ["An Phú", "Nam Long", "Hồng Loan"]
        for f in self.available_folders:
            f_norm = strip_vn_accents(f)
            if "an phu" in f_norm:
                all_projs.add("An Phú")
            elif "nam long" in f_norm:
                all_projs.add("Nam Long")
            elif "hong loan" in f_norm:
                all_projs.add("Hồng Loan")

        if not all_projs:
            all_projs = set(default_candidates)

        self.project_filter_list = ["Tất cả dự án"] + sorted(list(all_projs), key=natural_sort_key)

        # Trạng thái phân trang & bộ lọc (Hiển thị 5 khách hàng / trang)
        self.page_size = 5
        self.current_page = 0
        self.total_pages = 1
        self.filtered_records = []
        self._selected_real_idx = None

        # 1. Header
        header = tk.Frame(self, bg="#ffffff", padx=18, pady=12, highlightthickness=1, highlightbackground="#e2e8f0")
        header.pack(fill="x")

        lbl_title = tk.Label(
            header, text="📋 CHỌN HỒ SƠ XUẤT FILE HÀNG LOẠT (BIỂU MẪU WORD)",
            font=("Segoe UI", 12, "bold"), fg="#1e293b", bg="#ffffff"
        )
        lbl_title.pack(anchor="w")

        src_name = os.path.basename(excel_path)
        is_web_sync = "SYNC" in src_name.upper()
        tag_text = "🟢 Dữ liệu đồng bộ trực tiếp từ Web" if is_web_sync else f"📁 File Excel: {src_name}"
        lbl_sub = tk.Label(
            header, text=f"Nguồn: {tag_text}  •  Tìm thấy {len(self.records)} hồ sơ trong danh sách",
            font=("Segoe UI", 9), fg="#059669" if is_web_sync else "#475569", bg="#ffffff"
        )
        lbl_sub.pack(anchor="w", pady=(2, 0))

        # 2. Khung Tùy chọn Phạm vi Xuất (Radio Buttons & Chọn thư mục mẫu)
        opt_card = tk.Frame(self, bg="#ffffff", padx=16, pady=8, highlightthickness=1, highlightbackground="#e2e8f0")
        opt_card.pack(fill="x", padx=16, pady=(6, 3))

        self.scope_var = tk.StringVar(value="one" if len(self.records) > 0 else "all")

        row_radio = tk.Frame(opt_card, bg="#ffffff")
        row_radio.pack(fill="x")

        r_one = tk.Radiobutton(
            row_radio, text="👉 Chỉ xuất 01 HỒ SƠ ĐANG CHỌN trong bảng dưới đây",
            variable=self.scope_var, value="one",
            font=("Segoe UI", 9, "bold"), fg="#2563eb", bg="#ffffff",
            activebackground="#ffffff", activeforeground="#1d4ed8"
        )
        r_one.pack(side="left", padx=(0, 20))

        r_all = tk.Radiobutton(
            row_radio, text=f"📦 Xuất TOÀN BỘ ({len(self.records)} hồ sơ)",
            variable=self.scope_var, value="all",
            font=("Segoe UI", 9), fg="#475569", bg="#ffffff",
            activebackground="#ffffff", activeforeground="#0f172a"
        )
        r_all.pack(side="left")

        # Hàng hiển thị / chọn thư mục mẫu Word & Tình trạng hôn nhân
        row_folder = tk.Frame(opt_card, bg="#ffffff")
        row_folder.pack(fill="x", pady=(6, 2))

        tk.Label(
            row_folder, text="📁 Thư mục mẫu Word:",
            font=("Segoe UI", 9, "bold"), fg="#1e293b", bg="#ffffff"
        ).pack(side="left", padx=(0, 6))

        self.available_folders = get_case_template_folders()
        self.cb_case_folder = ttk.Combobox(
            row_folder, values=self.available_folders, state="readonly",
            width=32, font=("Segoe UI", 9)
        )
        self.cb_case_folder.pack(side="left", padx=(0, 10))
        self.cb_case_folder.bind("<<ComboboxSelected>>", self.on_folder_change)

        tk.Label(
            row_folder, text="💍 Hôn nhân:",
            font=("Segoe UI", 9, "bold"), fg="#1e293b", bg="#ffffff"
        ).pack(side="left", padx=(0, 4))

        self.cb_marital_state = ttk.Combobox(
            row_folder, values=["Có gia đình / Kết hôn", "Độc thân / Ly hôn"], state="readonly",
            width=20, font=("Segoe UI", 9)
        )
        self.cb_marital_state.pack(side="left", padx=(0, 10))
        self.cb_marital_state.bind("<<ComboboxSelected>>", self.on_marital_change)

        tk.Label(
            row_folder, text="💡 Tự động khớp theo Dự án & Hôn nhân",
            font=("Segoe UI", 8, "italic"), fg="#059669", bg="#ffffff"
        ).pack(side="left")

        self.lbl_selected_hint = tk.Label(
            opt_card, text="💡 Nhấp chuột vào một dòng trong bảng dưới đây để chọn hồ sơ cần xuất:",
            font=("Segoe UI", 8, "italic"), fg="#64748b", bg="#ffffff"
        )
        self.lbl_selected_hint.pack(anchor="w", pady=(4, 0))

        # 2.5 Khung Tìm Kiếm Khách Hàng & Lọc Dự Án (Search & Filter Bar)
        filter_card = tk.Frame(self, bg="#ffffff", padx=16, pady=7, highlightthickness=1, highlightbackground="#e2e8f0")
        filter_card.pack(fill="x", padx=16, pady=(3, 3))

        row_filter = tk.Frame(filter_card, bg="#ffffff")
        row_filter.pack(fill="x")

        # Ô tìm kiếm khách hàng
        tk.Label(
            row_filter, text="🔍 Tìm khách hàng:",
            font=("Segoe UI", 9, "bold"), fg="#1e293b", bg="#ffffff"
        ).pack(side="left", padx=(0, 6))

        self.search_kh_var = tk.StringVar()
        self.ent_search_kh = ttk.Entry(row_filter, textvariable=self.search_kh_var, width=28, font=("Segoe UI", 9))
        self.ent_search_kh.pack(side="left", padx=(0, 14))
        self.ent_search_kh.bind("<KeyRelease>", lambda e: self.apply_filter(reset_page=True))

        # Dropdown lọc theo dự án
        tk.Label(
            row_filter, text="🏢 Dự án:",
            font=("Segoe UI", 9, "bold"), fg="#1e293b", bg="#ffffff"
        ).pack(side="left", padx=(0, 6))

        self.cb_filter_proj = ttk.Combobox(
            row_filter, values=self.project_filter_list, state="readonly",
            width=22, font=("Segoe UI", 9)
        )
        self.cb_filter_proj.set("Tất cả dự án")
        self.cb_filter_proj.pack(side="left", padx=(0, 10))
        self.cb_filter_proj.bind("<<ComboboxSelected>>", lambda e: self.apply_filter(reset_page=True))

        # Nút Xóa lọc
        btn_clear_filter = tk.Button(
            row_filter, text="❌ Xóa Lọc",
            font=("Segoe UI", 8, "bold"), bg="#f1f5f9", fg="#475569",
            activebackground="#e2e8f0", activeforeground="#0f172a",
            relief="flat", padx=8, pady=2, cursor="hand2",
            highlightthickness=1, highlightbackground="#cbd5e1",
            command=self.clear_filter
        )
        btn_clear_filter.pack(side="left", padx=(0, 12))

        # Nhãn đếm kết quả
        self.lbl_filter_count = tk.Label(
            row_filter, text=f"📊 Có {len(self.records)} hồ sơ",
            font=("Segoe UI", 8, "italic"), fg="#059669", bg="#ffffff"
        )
        self.lbl_filter_count.pack(side="right")

        # 3. Bảng Danh Sách Hồ Sơ (Treeview - Hiển thị 5 dòng)
        table_frame = tk.Frame(self, bg="#f8fafc")
        table_frame.pack(fill="x", padx=16, pady=(2, 2))

        cols = ("stt", "hoten", "canho", "duan", "honnhan", "folder", "trangthai", "vonvay", "cccd")
        self.tree = ttk.Treeview(table_frame, columns=cols, show="headings", height=5, selectmode="browse")
        
        self.tree.heading("stt", text="STT")
        self.tree.heading("hoten", text="Họ và Tên Khách Hàng")
        self.tree.heading("canho", text="Căn Hộ")
        self.tree.heading("duan", text="Dự Án")
        self.tree.heading("honnhan", text="Hôn Nhân")
        self.tree.heading("folder", text="📁 Thư Mục Mẫu Word")
        self.tree.heading("trangthai", text="Trạng Thái GN")
        self.tree.heading("vonvay", text="Vốn Vay (VNĐ)")
        self.tree.heading("cccd", text="Số CCCD")

        self.tree.column("stt", width=42, anchor="center")
        self.tree.column("hoten", width=160, anchor="w")
        self.tree.column("canho", width=65, anchor="center")
        self.tree.column("duan", width=105, anchor="w")
        self.tree.column("honnhan", width=90, anchor="center")
        self.tree.column("folder", width=190, anchor="w")
        self.tree.column("trangthai", width=105, anchor="center")
        self.tree.column("vonvay", width=95, anchor="e")
        self.tree.column("cccd", width=95, anchor="center")

        scroll_y = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scroll_y.set)
        scroll_y.pack(side="right", fill="y")
        self.tree.pack(side="left", fill="both", expand=True)

        # 3.2 Thanh Phân Trang (Hiển thị 5 khách hàng / trang)
        self.page_bar = tk.Frame(self, bg="#ffffff", padx=14, pady=5, highlightthickness=1, highlightbackground="#e2e8f0")
        self.page_bar.pack(fill="x", padx=16, pady=(2, 4))

        self.lbl_page_info = tk.Label(
            self.page_bar, text="",
            font=("Segoe UI", 8, "bold"), fg="#475569", bg="#ffffff"
        )
        self.lbl_page_info.pack(side="left")

        # Nút phân trang bên phải
        nav_frame = tk.Frame(self.page_bar, bg="#ffffff")
        nav_frame.pack(side="right")

        self.btn_first = tk.Button(
            nav_frame, text="⏮ Đầu", font=("Segoe UI", 8, "bold"),
            bg="#f8fafc", fg="#334155", relief="flat", padx=8, pady=1, cursor="hand2",
            highlightthickness=1, highlightbackground="#cbd5e1", command=self.go_first_page
        )
        self.btn_first.pack(side="left", padx=(0, 4))

        self.btn_prev = tk.Button(
            nav_frame, text="◀ Trước", font=("Segoe UI", 8, "bold"),
            bg="#f8fafc", fg="#334155", relief="flat", padx=8, pady=1, cursor="hand2",
            highlightthickness=1, highlightbackground="#cbd5e1", command=self.go_prev_page
        )
        self.btn_prev.pack(side="left", padx=(0, 6))

        self.lbl_page_num = tk.Label(
            nav_frame, text="Trang 1 / 1",
            font=("Segoe UI", 8, "bold"), fg="#2563eb", bg="#eff6ff",
            padx=10, pady=2, highlightthickness=1, highlightbackground="#bfdbfe"
        )
        self.lbl_page_num.pack(side="left", padx=(0, 6))

        self.btn_next = tk.Button(
            nav_frame, text="Sau ▶", font=("Segoe UI", 8, "bold"),
            bg="#f8fafc", fg="#334155", relief="flat", padx=8, pady=1, cursor="hand2",
            highlightthickness=1, highlightbackground="#cbd5e1", command=self.go_next_page
        )
        self.btn_next.pack(side="left", padx=(0, 4))

        self.btn_last = tk.Button(
            nav_frame, text="Cuối ⏭", font=("Segoe UI", 8, "bold"),
            bg="#f8fafc", fg="#334155", relief="flat", padx=8, pady=1, cursor="hand2",
            highlightthickness=1, highlightbackground="#cbd5e1", command=self.go_last_page
        )
        self.btn_last.pack(side="left")

        # 3.5 Khung Hiển Thị Danh Sách Biểu Mẫu Word Của Thư Mục (Có Checkbox Tùy Chọn Mẫu)
        self.tmpl_frame = tk.LabelFrame(
            self, text=" 📋 DANH SÁCH BIỂU MẪU WORD SẼ XUẤT THEO THỨ TỰ (1 - 12) ",
            font=("Segoe UI", 9, "bold"), fg="#1e293b", bg="#ffffff",
            padx=12, pady=6, highlightthickness=1, highlightbackground="#cbd5e1"
        )
        self.tmpl_frame.pack(fill="both", expand=True, padx=16, pady=(4, 6))

        # Thanh công cụ chọn nhanh mẫu xuất
        tmpl_toolbar = tk.Frame(self.tmpl_frame, bg="#ffffff")
        tmpl_toolbar.pack(fill="x", pady=(0, 6))

        tk.Label(
            tmpl_toolbar, text="💡 Nhấp chuột (hoặc phím Space) vào dòng để Bật/Tắt mẫu xuất:",
            font=("Segoe UI", 8, "italic"), fg="#64748b", bg="#ffffff"
        ).pack(side="left", padx=(0, 8))

        btn_sel_all = tk.Button(
            tmpl_toolbar, text="☑ Chọn Tất Cả",
            font=("Segoe UI", 8, "bold"), bg="#eff6ff", fg="#1d4ed8",
            activebackground="#dbeafe", activeforeground="#1e40af",
            relief="flat", padx=8, pady=2, cursor="hand2",
            highlightthickness=1, highlightbackground="#bfdbfe",
            command=self.select_all_templates
        )
        btn_sel_all.pack(side="left", padx=(0, 4))

        btn_desel_all = tk.Button(
            tmpl_toolbar, text="☐ Bỏ Chọn Hết",
            font=("Segoe UI", 8), bg="#f8fafc", fg="#475569",
            activebackground="#f1f5f9", activeforeground="#0f172a",
            relief="flat", padx=8, pady=2, cursor="hand2",
            highlightthickness=1, highlightbackground="#cbd5e1",
            command=self.deselect_all_templates
        )
        btn_desel_all.pack(side="left", padx=(0, 4))

        btn_invert = tk.Button(
            tmpl_toolbar, text="🔄 Đảo Lựa Chọn",
            font=("Segoe UI", 8), bg="#f8fafc", fg="#475569",
            activebackground="#f1f5f9", activeforeground="#0f172a",
            relief="flat", padx=8, pady=2, cursor="hand2",
            highlightthickness=1, highlightbackground="#cbd5e1",
            command=self.invert_templates
        )
        btn_invert.pack(side="left", padx=(0, 10))

        self.lbl_tmpl_count = tk.Label(
            tmpl_toolbar, text="📊 Đã chọn: 0/0 biểu mẫu",
            font=("Segoe UI", 9, "bold"), fg="#059669", bg="#ffffff"
        )
        self.lbl_tmpl_count.pack(side="right")

        tmpl_cols = ("check", "stt", "fname", "status")
        self.tree_tmpl = ttk.Treeview(self.tmpl_frame, columns=tmpl_cols, show="headings", height=7, selectmode="browse")
        self.tree_tmpl.heading("check", text="☑ / ☐", command=self.toggle_all_templates)
        self.tree_tmpl.heading("stt", text="STT")
        self.tree_tmpl.heading("fname", text="Tên Tập Tin Biểu Mẫu Word (.docx)")
        self.tree_tmpl.heading("status", text="Trạng Thái Xuất")

        self.tree_tmpl.column("check", width=50, anchor="center")
        self.tree_tmpl.column("stt", width=45, anchor="center")
        self.tree_tmpl.column("fname", width=620, anchor="w")
        self.tree_tmpl.column("status", width=150, anchor="center")

        self.tree_tmpl.tag_configure("checked", foreground="#0f172a")
        self.tree_tmpl.tag_configure("unchecked", foreground="#94a3b8")

        self.tree_tmpl.bind("<Button-1>", self.on_template_tree_click)
        self.tree_tmpl.bind("<Double-1>", self.on_template_tree_click)
        self.tree_tmpl.bind("<space>", self.on_template_tree_space)

        scroll_tmpl_y = ttk.Scrollbar(self.tmpl_frame, orient="vertical", command=self.tree_tmpl.yview)
        self.tree_tmpl.configure(yscrollcommand=scroll_tmpl_y.set)
        scroll_tmpl_y.pack(side="right", fill="y")
        self.tree_tmpl.pack(side="left", fill="both", expand=True)

        self.populate_table()
        self.tree.bind("<<TreeviewSelect>>", self.on_select_row)

        # 4. Chân hộp thoại (Buttons)
        footer = tk.Frame(self, bg="#ffffff", padx=16, pady=10, highlightthickness=1, highlightbackground="#e2e8f0")
        footer.pack(fill="x", side="bottom")

        btn_other = tk.Button(
            footer, text="📁 Chọn File Excel Khác...",
            font=("Segoe UI", 9), bg="#f1f5f9", fg="#334155",
            relief="flat", padx=12, pady=6, cursor="hand2", command=self.on_choose_another
        )
        btn_other.pack(side="left")

        btn_cancel = tk.Button(
            footer, text="Hủy",
            font=("Segoe UI", 9), bg="#e2e8f0", fg="#475569",
            relief="flat", padx=14, pady=6, cursor="hand2", command=self.destroy
        )
        btn_cancel.pack(side="right", padx=(6, 0))

        self.btn_export = tk.Button(
            footer, text="🚀 XUẤT FILE HÀNG LOẠT",
            font=("Segoe UI", 9, "bold"), bg="#7c3aed", fg="white",
            activebackground="#6d28d9", activeforeground="white",
            relief="flat", padx=18, pady=6, cursor="hand2", command=self.on_export_click
        )
        self.btn_export.pack(side="right")

    def get_checked_form_keys(self):
        keys = set()
        for item in self.tree_tmpl.get_children():
            vals = self.tree_tmpl.item(item, "values")
            if vals and vals[0] == "☑":
                keys.add(get_form_key(vals[2]))
        return keys

    def update_template_list_preview(self, folder_name, keep_selection=False):
        saved_keys = self.get_checked_form_keys() if keep_selection else None
        self._current_preview_folder = folder_name
        for item in self.tree_tmpl.get_children():
            self.tree_tmpl.delete(item)
        files = get_sorted_templates_in_folder(folder_name)
        self.tmpl_frame.config(
            text=f" 📋 DANH SÁCH BIỂU MẪU WORD SẼ XUẤT - THƯ MỤC: [{folder_name}] ({len(files)} BIỂU MẪU TỪ 1 ĐẾN {len(files)}) "
        )
        for idx, fn in enumerate(files, 1):
            k = get_form_key(fn)
            if saved_keys is not None:
                is_checked = (k in saved_keys)
            else:
                is_checked = True
            check_str = "☑" if is_checked else "☐"
            status_str = "✓ Sẵn sàng xuất" if is_checked else "⏭ Bỏ qua (Không xuất)"
            tag_str = "checked" if is_checked else "unchecked"
            self.tree_tmpl.insert("", "end", values=(check_str, f"{idx:02d}", fn, status_str), tags=(tag_str,))
        self.update_tmpl_count_badge()

    def on_template_tree_click(self, event):
        region = self.tree_tmpl.identify_region(event.x, event.y)
        if region == "heading":
            col = self.tree_tmpl.identify_column(event.x)
            if col == "#1":
                self.toggle_all_templates()
                return "break"
        row_id = self.tree_tmpl.identify_row(event.y)
        if row_id:
            self.toggle_template_row(row_id)
            self.tree_tmpl.selection_set(row_id)
            return "break"

    def on_template_tree_space(self, event):
        selected = self.tree_tmpl.selection()
        for row_id in selected:
            self.toggle_template_row(row_id)
        return "break"

    def toggle_template_row(self, row_id):
        vals = list(self.tree_tmpl.item(row_id, "values"))
        if not vals:
            return
        if vals[0] == "☑":
            vals[0] = "☐"
            vals[3] = "⏭ Bỏ qua (Không xuất)"
            self.tree_tmpl.item(row_id, values=vals, tags=("unchecked",))
        else:
            vals[0] = "☑"
            vals[3] = "✓ Sẵn sàng xuất"
            self.tree_tmpl.item(row_id, values=vals, tags=("checked",))
        self.update_tmpl_count_badge()

    def select_all_templates(self):
        for item in self.tree_tmpl.get_children():
            vals = list(self.tree_tmpl.item(item, "values"))
            if vals:
                vals[0] = "☑"
                vals[3] = "✓ Sẵn sàng xuất"
                self.tree_tmpl.item(item, values=vals, tags=("checked",))
        self.update_tmpl_count_badge()

    def deselect_all_templates(self):
        for item in self.tree_tmpl.get_children():
            vals = list(self.tree_tmpl.item(item, "values"))
            if vals:
                vals[0] = "☐"
                vals[3] = "⏭ Bỏ qua (Không xuất)"
                self.tree_tmpl.item(item, values=vals, tags=("unchecked",))
        self.update_tmpl_count_badge()

    def invert_templates(self):
        for item in self.tree_tmpl.get_children():
            self.toggle_template_row(item)

    def toggle_all_templates(self):
        any_unchecked = False
        for item in self.tree_tmpl.get_children():
            vals = self.tree_tmpl.item(item, "values")
            if vals and vals[0] == "☐":
                any_unchecked = True
                break
        if any_unchecked:
            self.select_all_templates()
        else:
            self.deselect_all_templates()

    def update_tmpl_count_badge(self):
        total = 0
        checked = 0
        for item in self.tree_tmpl.get_children():
            total += 1
            vals = self.tree_tmpl.item(item, "values")
            if vals and vals[0] == "☑":
                checked += 1
        if hasattr(self, 'lbl_tmpl_count'):
            self.lbl_tmpl_count.config(
                text=f"📊 Đã chọn: {checked}/{total} biểu mẫu",
                fg="#059669" if checked > 0 else "#dc2626"
            )
        if hasattr(self, 'btn_export'):
            if checked == 0:
                self.btn_export.config(
                    text="⚠ Vui lòng chọn ít nhất 1 mẫu",
                    bg="#94a3b8"
                )
            elif checked == total:
                self.btn_export.config(
                    text=f"🚀 XUẤT TOÀN BỘ ({total} MẪU)",
                    bg="#7c3aed"
                )
            else:
                self.btn_export.config(
                    text=f"🚀 XUẤT {checked} MẪU ĐÃ CHỌN",
                    bg="#2563eb"
                )

    def load_records(self, excel_path):
        raw_records = []
        # 1. Tối ưu siêu tốc: Nếu là file đồng bộ Web, đọc trực tiếp từ database_records_rows206.json (< 1ms thay vì đợi openpyxl load)
        src_name = os.path.basename(excel_path)
        rows206_file = os.path.join(BASE_DIR, "database_records_rows206.json")
        if "SYNC" in src_name.upper():
            db_len = 0
            if os.path.exists(DB_FILE):
                try:
                    with open(DB_FILE, 'r', encoding='utf-8') as f:
                        f_data = json.load(f)
                        if isinstance(f_data, list):
                            db_len = len(f_data)
                except Exception:
                    pass

            if os.path.exists(rows206_file):
                try:
                    with open(rows206_file, 'r', encoding='utf-8') as rf:
                        data = json.load(rf)
                        if isinstance(data, list) and len(data) >= db_len and len(data) > 0:
                            raw_records = data
                except Exception:
                    pass

            if not raw_records:
                synced = sync_all_database_records()
                if synced:
                    raw_records = synced

        # 2. Nếu là file Excel khác do người dùng chọn: đọc bằng openpyxl
        if not raw_records:
            try:
                import openpyxl
                wb = openpyxl.load_workbook(excel_path, data_only=True)
                ws = wb['data'] if 'data' in wb.sheetnames else wb.active
                headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
                records = []
                for r_idx in range(2, ws.max_row + 1):
                    row_vals = [ws.cell(r_idx, c).value for c in range(1, len(headers) + 1)]
                    if any(row_vals):
                        rec = {}
                        for idx, h in enumerate(headers):
                            if h:
                                val = row_vals[idx] if idx < len(row_vals) else ""
                                rec[str(h).strip()] = val if val is not None else ""
                        records.append(rec)
                raw_records = records
            except Exception:
                raw_records = []

        # 3. Làm giàu dữ liệu: Luôn nhận diện 100% Dự án & Hôn nhân & Thư mục mẫu Word
        db_records_map = {}
        if os.path.exists(DB_FILE):
            try:
                with open(DB_FILE, 'r', encoding='utf-8') as f:
                    db_data = json.load(f)
                    if isinstance(db_data, list):
                        for item in db_data:
                            if isinstance(item, dict):
                                c = str(item.get("cccd") or item.get("CCCD") or "").strip().replace(' ', '').replace('.', '').replace('-', '')
                                n = str(item.get("hoTen") or "").strip().lower()
                                if c: db_records_map[c] = item
                                if n: db_records_map[n] = item
            except Exception:
                pass

        for r in raw_records:
            if not isinstance(r, dict):
                continue
            c_id = str(r.get("cccd") or r.get("CCCD") or r.get("soCCCD") or "").strip().replace(' ', '').replace('.', '').replace('-', '')
            n_id = str(r.get("hoTen") or r.get("Họ và tên KH thường") or r.get("Họ và tên KH hoa") or "").strip().lower()
            matched_db = db_records_map.get(c_id) or db_records_map.get(n_id)

            # Tự động nhận diện chuẩn xác dự án
            proj = detect_project_from_record(r, db_records_map)
            r["Dự án"] = proj
            r["project"] = proj
            r["duAn"] = proj

            # Đồng bộ tình trạng hôn nhân từ database nếu có
            if matched_db:
                db_hn = matched_db.get("tinhTrangHonNhan")
                if db_hn and (not r.get("tinhTrangHonNhan") or r.get("tinhTrangHonNhan") in ["None", ""]):
                    r["tinhTrangHonNhan"] = db_hn
                    r["Tình trạng hôn nhân"] = db_hn

            # Khớp thư mục mẫu Word
            avail = getattr(self, 'available_folders', None) or get_case_template_folders()
            auto_folder = resolve_template_folder(r, folders=avail, db_records_map=db_records_map)
            r["_resolved_folder"] = auto_folder
            r["case_folder"] = auto_folder

        return raw_records

    def clear_filter(self):
        self.search_kh_var.set("")
        self.cb_filter_proj.set("Tất cả dự án")
        self.apply_filter(reset_page=True)

    def apply_filter(self, reset_page=True):
        if reset_page:
            self.current_page = 0

        q = strip_vn_accents(self.search_kh_var.get().strip().lower()) if hasattr(self, 'search_kh_var') else ""
        selected_proj = self.cb_filter_proj.get().strip() if hasattr(self, 'cb_filter_proj') else "Tất cả dự án"

        matched = []
        for idx, r in enumerate(self.records):
            # 1. Lọc theo dự án
            if selected_proj and selected_proj != "Tất cả dự án":
                r_proj = str(r.get("Dự án") or r.get("project") or r.get("duAn") or "").strip().lower()
                norm_sel = strip_vn_accents(selected_proj.lower())
                norm_r = strip_vn_accents(r_proj)
                if norm_sel not in norm_r and norm_r not in norm_sel:
                    continue

            # 2. Lọc theo từ khóa khách hàng (Tên, CCCD, Căn hộ, SĐT)
            if q:
                ho_ten = strip_vn_accents(str(r.get("Họ và tên KH hoa") or r.get("Họ và tên KH thường") or r.get("hoTen") or "").lower())
                cccd = strip_vn_accents(str(r.get("CCCD") or r.get("cccd") or "").lower())
                can = strip_vn_accents(str(r.get("Căn hộ số") or r.get("canHoSo") or "").lower())
                sdt = strip_vn_accents(str(r.get("Số điện thoại KH") or r.get("sdtKH") or "").lower())
                if q not in ho_ten and q not in cccd and q not in can and q not in sdt:
                    continue

            matched.append((idx, r))

        self.filtered_records = matched
        self.total_pages = max(1, (len(self.filtered_records) + self.page_size - 1) // self.page_size)
        if self.current_page >= self.total_pages:
            self.current_page = max(0, self.total_pages - 1)

        self.update_pagination_ui()
        self.render_current_page()

    def update_pagination_ui(self):
        total_items = len(self.filtered_records)
        total_all = len(self.records)

        if hasattr(self, 'lbl_filter_count'):
            if total_items < total_all:
                self.lbl_filter_count.config(
                    text=f"🎯 Khớp: {total_items}/{total_all} hồ sơ",
                    fg="#2563eb", font=("Segoe UI", 8, "bold")
                )
            else:
                self.lbl_filter_count.config(
                    text=f"📊 Tổng số: {total_all} hồ sơ",
                    fg="#059669", font=("Segoe UI", 8, "italic")
                )

        if not hasattr(self, 'lbl_page_info'):
            return

        if total_items == 0:
            self.lbl_page_info.config(text="Hiển thị 0 / 0 hồ sơ")
            self.lbl_page_num.config(text="Trang 0 / 0")
            self.btn_first.config(state="disabled")
            self.btn_prev.config(state="disabled")
            self.btn_next.config(state="disabled")
            self.btn_last.config(state="disabled")
            return

        start = self.current_page * self.page_size + 1
        end = min((self.current_page + 1) * self.page_size, total_items)
        self.lbl_page_info.config(text=f"Hiển thị hồ sơ {start} - {end} trong tổng số {total_items} hồ sơ")
        self.lbl_page_num.config(text=f"Trang {self.current_page + 1} / {self.total_pages}")

        self.btn_first.config(state="normal" if self.current_page > 0 else "disabled")
        self.btn_prev.config(state="normal" if self.current_page > 0 else "disabled")
        self.btn_next.config(state="normal" if self.current_page < self.total_pages - 1 else "disabled")
        self.btn_last.config(state="normal" if self.current_page < self.total_pages - 1 else "disabled")

    def go_first_page(self):
        if self.current_page > 0:
            self.current_page = 0
            self.update_pagination_ui()
            self.render_current_page()

    def go_prev_page(self):
        if self.current_page > 0:
            self.current_page -= 1
            self.update_pagination_ui()
            self.render_current_page()

    def go_next_page(self):
        if self.current_page < self.total_pages - 1:
            self.current_page += 1
            self.update_pagination_ui()
            self.render_current_page()

    def go_last_page(self):
        if self.current_page < self.total_pages - 1:
            self.current_page = self.total_pages - 1
            self.update_pagination_ui()
            self.render_current_page()

    def populate_table(self):
        self.apply_filter(reset_page=True)

    def render_current_page(self):
        for item in self.tree.get_children():
            self.tree.delete(item)

        start = self.current_page * self.page_size
        end = min(start + self.page_size, len(self.filtered_records))
        page_items = self.filtered_records[start:end]

        for display_pos, (real_idx, r) in enumerate(page_items):
            r['_raw_idx'] = real_idx
            ho_ten = r.get("Họ và tên KH hoa") or r.get("Họ và tên KH thường") or r.get("hoTen") or f"Khách hàng {real_idx+1}"
            can = r.get("Căn hộ số") or r.get("canHoSo") or ""
            du_an = r.get("Dự án") or r.get("project") or detect_project_from_record(r)
            r["Dự án"] = du_an
            r["project"] = du_an
            von_vay = r.get("Vốn vay") or r.get("vonVay") or ""
            if isinstance(von_vay, (int, float)):
                von_txt = f"{von_vay:,.0f}".replace(',', '.') if von_vay > 0 else ""
            else:
                s_von = str(von_vay).strip()
                von_txt = "" if s_von in ["0", "0.0", "", "null", "none"] else s_von
            cccd = r.get("CCCD") or r.get("cccd") or ""

            if "_resolved_folder" not in r or not r.get("_resolved_folder"):
                r["_resolved_folder"] = resolve_template_folder(r, folders=self.available_folders)
            f_name = r["_resolved_folder"]

            status_gn = r.get("trangThaiGiaiNgan")
            if not status_gn:
                if r.get("soTBNgayTB3") or r.get("số tb - ngày tb3") or r.get("số tb - ngày tb lần 3"):
                    status_gn = "Đã GN đợt 3"
                elif r.get("soTBNgayTB2") or r.get("số tb - ngày tb2") or r.get("số tb - ngày tb lần 2"):
                    status_gn = "Đã GN đợt 2"
                elif r.get("soTBNgayTB") or r.get("số tb - ngày tb") or r.get("soThongBao"):
                    status_gn = "Đã GN đợt 1"
                else:
                    status_gn = "Chưa GN"

            is_s = is_single_profile(r)
            hon_nhan_txt = "Độc thân" if is_s else "Có gia đình"

            iid = str(real_idx)
            self.tree.insert("", "end", iid=iid, values=(real_idx+1, ho_ten, can, du_an, hon_nhan_txt, f_name, status_gn, von_txt, cccd))

        # Khôi phục hoặc tự động chọn 1 dòng trong trang
        children = self.tree.get_children()
        if children:
            if getattr(self, '_selected_real_idx', None) is not None and str(self._selected_real_idx) in children:
                self.tree.selection_set(str(self._selected_real_idx))
            else:
                first_iid = children[0]
                self.tree.selection_set(first_iid)
                self.on_select_row(None)
        else:
            self.lbl_selected_hint.config(
                text="⚠ Không tìm thấy hồ sơ nào phù hợp với bộ lọc tìm kiếm!",
                fg="#dc2626", font=("Segoe UI", 9, "italic")
            )

    def on_select_row(self, event):
        selected = self.tree.selection()
        if not selected: return
        try:
            real_idx = int(selected[0])
        except Exception:
            return

        self._selected_real_idx = real_idx
        if 0 <= real_idx < len(self.records):
            r = self.records[real_idx]
            ho_ten = r.get("Họ và tên KH hoa") or r.get("Họ và tên KH thường") or r.get("hoTen") or f"Khách hàng {real_idx+1}"
            can = r.get("Căn hộ số") or r.get("canHoSo") or ""
            f_name = r.get("_resolved_folder") or ""
            self.scope_var.set("one")
            if f_name in self.available_folders:
                self.cb_case_folder.set(f_name)
            is_s = is_single_profile(r)
            if hasattr(self, 'cb_marital_state'):
                self.cb_marital_state.set("Độc thân / Ly hôn" if is_s else "Có gia đình / Kết hôn")
            self.update_hint(real_idx, ho_ten, can, f_name)
            if f_name != getattr(self, '_current_preview_folder', None):
                self.update_template_list_preview(f_name, keep_selection=True)

    def on_folder_change(self, event):
        selected = self.tree.selection()
        if not selected: return
        try:
            real_idx = int(selected[0])
        except Exception:
            return

        new_folder = self.cb_case_folder.get()
        if 0 <= real_idx < len(self.records):
            self.records[real_idx]["_resolved_folder"] = new_folder
            self.records[real_idx]["case_folder"] = new_folder
            if "doc than" in new_folder.lower() or "độc thân" in new_folder.lower():
                self.records[real_idx]["tinhTrangHonNhan"] = "Độc thân"
                self.records[real_idx]["Tình trạng hôn nhân"] = "Độc thân"
                if hasattr(self, 'cb_marital_state'):
                    self.cb_marital_state.set("Độc thân / Ly hôn")
            
            vals = list(self.tree.item(selected[0], "values"))
            if vals:
                vals[4] = "Độc thân" if is_single_profile(self.records[real_idx]) else "Có gia đình"
                vals[5] = new_folder
                self.tree.item(selected[0], values=vals)
                ho_ten = vals[1]
                can = vals[2]
                self.update_hint(real_idx, ho_ten, can, new_folder)
        self.update_template_list_preview(new_folder, keep_selection=True)

    def on_marital_change(self, event):
        selected = self.tree.selection()
        if not selected: return
        try:
            real_idx = int(selected[0])
        except Exception:
            return

        val = self.cb_marital_state.get()
        is_single = "độc thân" in val.lower() or "ly hôn" in val.lower()
        if 0 <= real_idx < len(self.records):
            rec = self.records[real_idx]
            rec["tinhTrangHonNhan"] = "Độc thân" if is_single else "Đã kết hôn"
            rec["Tình trạng hôn nhân"] = rec["tinhTrangHonNhan"]
            
            # Tự động gợi ý lại thư mục tương ứng
            auto_fld = resolve_template_folder(rec, folders=self.available_folders)
            rec["_resolved_folder"] = auto_fld
            rec["case_folder"] = auto_fld
            if auto_fld in self.available_folders:
                self.cb_case_folder.set(auto_fld)
            
            vals = list(self.tree.item(selected[0], "values"))
            if vals:
                vals[4] = "Độc thân" if is_single else "Có gia đình"
                vals[5] = auto_fld
                self.tree.item(selected[0], values=vals)
                self.update_hint(real_idx, vals[1], vals[2], auto_fld)
            self.update_template_list_preview(auto_fld, keep_selection=True)

    def update_hint(self, idx, ho_ten, can, f_name=""):
        can_txt = f" (Căn: {can})" if can else ""
        folder_txt = f"  •  📁 Thư mục mẫu: {f_name}" if f_name else ""
        self.lbl_selected_hint.config(
            text=f"👉 Đang chọn xuất: [{idx+1}] {ho_ten}{can_txt}{folder_txt}",
            fg="#2563eb", font=("Segoe UI", 9, "bold")
        )

    def on_choose_another(self):
        self.destroy()
        self.on_browse_another()

    def on_export_click(self):
        if not self.records:
            messagebox.showwarning("Thông báo", "Không có dữ liệu hồ sơ để xuất.")
            return

        mode = self.scope_var.get()
        if mode == "all":
            # Nếu người dùng đang lọc hiển thị một nhóm hồ sơ, hỏi có muốn xuất nhóm đang lọc hay tất cả
            if len(self.filtered_records) < len(self.records) and len(self.filtered_records) > 0:
                msg = f"Bạn đang lọc hiển thị {len(self.filtered_records)} hồ sơ (trên tổng số {len(self.records)} hồ sơ).\n\n" \
                      f"• Bấm 'Yes' để xuất {len(self.filtered_records)} hồ sơ ĐANG LỌC.\n" \
                      f"• Bấm 'No' để xuất TOÀN BỘ {len(self.records)} hồ sơ trong file."
                res = messagebox.askyesnocancel("Xác nhận phạm vi xuất", msg)
                if res is None:
                    return
                elif res is True:
                    selected_indices = [idx for (idx, r) in self.filtered_records]
                else:
                    selected_indices = list(range(len(self.records)))
            else:
                selected_indices = list(range(len(self.records)))
        else:
            selected = self.tree.selection()
            if not selected:
                messagebox.showwarning("Chưa chọn", "Vui lòng nhấp chọn 1 hồ sơ trong bảng danh sách!")
                return
            try:
                real_idx = int(selected[0])
            except Exception:
                vals = self.tree.item(selected[0], "values")
                real_idx = int(vals[0]) - 1
            selected_indices = [real_idx]

        folder_mapping = {i: self.records[i].get("_resolved_folder", "") for i in selected_indices}
        marital_mapping = {i: self.records[i].get("tinhTrangHonNhan", "") for i in selected_indices}

        # Thu thập danh sách các biểu mẫu được đánh dấu chọn ☑
        checked_templates = []
        for item in self.tree_tmpl.get_children():
            vals = self.tree_tmpl.item(item, "values")
            if vals and vals[0] == "☑":
                checked_templates.append(vals[2])

        if not checked_templates:
            messagebox.showwarning("Chưa chọn mẫu", "Vui lòng tích chọn (☑) ít nhất 1 biểu mẫu Word để xuất!")
            return

        self.destroy()
        self.on_submit(self.excel_path, selected_indices, folder_mapping, checked_templates, marital_mapping)


class AppGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("NGÂN HÀNG CHÍNH SÁCH XÃ HỘI THÀNH PHỐ CẦN THƠ - HỆ THỐNG QUẢN LÝ HỒ SƠ NOXH")
        
        # Căn chỉnh cửa sổ chính xác vào giữa màn hình khi mở lên
        win_w, win_h = 860, 630
        self.root.update_idletasks()
        scr_w = self.root.winfo_screenwidth()
        scr_h = self.root.winfo_screenheight()
        pos_x = max(0, (scr_w - win_w) // 2)
        pos_y = max(0, (scr_h - win_h) // 2)
        self.root.geometry(f"{win_w}x{win_h}+{pos_x}+{pos_y}")
        self.root.minsize(780, 520)
        self.root.configure(bg="#f8fafc")

        # Khởi động server API nền
        start_local_server()
        # Nạp trước openpyxl trong luồng nền để khi bấm Xuất Word cửa sổ hiện lên tức thì
        threading.Thread(target=lambda: __import__('openpyxl'), daemon=True).start()

        # Style TTK
        style = ttk.Style()
        style.theme_use('clam')

        # -------------------------------------------------------------
        # 1. HEADER HIỆN ĐẠI: NỀN TRẮNG, LOGO VBSP, BORDER TINH TẾ
        # -------------------------------------------------------------
        header = tk.Frame(root, bg="#ffffff", padx=20, pady=12, highlightthickness=1, highlightbackground="#e2e8f0")
        header.pack(fill="x")

        # Logo VBSP
        try:
            self.logo_img = tk.PhotoImage(data=VBSP_LOGO_B64)
            lbl_logo = tk.Label(header, image=self.logo_img, bg="#ffffff")
            lbl_logo.pack(side="left", padx=(0, 16))
        except Exception:
            lbl_logo = tk.Label(header, text="🏛️ VBSP", font=("Segoe UI", 16, "bold"), fg="#be185d", bg="#ffffff")
            lbl_logo.pack(side="left", padx=(0, 16))

        # Khung chữ Tiêu đề
        title_frame = tk.Frame(header, bg="#ffffff")
        title_frame.pack(side="left", fill="y", expand=True)

        lbl_bank = tk.Label(
            title_frame, text="NGÂN HÀNG CHÍNH SÁCH XÃ HỘI THÀNH PHỐ CẦN THƠ (VBSP)",
            font=("Segoe UI", 13, "bold"), fg="#9d174d", bg="#ffffff"
        )
        lbl_bank.pack(anchor="w")

        lbl_title = tk.Label(
            title_frame, text="HỆ THỐNG QUẢN LÝ HỒ SƠ & XUẤT FILE HÀNG LOẠT VAY VỐN NOXH",
            font=("Segoe UI", 11, "bold"), fg="#1e293b", bg="#ffffff"
        )
        lbl_title.pack(anchor="w", pady=(2, 0))

        # Badge trạng thái bên phải
        badge_frame = tk.Frame(header, bg="#ecfdf5", padx=10, pady=4, highlightthickness=1, highlightbackground="#a7f3d0")
        badge_frame.pack(side="right")
        lbl_badge = tk.Label(badge_frame, text="● Sẵn sàng", font=("Segoe UI", 9, "bold"), fg="#047857", bg="#ecfdf5")
        lbl_badge.pack()

        # -------------------------------------------------------------
        # 1.5 THANH TRẠNG THÁI ĐỒNG BỘ DỮ LIỆU TỪ WEB (REAL-TIME)
        # -------------------------------------------------------------
        self.sync_card = tk.Frame(root, bg="#f0fdf4", padx=16, pady=8, highlightthickness=1, highlightbackground="#bbf7d0")
        self.sync_card.pack(fill="x", padx=16, pady=(10, 0))

        self.lbl_sync_status = tk.Label(
            self.sync_card,
            text="🟢 Đang kết nối dữ liệu Web...",
            font=("Segoe UI", 9, "bold"), fg="#166534", bg="#f0fdf4"
        )
        self.lbl_sync_status.pack(side="left")

        btn_refresh = tk.Button(
            self.sync_card, text="🔄 Làm mới",
            font=("Segoe UI", 8, "bold"), bg="#dcfce7", fg="#15803d",
            activebackground="#bbf7d0", activeforeground="#166534",
            relief="flat", padx=8, pady=2, cursor="hand2", command=self.update_sync_status
        )
        btn_refresh.pack(side="right")

        # -------------------------------------------------------------
        # 2. KHUNG ĐIỀU KHIỂN CHÍNH (CARD HIỆN ĐẠI, NÚT BẤM DỄ NHÌN)
        # -------------------------------------------------------------
        ctrl_card = tk.Frame(root, bg="#ffffff", padx=18, pady=14, highlightthickness=1, highlightbackground="#e2e8f0")
        ctrl_card.pack(fill="x", padx=16, pady=(12, 6))

        # Nút lớn: Mở giao diện Web trên trình duyệt
        btn_open = tk.Button(
            ctrl_card, text="🌐 MỞ HỆ THỐNG QUẢN LÝ (TRÌNH DUYỆT)",
            font=("Segoe UI", 11, "bold"), bg="#2563eb", fg="white",
            activebackground="#1d4ed8", activeforeground="white",
            relief="flat", padx=16, pady=10, cursor="hand2", command=self.on_open_web
        )
        btn_open.pack(fill="x", pady=(0, 10))

        # Hàng 3 nút chức năng cốt lõi
        row_actions = tk.Frame(ctrl_card, bg="#ffffff")
        row_actions.pack(fill="x", pady=2)

        btn_add_tmpl = tk.Button(
            row_actions, text="➕ Thêm Mẫu Word Vào template/",
            font=("Segoe UI", 9, "bold"), bg="#0284c7", fg="white",
            activebackground="#0369a1", activeforeground="white",
            relief="flat", padx=10, pady=8, cursor="hand2", command=self.on_add_template_files
        )
        btn_add_tmpl.pack(side="left", expand=True, fill="x", padx=(0, 4))

        btn_sync = tk.Button(
            row_actions, text="🔄 Đồng Bộ Mẫu Word (template/)",
            font=("Segoe UI", 9, "bold"), bg="#059669", fg="white",
            activebackground="#047857", activeforeground="white",
            relief="flat", padx=10, pady=8, cursor="hand2", command=self.on_sync_templates
        )
        btn_sync.pack(side="left", expand=True, fill="x", padx=(4, 4))

        btn_merge = tk.Button(
            row_actions, text="🚀 Xuất File Hàng Loạt",
            font=("Segoe UI", 9, "bold"), bg="#7c3aed", fg="white",
            activebackground="#6d28d9", activeforeground="white",
            relief="flat", padx=10, pady=8, cursor="hand2", command=self.on_mail_merge
        )
        btn_merge.pack(side="left", expand=True, fill="x", padx=(4, 0))

        # Hàng chức năng mở rộng: Gộp hồ sơ nhiều máy & Mở thư mục nhanh
        row_secondary = tk.Frame(ctrl_card, bg="#ffffff")
        row_secondary.pack(fill="x", pady=(8, 2))

        btn_combine = tk.Button(
            row_secondary, text="📦 Gộp Dữ Liệu Hồ Sơ (Nhiều Máy)",
            font=("Segoe UI", 9, "bold"), bg="#d97706", fg="white",
            activebackground="#b45309", activeforeground="white",
            relief="flat", padx=10, pady=6, cursor="hand2", command=self.on_combine_records
        )
        btn_combine.pack(side="left", expand=True, fill="x", padx=(0, 4))

        btn_open_tmpl = tk.Button(
            row_secondary, text="📁 Mở Thư Mục template/",
            font=("Segoe UI", 9), bg="#f8fafc", fg="#334155",
            activebackground="#f1f5f9", activeforeground="#0f172a",
            relief="flat", padx=8, pady=6, cursor="hand2",
            highlightthickness=1, highlightbackground="#cbd5e1",
            command=lambda: os.startfile(TEMPLATE_DIR) if os.path.exists(TEMPLATE_DIR) else messagebox.showinfo("Thông báo", "Thư mục template chưa tồn tại.")
        )
        btn_open_tmpl.pack(side="left", expand=True, fill="x", padx=(4, 4))

        btn_open_out = tk.Button(
            row_secondary, text="📁 Mở Thư Mục output_word/",
            font=("Segoe UI", 9), bg="#f8fafc", fg="#334155",
            activebackground="#f1f5f9", activeforeground="#0f172a",
            relief="flat", padx=8, pady=6, cursor="hand2",
            highlightthickness=1, highlightbackground="#cbd5e1",
            command=lambda: os.startfile(OUTPUT_DIR) if os.path.exists(OUTPUT_DIR) else messagebox.showinfo("Thông báo", "Thư mục output_word chưa tồn tại.")
        )
        btn_open_out.pack(side="left", expand=True, fill="x", padx=(4, 0))


        # -------------------------------------------------------------
        # 3. KHUNG NHẬT KÝ TIẾN TRÌNH (DARK SLATE CONSOLE HIỆN ĐẠI)
        # -------------------------------------------------------------
        log_frame = tk.LabelFrame(
            root, text=" 📋 Nhật Ký Hoạt Động & Tiến Trình ",
            font=("Segoe UI", 9, "bold"), fg="#1e293b", bg="#f8fafc",
            padx=10, pady=8, highlightthickness=1, highlightbackground="#e2e8f0"
        )
        log_frame.pack(fill="both", expand=True, padx=16, pady=(6, 12))

        self.txt_log = tk.Text(
            log_frame, wrap="word", font=("Consolas", 9),
            bg="#0f172a", fg="#f8fafc", insertbackground="#38bdf8",
            relief="flat", padx=10, pady=8
        )
        scroll = ttk.Scrollbar(log_frame, command=self.txt_log.yview)
        self.txt_log.configure(yscrollcommand=scroll.set)
        scroll.pack(side="right", fill="y")
        self.txt_log.pack(side="left", fill="both", expand=True)

        self.log(f"✓ Ứng dụng VBSP đã sẵn sàng. Thư mục hoạt động:\n  {BASE_DIR}")
        self.log(f"✓ Máy chủ nền đang chạy tại: http://127.0.0.1:{server_port}")
        self.log("💡 Hướng dẫn nhanh:")
        self.log("  1. Bấm 'MỞ HỆ THỐNG QUẢN LÝ' để nhập liệu hồ sơ và nạp mẫu biểu.")
        self.log("  2. Bấm 'Xuất File Excel 262 Cột' trên web để xuất file Excel dữ liệu.")
        self.log("  3. Bấm 'Xuất File Hàng Loạt' - hệ thống sẽ tự động lấy hồ sơ từ Web / database_records.json để tạo file.")
        self.update_sync_status()
        self.schedule_auto_sync_check()

    def schedule_auto_sync_check(self):
        self.update_sync_status()
        self.root.after(2500, self.schedule_auto_sync_check)

    def update_sync_status(self):
        try:
            if os.path.exists(DB_FILE):
                with open(DB_FILE, 'r', encoding='utf-8') as f:
                    recs = json.load(f)
                if isinstance(recs, list) and len(recs) > 0:
                    latest = recs[0]
                    name = latest.get('hoTen') or 'Khách hàng'
                    can = latest.get('canHoSo')
                    can_txt = f" (Căn: {can})" if can else ""
                    mtime = time.strftime('%H:%M:%S', time.localtime(os.path.getmtime(DB_FILE)))
                    self.lbl_sync_status.config(
                        text=f"🟢 Đã đồng bộ từ Web: {len(recs)} hồ sơ  |  Hồ sơ mới nhất: {name}{can_txt}  (Lúc {mtime})"
                    )
                    return
            self.lbl_sync_status.config(text="⚪ Chưa có hồ sơ nào từ Web (Vui lòng mở Web để nhập liệu)")
        except Exception:
            pass

    def log(self, text):
        self.txt_log.insert(tk.END, text + "\n")
        self.txt_log.see(tk.END)

    def on_open_web(self):
        self.log("\n🌐 Đang mở Hệ Thống Quản Lý Hồ Sơ trên trình duyệt...")
        open_web_app()

    def on_add_template_files(self):
        files = filedialog.askopenfilenames(
            title="Chọn các file mẫu Word (.docx) để thêm vào hệ thống",
            filetypes=[("Word Document (*.docx)", "*.docx")]
        )
        if not files: return
        
        def handle_success(count):
            self.log(f"🔄 Đang tự động quét và đồng bộ {count} mẫu Word mới vào hệ thống HTML...")
            success, msg = run_sync_templates(log_callback=self.log)
            if success:
                messagebox.showinfo("Thành công", f"Đã nạp {count} file vào thư mục template/ và đồng bộ xong!")
            else:
                messagebox.showwarning("Thông báo", msg)

        AddTemplateConfigDialog(self.root, files, handle_success)

    def on_sync_templates(self):
        def task():
            success, msg = run_sync_templates(log_callback=self.log)
            if success:
                messagebox.showinfo("Hoàn tất", msg)
            else:
                messagebox.showwarning("Thông báo", msg)
        threading.Thread(target=task, daemon=True).start()

    def on_combine_records(self):
        """Chức năng gộp dữ liệu database_records.json từ nhiều máy khác nhau."""
        files = filedialog.askopenfilenames(
            title="Chọn các file database_records.json cần gộp (từ các máy)",
            filetypes=[("Tệp JSON (*.json)", "*.json"), ("Tất cả tệp", "*.*")]
        )
        if not files:
            return

        choice_dedup = messagebox.askyesnocancel(
            "Tùy chọn lọc trùng lặp",
            f"Bạn đã chọn {len(files)} file dữ liệu để gộp vào hệ thống.\n\n"
            "Bạn có muốn TỰ ĐỘNG LỌC TRÙNG (dựa theo số CCCD / Tên khách hàng) không?\n\n"
            "• Chọn 'Yes' (Khuyên dùng): Bỏ qua hồ sơ trùng, giữ hồ sơ hiện có.\n"
            "• Chọn 'No': Ghi đè / cập nhật hồ sơ trùng bằng bản mới.\n"
            "• Chọn 'Cancel': Hủy tiến trình gộp."
        )
        if choice_dedup is None:
            return

        def task():
            self.log("\n" + "="*50)
            self.log(f"📦 BẮT ĐẦU TIẾN TRÌNH GỘP DỮ LIỆU TỪ {len(files)} FILE...")

            # 1. Đọc dữ liệu hiện tại trong hệ thống
            current_records = []
            current_rows206 = []
            if os.path.exists(DB_FILE):
                try:
                    with open(DB_FILE, 'r', encoding='utf-8') as f:
                        data = json.load(f)
                        if isinstance(data, list):
                            current_records = data
                except Exception as e:
                    self.log(f"⚠ Lỗi đọc DB_FILE hiện tại: {e}")

            rows206_file = os.path.join(BASE_DIR, "database_records_rows206.json")
            if os.path.exists(rows206_file):
                try:
                    with open(rows206_file, 'r', encoding='utf-8') as rf:
                        r206 = json.load(rf)
                        if isinstance(r206, list):
                            current_rows206 = r206
                except Exception:
                    pass

            # 2. Tự động sao lưu bản cũ trước khi gộp để đảm bảo an toàn tuyệt đối
            if current_records:
                backup_ts = time.strftime('%Y%m%d_%H%M%S')
                backup_path = os.path.join(BASE_DIR, f"database_records_backup_{backup_ts}.json")
                try:
                    shutil.copy2(DB_FILE, backup_path)
                    self.log(f"💾 Đã tạo bản sao lưu an toàn: {os.path.basename(backup_path)}")
                except Exception:
                    pass

            def get_rec_key(r):
                cccd = str(r.get('cccd') or r.get('soCCCD') or r.get('f_soCCCD') or '').strip().replace(' ', '').replace('.', '').replace('-', '')
                if cccd and len(cccd) >= 8:
                    return f"cccd_{cccd}"
                name = str(r.get('hoTen') or r.get('f_hoTen') or '').strip().lower()
                can = str(r.get('canHoSo') or r.get('canHo') or '').strip().lower()
                proj = str(r.get('project') or r.get('duAn') or '').strip().lower()
                if name:
                    return f"name_{name}_{can}_{proj}"
                return f"id_{r.get('_id', '')}"

            # Xây dựng danh bạ hồ sơ hiện tại
            record_map = {}
            for r in current_records:
                if isinstance(r, dict):
                    k = get_rec_key(r)
                    record_map[k] = r

            new_count = 0
            updated_count = 0
            duplicate_skipped = 0
            total_read = 0

            # 3. Lần lượt đọc từng file được chọn
            for fpath in files:
                fname = os.path.basename(fpath)
                self.log(f"📂 Đang xử lý: {fname}...")
                try:
                    with open(fpath, 'r', encoding='utf-8') as f:
                        file_data = json.load(f)
                except Exception as e:
                    self.log(f"   ❌ Không đọc được file: {e}")
                    continue

                records_in_file = []
                if isinstance(file_data, list):
                    records_in_file = file_data
                elif isinstance(file_data, dict):
                    if "records" in file_data and isinstance(file_data["records"], list):
                        records_in_file = file_data["records"]
                    else:
                        records_in_file = [file_data]

                if not isinstance(records_in_file, list):
                    self.log(f"   ⚠ Cấu trúc file {fname} không phải danh sách hồ sơ hợp lệ.")
                    continue

                self.log(f"   ✓ Tìm thấy {len(records_in_file)} hồ sơ trong file.")
                total_read += len(records_in_file)

                for r in records_in_file:
                    if not isinstance(r, dict):
                        continue
                    k = get_rec_key(r)
                    if k in record_map:
                        if choice_dedup:
                            duplicate_skipped += 1
                        else:
                            merged_item = dict(record_map[k])
                            merged_item.update(r)
                            record_map[k] = merged_item
                            updated_count += 1
                    else:
                        record_map[k] = r
                        new_count += 1

            # 4. Tạo danh sách hồ sơ tổng sau khi gộp
            merged_records = list(record_map.values())

            # 5. Lưu lại vào DB_FILE
            try:
                safe_unhide_file(DB_FILE)
                with open(DB_FILE, 'w', encoding='utf-8') as f:
                    json.dump(merged_records, f, ensure_ascii=False, indent=2)
                safe_rehide_file(DB_FILE)
                self.log(f"💾 Đã lưu {len(merged_records)} hồ sơ vào database_records.json")
            except Exception as e:
                self.log(f"❌ Lỗi ghi database_records.json: {e}")
                messagebox.showerror("Lỗi", f"Không thể ghi file database_records.json: {e}")
                return

            # 6. Đồng bộ 100% hồ sơ vào database_records_rows206.json và DATA_MAIL_MERGE_SYNC.xlsx
            try:
                merged_rows206 = sync_all_database_records(merged_records)
                self.log(f"📊 Đã đồng bộ {len(merged_rows206)} hồ sơ vào DATA_MAIL_MERGE_SYNC.xlsx & cache rows206")
            except Exception as ex:
                self.log(f"⚠ Không thể tạo file Excel sync: {ex}")

            # 7. Cập nhật trạng thái giao diện
            self.update_sync_status()
            self.log("\n" + "="*50)
            self.log("🎉 TỔNG KẾT GỘP DỮ LIỆU THÀNH CÔNG:")
            self.log(f"   • Số file đã gộp: {len(files)} file")
            self.log(f"   • Tổng hồ sơ đọc được: {total_read} hồ sơ")
            self.log(f"   • Hồ sơ mới thêm: +{new_count}")
            if choice_dedup:
                self.log(f"   • Hồ sơ trùng đã bỏ qua: {duplicate_skipped}")
            else:
                self.log(f"   • Hồ sơ trùng đã cập nhật: {updated_count}")
            self.log(f"   • TỔNG SỐ HỒ SƠ TRONG HỆ THỐNG: {len(merged_records)} HỒ SƠ")
            self.log("💡 Mẹo: Bấm 'MỞ HỆ THỐNG QUẢN LÝ' để xem danh sách tổng hoặc bấm 'Xuất File Hàng Loạt'!")

            msg_box_text = (
                f"Đã gộp thành công dữ liệu từ {len(files)} file!\n\n"
                f"• Hồ sơ mới thêm: +{new_count}\n"
                f"• Hồ sơ trùng lặp: {duplicate_skipped if choice_dedup else updated_count}\n"
                f"• Tổng số hồ sơ trong hệ thống: {len(merged_records)} hồ sơ\n\n"
                f"Dữ liệu đã sẵn sàng để Xuất File Hàng Loạt hoặc mở xem trên Web."
            )
            messagebox.showinfo("Gộp Dữ Liệu Thành Công", msg_box_text)

        threading.Thread(target=task, daemon=True).start()


    def on_mail_merge(self):
        # 1. Chỉ sử dụng dữ liệu thật từ database_records.json (không tự tạo hồ sơ giả)
        db_recs = []
        if os.path.exists(DB_FILE):
            try:
                with open(DB_FILE, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                    if isinstance(data, list):
                        db_recs = data
            except Exception:
                db_recs = []

        sync_file = os.path.join(BASE_DIR, "DATA_MAIL_MERGE_SYNC.xlsx")
        rows206_file = os.path.join(BASE_DIR, "database_records_rows206.json")

        if len(db_recs) == 0:
            if os.path.exists(sync_file):
                try: os.remove(sync_file)
                except Exception: pass
            if os.path.exists(rows206_file):
                try: os.remove(rows206_file)
                except Exception: pass

            messagebox.showinfo(
                "Chưa có dữ liệu hồ sơ",
                "Hệ thống hiện tại chưa có hồ sơ nào trong database_records.json.\n\n"
                "Vui lòng bấm nút 'MỞ HỆ THỐNG QUẢN LÝ' để nhập hồ sơ khách hàng thật trên Web trước khi xuất file hàng loạt."
            )
            return

        # 2. Luôn đảm bảo 100% hồ sơ trong database_records.json đều được đồng bộ tươi mới
        # vào DATA_MAIL_MERGE_SYNC.xlsx và database_records_rows206.json
        # Điều này đảm bảo khi người dùng sửa bất kỳ thông tin nào trên Web, xuất Word sẽ ăn ngay lập tức!
        try:
            sync_all_database_records(db_recs)
        except Exception as ex_sync:
            self.log(f"⚠️ Cảnh báo đồng bộ Excel: {ex_sync}")

        self.open_select_dialog(sync_file)

    def open_select_dialog(self, excel_path):
        def do_export(fpath, selected_indices, folder_mapping=None, selected_templates=None, marital_mapping=None):
            def task():
                success, msg_res = run_mail_merge(
                    excel_path=fpath,
                    selected_indices=selected_indices,
                    folder_mapping=folder_mapping,
                    marital_mapping=marital_mapping,
                    selected_templates=selected_templates,
                    log_callback=self.log
                )
                if success:
                    if messagebox.askyesno("Thành công", f"{msg_res}\n\nBạn có muốn mở ngay thư mục chứa các file Word đã xuất không?"):
                        try:
                            os.startfile(OUTPUT_DIR)
                        except Exception:
                            pass
                else:
                    messagebox.showwarning("Thông báo", msg_res)
            threading.Thread(target=task, daemon=True).start()

        SelectRecordsDialog(
            parent=self.root,
            excel_path=excel_path,
            on_submit=do_export,
            on_browse_another=self.browse_and_open_dialog
        )

    def browse_and_open_dialog(self):
        fpath = filedialog.askopenfilename(
            title="Chọn file Excel nguồn để trộn thư (sheet 'data')",
            initialdir=BASE_DIR,
            filetypes=[("Excel Files (*.xlsx, *.xlsm)", "*.xlsx;*.xlsm"), ("Tất cả tập tin", "*.*")]
        )
        if fpath:
            self.open_select_dialog(fpath)

def main():
    if len(sys.argv) > 1:
        arg = sys.argv[1].lower()
        if arg in ["--sync", "-s"]:
            run_sync_templates()
            return
        elif arg in ["--merge", "-m"]:
            run_mail_merge()
            return
        elif arg in ["--web", "-w"]:
            open_web_app()
            return

    root = tk.Tk()
    app = AppGUI(root)
    root.mainloop()

if __name__ == "__main__":
    main()
